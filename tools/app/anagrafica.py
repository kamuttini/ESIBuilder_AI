"""Reader/writer for `encoding_struct*.xlsx`, the project's registry of echos and probes.

Why not openpyxl for reading: the sheets are read on every request and the stdlib zip+xml
path is fast and dependency free. Writing a new row *does* use openpyxl, so the rest of the
workbook (formatting, other sheets) survives untouched, and a timestamped backup is taken
before saving.

Careful: empty cells are omitted in the XML, so every value is placed by its column letter
(`A2`, `G2`), never by position. Reading the FSS sheet positionally misaligns the columns.
"""

from __future__ import annotations

import re
import shutil
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from xml.etree import ElementTree as ET

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

SHEET_ECHO = "ECO"
SHEET_PROBE = "PROBE"
SHEET_PROBE_TYPE = "PROBE TYPE"
SHEET_NDG = "NDG"
SHEET_FSS = "FSS"

COL_ECHO_MODEL, COL_ECHO_ID = "A", "B"
COL_PROBE_MODEL, COL_PROBE_ID = "A", "B"
# FSS sheet: modello eco | modello sonda | modello ndg | ID ECHO | ID SONDA | ID NDG | tipo | FSS
FSS_COLUMNS = {
    "echo_model": "A",
    "probe_model": "B",
    "ndg_model": "C",
    "id_echo": "D",
    "id_probe": "E",
    "id_ndg": "F",
    "probe_type": "G",
    "fss": "H",
    "note": "I",
    "resolution": "K",
}

_COL_RE = re.compile(r"^([A-Z]+)")


def _column(ref: str) -> str:
    match = _COL_RE.match(ref or "")
    return match.group(1) if match else ""


def normalize(text: str) -> str:
    """Lowercase, drop punctuation, collapse spaces: 'Esaote, MyLab X8 eHD' -> 'esaote mylab x8 ehd'."""
    lowered = (text or "").lower()
    lowered = re.sub(r"[^a-z0-9]+", " ", lowered)
    return " ".join(lowered.split())


def _tokens(text: str) -> List[str]:
    return normalize(text).split()


def similarity(query: str, candidate: str) -> float:
    """Token overlap weighted towards the query, plus a bonus for a contained string."""
    q, c = _tokens(query), _tokens(candidate)
    if not q or not c:
        return 0.0
    qs, cs = set(q), set(c)
    overlap = len(qs & cs)
    score = overlap / len(qs)  # how much of what the user typed was found
    score *= 0.6 + 0.4 * (overlap / len(cs))  # penalise very long unrelated entries
    if normalize(query) and normalize(query) in normalize(candidate):
        score = max(score, 0.9)
    if normalize(candidate) and normalize(candidate) in normalize(query):
        score = max(score, 0.85)
    return round(min(score, 1.0), 4)


@dataclass
class Match:
    value: str
    id: Optional[int]
    score: float
    row: int

    def as_dict(self) -> Dict:
        return {"value": self.value, "id": self.id, "score": self.score, "row": self.row}


class Anagrafica:
    """The workbook, read once per instance."""

    def __init__(self, path: Path):
        self.path = Path(path)
        self._sheets: Dict[str, List[Dict[str, str]]] = {}
        self._load()

    # -- reading -----------------------------------------------------------
    def _load(self) -> None:
        with zipfile.ZipFile(self.path) as archive:
            shared: List[str] = []
            if "xl/sharedStrings.xml" in archive.namelist():
                root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
                shared = [
                    "".join(node.text or "" for node in si.iter(f"{NS}t"))
                    for si in root.iter(f"{NS}si")
                ]
            workbook = ET.fromstring(archive.read("xl/workbook.xml"))
            rels = {
                rel.get("Id"): rel.get("Target")
                for rel in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            }
            targets = {
                sheet.get("name"): "xl/" + rels[sheet.get(f"{REL}id")].lstrip("/").replace("xl/", "")
                for sheet in workbook.iter(f"{NS}sheet")
            }

            def cell_value(cell) -> str:  # noqa: ANN001
                value = cell.find(f"{NS}v")
                if cell.get("t") == "s" and value is not None:
                    index = int(value.text)
                    return shared[index] if index < len(shared) else ""
                if cell.get("t") == "inlineStr":
                    return "".join(node.text or "" for node in cell.iter(f"{NS}t"))
                return value.text if value is not None else ""

            for name, target in targets.items():
                if name not in (SHEET_ECHO, SHEET_PROBE, SHEET_PROBE_TYPE, SHEET_NDG, SHEET_FSS):
                    continue
                sheet = ET.fromstring(archive.read(target))
                rows: List[Dict[str, str]] = []
                for row in sheet.iter(f"{NS}row"):
                    cells = {
                        _column(cell.get("r", "")): (cell_value(cell) or "").strip()
                        for cell in row.iter(f"{NS}c")
                    }
                    cells["__row__"] = row.get("r", "")
                    rows.append(cells)
                self._sheets[name] = rows

    def _data_rows(self, sheet: str) -> List[Dict[str, str]]:
        return self._sheets.get(sheet, [])[1:]

    # -- lookups -----------------------------------------------------------
    def probe_types(self) -> Dict[int, str]:
        out: Dict[int, str] = {}
        for row in self._data_rows(SHEET_PROBE_TYPE):
            raw_id, description = row.get("A", ""), row.get("B", "")
            if raw_id.isdigit() and description:
                out[int(raw_id)] = description
        return out

    def _match_column(
        self, sheet: str, value_col: str, id_col: str, query: str, limit: int = 5
    ) -> List[Match]:
        if not (query or "").strip():
            return []
        matches: List[Match] = []
        for row in self._data_rows(sheet):
            value = row.get(value_col, "")
            if not value:
                continue
            score = similarity(query, value)
            if score <= 0:
                continue
            raw_id = row.get(id_col, "")
            matches.append(
                Match(
                    value=value,
                    id=int(raw_id) if raw_id.lstrip("-").isdigit() else None,
                    score=score,
                    row=int(row.get("__row__") or 0),
                )
            )
        matches.sort(key=lambda m: (-m.score, m.value))
        return matches[:limit]

    def match_echo(self, query: str, limit: int = 5) -> List[Match]:
        return self._match_column(SHEET_ECHO, COL_ECHO_MODEL, COL_ECHO_ID, query, limit)

    def match_probe(self, query: str, limit: int = 5) -> List[Match]:
        return self._match_column(SHEET_PROBE, COL_PROBE_MODEL, COL_PROBE_ID, query, limit)

    def fss_rows(self) -> List[Dict[str, str]]:
        out = []
        for row in self._data_rows(SHEET_FSS):
            entry = {key: row.get(col, "") for key, col in FSS_COLUMNS.items()}
            entry["__row__"] = row.get("__row__", "")
            if entry["echo_model"] or entry["probe_model"]:
                out.append(entry)
        return out

    def match_combination(self, echo_query: str, probe_query: str, limit: int = 5) -> List[Dict]:
        """Existing echo+probe setups: the richest hit, it carries every id at once."""
        scored = []
        for row in self.fss_rows():
            echo_score = similarity(echo_query, row["echo_model"]) if echo_query else 0.0
            probe_score = similarity(probe_query, row["probe_model"]) if probe_query else 0.0
            if echo_score <= 0 or probe_score <= 0:
                continue
            entry = dict(row)
            entry["score"] = round((echo_score + probe_score) / 2, 4)
            entry["echo_score"] = echo_score
            entry["probe_score"] = probe_score
            scored.append(entry)
        scored.sort(key=lambda item: -item["score"])
        return scored[:limit]

    def probe_by_id(self, probe_id: int) -> Dict:
        """From the probe ID (what the probe classifier predicts) to model and probe type.

        Two types for the same probe (3 and 4) mean a transrectal biplane: which of the two
        applies depends on the plane of this setup, and that is what the L/T stage decides.
        """
        model = ""
        for row in self._data_rows(SHEET_PROBE):
            raw = row.get(COL_PROBE_ID, "")
            if raw.lstrip("-").isdigit() and int(raw) == int(probe_id):
                model = row.get(COL_PROBE_MODEL, "")
                break

        types: Dict[int, int] = {}
        echo_models: List[str] = []
        for row in self.fss_rows():
            raw_id = row.get("id_probe", "")
            if not (raw_id.lstrip("-").isdigit() and int(raw_id) == int(probe_id)):
                continue
            raw_type = row.get("probe_type", "")
            if raw_type.lstrip("-").isdigit():
                types[int(raw_type)] = types.get(int(raw_type), 0) + 1
            if row.get("echo_model"):
                echo_models.append(row["echo_model"])

        ordered = sorted(types.items(), key=lambda item: -item[1])
        return {
            "probe_id": int(probe_id),
            "model": model,
            "probe_types": [t for t, _ in ordered],
            "probe_type_counts": types,
            "biplane": set(types) >= {3, 4},
            "echo_models": echo_models[:8],
        }

    def combinations_for_probe(
        self, probe_id: int, probe_type: Optional[int] = None
    ) -> List[Dict]:
        """Registered setups for this probe: the list the user picks the machine from.

        The vendor network recognises the *brand*, not the machine model, and the same probe
        appears with dozens of different ID ECHO (68 rows and 30 distinct ID ECHO for TLC3-13).
        So #02 cannot be derived from brand + probe: it has to be chosen.
        """
        out = []
        for row in self.fss_rows():
            raw_id = row.get("id_probe", "")
            if not (raw_id.lstrip("-").isdigit() and int(raw_id) == int(probe_id)):
                continue
            raw_type = row.get("probe_type", "")
            if probe_type is not None and raw_type.lstrip("-").isdigit():
                if int(raw_type) != int(probe_type):
                    continue
            out.append(
                {
                    "echo_model": row.get("echo_model", ""),
                    "probe_model": row.get("probe_model", ""),
                    "id_echo": row.get("id_echo", ""),
                    "id_ndg": row.get("id_ndg", ""),
                    "id_setup": row.get("fss", ""),
                    "probe_type": raw_type,
                    "resolution": row.get("resolution", ""),
                    "row": row.get("__row__", ""),
                }
            )
        return sorted(out, key=lambda item: item["echo_model"].lower())

    def next_id(self, sheet: str, id_col: str) -> int:
        used = [
            int(row[id_col])
            for row in self._data_rows(sheet)
            if row.get(id_col, "").lstrip("-").isdigit()
        ]
        return (max(used) + 1) if used else 0

    def resolve(self, echo_query: str, probe_query: str, min_score: float = 0.6) -> Dict:
        """What the codes page should propose, and what is still missing.

        Never invents an id: when nothing matches it reports the id a new row would get, and
        the caller decides whether to write it.
        """
        combinations = self.match_combination(echo_query, probe_query)
        echo_matches = self.match_echo(echo_query)
        probe_matches = self.match_probe(probe_query)

        # A single best score is not enough: if the top rows disagree on ID ECHO the answer
        # is "choose", not a guess.
        top_score = combinations[0]["score"] if combinations else 0.0
        tied = [row for row in combinations if abs(row["score"] - top_score) < 1e-9]
        distinct_echo = {row.get("id_echo", "") for row in tied}
        ambiguous = len(distinct_echo) > 1
        if not combinations and echo_matches and len(echo_matches) > 1:
            if (echo_matches[0].score - echo_matches[1].score) < 0.10:
                ambiguous = True
        best_combo = (
            combinations[0]
            if combinations and top_score >= min_score and not ambiguous
            else None
        )
        # Same rule for the single-sheet fallback: a bare brand ("Esaote") scores high against
        # many ECO rows, so a winner without a clear margin is not an answer.
        def _clear_winner(matches: List[Match], margin: float = 0.10) -> Optional[Match]:
            if not matches or matches[0].score < min_score:
                return None
            if len(matches) > 1 and (matches[0].score - matches[1].score) < margin:
                return None
            return matches[0]

        best_echo = _clear_winner(echo_matches)
        best_probe = _clear_winner(probe_matches)

        proposal: Dict[str, Optional[int]] = {"id_echo": None, "id_probe": None, "id_ndg": None,
                                              "probe_type": None, "id_setup": None}
        sources: Dict[str, str] = {}
        if best_combo:
            for key, field in (("id_echo", "id_echo"), ("id_probe", "id_probe"),
                               ("id_ndg", "id_ndg"), ("probe_type", "probe_type"),
                               ("id_setup", "fss")):
                raw = best_combo.get(field, "")
                if raw.lstrip("-").isdigit():
                    proposal[key] = int(raw)
                    sources[key] = f"FSS riga {best_combo['__row__']}"
        if proposal["id_echo"] is None and best_echo and best_echo.id is not None:
            proposal["id_echo"] = best_echo.id
            sources["id_echo"] = f"ECO riga {best_echo.row}"
        if proposal["id_probe"] is None and best_probe and best_probe.id is not None:
            proposal["id_probe"] = best_probe.id
            sources["id_probe"] = f"PROBE riga {best_probe.row}"

        missing = []
        if proposal["id_echo"] is None:
            missing.append({"sheet": SHEET_ECHO, "field": "id_echo", "query": echo_query,
                            "next_id": self.next_id(SHEET_ECHO, COL_ECHO_ID)})
        if proposal["id_probe"] is None:
            missing.append({"sheet": SHEET_PROBE, "field": "id_probe", "query": probe_query,
                            "next_id": self.next_id(SHEET_PROBE, COL_PROBE_ID)})

        return {
            "proposal": proposal,
            "sources": sources,
            "missing": missing,
            "ambiguous": ambiguous,
            "ambiguous_ids": sorted(distinct_echo) if ambiguous else [],
            "combination_matches": combinations,
            "echo_matches": [m.as_dict() for m in echo_matches],
            "probe_matches": [m.as_dict() for m in probe_matches],
        }

    # -- writing -----------------------------------------------------------
    def backup(self) -> Path:
        """Timestamped copy. Two writes in the same second must not overwrite each other."""
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        target = self.path.with_name(f"{self.path.stem}.backup_{stamp}{self.path.suffix}")
        counter = 2
        while target.exists():
            target = self.path.with_name(
                f"{self.path.stem}.backup_{stamp}_{counter}{self.path.suffix}"
            )
            counter += 1
        shutil.copy2(self.path, target)
        return target

    def append_row(self, sheet: str, values: Dict[str, str]) -> Dict:
        """Append a row keyed by column letter. Takes a backup first, then reloads."""
        try:
            from openpyxl import load_workbook
        except ImportError as error:  # pragma: no cover - optional dependency
            raise RuntimeError(
                "per scrivere in encoding_struct serve openpyxl (pip install openpyxl)"
            ) from error

        backup = self.backup()
        workbook = load_workbook(self.path)
        if sheet not in workbook.sheetnames:
            raise KeyError(f"foglio assente: {sheet}")
        worksheet = workbook[sheet]

        # First fully empty row, so the pre-allocated blank rows of the sheet get reused.
        target_row = worksheet.max_row + 1
        for index in range(2, worksheet.max_row + 2):
            if all(worksheet.cell(row=index, column=col).value in (None, "") for col in range(1, 9)):
                target_row = index
                break

        for column, value in values.items():
            worksheet[f"{column}{target_row}"] = value
        workbook.save(self.path)
        self._sheets.clear()
        self._load()
        return {"sheet": sheet, "row": target_row, "backup": str(backup), "values": values}

    def add_echo(self, model: str) -> Dict:
        new_id = self.next_id(SHEET_ECHO, COL_ECHO_ID)
        result = self.append_row(SHEET_ECHO, {COL_ECHO_MODEL: model, COL_ECHO_ID: str(new_id)})
        result["id"] = new_id
        return result

    def add_probe(self, model: str) -> Dict:
        new_id = self.next_id(SHEET_PROBE, COL_PROBE_ID)
        result = self.append_row(SHEET_PROBE, {COL_PROBE_MODEL: model, COL_PROBE_ID: str(new_id)})
        result["id"] = new_id
        return result

    def add_combination(self, entry: Dict[str, str]) -> Dict:
        values = {
            FSS_COLUMNS[key]: str(entry[key])
            for key in FSS_COLUMNS
            if entry.get(key) not in (None, "")
        }
        return self.append_row(SHEET_FSS, values)


def default_path(repo_root: Path) -> Optional[Path]:
    """The newest `encoding_struct*.xlsx` in the repo root."""
    candidates = sorted(Path(repo_root).glob("encoding_struct*.xlsx"))
    candidates = [c for c in candidates if ".backup_" not in c.name]
    return candidates[-1] if candidates else None
