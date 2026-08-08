#!/usr/bin/env python3
"""Predict key .fss head lines from raw acquisition folders.

Target lines:
- #01 VERSION (fixed to 4.0)
- #02 ID_ECHO (resolved from predicted vendor + historical mapping)
- #03 ID_PROBE (predicted from probe classifier; temporary stage before dedicated probe-recognition pipeline)
- #06..#10 VIDEO metadata (from acquisition filename pattern)
- #11 RECT_ECHO (predicted by rect detector)
- #12 GROUP_ORIENTATION (symbol/depth group selector; currently emitted as symbol with orientation/depth evidence)
- #13 RECT_NAME_ECHO (temporary: resolved from historical .fss by vendor/video, with interactive fallback)
- #14 RECT_NAME_PROBE (temporary: resolved from historical .fss by vendor/probe/video, with interactive fallback)
- #16 RECT_ORIENTATION (official orientation evidence from SU/GIU + LR marker)

Additional official folder-image stages:
- SU/GIU orientation classification on rect crops (one prediction per frame crop).
- LR marker orientation recognition (folder-level #16 plus per-frame evidence).
- L/T orientation classification on rect crops (one prediction per frame crop).
- RECT_DEPTH autonomous recognition (OCR/classical/ranker module, one prediction per frame).
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import statistics
import subprocess
import sys
import tempfile
import time
import types
import zipfile
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np
import torch
import torch.nn as nn
import torch.nn.functional as F
from PIL import Image
from torchvision.models import resnet18
from torchvision.transforms import InterpolationMode
from torchvision.transforms import functional as TF

try:
    from rect_red_pipeline import compute_rect_red_pipeline
except ModuleNotFoundError:
    from tools.ultrasound.rect_red_pipeline import compute_rect_red_pipeline
try:
    from probe_type_router import ProbeTypeRouter
except ModuleNotFoundError:
    from tools.ultrasound.probe_type_router import ProbeTypeRouter
from train_ultrasound_probe_classifier import ProbeClassifier
from train_ultrasound_rect_net import RectRegressor, normalize_box_order
from train_ultrasound_vendor_classifier import VendorClassifier, choose_device

SCRIPT_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parent.parent
DEFAULT_LR_MARKER_VENDOR_LIBRARY = REPO_ROOT / "artifacts/10_active_pipeline/pipeline_fss_head/lr_marker_vendor_template_library"
DEFAULT_LR_MARKER_REVIEW_FILE = DEFAULT_LR_MARKER_VENDOR_LIBRARY / "review_decisions.json"
DEFAULT_ORIENTATION_MARKER_BUNDLE_ZIP = Path("/Users/Shared/41_orientation_marker_detector_bundle.zip")
DEFAULT_ORIENTATION_MARKER_BUNDLE_DIR = REPO_ROOT / "artifacts/41_orientation_marker_detector_bundle"
DEFAULT_ORIENTATION_MARKER_BUNDLE_LIBRARY_ROOT = (
    DEFAULT_ORIENTATION_MARKER_BUNDLE_DIR / "orientation_marker_detector" / "templates"
)
ORIENTATION_TOOL_DIR = Path(__file__).resolve().parent.parent / "orientation"
if ORIENTATION_TOOL_DIR.is_dir() and ORIENTATION_TOOL_DIR.as_posix() not in sys.path:
    sys.path.insert(0, ORIENTATION_TOOL_DIR.as_posix())
try:
    from prepare_lr_marker_dataset_sugiu_v2 import (  # type: ignore
        SKIP_DIRS as LR_MARKER_SKIP_DIRS,
        TEMPLATE_RE as LR_MARKER_TEMPLATE_RE,
        _expand_rect as _lr_marker_expand_rect,
        _load_template as _lr_marker_load_template,
        _locate_match_in_abs_rect as _lr_marker_locate_match_in_abs_rect,
        _lr_from_detected_side as _lr_marker_label_from_side,
        _parse_float_steps as _lr_marker_parse_float_steps,
        _roi_from_sugiu as _lr_marker_roi_from_sugiu,
        _score_of as _lr_marker_score_of,
        infer_manufacturer as _lr_marker_infer_manufacturer,
    )
except Exception:  # pragma: no cover - optional runtime feature
    LR_MARKER_SKIP_DIRS = set()
    LR_MARKER_TEMPLATE_RE = None
    _lr_marker_expand_rect = None
    _lr_marker_load_template = None
    _lr_marker_locate_match_in_abs_rect = None
    _lr_marker_label_from_side = None
    _lr_marker_parse_float_steps = None
    _lr_marker_roi_from_sugiu = None
    _lr_marker_score_of = None
    _lr_marker_infer_manufacturer = None


IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp", ".gif"}
CAPTURE_INPUT_TOKENS = {"hdmi", "vga"}

# Opt-in progress channel for the review tool (`--stage-events`). Off by default: with the
# flag absent the pipeline prints exactly what it printed before.
_STAGE_EVENTS_ENABLED = False


def _emit_stage_event(stage: str, folder_name: str, **payload: object) -> None:
    """Print one ``##STAGE {json}`` line per stage per folder, when enabled.

    The review tool reads these while the folder is still running, so the user sees vendor /
    probe / rect / depth / scala appear one after the other instead of waiting for the CSV,
    which the pipeline only writes when every folder is done.
    """
    if not _STAGE_EVENTS_ENABLED:
        return
    event = {
        "stage": stage,
        "folder": folder_name,
        "ts_utc": datetime.now(timezone.utc).isoformat(),
    }
    event.update(payload)
    try:
        text = json.dumps(event, ensure_ascii=False, default=str)
    except Exception:  # noqa: BLE001 - a progress line must never break a run
        return
    print(f"##STAGE {text}", flush=True)
CAPTURE_RES_TOKEN_RE = re.compile(r"^(\d{3,5})[xX](\d{3,5})$")
CAPTURE_RES_ANY_RE = re.compile(r"(\d{3,5})[xX](\d{3,5})")
RECT_ECHO_RE = re.compile(r"^\s*(\d+)\|(\d+)\|(\d+)\|(\d+)\|\s*$")
RECT_NAME_ECHO_PREFIX_RE = re.compile(r"^\s*(\d+)\|(\d+)\|(\d+)\|(\d+)\|")
LINE13_DISABLED_TAIL = "0|0:0.000000:0:0:0:0:0|0|"
LINE16_PENDING_MATCH_TAIL = "1|0:0.000000:0:0:0:0:0|0|"
LR_MARKER_RELIABLE_MATCH_SCORE = 0.62
LR_MARKER_MIN_TEXTURE_STD = 3.0
LR_MARKER_MIN_HISTORICAL_TEMPLATE_AREA = 256
LR_MARKER_MIN_HISTORICAL_TEMPLATE_SHORT_SIDE = 12
LR_MARKER_MIN_RELIABLE_ROW_RATIO = 0.55
LR_MARKER_MAX_BLANK_ROW_RATIO = 0.25
LR_MARKER_MAX_LOW_TEXTURE_ROW_RATIO = 0.25
LR_MARKER_MIN_DERIVED_TEMPLATE_AREA = 160
LR_MARKER_MIN_DERIVED_TEMPLATE_SHORT_SIDE = 8
LR_MARKER_MIN_MANUAL_TEMPLATE_AREA = 80
LR_MARKER_MIN_MANUAL_TEMPLATE_SHORT_SIDE = 6
LR_MARKER_DUAL_SUGIU_SEARCH_CONFIDENCE = 0.90
# Expected acquisition filename pattern: timestamp_<vga|hdmi>_<WxH>[...]
CAPTURE_FILENAME_PATTERN_RE = re.compile(
    r"^.+_(vga|hdmi)_(\d{3,5})[xX](\d{3,5})(?:[_\-].*)?$",
    flags=re.IGNORECASE,
)
GROUP_ORIENTATION_SYMBOL = 4
GROUP_ORIENTATION_DEPTH = 5
GROUP_ORIENTATION_LABEL = {
    GROUP_ORIENTATION_SYMBOL: "symbol",
    GROUP_ORIENTATION_DEPTH: "depth",
}
ROTATION_VALUES_CLOCKWISE = (0, 90, 180, 270)
TESSERACT_OSD_ROTATE_RE = re.compile(r"Rotate:\s*(\d+)")
TESSERACT_OSD_CONF_RE = re.compile(r"Orientation confidence:\s*([0-9]*\.?[0-9]+)")
VENDOR_OCR_GENERIC_TOKENS = {
    "ultrasound",
    "medical",
    "imaging",
    "system",
    "scanner",
    "echo",
    "eco",
    "model",
    "probe",
}
VENDOR_OCR_SHORT_ALIAS_ALLOWLIST = {
    "bk",
}
SU_GIU_LABELS = ("su", "giu")
LT_LABELS = ("L", "T")


@dataclass(frozen=True)
class FolderPrediction:
    folder_path: Path
    folder_name: str
    images_total: int
    images_total_raw: int
    images_duplicates_removed: int
    images_used_cls: int
    images_used_rect: int
    line_01_version: str
    line_02_id_echo: str
    line_02_source: str
    line_02_support: float
    line_03_id_probe: str
    line_03_probe_conf: float
    line_04_probe_type: str
    line_04_probe_type_source: str
    line_04_probe_type_strategy: str
    line_04_probe_type_needs_secondary_model: str
    line_04_probe_type_candidate_ints: str
    line_06_video_input: str
    line_06_video_input_code: str
    line_07_video_input_size_x: str
    line_08_video_input_size_y: str
    line_09_video_x_size: str
    line_10_video_y_size: str
    line_11_rect_echo: str
    line_11_top: str
    line_11_left: str
    line_11_bottom: str
    line_11_right: str
    line_11_source: str
    line_11_model_checkpoint: str
    line_11_method: str
    line_11_rect_red_margin_pct: float
    line_11_rect_red_winner_group: str
    line_12_group_orientation: str
    line_12_group_orientation_label: str
    line_12_source: str
    line_12_support: float
    su_giu_images_predicted: int
    su_giu_majority_label: str
    su_giu_majority_vote_ratio: float
    su_giu_mean_confidence: float
    su_giu_mean_prob_su: float
    su_giu_mean_prob_giu: float
    su_giu_source: str
    su_giu_checkpoint: str
    lr_marker_images_predicted: int
    lr_marker_majority_label: str
    lr_marker_majority_vote_ratio: float
    lr_marker_best_label: str
    lr_marker_best_score: float
    lr_marker_best_template_path: str
    lr_marker_best_search_strategy: str
    lr_marker_source: str
    lr_marker_method: str
    line_16_rect_orientation: str
    line_16_source: str
    line_16_marker_boxes_count: int
    line_16_groups_json: str
    rect_depth_status: str
    rect_depth_images_predicted: int
    rect_depth_accepted_count: int
    rect_depth_review_count: int
    rect_depth_reject_count: int
    rect_depth_missing_count: int
    rect_depth_acceptance_ratio: float
    rect_depth_majority_mode: str
    rect_depth_unique_depths_json: str
    rect_depth_source: str
    rect_depth_output_dir: str
    rect_depth_predictions_csv: str
    rect_depth_summary_json: str
    lt_images_predicted: int
    lt_majority_label: str
    lt_majority_vote_ratio: float
    lt_mean_confidence: float
    lt_mean_prob_l: float
    lt_mean_prob_t: float
    lt_source: str
    lt_checkpoint: str
    line_13_rect_name_echo: str
    line_13_top: str
    line_13_left: str
    line_13_bottom: str
    line_13_right: str
    line_13_source: str
    line_13_support: float
    line_13_model_checkpoint: str
    line_13_model_route: str
    line_13_pre_dark_trim_rect_name_echo: str
    line_14_rect_name_probe: str
    line_14_top: str
    line_14_left: str
    line_14_bottom: str
    line_14_right: str
    line_14_source: str
    line_14_support: float
    vendor_pred: str
    vendor_conf: float
    vendor_margin: float
    vendor_source: str
    vendor_decision_reason: str
    vendor_top3_json: str
    vendor_threshold: float
    vendor_ocr_used: str
    vendor_ocr_candidate: str
    vendor_ocr_candidate_score: str
    vendor_ocr_candidate_hits: str
    vendor_ocr_delta_score: str
    vendor_ocr_text: str
    vendor_ocr_top3_json: str
    vendor_ocr_samples_checked: str
    vendor_ocr_reason: str
    probe_source: str
    probe_decision_reason: str
    probe_top3_json: str
    probe_threshold: float
    capture_metadata_support: int
    rotation_deg_clockwise: int
    rotation_source: str
    rotation_vote_ratio: float
    rotation_votes_total: int
    rotation_samples_checked: int
    rotation_decision_reason: str
    rotation_osd_debug_json: str
    rotation_ocr_debug_json: str
    status: str
    review_reasons: str
    # Scale stage (#18-#21). These carry defaults so the dataclass stays constructible when
    # the stage is disabled, and so adding a field never breaks the positional call.
    line_18_vect_depth: str = ""
    line_19_pixel_ratio_x: str = ""
    line_20_pixel_ratio_y: str = ""
    line_21_scale_line: str = ""
    scale_status: str = ""
    scale_source: str = ""
    scale_profile: str = ""
    scale_frames_studied: int = 0
    scale_depths_total: int = 0
    scale_depths_accepted: int = 0
    scale_depths_review: int = 0
    scale_depths_reject: int = 0
    scale_depths_interpolated: int = 0
    scale_acceptance_ratio: float = 0.0
    scale_ruler_x: str = ""
    scale_output_dir: str = ""
    scale_per_image_csv: str = ""
    scale_per_depth_csv: str = ""


def _video_input_code_0hdmi_1vga(value: Optional[str]) -> Optional[int]:
    if value is None:
        return None
    lowered = value.strip().lower()
    if lowered == "hdmi":
        return 0
    if lowered == "vga":
        return 1
    return None


def _safe_int_value(value: object, default: int = 0) -> int:
    try:
        return int(float(str(value).strip()))
    except Exception:
        return int(default)


def _parse_capture_metadata_from_filename(path: Path) -> Tuple[Optional[str], Optional[int], Optional[int]]:
    stem = path.stem
    tokens = [tok for tok in re.split(r"[_\s-]+", stem) if tok]

    input_label: Optional[str] = None
    input_idx = -1
    for idx, tok in enumerate(tokens):
        lowered = tok.lower()
        if lowered in CAPTURE_INPUT_TOKENS:
            input_label = lowered
            input_idx = idx
            break

    width: Optional[int] = None
    height: Optional[int] = None
    if input_label is not None and (input_idx + 1) < len(tokens):
        match = CAPTURE_RES_TOKEN_RE.match(tokens[input_idx + 1])
        if match:
            width = int(match.group(1))
            height = int(match.group(2))

    if width is None or height is None:
        match_any = CAPTURE_RES_ANY_RE.search(stem)
        if match_any:
            width = int(match_any.group(1))
            height = int(match_any.group(2))

    return input_label, width, height


def _infer_capture_metadata_majority(
    image_paths: Iterable[Path],
) -> Tuple[Optional[str], Optional[int], Optional[int], int]:
    histogram: Dict[Tuple[str, int, int], int] = defaultdict(int)
    for path in image_paths:
        input_label, width, height = _parse_capture_metadata_from_filename(path)
        if input_label is None and (width is None or height is None):
            continue
        key = (input_label or "", int(width or 0), int(height or 0))
        histogram[key] += 1

    if not histogram:
        return None, None, None, 0

    best_key, count = max(
        histogram.items(),
        key=lambda item: (item[1], item[0][0] == "hdmi", item[0][1] * item[0][2]),
    )
    input_label = best_key[0] if best_key[0] else None
    width = best_key[1] if best_key[1] > 0 else None
    height = best_key[2] if best_key[2] > 0 else None
    return input_label, width, height, count


def _safe_line(lines: Sequence[str], line_num_1based: int) -> Optional[str]:
    if 1 <= line_num_1based <= len(lines):
        return lines[line_num_1based - 1].strip()
    return None


def _extract_rect_echo_line_num(lines: Sequence[str]) -> Optional[int]:
    for candidate in (11, 10):
        value = _safe_line(lines, candidate)
        if value and RECT_ECHO_RE.match(value):
            return candidate
    for idx in range(1, min(len(lines), 20) + 1):
        value = _safe_line(lines, idx)
        if value and RECT_ECHO_RE.match(value):
            return idx
    return None


def _extract_fss_line_with_rect_offset(fss_path: Path, target_line_1based: int) -> Optional[str]:
    try:
        lines = fss_path.read_text(encoding="utf-8", errors="replace").splitlines()
    except OSError:
        return None
    rect_line = _extract_rect_echo_line_num(lines)
    if rect_line is None:
        return None
    offset = rect_line - 11
    return _safe_line(lines, target_line_1based + offset)


def _parse_rect_name_echo_coords(line_13: str) -> Optional[Tuple[int, int, int, int]]:
    match = RECT_NAME_ECHO_PREFIX_RE.match(line_13.strip())
    if not match:
        return None
    top, left, bottom, right = (int(match.group(i)) for i in range(1, 5))
    return top, left, bottom, right


def _format_line13_with_disabled_tail(top: int, left: int, bottom: int, right: int) -> str:
    return f"{int(top)}|{int(left)}|{int(bottom)}|{int(right)}|{LINE13_DISABLED_TAIL}"


def _force_line13_disabled_tail(line_13_value: str) -> str:
    coords = _parse_rect_name_echo_coords(str(line_13_value or ""))
    if coords is None:
        return str(line_13_value or "")
    top, left, bottom, right = coords
    return _format_line13_with_disabled_tail(top, left, bottom, right)


def _rect_name_tail_from_template(rect_name_template: str) -> str:
    parts = rect_name_template.strip().split("|")
    if len(parts) < 8:
        return "1|7:0.800000:0:0:0:0:0|5|"
    tail = "|".join(parts[4:])
    if not tail.endswith("|"):
        tail += "|"
    return tail


def _normalize_rect_name_user_value(user_value: str, fallback_template: str) -> str:
    value = user_value.strip()
    if _parse_rect_name_echo_coords(value) is not None and value.count("|") >= 7:
        return value
    if RECT_ECHO_RE.match(value):
        parts = value.split("|")
        if len(parts) >= 4:
            top, left, bottom, right = parts[0], parts[1], parts[2], parts[3]
            tail = _rect_name_tail_from_template(fallback_template)
            return f"{top}|{left}|{bottom}|{right}|{tail}"
    return value


def _collect_acquisition_images(folder: Path) -> List[Path]:
    out_all: List[Path] = []
    out_pattern: List[Path] = []
    for path in folder.rglob("*"):
        if not path.is_file():
            continue
        if path.suffix.lower() not in IMAGE_SUFFIXES:
            continue
        out_all.append(path)
        if CAPTURE_FILENAME_PATTERN_RE.match(path.stem):
            out_pattern.append(path)
    # Preferred path: strict acquisition naming pattern.
    # Fallback: if pattern is missing (legacy/preprocessed sets), use all image files.
    selected = out_pattern if out_pattern else out_all
    selected.sort()
    return selected


def _sha1_file(path: Path, chunk_size: int = 1 << 20) -> str:
    hasher = hashlib.sha1()
    with path.open("rb") as fh:
        while True:
            chunk = fh.read(chunk_size)
            if not chunk:
                break
            hasher.update(chunk)
    return hasher.hexdigest()


def _deduplicate_exact_images(paths: Sequence[Path]) -> Tuple[List[Path], int]:
    """Drop exact duplicate files (same byte content), preserving order."""
    if len(paths) <= 1:
        return list(paths), 0

    sizes: List[Optional[int]] = []
    size_hist: Counter[int] = Counter()
    for path in paths:
        try:
            sz = int(path.stat().st_size)
        except OSError:
            sz = None
        sizes.append(sz)
        if sz is not None:
            size_hist[sz] += 1

    unique_paths: List[Path] = []
    seen_hashes: set[Tuple[int, str]] = set()
    duplicates_removed = 0

    for path, sz in zip(paths, sizes):
        if sz is None or size_hist[sz] <= 1:
            unique_paths.append(path)
            continue
        try:
            digest = _sha1_file(path)
        except OSError:
            # Keep unreadable files in the stream; downstream loaders handle them.
            unique_paths.append(path)
            continue
        key = (sz, digest)
        if key in seen_hashes:
            duplicates_removed += 1
            continue
        seen_hashes.add(key)
        unique_paths.append(path)

    return unique_paths, duplicates_removed


def _select_uniform_subset(paths: Sequence[Path], limit: int) -> List[Path]:
    if limit <= 0 or len(paths) <= limit:
        return list(paths)
    if limit == 1:
        return [paths[len(paths) // 2]]

    total = len(paths)
    selected_idx: List[int] = []
    seen = set()
    for i in range(limit):
        idx = round(i * (total - 1) / (limit - 1))
        if idx not in seen:
            selected_idx.append(idx)
            seen.add(idx)
    if len(selected_idx) < limit:
        for idx in range(total):
            if idx in seen:
                continue
            selected_idx.append(idx)
            seen.add(idx)
            if len(selected_idx) >= limit:
                break
    selected_idx.sort()
    return [paths[i] for i in selected_idx]


def _normalize_rotation_deg_clockwise(value: int) -> int:
    value_i = int(value) % 360
    if value_i in ROTATION_VALUES_CLOCKWISE:
        return value_i
    return 0


def _run_tesseract_osd_rotation(
    image_path: Path,
    timeout_sec: float,
) -> Tuple[Optional[int], Optional[float], bool]:
    try:
        proc = subprocess.run(
            [
                "tesseract",
                str(image_path),
                "stdout",
                "--psm",
                "0",
                "-l",
                "osd",
            ],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            timeout=timeout_sec,
            check=False,
        )
    except FileNotFoundError:
        return None, None, True
    except Exception:
        return None, None, False

    text = f"{proc.stdout or ''}\n{proc.stderr or ''}"
    rotate_match = TESSERACT_OSD_ROTATE_RE.search(text)
    conf_match = TESSERACT_OSD_CONF_RE.search(text)
    conf_value: Optional[float] = None
    if conf_match:
        try:
            conf_value = float(conf_match.group(1))
        except ValueError:
            conf_value = None
    if not rotate_match:
        return None, conf_value, False
    rotate_deg = _normalize_rotation_deg_clockwise(int(rotate_match.group(1)))
    return rotate_deg, conf_value, False


def _estimate_folder_rotation_with_osd(
    image_paths: Sequence[Path],
    sample_limit: int,
    timeout_sec: float,
    min_confidence: float,
    min_votes: int,
    min_ratio: float,
) -> Tuple[int, str, float, int, int, bool, Dict[str, object]]:
    sampled_paths = _select_uniform_subset(image_paths, sample_limit)
    votes: Counter[int] = Counter()
    samples_checked = 0
    sample_records: List[Dict[str, object]] = []
    votes_map: Dict[str, int] = {str(deg): 0 for deg in ROTATION_VALUES_CLOCKWISE}

    for path in sampled_paths:
        rotate_deg, confidence, tesseract_missing = _run_tesseract_osd_rotation(
            image_path=path,
            timeout_sec=timeout_sec,
        )
        samples_checked += 1
        accepted = bool(
            rotate_deg is not None
            and (confidence is None or float(confidence) >= float(min_confidence))
        )
        sample_records.append(
            {
                "image_name": str(path.name),
                "rotate_deg": int(rotate_deg) if rotate_deg is not None else None,
                "osd_confidence": float(confidence) if confidence is not None else None,
                "accepted": bool(accepted),
            }
        )
        if tesseract_missing:
            return (
                0,
                "osd_unavailable",
                0.0,
                int(sum(votes.values())),
                samples_checked,
                True,
                {
                    "sample_limit": int(sample_limit),
                    "min_confidence": float(min_confidence),
                    "min_votes": int(min_votes),
                    "min_ratio": float(min_ratio),
                    "votes": votes_map,
                    "samples": sample_records,
                    "best_rotation": 0,
                    "best_count": 0,
                    "ratio": 0.0,
                },
            )
        if rotate_deg is None:
            continue
        if confidence is not None and confidence < min_confidence:
            continue
        votes[int(rotate_deg)] += 1
        votes_map[str(int(rotate_deg))] = int(votes[int(rotate_deg)])

    total_votes = int(sum(votes.values()))
    if total_votes <= 0:
        return (
            0,
            "osd_no_votes",
            0.0,
            total_votes,
            samples_checked,
            False,
            {
                "sample_limit": int(sample_limit),
                "min_confidence": float(min_confidence),
                "min_votes": int(min_votes),
                "min_ratio": float(min_ratio),
                "votes": votes_map,
                "samples": sample_records,
                "best_rotation": 0,
                "best_count": 0,
                "ratio": 0.0,
            },
        )

    best_rotation, best_count = max(
        votes.items(),
        key=lambda item: (item[1], item[0] == 0, -item[0]),
    )
    ratio = float(best_count / max(1, total_votes))
    diagnostics = {
        "sample_limit": int(sample_limit),
        "min_confidence": float(min_confidence),
        "min_votes": int(min_votes),
        "min_ratio": float(min_ratio),
        "votes": votes_map,
        "samples": sample_records,
        "best_rotation": int(best_rotation),
        "best_count": int(best_count),
        "ratio": float(ratio),
    }
    if best_count >= int(max(1, min_votes)) and ratio >= float(min_ratio):
        return int(best_rotation), "osd_majority", ratio, total_votes, samples_checked, False, diagnostics
    return 0, "osd_low_support", ratio, total_votes, samples_checked, False, diagnostics


def _run_tesseract_ocr_word_score(
    image_path: Path,
    rotate_deg_clockwise: int,
    timeout_sec: float,
    lang: str,
    psm: int,
    min_word_conf: float,
) -> Tuple[float, int, bool]:
    tokens, tesseract_missing = _run_tesseract_ocr_tokens(
        image_path=image_path,
        rotate_deg_clockwise=rotate_deg_clockwise,
        timeout_sec=timeout_sec,
        lang=lang,
        psm=psm,
        min_word_conf=min_word_conf,
    )
    if tesseract_missing:
        return 0.0, 0, True
    conf_sum = float(sum(float(item.get("conf", 0.0)) for item in tokens))
    words = int(len(tokens))
    return conf_sum, words, False


def _normalize_ocr_token(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(text or "").lower())


def _build_vendor_ocr_alias_map(vendor_classes: Sequence[str]) -> Dict[str, List[str]]:
    alias_map: Dict[str, List[str]] = {}
    for vendor_name in vendor_classes:
        vendor = str(vendor_name or "").strip()
        if not vendor:
            continue
        aliases = set()
        full_norm = _normalize_ocr_token(vendor)
        if len(full_norm) >= 3 or full_norm in VENDOR_OCR_SHORT_ALIAS_ALLOWLIST:
            aliases.add(full_norm)
        for token in re.split(r"[^A-Za-z0-9]+", vendor):
            tok = _normalize_ocr_token(token)
            if len(tok) < 2:
                continue
            if tok in VENDOR_OCR_GENERIC_TOKENS:
                continue
            if len(tok) < 3 and tok not in VENDOR_OCR_SHORT_ALIAS_ALLOWLIST:
                continue
            aliases.add(tok)
        if not aliases and full_norm:
            if len(full_norm) >= 3 or full_norm in VENDOR_OCR_SHORT_ALIAS_ALLOWLIST:
                aliases.add(full_norm)
        alias_map[vendor] = sorted(aliases, key=lambda x: (-len(x), x))
    return alias_map


def _run_tesseract_ocr_tokens(
    image_path: Path,
    rotate_deg_clockwise: int,
    timeout_sec: float,
    lang: str,
    psm: int,
    min_word_conf: float,
) -> Tuple[List[Dict[str, object]], bool]:
    rotate = _normalize_rotation_deg_clockwise(rotate_deg_clockwise)
    tmp_path: Optional[Path] = None
    source_path = image_path
    try:
        if rotate != 0:
            with Image.open(image_path) as img:
                rotated = img.convert("RGB").rotate(-rotate, expand=True)
                with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
                    tmp_path = Path(tmp.name)
                rotated.save(tmp_path, format="PNG")
            source_path = tmp_path

        proc = subprocess.run(
            [
                "tesseract",
                str(source_path),
                "stdout",
                "--oem",
                "1",
                "--psm",
                str(psm),
                "-l",
                lang,
                "tsv",
            ],
            stdout=subprocess.PIPE,
            stderr=subprocess.DEVNULL,
            text=True,
            timeout=timeout_sec,
            check=False,
        )
    except FileNotFoundError:
        return [], True
    except Exception:
        return [], False
    finally:
        if tmp_path is not None:
            try:
                tmp_path.unlink(missing_ok=True)
            except Exception:
                pass

    out: List[Dict[str, object]] = []
    for line in (proc.stdout or "").splitlines()[1:]:
        parts = line.split("\t")
        if len(parts) < 12:
            continue
        raw_text = str(parts[11] or "").strip()
        if not raw_text:
            continue
        if not any(ch.isalnum() for ch in raw_text):
            continue
        try:
            conf = float(parts[10])
        except ValueError:
            continue
        if conf < float(min_word_conf):
            continue
        token = _normalize_ocr_token(raw_text)
        if len(token) < 2:
            continue
        out.append(
            {
                "token": token,
                "text": raw_text,
                "conf": float(conf),
            }
        )
    return out, False


def _resolve_vendor_with_ocr(
    *,
    image_paths: Sequence[Path],
    rotate_deg_clockwise: int,
    vendor_alias_map: Dict[str, List[str]],
    sample_limit: int,
    timeout_sec: float,
    lang: str,
    psm: int,
    min_word_conf: float,
    min_score: float,
    min_score_delta: float,
    min_hits: int,
    cnn_vendor_hint: str,
) -> Dict[str, object]:
    sampled_paths = _select_uniform_subset(image_paths, sample_limit)
    diagnostics: Dict[str, object] = {
        "tesseract_missing": False,
        "samples_checked": int(len(sampled_paths)),
        "candidate": "",
        "candidate_score": 0.0,
        "candidate_hits": 0,
        "delta_score": 0.0,
        "accepted": False,
        "reason": "no_samples",
        "top3": [],
        "text_excerpt": "",
        "confidence_proxy": 0.0,
    }
    if not sampled_paths:
        return diagnostics
    if not vendor_alias_map:
        diagnostics["reason"] = "vendor_alias_map_empty"
        return diagnostics

    score_map: Dict[str, float] = {vendor: 0.0 for vendor in vendor_alias_map.keys()}
    hit_map: Dict[str, int] = {vendor: 0 for vendor in vendor_alias_map.keys()}
    raw_words: List[str] = []
    sample_debug: List[Dict[str, object]] = []
    for path in sampled_paths:
        tokens, tesseract_missing = _run_tesseract_ocr_tokens(
            image_path=path,
            rotate_deg_clockwise=int(rotate_deg_clockwise),
            timeout_sec=timeout_sec,
            lang=lang,
            psm=psm,
            min_word_conf=min_word_conf,
        )
        if tesseract_missing:
            diagnostics["tesseract_missing"] = True
            diagnostics["reason"] = "tesseract_unavailable"
            return diagnostics

        sample_tokens = [str(item.get("text", "")) for item in tokens]
        if sample_tokens:
            raw_words.extend(sample_tokens)
        sample_debug.append(
            {
                "image_name": str(path.name),
                "tokens_count": int(len(tokens)),
            }
        )

        for item in tokens:
            token = str(item.get("token", "")).strip()
            conf = float(item.get("conf", 0.0) or 0.0)
            if not token:
                continue
            for vendor, aliases in vendor_alias_map.items():
                best_local = 0.0
                for alias in aliases:
                    alias_txt = str(alias or "").strip()
                    if len(alias_txt) < 2:
                        continue
                    if token == alias_txt:
                        score = 1.00 + (conf / 100.0)
                    elif len(alias_txt) >= 4 and alias_txt in token:
                        score = 0.40 + (conf / 220.0)
                    elif len(token) >= 4 and token in alias_txt:
                        score = 0.24 + (conf / 260.0)
                    else:
                        score = 0.0
                    if score > best_local:
                        best_local = float(score)
                if best_local > 0.0:
                    score_map[vendor] += float(best_local)
                    hit_map[vendor] += 1

    ranking = sorted(
        list(score_map.keys()),
        key=lambda name: (
            -float(score_map.get(name, 0.0)),
            -int(hit_map.get(name, 0)),
            0 if str(name) == str(cnn_vendor_hint) else 1,
            str(name),
        ),
    )
    if not ranking:
        diagnostics["reason"] = "no_vendor_scores"
        diagnostics["sample_debug"] = sample_debug
        return diagnostics

    best_vendor = str(ranking[0])
    best_score = float(score_map.get(best_vendor, 0.0))
    best_hits = int(hit_map.get(best_vendor, 0))
    second_score = float(score_map.get(ranking[1], 0.0)) if len(ranking) > 1 else 0.0
    delta_score = float(best_score - second_score)
    top3_raw = [(str(name), float(score_map.get(name, 0.0))) for name in ranking[:3]]
    top3_score_sum = float(sum(max(0.0, score) for _, score in top3_raw))
    top3 = []
    for name, score in top3_raw:
        conf_proxy = float(score / top3_score_sum) if top3_score_sum > 1e-8 else 0.0
        top3.append({"label": name, "confidence": conf_proxy, "score": float(score), "hits": int(hit_map.get(name, 0))})

    accepted = bool(
        best_score >= float(min_score)
        and delta_score >= float(min_score_delta)
        and best_hits >= int(max(1, min_hits))
    )
    confidence_proxy = min(
        0.99,
        max(
            0.0,
            0.45 + 0.08 * min(float(best_hits), 5.0) + 0.04 * min(float(delta_score), 5.0),
        ),
    )

    diagnostics.update(
        {
            "candidate": best_vendor,
            "candidate_score": float(best_score),
            "candidate_hits": int(best_hits),
            "delta_score": float(delta_score),
            "accepted": bool(accepted),
            "reason": "accepted" if accepted else "below_ocr_threshold",
            "top3": top3,
            "text_excerpt": " | ".join(raw_words[:20])[:400],
            "confidence_proxy": float(confidence_proxy),
            "sample_debug": sample_debug,
            "min_score": float(min_score),
            "min_score_delta": float(min_score_delta),
            "min_hits": int(min_hits),
        }
    )
    return diagnostics


def _refine_folder_rotation_with_ocr(
    image_paths: Sequence[Path],
    candidate_rotate_deg: int,
    base_source: str,
    sample_limit: int,
    timeout_sec: float,
    lang: str,
    psm: int,
    min_word_conf: float,
    score_word_bonus: float,
    min_words: int,
    min_score_delta: float,
) -> Tuple[int, str, bool, Dict[str, object]]:
    candidate = _normalize_rotation_deg_clockwise(candidate_rotate_deg)
    sampled_paths = _select_uniform_subset(image_paths, sample_limit)
    if not sampled_paths:
        return (
            candidate,
            f"{base_source}_ocr_no_samples",
            False,
            {
                "candidate_rotation": int(candidate),
                "selected_rotation": int(candidate),
                "scores": {str(d): 0.0 for d in ROTATION_VALUES_CLOCKWISE},
                "word_counts": {str(d): 0 for d in ROTATION_VALUES_CLOCKWISE},
                "delta_score": 0.0,
                "min_words": int(min_words),
                "min_score_delta": float(min_score_delta),
                "sample_count": 0,
            },
        )

    scores: Dict[int, float] = {}
    word_counts: Dict[int, int] = {}
    per_angle_details: Dict[str, Dict[str, object]] = {}
    for rotate_deg in ROTATION_VALUES_CLOCKWISE:
        conf_sum_total = 0.0
        words_total = 0
        sample_scores: List[Dict[str, object]] = []
        for path in sampled_paths:
            conf_sum, words, tesseract_missing = _run_tesseract_ocr_word_score(
                image_path=path,
                rotate_deg_clockwise=rotate_deg,
                timeout_sec=timeout_sec,
                lang=lang,
                psm=psm,
                min_word_conf=min_word_conf,
            )
            if tesseract_missing:
                return (
                    candidate,
                    f"{base_source}_ocr_unavailable",
                    True,
                    {
                        "candidate_rotation": int(candidate),
                        "selected_rotation": int(candidate),
                        "scores": {str(d): float(scores.get(int(d), 0.0)) for d in ROTATION_VALUES_CLOCKWISE},
                        "word_counts": {str(d): int(word_counts.get(int(d), 0)) for d in ROTATION_VALUES_CLOCKWISE},
                        "delta_score": 0.0,
                        "min_words": int(min_words),
                        "min_score_delta": float(min_score_delta),
                        "sample_count": len(sampled_paths),
                        "per_angle": per_angle_details,
                    },
                )
            conf_sum_total += conf_sum
            words_total += words
            sample_scores.append(
                {
                    "image_name": str(path.name),
                    "conf_sum": float(conf_sum),
                    "words": int(words),
                }
            )
        scores[rotate_deg] = float(conf_sum_total + float(words_total) * float(score_word_bonus))
        word_counts[rotate_deg] = int(words_total)
        per_angle_details[str(int(rotate_deg))] = {
            "score": float(scores[rotate_deg]),
            "words": int(words_total),
            "samples": sample_scores,
        }

    best_deg = max(
        ROTATION_VALUES_CLOCKWISE,
        key=lambda deg: (
            scores.get(int(deg), 0.0),
            word_counts.get(int(deg), 0),
            int(deg) == candidate,
            int(deg) == 0,
            -int(deg),
        ),
    )
    best_words = int(word_counts.get(int(best_deg), 0))
    best_score = float(scores.get(int(best_deg), 0.0))
    candidate_score = float(scores.get(candidate, 0.0))
    delta = float(best_score - candidate_score)
    diagnostics = {
        "candidate_rotation": int(candidate),
        "selected_rotation": int(best_deg),
        "scores": {str(int(d)): float(scores.get(int(d), 0.0)) for d in ROTATION_VALUES_CLOCKWISE},
        "word_counts": {str(int(d)): int(word_counts.get(int(d), 0)) for d in ROTATION_VALUES_CLOCKWISE},
        "delta_score": float(delta),
        "min_words": int(min_words),
        "min_score_delta": float(min_score_delta),
        "sample_count": len(sampled_paths),
        "per_angle": per_angle_details,
    }

    if best_words < int(max(0, min_words)):
        return candidate, f"{base_source}_ocr_no_text_support", False, diagnostics
    if int(best_deg) != candidate and delta >= float(min_score_delta):
        return int(best_deg), f"{base_source}_ocr_override", False, diagnostics
    if int(best_deg) != candidate:
        return candidate, f"{base_source}_ocr_weak_delta", False, diagnostics
    return candidate, f"{base_source}_ocr_confirm", False, diagnostics


def _inverse_rotate_point_clockwise(
    x_rot: float,
    y_rot: float,
    original_width: int,
    original_height: int,
    rotate_deg_clockwise: int,
) -> Tuple[float, float]:
    rotate = _normalize_rotation_deg_clockwise(rotate_deg_clockwise)
    ow = float(original_width)
    oh = float(original_height)

    if rotate == 0:
        return x_rot, y_rot
    if rotate == 90:
        return y_rot, oh - x_rot
    if rotate == 180:
        return ow - x_rot, oh - y_rot
    if rotate == 270:
        return ow - y_rot, x_rot
    return x_rot, y_rot


def _rotate_box_back_to_original_coords(
    box_rotated: Tuple[float, float, float, float],
    original_width: int,
    original_height: int,
    rotate_deg_clockwise: int,
) -> Tuple[float, float, float, float]:
    rotate = _normalize_rotation_deg_clockwise(rotate_deg_clockwise)
    if rotate == 0:
        return box_rotated

    x1, y1, x2, y2 = box_rotated
    points_rot = ((x1, y1), (x2, y1), (x1, y2), (x2, y2))
    points_original = [
        _inverse_rotate_point_clockwise(
            x_rot=px,
            y_rot=py,
            original_width=original_width,
            original_height=original_height,
            rotate_deg_clockwise=rotate,
        )
        for px, py in points_rot
    ]
    xs = [p[0] for p in points_original]
    ys = [p[1] for p in points_original]
    x1_o = min(xs)
    y1_o = min(ys)
    x2_o = max(xs)
    y2_o = max(ys)

    x1_o = max(0.0, min(x1_o, float(original_width) - 1.0))
    y1_o = max(0.0, min(y1_o, float(original_height) - 1.0))
    x2_o = max(x1_o + 1.0, min(x2_o, float(original_width)))
    y2_o = max(y1_o + 1.0, min(y2_o, float(original_height)))
    return x1_o, y1_o, x2_o, y2_o


class SuGiuRectClassifier(nn.Module):
    """Binary classifier (su/giu) trained on rect crops."""

    def __init__(self) -> None:
        super().__init__()
        model = resnet18(weights=None)
        in_features = model.fc.in_features
        model.fc = nn.Linear(in_features, 2)
        self.model = model

    def forward(self, x: torch.Tensor) -> torch.Tensor:  # type: ignore[override]
        return self.model(x)


class LTRectClassifier(nn.Module):
    """Binary classifier (L/T) trained on ultrasound rect crops."""

    def __init__(self) -> None:
        super().__init__()
        backbone = resnet18(weights=None)
        in_features = backbone.fc.in_features
        backbone.fc = nn.Identity()
        self.backbone = backbone
        self.head = nn.Sequential(
            nn.Linear(in_features, 256),
            nn.ReLU(inplace=True),
            nn.Dropout(p=0.30),
            nn.Linear(256, 2),
        )

    def forward(self, x: torch.Tensor) -> torch.Tensor:  # type: ignore[override]
        feats = self.backbone(x)
        return self.head(feats)


def _clamp_crop_box(
    box: Tuple[float, float, float, float],
    width: int,
    height: int,
) -> Optional[Tuple[int, int, int, int]]:
    if width <= 1 or height <= 1:
        return None
    x1, y1, x2, y2 = box
    left = int(round(max(0.0, min(x1, float(width - 1)))))
    top = int(round(max(0.0, min(y1, float(height - 1)))))
    right = int(round(max(float(left + 1), min(x2, float(width)))))
    bottom = int(round(max(float(top + 1), min(y2, float(height)))))
    if right <= left:
        right = min(width, left + 1)
    if bottom <= top:
        bottom = min(height, top + 1)
    if right <= left or bottom <= top:
        return None
    return top, left, bottom, right


def _load_rect_crop_classification_tensor(
    path: Path,
    crop_box: Tuple[float, float, float, float],
    image_size: int,
    rotate_deg_clockwise: int = 0,
) -> Tuple[torch.Tensor, Tuple[int, int, int, int]]:
    with Image.open(path) as img:
        image = img.convert("RGB")
        width, height = image.size
        crop_tlbr = _clamp_crop_box(crop_box, width=width, height=height)
        if crop_tlbr is None:
            raise ValueError("invalid_crop_box")
        top, left, bottom, right = crop_tlbr
        crop = image.crop((left, top, right, bottom))
        rotate = _normalize_rotation_deg_clockwise(rotate_deg_clockwise)
        if rotate:
            crop = crop.rotate(-rotate, expand=True)
        crop = TF.resize(
            crop,
            size=[image_size, image_size],
            interpolation=InterpolationMode.BILINEAR,
            antialias=True,
        )
        tensor = TF.to_tensor(crop)
        tensor = TF.normalize(
            tensor,
            mean=(0.485, 0.456, 0.406),
            std=(0.229, 0.224, 0.225),
        )
    return tensor, crop_tlbr


def _predict_su_giu_on_rect_crops(
    model: torch.nn.Module,
    class_names: Sequence[str],
    image_paths: Sequence[Path],
    rect_boxes_abs: Sequence[Tuple[float, float, float, float]],
    fallback_rect_box_abs: Optional[Tuple[float, float, float, float]],
    image_size: int,
    batch_size: int,
    device: torch.device,
    rotate_deg_clockwise: int = 0,
) -> List[Dict[str, object]]:
    if not image_paths:
        return []
    if len(class_names) < 2:
        return []

    class_names_list = [str(name).strip().lower() for name in class_names]
    su_idx = class_names_list.index("su") if "su" in class_names_list else 0
    giu_idx = class_names_list.index("giu") if "giu" in class_names_list else min(1, len(class_names_list) - 1)

    jobs: List[Tuple[int, Path, Tuple[float, float, float, float], str]] = []
    for idx, path in enumerate(image_paths):
        if idx < len(rect_boxes_abs):
            jobs.append((idx, path, rect_boxes_abs[idx], "per_image_rect"))
        elif fallback_rect_box_abs is not None:
            jobs.append((idx, path, fallback_rect_box_abs, "global_rect_fallback"))

    if not jobs:
        return []

    out_rows: List[Dict[str, object]] = []
    model.eval()
    with torch.inference_mode():
        for start in range(0, len(jobs), batch_size):
            batch_jobs = jobs[start : start + batch_size]
            tensors: List[torch.Tensor] = []
            metas: List[Tuple[int, Path, Tuple[int, int, int, int], str]] = []
            for image_idx, image_path, crop_box, crop_source in batch_jobs:
                try:
                    tensor, crop_tlbr = _load_rect_crop_classification_tensor(
                        path=image_path,
                        crop_box=crop_box,
                        image_size=image_size,
                        rotate_deg_clockwise=rotate_deg_clockwise,
                    )
                except Exception:
                    continue
                tensors.append(tensor)
                metas.append((image_idx, image_path, crop_tlbr, crop_source))
            if not tensors:
                continue

            images = torch.stack(tensors, dim=0).to(device, non_blocking=True)
            probs = torch.softmax(model(images), dim=1).detach().cpu().numpy()
            pred_idx = np.argmax(probs, axis=1).tolist()

            for i, (image_idx, image_path, crop_tlbr, crop_source) in enumerate(metas):
                pred_i = int(pred_idx[i])
                pred_label = (
                    str(class_names[pred_i]).strip().lower()
                    if 0 <= pred_i < len(class_names)
                    else str(pred_i)
                )
                p_su = float(probs[i][su_idx]) if 0 <= su_idx < probs.shape[1] else 0.0
                p_giu = float(probs[i][giu_idx]) if 0 <= giu_idx < probs.shape[1] else 0.0
                conf = float(probs[i][pred_i]) if 0 <= pred_i < probs.shape[1] else 0.0
                top, left, bottom, right = crop_tlbr
                out_rows.append(
                    {
                        "image_index": int(image_idx),
                        "image_path": image_path.as_posix(),
                        "pred_label": pred_label,
                        "pred_idx": pred_i,
                        "confidence": conf,
                        "prob_su": p_su,
                        "prob_giu": p_giu,
                        "crop_top": int(top),
                        "crop_left": int(left),
                        "crop_bottom": int(bottom),
                        "crop_right": int(right),
                        "crop_source": crop_source,
                    }
                )

    out_rows.sort(key=lambda row: int(row.get("image_index", 0)))
    return out_rows


def _safe_dir_name(value: str, fallback: str = "folder") -> str:
    txt = str(value or "").strip()
    out = re.sub(r"[^A-Za-z0-9_.-]+", "_", txt).strip("._")
    return out or fallback


def _line11_to_rect_depth_arg(line_11_rect_echo: str) -> str:
    """Convert line #11 top|left|bottom|right| to depth CLI left,top,right,bottom."""
    match = RECT_ECHO_RE.match(str(line_11_rect_echo or "").strip())
    if not match:
        return ""
    top, left, bottom, right = [int(match.group(i)) for i in range(1, 5)]
    if bottom <= top or right <= left:
        return ""
    return f"{left},{top},{right},{bottom}"


def _rect_depth_unique_depths_json(rows: Sequence[Dict[str, object]]) -> str:
    values: List[float] = []
    for row in rows:
        status = str(row.get("status", "") or "").strip().lower()
        if status not in {"accepted", "review"}:
            continue
        try:
            depth = float(row.get("depth_mm", "") or 0.0)
        except (TypeError, ValueError):
            continue
        if depth > 0 and np.isfinite(depth):
            values.append(round(float(depth), 3))
    unique = sorted(set(values))
    return json.dumps(unique, ensure_ascii=False, separators=(",", ":"))


def _summarize_rect_depth_predictions(rows: Sequence[Dict[str, object]]) -> Dict[str, object]:
    status_counts: Counter[str] = Counter(str(row.get("status", "") or "missing").strip().lower() for row in rows)
    mode_counts: Counter[str] = Counter(str(row.get("mode", "") or "").strip() for row in rows if str(row.get("mode", "") or "").strip())
    total = int(len(rows))
    accepted = int(status_counts.get("accepted", 0))
    majority_mode = mode_counts.most_common(1)[0][0] if mode_counts else ""
    return {
        "images_predicted": total,
        "accepted_count": accepted,
        "review_count": int(status_counts.get("review", 0)),
        "reject_count": int(status_counts.get("reject", 0)),
        "missing_count": int(status_counts.get("missing", 0)),
        "acceptance_ratio": float(accepted / max(1, total)),
        "majority_mode": majority_mode,
        "unique_depths_json": _rect_depth_unique_depths_json(rows),
        "status_counts": dict(status_counts),
        "mode_counts": dict(mode_counts),
    }


def _run_rect_depth_autonomous_stage(
    *,
    folder: Path,
    output_dir: Path,
    folder_index: int,
    python_bin: str,
    vendor_pred: str,
    probe_id: str,
    line_11_rect_echo: str,
    video_x: Optional[int],
    video_y: Optional[int],
    rotation_deg_clockwise: int,
    max_images: int,
    max_candidates_per_sample: int,
    ocr_timeout: float,
    scale_side_preference: str,
    subprocess_timeout_sec: float,
    min_accepted_ratio: float,
) -> Tuple[Dict[str, object], List[Dict[str, object]]]:
    script = REPO_ROOT / "tools/depth/predict_rect_depth_autonomous.py"
    folder_slug = _safe_dir_name(folder.name, fallback=f"folder_{folder_index:04d}")
    folder_hash = hashlib.sha1(folder.as_posix().encode("utf-8")).hexdigest()[:8]
    stage_dir = output_dir / "rect_depth_autonomous" / f"{folder_index:04d}_{folder_slug}_{folder_hash}"
    summary_path = stage_dir / "summary.json"
    predictions_csv = stage_dir / "rect_depth_autonomous_predictions.csv"
    context_json = stage_dir / "pipeline_context.json"
    log_path = stage_dir / "pipeline_subprocess.log"
    out: Dict[str, object] = {
        "status": "review",
        "source": "not_run",
        "output_dir": stage_dir.as_posix(),
        "predictions_csv": predictions_csv.as_posix(),
        "summary_json": summary_path.as_posix(),
        "log_path": log_path.as_posix(),
        "images_predicted": 0,
        "accepted_count": 0,
        "review_count": 0,
        "reject_count": 0,
        "missing_count": 0,
        "acceptance_ratio": 0.0,
        "majority_mode": "",
        "unique_depths_json": "[]",
        "error": "",
    }
    if not script.is_file():
        out["source"] = "script_missing"
        out["error"] = f"script non trovato: {script}"
        return out, []

    stage_dir.mkdir(parents=True, exist_ok=True)
    rect_echo_arg = _line11_to_rect_depth_arg(line_11_rect_echo)
    context = {
        "folder_path": folder.as_posix(),
        "vendor_predicted": vendor_pred,
        "manufacturer": vendor_pred,
        "predicted_probe_id": probe_id,
        "predicted_probe": probe_id,
        "line_11_rect_echo": line_11_rect_echo,
        "rect_echo_cli_left_top_right_bottom": rect_echo_arg,
        "video_x": int(video_x or 0),
        "video_y": int(video_y or 0),
        "rotation_deg_clockwise": int(rotation_deg_clockwise),
        "pipeline_stage": "official_rect_depth_autonomous",
    }
    context_json.write_text(json.dumps(context, ensure_ascii=False, indent=2), encoding="utf-8")

    cmd = [
        python_bin,
        script.as_posix(),
        "--folder",
        folder.as_posix(),
        "--output-dir",
        stage_dir.as_posix(),
        "--context-json",
        context_json.as_posix(),
        "--max-images",
        str(int(max_images)),
        "--max-candidates-per-sample",
        str(int(max_candidates_per_sample)),
        "--ocr-timeout",
        str(float(ocr_timeout)),
    ]
    if vendor_pred:
        cmd.extend(["--vendor", str(vendor_pred)])
    if probe_id:
        cmd.extend(["--probe", str(probe_id)])
    if rect_echo_arg:
        cmd.extend(["--rect-echo", rect_echo_arg])
    if scale_side_preference:
        cmd.extend(["--scale-side-preference", str(scale_side_preference)])

    try:
        proc = subprocess.run(
            cmd,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            timeout=(float(subprocess_timeout_sec) if float(subprocess_timeout_sec) > 0 else None),
        )
        log_path.write_text(proc.stdout or "", encoding="utf-8")
    except subprocess.TimeoutExpired as exc:
        out["source"] = "autonomous_depth_timeout"
        out["error"] = f"timeout dopo {subprocess_timeout_sec:.1f}s"
        log_path.write_text(str(exc.stdout or exc), encoding="utf-8")
        return out, []
    except Exception as exc:
        out["source"] = "autonomous_depth_exception"
        out["error"] = str(exc)
        log_path.write_text(str(exc), encoding="utf-8")
        return out, []

    if proc.returncode != 0:
        out["source"] = "autonomous_depth_failed"
        out["error"] = f"returncode={proc.returncode}"
        return out, []
    if not predictions_csv.is_file():
        out["source"] = "autonomous_depth_no_predictions_csv"
        out["error"] = f"CSV predizioni non trovato: {predictions_csv}"
        return out, []

    rows: List[Dict[str, object]] = []
    try:
        with predictions_csv.open("r", encoding="utf-8", newline="") as fh:
            reader = csv.DictReader(fh)
            for image_index, row in enumerate(reader):
                clean = {str(k): ("" if v is None else str(v)) for k, v in row.items()}
                clean["folder_path"] = folder.as_posix()
                clean["folder_name"] = folder.name
                clean["folder_index"] = str(int(folder_index))
                clean["image_index"] = str(int(image_index))
                clean["rect_depth_run_dir"] = stage_dir.as_posix()
                rows.append(clean)
    except Exception as exc:
        out["source"] = "autonomous_depth_read_error"
        out["error"] = str(exc)
        return out, []

    summary = _summarize_rect_depth_predictions(rows)
    out.update(summary)
    out["source"] = "autonomous_ocr_classical_ranker"
    if int(out["images_predicted"]) <= 0:
        out["status"] = "review"
        out["error"] = "nessuna predizione depth valida"
    elif float(out["acceptance_ratio"]) >= float(min_accepted_ratio):
        out["status"] = "ok"
    else:
        out["status"] = "review"
        out["error"] = (
            f"accepted_ratio={float(out['acceptance_ratio']):.3f} "
            f"< min={float(min_accepted_ratio):.3f}"
        )
    return out, rows


def _build_scale_frames_context(
    *,
    images: Sequence[Path],
    su_giu_rows: Sequence[Dict[str, object]],
    lr_marker_rows: Sequence[Dict[str, object]],
    rect_depth_rows: Sequence[Dict[str, object]],
    line_11_rect_echo: str,
) -> List[Dict[str, object]]:
    """Weave what the earlier stages already know about each frame, for the scale stage.

    Nothing is recomputed here: the up/down comes from the LR-marker rows, whose
    ``su_giu_pred`` is *already* the marker's verdict over the net's prior (the marker can
    overturn it, and ``bundle_vertical_correction`` records when it did); the rect comes from
    the same rows; the depth and its evidence come from the RECT_DEPTH stage. A frame with no
    depth still travels, because it can still help agree the ruler's column folder-wide.

    Everything is joined by image path, i.e. by identity of the file. The depth stage picks its
    own frames and may well have looked at others; when it did, that frame simply carries no
    depth instead of being paired with a different acquisition because two names looked alike.
    """
    sugiu_by_path: Dict[str, Dict[str, object]] = {}
    for row in su_giu_rows:
        path = str(row.get("image_path", "") or "")
        if path:
            sugiu_by_path[path] = row
    marker_by_path: Dict[str, Dict[str, object]] = {}
    for row in lr_marker_rows:
        path = str(row.get("image_path", "") or "")
        if path:
            marker_by_path[path] = row
    depth_by_path: Dict[str, Dict[str, object]] = {}
    for row in rect_depth_rows:
        path = str(row.get("image_path", "") or "")
        if not path:
            continue
        # The depth stage emits one row per image; keep the strongest evidence if it ever
        # emits more, ranked as the reference study ranks it: accepted over review, and a
        # value read in the interface over one inferred from the scale.
        mode = str(row.get("mode", "") or "").strip()
        status = str(row.get("status", "") or "").strip().lower()
        rank = ({"accepted": 2, "review": 1}.get(status, 0), 1 if mode == "direct_label" else 0)
        previous = depth_by_path.get(path)
        if previous is None or rank > previous.get("_rank", (0, 0)):
            depth_by_path[path] = {**row, "_rank": rank}

    line11 = _line11_to_rect_depth_arg(line_11_rect_echo)
    folder_rect: Optional[List[float]] = None
    if line11:
        folder_rect = [float(v) for v in line11.split(",")]

    frames: List[Dict[str, object]] = []
    for path in images:
        key = path.as_posix()
        marker = marker_by_path.get(key, {})
        sugiu_row = sugiu_by_path.get(key, {})
        depth_row = depth_by_path.get(key, {})

        # The marker's verdict when there is one, the net's label otherwise.
        sugiu = str(marker.get("su_giu_pred", "") or sugiu_row.get("pred_label", "") or "")
        if marker:
            corrected = str(marker.get("bundle_vertical_correction", "") or "") == "corrected"
            source = "marker (ha corretto la rete)" if corrected else "marker"
            if not str(marker.get("bundle_vertical_source", "") or ""):
                source = "rete"
        else:
            source = "rete" if sugiu else ""

        rect_ltrb: Optional[List[float]] = None
        try:
            rect_ltrb = [
                float(marker["echo_rect_left_abs"]), float(marker["echo_rect_top_abs"]),
                float(marker["echo_rect_right_abs"]), float(marker["echo_rect_bottom_abs"]),
            ]
        except (KeyError, TypeError, ValueError):
            rect_ltrb = None
        if rect_ltrb is None:
            try:  # the su/giu rows carry the crop the classifier actually used
                rect_ltrb = [
                    float(sugiu_row["crop_left"]), float(sugiu_row["crop_top"]),
                    float(sugiu_row["crop_right"]), float(sugiu_row["crop_bottom"]),
                ]
            except (KeyError, TypeError, ValueError):
                rect_ltrb = None
        if rect_ltrb is None and folder_rect:
            rect_ltrb = list(folder_rect)

        depth_mm = None
        try:
            value = float(depth_row.get("depth_mm", "") or 0.0)
            depth_mm = value if value > 0 else None
        except (TypeError, ValueError):
            depth_mm = None
        depth_mode = str(depth_row.get("mode", "") or "")
        depth_status = str(depth_row.get("status", "") or "")
        # Only a depth actually printed in the interface is independent evidence: one the
        # module inferred *from the scale* would be compared with itself.
        from_interface = depth_mode == "direct_label"
        if depth_status.strip().lower() == "reject":
            depth_mm = None

        frames.append(
            {
                "image_path": key,
                "sugiu": sugiu,
                "sugiu_conf": _safe_float_or_none(sugiu_row.get("confidence")) or 0.0,
                "sugiu_source": source,
                "marker_score": _safe_float_or_none(marker.get("match_score")),
                "orientation_group": str(marker.get("orientation_group", "") or ""),
                "rect_ltrb": rect_ltrb,
                "depth_mm": depth_mm,
                "depth_from_interface": bool(from_interface),
                "depth_mode": depth_mode,
                "depth_status": depth_status,
                "depth_box": [
                    _safe_float_or_none(depth_row.get(k)) for k in ("left", "top", "right", "bottom")
                ] if depth_row else None,
                "depth_ocr_text": str(depth_row.get("ocr_text", "") or ""),
            }
        )
    return frames


def _safe_float_or_none(value: object) -> Optional[float]:
    try:
        out = float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return None
    return out if out == out else None


def _run_scale_stage(
    *,
    folder: Path,
    output_dir: Path,
    folder_index: int,
    python_bin: str,
    vendor_pred: str,
    vendor_conf: float,
    probe_id: str,
    line_11_rect_echo: str,
    video_x: Optional[int],
    video_y: Optional[int],
    rotation_deg_clockwise: int,
    frames: Sequence[Dict[str, object]],
    max_frames: int,
    subprocess_timeout_sec: float,
    min_accepted_ratio: float,
    corrections_path: Optional[Path] = None,
) -> Tuple[Dict[str, object], List[Dict[str, object]]]:
    """Run the scale study as a subprocess, the same way the depth stage is run.

    A subprocess and not an import: the scale detector needs cv2 and Tesseract, and a folder
    whose OCR wedges or whose cv2 build dies must cost one folder marked ``review``, not the
    whole batch. The context JSON left on disk is also what makes the run reproducible by
    hand and what the review page reads.
    """
    script = REPO_ROOT / "tools/scale/predict_scale_from_pipeline.py"
    folder_slug = _safe_dir_name(folder.name, fallback=f"folder_{folder_index:04d}")
    folder_hash = hashlib.sha1(folder.as_posix().encode("utf-8")).hexdigest()[:8]
    stage_dir = output_dir / "scale" / f"{folder_index:04d}_{folder_slug}_{folder_hash}"
    per_image_csv = stage_dir / "scale_per_image.csv"
    per_depth_csv = stage_dir / "scale_per_depth.csv"
    summary_path = stage_dir / "summary.json"
    context_json = stage_dir / "pipeline_context.json"
    log_path = stage_dir / "pipeline_subprocess.log"
    out: Dict[str, object] = {
        "status": "review",
        "source": "not_run",
        "output_dir": stage_dir.as_posix(),
        "per_image_csv": per_image_csv.as_posix(),
        "per_depth_csv": per_depth_csv.as_posix(),
        "summary_json": summary_path.as_posix(),
        "log_path": log_path.as_posix(),
        "profile": "",
        "frames_studied": 0,
        "depths_total": 0,
        "depths_accepted": 0,
        "depths_review": 0,
        "depths_reject": 0,
        "depths_interpolated": 0,
        "acceptance_ratio": 0.0,
        "ruler_x": "",
        "line_18_vect_depth": "",
        "line_19_pixel_ratio_x": "",
        "line_20_pixel_ratio_y": "",
        "line_21_scale_line": "",
        "review_reasons": [],
        "error": "",
    }
    if not script.is_file():
        out["source"] = "script_missing"
        out["error"] = f"script non trovato: {script}"
        return out, []

    stage_dir.mkdir(parents=True, exist_ok=True)
    context = {
        "folder_path": folder.as_posix(),
        "vendor_predicted": vendor_pred,
        "vendor_confidence": float(vendor_conf),
        "predicted_probe_id": probe_id,
        "line_11_rect_echo": line_11_rect_echo,
        "video_x": int(video_x or 0),
        "video_y": int(video_y or 0),
        "rotation_deg_clockwise": int(rotation_deg_clockwise),
        "pipeline_stage": "official_scale_line21",
        "frames": list(frames),
    }
    context_json.write_text(json.dumps(context, ensure_ascii=False, indent=1), encoding="utf-8")

    cmd = [
        python_bin,
        script.as_posix(),
        "--context-json",
        context_json.as_posix(),
        "--output-dir",
        stage_dir.as_posix(),
        "--max-frames",
        str(int(max_frames)),
    ]
    if corrections_path is not None and corrections_path.is_file():
        cmd.extend(["--corrections", corrections_path.as_posix()])
    try:
        proc = subprocess.run(
            cmd,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            timeout=(float(subprocess_timeout_sec) if float(subprocess_timeout_sec) > 0 else None),
        )
        log_path.write_text(proc.stdout or "", encoding="utf-8")
    except subprocess.TimeoutExpired as exc:
        out["source"] = "scale_timeout"
        out["error"] = f"timeout dopo {subprocess_timeout_sec:.1f}s"
        log_path.write_text(str(exc.stdout or exc), encoding="utf-8")
        return out, []
    except Exception as exc:  # noqa: BLE001
        out["source"] = "scale_exception"
        out["error"] = str(exc)
        log_path.write_text(str(exc), encoding="utf-8")
        return out, []

    if proc.returncode != 0:
        out["source"] = "scale_failed"
        out["error"] = f"returncode={proc.returncode}"
        return out, []
    if not summary_path.is_file():
        out["source"] = "scale_no_summary"
        out["error"] = f"summary non trovato: {summary_path}"
        return out, []

    try:
        summary = json.loads(summary_path.read_text(encoding="utf-8"))
    except Exception as exc:  # noqa: BLE001
        out["source"] = "scale_read_error"
        out["error"] = str(exc)
        return out, []

    for key in (
        "status", "source", "profile", "frames_studied", "depths_total", "depths_accepted",
        "depths_review", "depths_reject", "depths_interpolated", "acceptance_ratio",
        "line_18_vect_depth", "line_19_pixel_ratio_x", "line_20_pixel_ratio_y",
        "line_21_scale_line", "review_reasons",
    ):
        if key in summary:
            out[key] = summary[key]
    zone = summary.get("zone") or {}
    out["ruler_x"] = "" if not isinstance(zone, dict) else str(zone.get("x", "") or "")

    rows: List[Dict[str, object]] = []
    if per_image_csv.is_file():
        try:
            with per_image_csv.open("r", encoding="utf-8", newline="") as handle:
                for row in csv.DictReader(handle):
                    clean = {str(k): ("" if v is None else str(v)) for k, v in row.items()}
                    clean["folder_path"] = folder.as_posix()
                    clean["folder_name"] = folder.name
                    clean["folder_index"] = str(int(folder_index))
                    clean["scale_run_dir"] = stage_dir.as_posix()
                    rows.append(clean)
        except Exception as exc:  # noqa: BLE001
            out["error"] = f"lettura per-immagine fallita: {exc}"

    depths_total = int(out.get("depths_total", 0) or 0)
    if depths_total <= 0:
        out["status"] = "review"
        if not out["error"]:
            out["error"] = "nessuna depth con una risposta di scala"
    elif float(out.get("acceptance_ratio", 0.0) or 0.0) < float(min_accepted_ratio):
        out["status"] = "review"
        out["error"] = out["error"] or (
            f"accepted_ratio={float(out['acceptance_ratio']):.3f} "
            f"< min={float(min_accepted_ratio):.3f}"
        )
    return out, rows


def _lr_marker_tooling_available() -> bool:
    return all(
        item is not None
        for item in (
            LR_MARKER_TEMPLATE_RE,
            _lr_marker_load_template,
            _lr_marker_locate_match_in_abs_rect,
            _lr_marker_label_from_side,
            _lr_marker_roi_from_sugiu,
            _lr_marker_score_of,
            _lr_marker_infer_manufacturer,
        )
    )


def _iter_lr_marker_config_folders(dataset_roots: Sequence[Path], max_depth: int) -> Iterable[Path]:
    skip_dirs = set(LR_MARKER_SKIP_DIRS or set())
    for root in dataset_roots:
        root = root.expanduser().resolve()
        if not root.is_dir():
            continue
        root_depth = len(root.parts)
        for dirpath, dirnames, _filenames in os.walk(root):
            depth = max(0, len(Path(dirpath).resolve().parts) - root_depth)
            dirnames[:] = [name for name in dirnames if name not in skip_dirs and not name.startswith(".")]
            if max_depth >= 0 and depth > max_depth:
                dirnames[:] = []
                continue
            names = set(dirnames)
            if {"DB_setup", "DB_echo", "image_samples"}.issubset(names):
                folder = Path(dirpath).resolve()
                yield folder
                dirnames[:] = [name for name in dirnames if name not in {"DB_setup", "DB_echo", "image_samples"}]


def _template_digest_for_dedup(path: Path) -> str:
    with Image.open(path) as img:
        gray = img.convert("L")
        payload = gray.tobytes() + str(gray.size).encode("ascii")
    return hashlib.sha1(payload).hexdigest()


def _normalize_lr_marker_review_vendor_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())


def _load_lr_marker_template_review(path: Optional[Path]) -> Dict[str, object]:
    if path is None:
        return {"available": False, "error": "missing_review_file_path"}
    review_path = path.expanduser()
    if not review_path.is_absolute():
        review_path = (REPO_ROOT / review_path).resolve()
    else:
        review_path = review_path.resolve()
    if not review_path.is_file():
        return {
            "available": False,
            "path": review_path.as_posix(),
            "error": "review_file_not_found",
        }
    try:
        payload = json.loads(review_path.read_text(encoding="utf-8"))
    except Exception as exc:
        return {
            "available": False,
            "path": review_path.as_posix(),
            "error": f"review_file_invalid:{exc}",
        }
    if not isinstance(payload, dict):
        return {
            "available": False,
            "path": review_path.as_posix(),
            "error": "review_file_not_object",
        }
    raw_vendors = payload.get("vendors", {})
    if not isinstance(raw_vendors, dict):
        raw_vendors = {}
    vendors: Dict[str, Dict[str, object]] = {}
    for vendor_name, entry_raw in raw_vendors.items():
        if not isinstance(entry_raw, dict):
            continue
        key = _normalize_lr_marker_review_vendor_key(vendor_name)
        if not key:
            continue
        accepted = entry_raw.get("accepted", [])
        rejected = entry_raw.get("rejected", [])
        if not isinstance(accepted, list):
            accepted = []
        if not isinstance(rejected, list):
            rejected = []
        vendors[key] = {
            "vendor": str(entry_raw.get("vendor", vendor_name) or vendor_name),
            "min_area": int(float(entry_raw.get("min_area", 0) or 0)),
            "min_short_side": int(float(entry_raw.get("min_short_side", 0) or 0)),
            "accepted": {str(x).replace("\\", "/").lstrip("/") for x in accepted if str(x).strip()},
            "rejected": {str(x).replace("\\", "/").lstrip("/") for x in rejected if str(x).strip()},
        }
    library_root_raw = str(payload.get("library_root", "") or "").strip()
    library_root = Path(library_root_raw).expanduser() if library_root_raw else review_path.parent
    if not library_root.is_absolute():
        library_root = (REPO_ROOT / library_root).resolve()
    else:
        library_root = library_root.resolve()
    return {
        "available": True,
        "path": review_path.as_posix(),
        "library_root": library_root.as_posix(),
        "vendors": vendors,
    }


def _lr_marker_review_vendor_entry(review: Dict[str, object], vendor_name: str) -> Optional[Dict[str, object]]:
    vendors = review.get("vendors", {})
    if not isinstance(vendors, dict):
        return None
    key = _normalize_lr_marker_review_vendor_key(vendor_name)
    entry = vendors.get(key)
    return entry if isinstance(entry, dict) else None


def _lr_marker_review_rel_path(path: Path, library_root: Path) -> str:
    try:
        return path.expanduser().resolve().relative_to(library_root.expanduser().resolve()).as_posix()
    except Exception:
        return path.name


def _lr_marker_review_template_decision(
    *,
    review: Dict[str, object],
    vendor_name: str,
    template_path: Path,
    width: int,
    height: int,
) -> str:
    entry = _lr_marker_review_vendor_entry(review, vendor_name)
    if not entry:
        return ""
    library_root_raw = str(review.get("library_root", "") or "").strip()
    library_root = Path(library_root_raw).expanduser().resolve() if library_root_raw else template_path.parent.parent
    rel_path = _lr_marker_review_rel_path(template_path, library_root)
    accepted = entry.get("accepted", set())
    rejected = entry.get("rejected", set())
    if not isinstance(accepted, set):
        accepted = {str(x).replace("\\", "/").lstrip("/") for x in accepted} if isinstance(accepted, list) else set()
    if not isinstance(rejected, set):
        rejected = {str(x).replace("\\", "/").lstrip("/") for x in rejected} if isinstance(rejected, list) else set()
    path_keys = {
        rel_path,
        template_path.name,
        template_path.expanduser().resolve().as_posix(),
    }
    if path_keys.intersection(rejected):
        return "rejected"
    if path_keys.intersection(accepted):
        return "accepted"
    if accepted or rejected:
        return "rejected"
    try:
        min_area = int(entry.get("min_area", 0) or 0)
        min_short = int(entry.get("min_short_side", 0) or 0)
    except (TypeError, ValueError):
        min_area = 0
        min_short = 0
    if min_area > 0 or min_short > 0:
        area = int(max(0, width) * max(0, height))
        short_side = int(min(max(0, width), max(0, height)))
        if area < max(0, min_area) or short_side < max(0, min_short):
            return "rejected"
        return "accepted"
    return ""


def _load_lr_marker_templates_by_vendor(
    *,
    roots: Sequence[Path],
    vendor_name: str,
    device: torch.device,
    max_depth: int,
    blank_template_max_value: int,
    exclude_vendors: Sequence[str],
    review_file: Optional[Path] = None,
) -> Tuple[List[object], Dict[str, object]]:
    if not _lr_marker_tooling_available():
        return [], {"available": False, "error": "lr_marker_tooling_unavailable"}
    vendor_norm = str(vendor_name or "").strip().lower()
    if not vendor_norm:
        return [], {"available": False, "error": "missing_vendor"}
    excluded = {str(v).strip().lower() for v in exclude_vendors if str(v).strip()}
    if vendor_norm in excluded:
        return [], {"available": False, "error": f"vendor_excluded:{vendor_name}"}

    assert LR_MARKER_TEMPLATE_RE is not None
    assert _lr_marker_load_template is not None
    assert _lr_marker_infer_manufacturer is not None

    review = _load_lr_marker_template_review(review_file)
    review_entry = _lr_marker_review_vendor_entry(review, vendor_name) if bool(review.get("available", False)) else None
    review_active_for_vendor = bool(review_entry)
    review_library_root_raw = str(review.get("library_root", "") or "").strip()
    review_library_root = Path(review_library_root_raw).expanduser().resolve() if review_library_root_raw else None

    templates: List[object] = []
    seen_digests: set[str] = set()
    configs_seen = 0
    files_seen = 0
    blank_or_unusable = 0
    library_files_seen = 0
    review_accepted = 0
    review_rejected = 0
    review_unmentioned = 0
    for root in roots:
        root = root.expanduser().resolve()
        if not root.is_dir() or not (root / "summary.json").is_file():
            continue
        try:
            vendor_dirs = sorted([p for p in root.iterdir() if p.is_dir()], key=lambda p: p.name.lower())
        except Exception:
            continue
        for vendor_dir in vendor_dirs:
            if vendor_dir.name.lower() != vendor_norm:
                continue
            for template_path in sorted(vendor_dir.glob("*.png")):
                if not template_path.is_file():
                    continue
                library_files_seen += 1
                try:
                    with Image.open(template_path) as img_for_review:
                        width_for_review, height_for_review = img_for_review.size
                except Exception:
                    width_for_review, height_for_review = 0, 0
                if review_active_for_vendor and review_library_root is not None:
                    decision = _lr_marker_review_template_decision(
                        review=review,
                        vendor_name=vendor_name,
                        template_path=template_path,
                        width=int(width_for_review),
                        height=int(height_for_review),
                    )
                    if decision == "rejected":
                        review_rejected += 1
                        continue
                    if decision == "accepted":
                        review_accepted += 1
                    else:
                        review_unmentioned += 1
                try:
                    digest = _template_digest_for_dedup(template_path)
                except Exception:
                    blank_or_unusable += 1
                    continue
                if digest in seen_digests:
                    continue
                seen_digests.add(digest)
                template = _lr_marker_load_template(
                    template_path,
                    device=device,
                    blank_template_max_value=int(blank_template_max_value),
                )
                if template is None:
                    blank_or_unusable += 1
                    continue
                templates.append(template)

    config_scan_skipped_by_review = bool(review_active_for_vendor)
    if not config_scan_skipped_by_review:
        for config_dir in _iter_lr_marker_config_folders(roots, max_depth=max_depth):
            manufacturer = str(_lr_marker_infer_manufacturer(config_dir.name) or "").strip()
            if manufacturer.lower() != vendor_norm:
                continue
            configs_seen += 1
            for template_path in sorted((config_dir / "DB_echo").rglob("*")):
                if (
                    not template_path.is_file()
                    or not LR_MARKER_TEMPLATE_RE.match(template_path.name)
                    or template_path.stem.lower() != "orientation_0"
                ):
                    continue
                files_seen += 1
                try:
                    digest = _template_digest_for_dedup(template_path)
                except Exception:
                    blank_or_unusable += 1
                    continue
                if digest in seen_digests:
                    continue
                seen_digests.add(digest)
                template = _lr_marker_load_template(
                    template_path,
                    device=device,
                    blank_template_max_value=int(blank_template_max_value),
                )
                if template is None:
                    blank_or_unusable += 1
                    continue
                templates.append(template)

    meta = {
        "available": bool(templates),
        "error": "" if templates else "no_usable_vendor_templates",
        "vendor": vendor_name,
        "configs_seen": int(configs_seen),
        "template_files_seen": int(files_seen),
        "prepared_library_files_seen": int(library_files_seen),
        "unique_usable_templates": int(len(templates)),
        "blank_or_unusable_templates": int(blank_or_unusable),
        "review_file": str(review.get("path", "") or ""),
        "review_available": bool(review.get("available", False)),
        "review_error": str(review.get("error", "") or ""),
        "review_active_for_vendor": bool(review_active_for_vendor),
        "review_library_root": str(review.get("library_root", "") or ""),
        "review_accepted_templates": int(review_accepted),
        "review_rejected_templates": int(review_rejected),
        "review_unmentioned_templates": int(review_unmentioned),
        "config_scan_skipped_by_review": bool(config_scan_skipped_by_review),
    }
    if review_active_for_vendor and not templates:
        meta["error"] = "no_usable_vendor_templates_after_review"
        meta["available"] = False
    return templates, meta


def _lr_marker_darkness_pct_for_sugiu_row(row_sg: Dict[str, object]) -> float:
    image_path_s = str(row_sg.get("image_path", "") or "").strip()
    if not image_path_s:
        return -1.0
    try:
        image_path = Path(image_path_s).expanduser().resolve()
        with Image.open(image_path) as img:
            gray = img.convert("L")
            width, height = gray.size
            top = max(0, min(_safe_int_value(str(row_sg.get("crop_top", "0") or "0"), 0), height))
            left = max(0, min(_safe_int_value(str(row_sg.get("crop_left", "0") or "0"), 0), width))
            bottom = max(0, min(_safe_int_value(str(row_sg.get("crop_bottom", "0") or "0"), 0), height))
            right = max(0, min(_safe_int_value(str(row_sg.get("crop_right", "0") or "0"), 0), width))
            if bottom <= top or right <= left:
                top, left, bottom, right = 0, 0, height, width
            region = gray.crop((left, top, right, bottom))
            total = int(region.size[0] * region.size[1])
            if total <= 0:
                return -1.0
            hist = region.histogram()
            dark = int(sum(hist[:33]))
            return (float(dark) / float(total)) * 100.0
    except Exception:
        return -1.0


def _extract_expanded_lr_marker_patch(
    *,
    image_path: Path,
    marker_box_abs: Tuple[int, int, int, int],
    blank_template_max_value: int,
) -> Optional[Image.Image]:
    """Extract a tight folder-local marker template around a historical match."""
    try:
        with Image.open(image_path) as sample_img:
            rgb = sample_img.convert("RGB")
            gray = sample_img.convert("L")
            width, height = gray.size
            marker_top, marker_left, marker_bottom, marker_right = marker_box_abs
            marker_top = max(0, min(height - 1, int(marker_top)))
            marker_bottom = max(0, min(height - 1, int(marker_bottom)))
            marker_left = max(0, min(width - 1, int(marker_left)))
            marker_right = max(0, min(width - 1, int(marker_right)))
            if marker_bottom <= marker_top or marker_right <= marker_left:
                return None

            seed_w = marker_right - marker_left + 1
            seed_h = marker_bottom - marker_top + 1
            seed_side = max(seed_w, seed_h)
            margin = min(24, max(6, int(round(seed_side * 0.45))))
            local_left = max(0, marker_left - margin)
            local_top = max(0, marker_top - margin)
            local_right = min(width - 1, marker_right + margin)
            local_bottom = min(height - 1, marker_bottom + margin)
            local = np.asarray(
                gray.crop((local_left, local_top, local_right + 1, local_bottom + 1)),
                dtype=np.float32,
            )
            if local.ndim != 2 or local.size <= 0:
                return None

            border = np.concatenate([local[0, :], local[-1, :], local[:, 0], local[:, -1]])
            bg = float(np.median(border)) if border.size else 0.0
            contrast_mask = np.abs(local - bg) >= 24.0
            bright_mask = local >= max(float(blank_template_max_value) + 8.0, bg + 18.0)
            mask = contrast_mask | bright_mask
            if bg >= 48.0:
                mask |= local <= max(0.0, bg - 18.0)
            if not bool(mask.any()):
                return None

            # Join nearby strokes touching the historical match, without pulling
            # distant ruler ticks or text into the derived folder template.
            radius = 1 if seed_side < 28 else 2
            dilated = mask.copy()
            for _ in range(radius):
                expanded = dilated.copy()
                for dy in (-1, 0, 1):
                    for dx in (-1, 0, 1):
                        if dx == 0 and dy == 0:
                            continue
                        src_y0 = max(0, -dy)
                        src_y1 = dilated.shape[0] - max(0, dy)
                        src_x0 = max(0, -dx)
                        src_x1 = dilated.shape[1] - max(0, dx)
                        dst_y0 = max(0, dy)
                        dst_y1 = expanded.shape[0] - max(0, -dy)
                        dst_x0 = max(0, dx)
                        dst_x1 = expanded.shape[1] - max(0, -dx)
                        if src_y1 > src_y0 and src_x1 > src_x0:
                            expanded[dst_y0:dst_y1, dst_x0:dst_x1] |= dilated[src_y0:src_y1, src_x0:src_x1]
                dilated = expanded

            seed_y0 = max(0, marker_top - local_top)
            seed_y1 = min(dilated.shape[0], marker_bottom - local_top + 1)
            seed_x0 = max(0, marker_left - local_left)
            seed_x1 = min(dilated.shape[1], marker_right - local_left + 1)
            seed_points = np.argwhere(dilated[seed_y0:seed_y1, seed_x0:seed_x1])
            if seed_points.size <= 0:
                patch = rgb.crop((marker_left, marker_top, marker_right + 1, marker_bottom + 1))
                patch = _trim_uniform_dark_border(patch, dark_max=50, keep_px=0)
                extrema = patch.convert("L").getextrema()
                if extrema is None or int(extrema[1]) <= int(blank_template_max_value):
                    return None
                return patch

            seen = np.zeros_like(dilated, dtype=bool)
            stack: List[Tuple[int, int]] = []
            for yy, xx in seed_points:
                y = int(seed_y0 + yy)
                x = int(seed_x0 + xx)
                if not seen[y, x]:
                    seen[y, x] = True
                    stack.append((y, x))
            while stack:
                y, x = stack.pop()
                for yy in range(max(0, y - 1), min(dilated.shape[0], y + 2)):
                    for xx in range(max(0, x - 1), min(dilated.shape[1], x + 2)):
                        if not seen[yy, xx] and bool(dilated[yy, xx]):
                            seen[yy, xx] = True
                            stack.append((yy, xx))

            ys, xs = np.where(seen & mask)
            if len(xs) <= 0 or len(ys) <= 0:
                return None
            max_extra = min(24, max(4, int(round(seed_side * 0.35))))
            comp_left = max(marker_left - max_extra, min(marker_left, int(xs.min()) + local_left))
            comp_top = max(marker_top - max_extra, min(marker_top, int(ys.min()) + local_top))
            comp_right = min(marker_right + max_extra, max(marker_right, int(xs.max()) + local_left))
            comp_bottom = min(marker_bottom + max_extra, max(marker_bottom, int(ys.max()) + local_top))
            pad = min(3, max(1, int(round(seed_side * 0.05))))
            crop_left = max(0, comp_left - pad)
            crop_top = max(0, comp_top - pad)
            crop_right = min(width - 1, comp_right + pad)
            crop_bottom = min(height - 1, comp_bottom + pad)
            patch = rgb.crop((crop_left, crop_top, crop_right + 1, crop_bottom + 1))
            patch = _trim_uniform_dark_border(patch, dark_max=50, keep_px=0)
            extrema = patch.convert("L").getextrema()
            if extrema is None or int(extrema[1]) <= int(blank_template_max_value):
                return None
            return patch
    except Exception:
        return None


def _trim_uniform_dark_border(patch: Image.Image, dark_max: int = 18, keep_px: int = 1) -> Image.Image:
    try:
        gray = np.asarray(patch.convert("L"), dtype=np.uint8)
        if gray.ndim != 2 or gray.size <= 0:
            return patch
        top = 0
        bottom = gray.shape[0] - 1
        left = 0
        right = gray.shape[1] - 1
        while top < bottom and int(gray[top, :].max()) <= int(dark_max):
            top += 1
        while bottom > top and int(gray[bottom, :].max()) <= int(dark_max):
            bottom -= 1
        while left < right and int(gray[:, left].max()) <= int(dark_max):
            left += 1
        while right > left and int(gray[:, right].max()) <= int(dark_max):
            right -= 1
        top = max(0, top - int(keep_px))
        left = max(0, left - int(keep_px))
        bottom = min(gray.shape[0] - 1, bottom + int(keep_px))
        right = min(gray.shape[1] - 1, right + int(keep_px))
        if bottom <= top or right <= left:
            return patch
        if top == 0 and left == 0 and bottom == gray.shape[0] - 1 and right == gray.shape[1] - 1:
            return patch
        return patch.crop((left, top, right + 1, bottom + 1))
    except Exception:
        return patch


def _parse_lr_marker_manual_rect(value: object) -> Optional[Tuple[int, int, int, int]]:
    if isinstance(value, (list, tuple)) and len(value) >= 4:
        raw_parts = list(value[:4])
    else:
        raw_parts = re.split(r"[|,;\s]+", str(value or "").strip())
    vals: List[int] = []
    for part in raw_parts:
        try:
            vals.append(int(round(float(str(part).strip()))))
        except (TypeError, ValueError):
            continue
    if len(vals) < 4:
        return None
    top, left, bottom, right = vals[:4]
    if bottom <= top or right <= left:
        return None
    return top, left, bottom, right


def _load_lr_marker_manual_seed_rows(path: Optional[Path]) -> List[Dict[str, object]]:
    if path is None:
        return []
    seed_path = path.expanduser().resolve()
    if not seed_path.is_file():
        return []
    try:
        payload = json.loads(seed_path.read_text(encoding="utf-8"))
    except Exception:
        return []
    raw_items: object
    if isinstance(payload, dict):
        raw_items = payload.get("seeds", payload.get("lr_marker_manual_seeds", []))
    else:
        raw_items = payload
    if not isinstance(raw_items, list):
        return []
    rows: List[Dict[str, object]] = []
    for idx, item in enumerate(raw_items, start=1):
        if not isinstance(item, dict):
            continue
        rect = _parse_lr_marker_manual_rect(item.get("rect", item.get("marker_rect", "")))
        if rect is None:
            continue
        flag = str(item.get("flag", "") or "").strip().lower()
        correction = str(item.get("correction", "") or "").strip()
        if flag == "exclude" or correction.lower() == "exclude":
            continue
        row = dict(item)
        row["seed_index"] = int(item.get("seed_index", idx) or idx)
        row["rect"] = ",".join(str(v) for v in rect)
        rows.append(row)
    return rows


def _resolve_lr_marker_manual_seed_image(
    seed: Dict[str, object],
    *,
    folder: Path,
    su_giu_rows: Sequence[Dict[str, object]],
) -> Optional[Path]:
    candidates: List[Path] = []
    for key in ("image_path", "source_image_path"):
        value = str(seed.get(key, "") or "").strip()
        if value:
            p = Path(value).expanduser()
            candidates.append(p if p.is_absolute() else folder / p)

    rel_value = str(seed.get("image_rel", "") or "").strip().replace("\\", "/")
    if rel_value:
        rel = Path(rel_value)
        candidates.append(folder / rel)
        candidates.append(folder.parent / rel)
        parts = [part for part in rel_value.split("/") if part]
        for marker in ("input_ref", "input"):
            if marker in parts:
                pos = parts.index(marker)
                tail = parts[pos + 1 :]
                if tail and tail[0] == folder.name:
                    candidates.append(folder / Path(*tail[1:]))
                elif tail:
                    candidates.append(folder.parent / Path(*tail))
        if parts and parts[0] == folder.name:
            candidates.append(folder / Path(*parts[1:]))
        if folder.name in parts:
            pos = parts.index(folder.name)
            if pos + 1 < len(parts):
                candidates.append(folder / Path(*parts[pos + 1 :]))

    for candidate in candidates:
        try:
            resolved = candidate.expanduser().resolve()
            if resolved.is_file():
                return resolved
        except Exception:
            continue

    try:
        seed_index = int(seed.get("image_index", -999999) or -999999)
    except (TypeError, ValueError):
        seed_index = -999999
    if seed_index != -999999:
        for row in su_giu_rows:
            try:
                if int(row.get("image_index", -999999) or -999999) != seed_index:
                    continue
            except (TypeError, ValueError):
                continue
            row_path = str(row.get("image_path", "") or "").strip()
            if row_path:
                p = Path(row_path).expanduser().resolve()
                if p.is_file():
                    return p

    wanted_name = Path(rel_value).name if rel_value else ""
    if wanted_name:
        for row in su_giu_rows:
            row_path = str(row.get("image_path", "") or "").strip()
            if row_path and Path(row_path).name == wanted_name:
                p = Path(row_path).expanduser().resolve()
                if p.is_file():
                    return p
    return None


def _extract_manual_lr_marker_seed_patch(
    *,
    image_path: Path,
    marker_box_abs: Tuple[int, int, int, int],
    blank_template_max_value: int,
) -> Optional[Image.Image]:
    patch = _extract_expanded_lr_marker_patch(
        image_path=image_path,
        marker_box_abs=marker_box_abs,
        blank_template_max_value=blank_template_max_value,
    )
    if patch is None:
        try:
            with Image.open(image_path) as img:
                rgb = img.convert("RGB")
                width, height = rgb.size
                top, left, bottom, right = marker_box_abs
                top = max(0, min(height - 1, int(top)))
                bottom = max(0, min(height - 1, int(bottom)))
                left = max(0, min(width - 1, int(left)))
                right = max(0, min(width - 1, int(right)))
                if bottom <= top or right <= left:
                    return None
                patch = rgb.crop((left, top, right + 1, bottom + 1))
                patch = _trim_uniform_dark_border(patch, dark_max=50, keep_px=0)
        except Exception:
            return None
    patch_w, patch_h = int(patch.size[0]), int(patch.size[1])
    if (
        patch_w * patch_h < int(LR_MARKER_MIN_MANUAL_TEMPLATE_AREA)
        or min(patch_w, patch_h) < int(LR_MARKER_MIN_MANUAL_TEMPLATE_SHORT_SIDE)
    ):
        return None
    extrema = patch.convert("L").getextrema()
    if extrema is None or int(extrema[1]) <= int(blank_template_max_value):
        return None
    return patch


def _build_lr_marker_manual_seed_templates(
    *,
    seeds: Sequence[Dict[str, object]],
    su_giu_rows: Sequence[Dict[str, object]],
    folder: Path,
    output_dir: Path,
    vendor_name: str,
    device: torch.device,
    blank_template_max_value: int,
) -> Tuple[List[object], List[Dict[str, object]]]:
    if not seeds or not _lr_marker_tooling_available():
        return [], []
    assert _lr_marker_load_template is not None
    template_dir = output_dir / "lr_marker_manual_seed_templates"
    templates: List[object] = []
    meta_rows: List[Dict[str, object]] = []
    safe_vendor = re.sub(r"[^A-Za-z0-9_.-]+", "_", vendor_name.strip() or "vendor").strip("_") or "vendor"
    for idx, seed in enumerate(seeds, start=1):
        rect = _parse_lr_marker_manual_rect(seed.get("rect", ""))
        meta: Dict[str, object] = {
            "seed_index": int(seed.get("seed_index", idx) or idx),
            "review_key": str(seed.get("review_key", "") or ""),
            "image_rel": str(seed.get("image_rel", "") or ""),
            "rect": str(seed.get("rect", "") or ""),
            "correction": str(seed.get("correction", "") or ""),
            "status": "rejected",
            "reason": "",
        }
        if rect is None:
            meta["reason"] = "invalid_rect"
            meta_rows.append(meta)
            continue
        image_path = _resolve_lr_marker_manual_seed_image(seed, folder=folder, su_giu_rows=su_giu_rows)
        if image_path is None:
            meta["reason"] = "image_not_found"
            meta_rows.append(meta)
            continue
        patch = _extract_manual_lr_marker_seed_patch(
            image_path=image_path,
            marker_box_abs=rect,
            blank_template_max_value=int(blank_template_max_value),
        )
        if patch is None:
            meta["reason"] = "unusable_patch"
            meta["image_path"] = image_path.as_posix()
            meta_rows.append(meta)
            continue
        try:
            template_dir.mkdir(parents=True, exist_ok=True)
            seed_hash = hashlib.sha1((image_path.as_posix() + str(rect)).encode("utf-8")).hexdigest()[:10]
            out_path = template_dir / f"{safe_vendor}_manual_seed_{idx:03d}_{seed_hash}.png"
            patch.save(out_path)
            template = _lr_marker_load_template(
                out_path,
                device=device,
                blank_template_max_value=int(blank_template_max_value),
            )
            if template is None:
                meta["reason"] = "load_template_rejected"
                meta["image_path"] = image_path.as_posix()
                meta["template_path"] = out_path.as_posix()
                meta_rows.append(meta)
                continue
            templates.append(template)
            meta.update(
                {
                    "status": "accepted",
                    "reason": "",
                    "image_path": image_path.as_posix(),
                    "template_path": out_path.resolve().as_posix(),
                    "template_width": int(patch.size[0]),
                    "template_height": int(patch.size[1]),
                    "template_area": int(patch.size[0] * patch.size[1]),
                }
            )
            meta_rows.append(meta)
        except Exception as exc:
            meta["reason"] = f"exception:{exc}"
            meta["image_path"] = image_path.as_posix()
            meta_rows.append(meta)
            continue
    return templates, meta_rows


def _build_lr_marker_bundle_manual_seed_library(
    *,
    seeds: Sequence[Dict[str, object]],
    su_giu_rows: Sequence[Dict[str, object]],
    folder: Path,
    output_dir: Path,
    vendor_name: str,
    blank_template_max_value: int,
) -> Tuple[Optional[Path], List[Dict[str, object]]]:
    if not seeds:
        return None, []
    library_root = output_dir / "lr_marker_bundle_manual_seed_templates"
    safe_vendor = re.sub(r"[^A-Za-z0-9_.-]+", "_", vendor_name.strip() or "vendor").strip("_") or "vendor"
    vendor_dir = library_root / safe_vendor
    accepted_rels: List[str] = []
    meta_rows: List[Dict[str, object]] = []
    for idx, seed in enumerate(seeds, start=1):
        rect = _parse_lr_marker_manual_rect(seed.get("rect", ""))
        meta: Dict[str, object] = {
            "seed_index": int(seed.get("seed_index", idx) or idx),
            "review_key": str(seed.get("review_key", "") or ""),
            "image_rel": str(seed.get("image_rel", "") or ""),
            "image_index": str(seed.get("image_index", "") or ""),
            "rect": str(seed.get("rect", "") or ""),
            "correction": str(seed.get("correction", "") or "").strip().upper(),
            "status": "rejected",
            "reason": "",
            "source": "bundle_manual_seed",
        }
        if rect is None:
            meta["reason"] = "invalid_rect"
            meta_rows.append(meta)
            continue
        correction = str(meta.get("correction", "") or "").strip().upper()
        if correction and correction not in {"NF", "LR", "UD", "LRUD"}:
            meta["reason"] = "invalid_correction"
            meta_rows.append(meta)
            continue
        image_path = _resolve_lr_marker_manual_seed_image(seed, folder=folder, su_giu_rows=su_giu_rows)
        if image_path is None:
            meta["reason"] = "image_not_found"
            meta_rows.append(meta)
            continue
        patch = _extract_manual_lr_marker_seed_patch(
            image_path=image_path,
            marker_box_abs=rect,
            blank_template_max_value=int(blank_template_max_value),
        )
        if patch is None:
            meta["reason"] = "unusable_patch"
            meta["image_path"] = image_path.as_posix()
            meta_rows.append(meta)
            continue
        try:
            vendor_dir.mkdir(parents=True, exist_ok=True)
            out_path = vendor_dir / f"marker_{len(accepted_rels) + 1:03d}.png"
            patch.save(out_path)
            rel = f"{safe_vendor}/{out_path.name}"
            accepted_rels.append(rel)
            meta.update(
                {
                    "status": "accepted",
                    "reason": "",
                    "image_path": image_path.as_posix(),
                    "template_path": out_path.resolve().as_posix(),
                    "template_rel": rel,
                    "template_width": int(patch.size[0]),
                    "template_height": int(patch.size[1]),
                    "template_area": int(patch.size[0] * patch.size[1]),
                }
            )
            meta_rows.append(meta)
        except Exception as exc:
            meta["reason"] = f"exception:{exc}"
            meta["image_path"] = image_path.as_posix()
            meta_rows.append(meta)
            continue

    if not accepted_rels:
        return None, meta_rows
    decisions = {
        "source": "pipeline_manual_lr_marker_seeds",
        "vendors": {
            safe_vendor: {
                "vendor": safe_vendor,
                "accepted": accepted_rels,
                "rejected": [],
            }
        },
    }
    try:
        library_root.mkdir(parents=True, exist_ok=True)
        (library_root / "review_decisions.json").write_text(json.dumps(decisions, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception as exc:
        for meta in meta_rows:
            if str(meta.get("status", "") or "") == "accepted":
                meta["status"] = "rejected"
                meta["reason"] = f"review_decisions_write_failed:{exc}"
        return None, meta_rows
    return library_root, meta_rows


def _lr_marker_patch_stats_from_gray(
    gray: np.ndarray,
    marker_box_abs: Tuple[int, int, int, int],
    *,
    dark_threshold: int = 30,
) -> Dict[str, object]:
    try:
        if gray.ndim != 2 or gray.size <= 0:
            return {"blank": True, "dark_pct": 100.0, "max_value": 0, "std": 0.0}
        height, width = gray.shape
        top, left, bottom, right = marker_box_abs
        top = max(0, min(height - 1, int(top)))
        bottom = max(0, min(height - 1, int(bottom)))
        left = max(0, min(width - 1, int(left)))
        right = max(0, min(width - 1, int(right)))
        if bottom <= top or right <= left:
            return {"blank": True, "dark_pct": 100.0, "max_value": 0, "std": 0.0}
        patch = gray[top : bottom + 1, left : right + 1]
        if patch.size <= 0:
            return {"blank": True, "dark_pct": 100.0, "max_value": 0, "std": 0.0}
        max_value = int(patch.max())
        std_value = float(patch.std())
        dark_pct = float((patch <= int(dark_threshold)).mean() * 100.0)
        blank = bool(max_value <= 18 or (dark_pct >= 98.0 and std_value <= 8.0))
        return {
            "blank": blank,
            "dark_pct": dark_pct,
            "max_value": max_value,
            "std": std_value,
        }
    except Exception:
        return {"blank": True, "dark_pct": 100.0, "max_value": 0, "std": 0.0}


def _lr_marker_row_matches_target_size(
    row: Dict[str, object],
    *,
    target_image_width: int = 0,
    target_image_height: int = 0,
) -> bool:
    if int(target_image_width or 0) <= 0 or int(target_image_height or 0) <= 0:
        return True
    try:
        image_w = int(row.get("image_width", 0) or 0)
        image_h = int(row.get("image_height", 0) or 0)
    except (TypeError, ValueError):
        return True
    if image_w <= 0 or image_h <= 0:
        return True
    return image_w == int(target_image_width) and image_h == int(target_image_height)


def _lr_marker_truthy_flag(value: object) -> bool:
    if isinstance(value, str):
        return value.strip().lower() in {"1", "true", "yes", "y", "on"}
    try:
        return bool(int(value))  # type: ignore[arg-type]
    except Exception:
        return bool(value)


def _lr_marker_row_source_orientation_group(row: Dict[str, object]) -> str:
    return _infer_lr_marker_orientation_group(str(row.get("image_path", "") or ""))


def _lr_marker_row_has_textured_match(row: Dict[str, object]) -> bool:
    if _lr_marker_truthy_flag(row.get("match_patch_is_blank", False)):
        return False
    try:
        return float(row.get("match_patch_std", 0.0) or 0.0) >= float(LR_MARKER_MIN_TEXTURE_STD)
    except (TypeError, ValueError):
        return False


def _lr_marker_template_size(template: object) -> Tuple[int, int]:
    try:
        width = int(getattr(template, "width", 0) or 0)
        height = int(getattr(template, "height", 0) or 0)
    except (TypeError, ValueError):
        return 0, 0
    return max(0, width), max(0, height)


def _lr_marker_preferred_historical_templates(templates: Sequence[object]) -> List[object]:
    items = list(templates)
    preferred: List[object] = []
    for template in items:
        width, height = _lr_marker_template_size(template)
        if (
            width * height >= int(LR_MARKER_MIN_HISTORICAL_TEMPLATE_AREA)
            and min(width, height) >= int(LR_MARKER_MIN_HISTORICAL_TEMPLATE_SHORT_SIDE)
        ):
            preferred.append(template)
    return preferred or items


def _lr_marker_expected_label_for_orientation_group(group: str) -> str:
    group_norm = str(group or "").strip().upper()
    if group_norm in {"NF", "UD"}:
        return "not_lr_flipped"
    if group_norm in {"LR", "LRUD"}:
        return "lr_flipped"
    return ""


def _lr_marker_expected_label_for_source_row(row: Dict[str, object]) -> str:
    return _lr_marker_expected_label_for_orientation_group(_lr_marker_row_source_orientation_group(row))


def _lr_marker_row_template_area(row: Dict[str, object]) -> int:
    try:
        width = int(row.get("template_width", 0) or 0)
        height = int(row.get("template_height", 0) or 0)
    except (TypeError, ValueError):
        return 0
    return max(0, width) * max(0, height)


def _lr_marker_template_candidate_rank_rows(rows: Sequence[Dict[str, object]], limit: int = 20) -> List[Dict[str, object]]:
    grouped: Dict[str, List[Dict[str, object]]] = {}
    for row in rows:
        path = str(row.get("template_path", "") or "").strip()
        if path:
            grouped.setdefault(path, []).append(dict(row))
    ranked: List[Dict[str, object]] = []
    for path, group_rows in grouped.items():
        scores: List[float] = []
        expected_total = 0
        expected_correct = 0
        expected_groups: set[str] = set()
        labels: Counter[str] = Counter()
        max_area = 0
        template_width = 0
        template_height = 0
        for item in group_rows:
            try:
                scores.append(float(item.get("match_score", 0.0) or 0.0))
            except (TypeError, ValueError):
                scores.append(0.0)
            label = str(item.get("lr_label", "") or "").strip()
            if label:
                labels[label] += 1
            area = _lr_marker_row_template_area(item)
            if area > max_area:
                max_area = area
                try:
                    template_width = int(item.get("template_width", 0) or 0)
                    template_height = int(item.get("template_height", 0) or 0)
                except (TypeError, ValueError):
                    template_width = 0
                    template_height = 0
            source_group = _lr_marker_row_source_orientation_group(item)
            expected = _lr_marker_expected_label_for_orientation_group(source_group)
            if expected:
                expected_total += 1
                expected_groups.add(source_group)
                if label == expected:
                    expected_correct += 1
        mean_score = sum(scores) / max(1, len(scores))
        median_score = float(statistics.median(scores)) if scores else 0.0
        expected_ratio = expected_correct / max(1, expected_total) if expected_total else 0.0
        ranked.append(
            {
                "template_path": path,
                "support": int(len(group_rows)),
                "mean_score": float(mean_score),
                "median_score": float(median_score),
                "expected_total": int(expected_total),
                "expected_correct": int(expected_correct),
                "expected_ratio": float(expected_ratio),
                "expected_groups": sorted(expected_groups),
                "expected_groups_count": int(len(expected_groups)),
                "label_counts": dict(labels),
                "template_width": int(template_width),
                "template_height": int(template_height),
                "template_area": int(max_area),
            }
        )
    ranked.sort(
        key=lambda item: (
            -float(item.get("expected_ratio", 0.0) or 0.0),
            -int(item.get("expected_correct", 0) or 0),
            -int(item.get("expected_groups_count", 0) or 0),
            -int(item.get("support", 0) or 0),
            -float(item.get("median_score", 0.0) or 0.0),
            -float(item.get("mean_score", 0.0) or 0.0),
            -int(item.get("template_area", 0) or 0),
            str(item.get("template_path", "")),
        )
    )
    return ranked[: max(1, int(limit))]


def _select_lr_marker_folder_template_candidate(rows: Sequence[Dict[str, object]]) -> Dict[str, object]:
    grouped: Dict[str, List[Dict[str, object]]] = {}
    for row in rows:
        path = str(row.get("template_path", "") or "").strip()
        if not path:
            continue
        grouped.setdefault(path, []).append(dict(row))
    if not grouped:
        return max(rows, key=lambda item: float(item.get("match_score", 0.0) or 0.0))

    def _group_key(group_rows: List[Dict[str, object]]) -> Tuple[float, float, float, float, float, float]:
        scores: List[float] = []
        expected_total = 0
        expected_correct = 0
        expected_groups: set[str] = set()
        max_area = 0
        for item in group_rows:
            try:
                scores.append(float(item.get("match_score", 0.0) or 0.0))
            except (TypeError, ValueError):
                scores.append(0.0)
            max_area = max(max_area, _lr_marker_row_template_area(item))
            source_group = _lr_marker_row_source_orientation_group(item)
            expected = _lr_marker_expected_label_for_orientation_group(source_group)
            if expected:
                expected_total += 1
                expected_groups.add(source_group)
                if str(item.get("lr_label", "") or "").strip() == expected:
                    expected_correct += 1
        support = len(group_rows)
        mean_score = sum(scores) / max(1, len(scores))
        median_score = float(statistics.median(scores)) if scores else 0.0
        expected_ratio = expected_correct / max(1, expected_total) if expected_total else 0.0
        if expected_total:
            return (
                1.0,
                expected_ratio,
                float(expected_correct),
                float(len(expected_groups)),
                float(support),
                median_score,
                mean_score,
                float(max_area),
            )
        return (0.0, 0.0, 0.0, 0.0, float(support), median_score, mean_score, float(max_area))

    best_group_rows = max(grouped.values(), key=_group_key)
    return max(
        best_group_rows,
        key=lambda item: (
            float(item.get("match_score", 0.0) or 0.0),
            float(_lr_marker_row_template_area(item)),
        ),
    )


def _lr_marker_orientation_group_from_quadrant(side: str, vertical: str) -> str:
    side_norm = str(side or "").strip().lower()
    vertical_norm = str(vertical or "").strip().lower()
    if side_norm == "left" and vertical_norm == "su":
        return "NF"
    if side_norm == "right" and vertical_norm == "su":
        return "LR"
    if side_norm == "left" and vertical_norm == "giu":
        return "UD"
    if side_norm == "right" and vertical_norm == "giu":
        return "LRUD"
    return ""


def _lr_marker_quadrant_validation_fields(row: Dict[str, object]) -> Dict[str, object]:
    fields: Dict[str, object] = {
        "quadrant_valid": 0,
        "quadrant_status": "unknown",
        "quadrant_reason": "",
        "quadrant_expected": "",
        "quadrant_center": "",
        "quadrant_group": "",
        "quadrant_center_group": "",
        "echo_mid_x_abs": "",
        "echo_mid_y_abs": "",
    }
    side = str(row.get("detected_marker_side", "") or "").strip().lower()
    sugiu = str(row.get("su_giu_pred", "") or "").strip().lower()
    try:
        echo_top = int(row.get("echo_rect_top_abs", 0) or 0)
        echo_left = int(row.get("echo_rect_left_abs", 0) or 0)
        echo_bottom = int(row.get("echo_rect_bottom_abs", 0) or 0)
        echo_right = int(row.get("echo_rect_right_abs", 0) or 0)
        marker_top = int(row.get("marker_top_abs", 0) or 0)
        marker_left = int(row.get("marker_left_abs", 0) or 0)
        marker_bottom = int(row.get("marker_bottom_abs", 0) or 0)
        marker_right = int(row.get("marker_right_abs", 0) or 0)
    except (TypeError, ValueError):
        fields["quadrant_reason"] = "invalid_quadrant_coordinates"
        return fields
    if echo_bottom <= echo_top or echo_right <= echo_left or marker_bottom <= marker_top or marker_right <= marker_left:
        fields["quadrant_reason"] = "invalid_quadrant_rect"
        return fields

    mid_x = (echo_left + echo_right) / 2.0
    mid_y = (echo_top + echo_bottom) / 2.0
    marker_cx = (marker_left + marker_right) / 2.0
    marker_cy = (marker_top + marker_bottom) / 2.0
    center_side = "left" if marker_cx < mid_x else "right"
    center_vertical = "su" if marker_cy < mid_y else "giu"
    expected = f"{side}_{sugiu}" if side in {"left", "right"} and sugiu in {"su", "giu"} else ""
    center = f"{center_side}_{center_vertical}"
    fields.update(
        {
            "quadrant_expected": expected,
            "quadrant_center": center,
            "quadrant_group": _lr_marker_orientation_group_from_quadrant(side, sugiu),
            "quadrant_center_group": _lr_marker_orientation_group_from_quadrant(center_side, center_vertical),
            "echo_mid_x_abs": f"{mid_x:.3f}",
            "echo_mid_y_abs": f"{mid_y:.3f}",
        }
    )

    reasons: List[str] = []
    if side not in {"left", "right"}:
        reasons.append("missing_marker_side")
    if sugiu not in {"su", "giu"}:
        reasons.append("missing_sugiu_pred")
    if side == "left" and marker_right > mid_x:
        reasons.append("marker_crosses_vertical_median")
    elif side == "right" and marker_left < mid_x:
        reasons.append("marker_crosses_vertical_median")
    if sugiu == "su" and marker_bottom > mid_y:
        reasons.append("marker_crosses_horizontal_median")
    elif sugiu == "giu" and marker_top < mid_y:
        reasons.append("marker_crosses_horizontal_median")

    if reasons:
        fields["quadrant_valid"] = 0
        fields["quadrant_status"] = "invalid"
        fields["quadrant_reason"] = ";".join(dict.fromkeys(reasons))
    else:
        fields["quadrant_valid"] = 1
        fields["quadrant_status"] = "ok"
        fields["quadrant_reason"] = ""
    return fields


def _lr_marker_row_quadrant_valid_or_unknown(row: Dict[str, object]) -> bool:
    raw = row.get("quadrant_valid", "")
    if str(raw).strip() != "":
        return _lr_marker_truthy_flag(raw)
    fields = _lr_marker_quadrant_validation_fields(row)
    return str(fields.get("quadrant_status", "") or "") != "invalid"


def _lr_marker_reliable_rows(
    rows: Sequence[Dict[str, object]],
    *,
    min_match_score: float,
    target_image_width: int = 0,
    target_image_height: int = 0,
) -> List[Dict[str, object]]:
    reliable: List[Dict[str, object]] = []
    for row in rows:
        try:
            score = float(row.get("match_score", 0.0) or 0.0)
        except (TypeError, ValueError):
            score = 0.0
        if score < float(min_match_score):
            continue
        if str(row.get("lr_marker_method", "") or "") == "bundle" and str(row.get("status", "") or "") != "ok":
            continue
        if str(row.get("search_strategy", "") or "") == "spatial_consensus_forced":
            continue
        if not _lr_marker_row_matches_target_size(
            row,
            target_image_width=int(target_image_width or 0),
            target_image_height=int(target_image_height or 0),
        ):
            continue
        if not _lr_marker_row_quadrant_valid_or_unknown(dict(row)):
            continue
        review_parts = set(str(row.get("review_reason", "") or "").split(";"))
        if _lr_marker_truthy_flag(row.get("match_patch_is_blank", False)) or "blank_marker_match" in review_parts:
            continue
        if not _lr_marker_row_has_textured_match(row) or "low_texture_marker_match" in review_parts:
            continue
        reliable.append(dict(row))
    return reliable


def _lr_marker_should_use_derived_fallback(
    rows: Sequence[Dict[str, object]],
    *,
    min_match_score: float,
    target_image_width: int = 0,
    target_image_height: int = 0,
) -> Tuple[bool, str]:
    if not rows:
        return True, "historical_no_rows"
    scores: List[float] = []
    comparable_rows: List[Dict[str, object]] = []
    for row in rows:
        try:
            scores.append(float(row.get("match_score", 0.0) or 0.0))
        except (TypeError, ValueError):
            scores.append(0.0)
        if _lr_marker_row_matches_target_size(
            row,
            target_image_width=int(target_image_width or 0),
            target_image_height=int(target_image_height or 0),
        ):
            comparable_rows.append(dict(row))
    best_score = max(scores, default=0.0)
    if best_score < float(min_match_score):
        return True, f"historical_best_score_below_min:{best_score:.4f}"
    reliable = _lr_marker_reliable_rows(
        rows,
        min_match_score=float(min_match_score),
        target_image_width=int(target_image_width or 0),
        target_image_height=int(target_image_height or 0),
    )
    comparable_total = len(comparable_rows) if comparable_rows else len(rows)
    if comparable_total >= 8:
        blank_count = 0
        low_texture_count = 0
        for row in comparable_rows:
            review_parts = set(str(row.get("review_reason", "") or "").split(";"))
            if _lr_marker_truthy_flag(row.get("match_patch_is_blank", False)) or "blank_marker_match" in review_parts:
                blank_count += 1
            if not _lr_marker_row_has_textured_match(row) or "low_texture_marker_match" in review_parts:
                low_texture_count += 1
        blank_ratio = float(blank_count) / float(comparable_total)
        low_texture_ratio = float(low_texture_count) / float(comparable_total)
        if blank_ratio > float(LR_MARKER_MAX_BLANK_ROW_RATIO):
            return True, f"historical_blank_rows_above_max:{blank_count}/{comparable_total}"
        if low_texture_ratio > float(LR_MARKER_MAX_LOW_TEXTURE_ROW_RATIO):
            return True, f"historical_low_texture_rows_above_max:{low_texture_count}/{comparable_total}"
        min_reliable = max(2, int(round(float(comparable_total) * float(LR_MARKER_MIN_RELIABLE_ROW_RATIO))))
        if len(reliable) < min_reliable:
            return True, f"historical_reliable_rows_below_min:{len(reliable)}/{comparable_total}"
    elif not reliable:
        return True, "historical_no_reliable_rows"
    return False, ""


def _predict_lr_marker_on_su_giu_rows(
    *,
    su_giu_rows: Sequence[Dict[str, object]],
    templates: Sequence[object],
    vendor_name: str,
    min_match_score: float,
    full_crop_fallback_threshold: float,
    expanded_search_threshold: float,
    expanded_search_steps: Sequence[float],
    min_sugiu_confidence: float,
    device: torch.device,
    force_single_template: bool = True,
    fixed_template_path: str = "",
    fixed_template_selection_score: float = 0.0,
    derived_template_dir: Optional[Path] = None,
    blank_template_max_value: int = 12,
    folder_template_seed_image_path: str = "",
    folder_template_seed_darkness_pct: float = -1.0,
    folder_template_seed_db_template_path: str = "",
    folder_template_rank_json: str = "",
    folder_template_policy: str = "",
    single_template_policy: str = "fixed_historical_best_template",
    target_image_width: int = 0,
    target_image_height: int = 0,
    prefer_large_historical_templates: bool = True,
) -> List[Dict[str, object]]:
    if not su_giu_rows or not templates:
        return []
    assert _lr_marker_locate_match_in_abs_rect is not None
    assert _lr_marker_roi_from_sugiu is not None
    assert _lr_marker_score_of is not None
    assert _lr_marker_expand_rect is not None
    assert _lr_marker_label_from_side is not None

    def _template_path_text(template: object) -> str:
        path_obj = getattr(template, "path", "")
        try:
            return Path(path_obj).expanduser().resolve().as_posix()
        except Exception:
            return str(path_obj or "")

    selected_template_path = str(fixed_template_path or "").strip()
    selected_template_score = float(fixed_template_selection_score or 0.0)
    if bool(force_single_template) and not selected_template_path:
        selection_templates = (
            _lr_marker_preferred_historical_templates(templates)
            if bool(prefer_large_historical_templates)
            else list(templates)
        )
        candidate_rows = _predict_lr_marker_on_su_giu_rows(
            su_giu_rows=su_giu_rows,
            templates=selection_templates,
            vendor_name=vendor_name,
            min_match_score=min_match_score,
            full_crop_fallback_threshold=full_crop_fallback_threshold,
            expanded_search_threshold=expanded_search_threshold,
            expanded_search_steps=expanded_search_steps,
            min_sugiu_confidence=min_sugiu_confidence,
            device=device,
            force_single_template=False,
            derived_template_dir=derived_template_dir,
            blank_template_max_value=blank_template_max_value,
            target_image_width=int(target_image_width or 0),
            target_image_height=int(target_image_height or 0),
            prefer_large_historical_templates=bool(prefer_large_historical_templates),
        )
        if not candidate_rows:
            return []
        canonical_candidate_rows = [
            row for row in candidate_rows
            if _lr_marker_row_matches_target_size(
                dict(row),
                target_image_width=int(target_image_width or 0),
                target_image_height=int(target_image_height or 0),
            )
            and _lr_marker_row_has_textured_match(dict(row))
            and _lr_marker_row_quadrant_valid_or_unknown(dict(row))
        ]
        if not canonical_candidate_rows:
            return []
        orientation_candidate_rows = [
            row for row in canonical_candidate_rows
            if _lr_marker_row_source_orientation_group(dict(row)) in {"NF", "LR", "UD", "LRUD"}
        ]
        candidate_rows_for_selection = orientation_candidate_rows or canonical_candidate_rows
        candidate_rank_rows = _lr_marker_template_candidate_rank_rows(candidate_rows_for_selection, limit=25)

        def _sample_is_canonical(item: Dict[str, object]) -> bool:
            if int(target_image_width or 0) <= 0 or int(target_image_height or 0) <= 0:
                return True
            image_path_s = str(item.get("image_path", "") or "").strip()
            if not image_path_s:
                return False
            try:
                with Image.open(Path(image_path_s).expanduser().resolve()) as img:
                    width, height = img.size
                return int(width) == int(target_image_width) and int(height) == int(target_image_height)
            except Exception:
                return False

        orientation_sample_rows = [
            row for row in su_giu_rows
            if _lr_marker_row_source_orientation_group(dict(row)) in {"NF", "LR", "UD", "LRUD"}
        ]
        sample_rows = sorted(
            list(orientation_sample_rows or su_giu_rows),
            key=lambda item: (
                0 if _sample_is_canonical(dict(item)) else 1,
                -_lr_marker_darkness_pct_for_sugiu_row(dict(item)),
                int(item.get("image_index", 0) or 0),
            ),
        )
        sample_path = str(sample_rows[0].get("image_path", "") or "").strip() if sample_rows else ""
        sample_darkness = _lr_marker_darkness_pct_for_sugiu_row(dict(sample_rows[0])) if sample_rows else -1.0
        best_candidate = _select_lr_marker_folder_template_candidate(candidate_rows_for_selection)
        selected_template_path = str(best_candidate.get("template_path", "") or "").strip()
        selected_template_score = float(best_candidate.get("match_score", 0.0) or 0.0)
        selected_templates = [t for t in selection_templates if _template_path_text(t) == selected_template_path]
        selected_template_policy = str(single_template_policy or "fixed_historical_best_template") if selected_template_path else ""
        seed_candidate = best_candidate
        if derived_template_dir is not None:
            safe_vendor = re.sub(r"[^A-Za-z0-9_.-]+", "_", vendor_name.strip() or "vendor").strip("_") or "vendor"
            best_candidate_idx = int(best_candidate.get("image_index", -999999) or -999999)
            best_candidate_name = Path(str(best_candidate.get("image_path", "") or "")).name
            sample_rows_for_derivation = sorted(
                sample_rows,
                key=lambda item: (
                    0
                    if int(item.get("image_index", -999999) or -999999) == best_candidate_idx
                    or Path(str(item.get("image_path", "") or "")).name == best_candidate_name
                    else 1,
                    0 if _sample_is_canonical(dict(item)) else 1,
                    -_lr_marker_darkness_pct_for_sugiu_row(dict(item)),
                    int(item.get("image_index", 0) or 0),
                ),
            )
            for sample_row in sample_rows_for_derivation:
                candidate_sample_path = str(sample_row.get("image_path", "") or "").strip()
                candidate_sample_idx = int(sample_row.get("image_index", 0) or 0)
                sample_candidates = [
                    row for row in candidate_rows
                    if int(row.get("image_index", -999999) or -999999) == candidate_sample_idx
                    or Path(str(row.get("image_path", "") or "")).name == Path(candidate_sample_path).name
                ]
                if not sample_candidates or not candidate_sample_path:
                    continue
                sample_candidates = [
                    row for row in sample_candidates
                    if _lr_marker_row_quadrant_valid_or_unknown(dict(row))
                    and _lr_marker_row_has_textured_match(dict(row))
                ]
                if not sample_candidates:
                    continue
                selected_template_candidates = [
                    row for row in sample_candidates
                    if selected_template_path and str(row.get("template_path", "") or "").strip() == selected_template_path
                ]
                sample_candidates_for_patch = selected_template_candidates or sample_candidates
                usable_sample_candidates = [
                    row for row in sample_candidates_for_patch
                    if not _lr_marker_truthy_flag(row.get("match_patch_is_blank", False))
                ] or sample_candidates_for_patch
                candidate = max(usable_sample_candidates, key=lambda item: float(item.get("match_score", 0.0) or 0.0))
                try:
                    marker_top = int(candidate.get("marker_top_abs", 0) or 0)
                    marker_left = int(candidate.get("marker_left_abs", 0) or 0)
                    marker_bottom = int(candidate.get("marker_bottom_abs", 0) or 0)
                    marker_right = int(candidate.get("marker_right_abs", 0) or 0)
                    patch = _extract_expanded_lr_marker_patch(
                        image_path=Path(candidate_sample_path).expanduser().resolve(),
                        marker_box_abs=(marker_top, marker_left, marker_bottom, marker_right),
                        blank_template_max_value=int(blank_template_max_value),
                    )
                    if patch is None:
                        continue
                    patch_w, patch_h = int(patch.size[0]), int(patch.size[1])
                    if (
                        patch_w * patch_h < int(LR_MARKER_MIN_DERIVED_TEMPLATE_AREA)
                        or min(patch_w, patch_h) < int(LR_MARKER_MIN_DERIVED_TEMPLATE_SHORT_SIDE)
                    ):
                        continue
                    derived_template_dir.mkdir(parents=True, exist_ok=True)
                    sample_hash = hashlib.sha1(candidate_sample_path.encode("utf-8")).hexdigest()[:10]
                    derived_path = derived_template_dir / f"{safe_vendor}_frame_{candidate_sample_idx:04d}_{sample_hash}.png"
                    patch.save(derived_path)
                    derived_template = _lr_marker_load_template(
                        derived_path,
                        device=device,
                        blank_template_max_value=int(blank_template_max_value),
                    )
                    if derived_template is None:
                        continue
                    selected_template_path = derived_path.resolve().as_posix()
                    selected_template_score = float(candidate.get("match_score", 0.0) or 0.0)
                    selected_templates = [derived_template]
                    selected_template_policy = "fixed_derived_folder_template"
                    sample_path = candidate_sample_path
                    sample_darkness = _lr_marker_darkness_pct_for_sugiu_row(dict(sample_row))
                    seed_candidate = candidate
                    break
                except Exception:
                    continue
        if derived_template_dir is not None and selected_template_policy != "fixed_derived_folder_template":
            return []
        if not selected_templates:
            return []
        rank_rows: List[Dict[str, object]] = []
        sample_row_for_rank = next(
            (
                row for row in sample_rows
                if str(row.get("image_path", "") or "").strip() == sample_path
            ),
            sample_rows[0] if sample_rows else None,
        )
        if sample_row_for_rank is not None:
            for template in selection_templates:
                try:
                    ranked = _predict_lr_marker_on_su_giu_rows(
                        su_giu_rows=[sample_row_for_rank],
                        templates=[template],
                        vendor_name=vendor_name,
                        min_match_score=min_match_score,
                        full_crop_fallback_threshold=full_crop_fallback_threshold,
                        expanded_search_threshold=expanded_search_threshold,
                        expanded_search_steps=expanded_search_steps,
                        min_sugiu_confidence=min_sugiu_confidence,
                        device=device,
                        force_single_template=False,
                        target_image_width=int(target_image_width or 0),
                        target_image_height=int(target_image_height or 0),
                    )
                    if not ranked:
                        continue
                    ranked_row = ranked[0]
                    rank_rows.append(
                        {
                            "template_path": str(ranked_row.get("template_path", "") or ""),
                            "score": float(ranked_row.get("match_score", 0.0) or 0.0),
                            "template_width": int(ranked_row.get("template_width", 0) or 0),
                            "template_height": int(ranked_row.get("template_height", 0) or 0),
                        }
                    )
                except Exception:
                    continue
        rank_rows.sort(key=lambda item: (-float(item.get("score", 0.0) or 0.0), str(item.get("template_path", ""))))
        rank_json = json.dumps(
            {
                "selection_candidates": candidate_rank_rows,
                "sample_scores": rank_rows[:25],
                "selected_template_path": selected_template_path,
                "selected_template_policy": selected_template_policy,
            },
            ensure_ascii=False,
            separators=(",", ":"),
        )
        return _predict_lr_marker_on_su_giu_rows(
            su_giu_rows=su_giu_rows,
            templates=selected_templates,
            vendor_name=vendor_name,
            min_match_score=min_match_score,
            full_crop_fallback_threshold=full_crop_fallback_threshold,
            expanded_search_threshold=expanded_search_threshold,
            expanded_search_steps=expanded_search_steps,
            min_sugiu_confidence=min_sugiu_confidence,
            device=device,
            force_single_template=False,
            fixed_template_path=selected_template_path,
            fixed_template_selection_score=selected_template_score,
            derived_template_dir=derived_template_dir,
            blank_template_max_value=blank_template_max_value,
            folder_template_seed_image_path=sample_path,
            folder_template_seed_darkness_pct=sample_darkness,
            folder_template_seed_db_template_path=str(seed_candidate.get("template_path", "") or "").strip(),
            folder_template_rank_json=rank_json,
            folder_template_policy=selected_template_policy,
            single_template_policy=single_template_policy,
            target_image_width=int(target_image_width or 0),
            target_image_height=int(target_image_height or 0),
        )

    out_rows: List[Dict[str, object]] = []
    for row_sg in su_giu_rows:
        image_path_s = str(row_sg.get("image_path", "") or "").strip()
        if not image_path_s:
            continue
        image_path = Path(image_path_s).expanduser().resolve()
        if not image_path.is_file():
            continue
        try:
            with Image.open(image_path) as img:
                rgb = img.convert("RGB")
                width, height = rgb.size
                image_is_canonical = (
                    int(target_image_width or 0) <= 0
                    or int(target_image_height or 0) <= 0
                    or (int(width) == int(target_image_width) and int(height) == int(target_image_height))
                )
                crop_top_ex = _safe_int_value(str(row_sg.get("crop_top", "0") or "0"), 0)
                crop_left_ex = _safe_int_value(str(row_sg.get("crop_left", "0") or "0"), 0)
                crop_bottom_ex = _safe_int_value(str(row_sg.get("crop_bottom", "0") or "0"), 0)
                crop_right_ex = _safe_int_value(str(row_sg.get("crop_right", "0") or "0"), 0)
                top = max(0, min(crop_top_ex, height - 1))
                left = max(0, min(crop_left_ex, width - 1))
                bottom = max(0, min(crop_bottom_ex - 1, height - 1))
                right = max(0, min(crop_right_ex - 1, width - 1))
                if bottom <= top or right <= left:
                    continue
                crop_width = right - left + 1
                crop_height = bottom - top + 1
                sugiu_model_pred = str(row_sg.get("pred_label", "") or "").strip().lower()
                sugiu_pred = sugiu_model_pred
                if sugiu_pred not in {"su", "giu"}:
                    continue
                roi_top, roi_left, roi_bottom, roi_right = _lr_marker_roi_from_sugiu(crop_width, crop_height, sugiu_pred)
                rect = (top, left, bottom, right)
                gray_uint8 = np.asarray(rgb.convert("L"), dtype=np.uint8)
                gray_full = torch.from_numpy(gray_uint8.astype(np.float32) / 255.0).to(
                    device=device,
                    dtype=torch.float32,
                )
                initial_search_rect_abs = (top + roi_top, left + roi_left, top + roi_bottom, left + roi_right)
                initial_loc = _lr_marker_locate_match_in_abs_rect(
                    gray_full=gray_full,
                    search_rect_abs=initial_search_rect_abs,
                    echo_rect_abs=rect,
                    templates=templates,
                    search_strategy="sugiu_roi",
                    search_scope="predicted_half",
                    search_margin_px=0,
                )
                best_loc = initial_loc
                dual_sugiu_search_used = False
                opposite_sugiu_match_score = ""
                try:
                    sugiu_conf_value = float(row_sg.get("confidence", 0.0) or 0.0)
                except (TypeError, ValueError):
                    sugiu_conf_value = 0.0
                if sugiu_conf_value < float(LR_MARKER_DUAL_SUGIU_SEARCH_CONFIDENCE):
                    opposite_sugiu = "giu" if sugiu_pred == "su" else "su"
                    opp_roi_top, opp_roi_left, opp_roi_bottom, opp_roi_right = _lr_marker_roi_from_sugiu(
                        crop_width,
                        crop_height,
                        opposite_sugiu,
                    )
                    opposite_search_rect_abs = (
                        top + opp_roi_top,
                        left + opp_roi_left,
                        top + opp_roi_bottom,
                        left + opp_roi_right,
                    )
                    opposite_loc = _lr_marker_locate_match_in_abs_rect(
                        gray_full=gray_full,
                        search_rect_abs=opposite_search_rect_abs,
                        echo_rect_abs=rect,
                        templates=templates,
                        search_strategy="opposite_sugiu_roi_low_conf",
                        search_scope="opposite_predicted_half",
                        search_margin_px=0,
                    )
                    opposite_sugiu_match_score = "" if opposite_loc is None else float(opposite_loc.score)
                    if _lr_marker_score_of(opposite_loc) > _lr_marker_score_of(best_loc):
                        best_loc = opposite_loc
                        sugiu_pred = opposite_sugiu
                    dual_sugiu_search_used = True
                full_crop_loc = None
                if _lr_marker_score_of(best_loc) < float(full_crop_fallback_threshold):
                    full_crop_loc = _lr_marker_locate_match_in_abs_rect(
                        gray_full=gray_full,
                        search_rect_abs=rect,
                        echo_rect_abs=rect,
                        templates=templates,
                        search_strategy="full_crop_low_score_fallback",
                        search_scope="full_echo_crop",
                        search_margin_px=0,
                    )
                    if _lr_marker_score_of(full_crop_loc) > _lr_marker_score_of(best_loc):
                        best_loc = full_crop_loc
                if _lr_marker_score_of(best_loc) < float(expanded_search_threshold):
                    for step in expanded_search_steps:
                        margin_px = max(1, int(round(max(crop_width, crop_height) * float(step))))
                        expanded_rect = _lr_marker_expand_rect(
                            initial_search_rect_abs,
                            width=width,
                            height=height,
                            margin_px=margin_px,
                        )
                        if expanded_rect is None:
                            continue
                        expanded_loc = _lr_marker_locate_match_in_abs_rect(
                            gray_full=gray_full,
                            search_rect_abs=expanded_rect,
                            echo_rect_abs=rect,
                            templates=templates,
                            search_strategy="expanded_rect_fallback",
                            search_scope="expanded_predicted_half",
                            search_margin_px=margin_px,
                        )
                        if _lr_marker_score_of(expanded_loc) > _lr_marker_score_of(best_loc):
                            best_loc = expanded_loc
                        if _lr_marker_score_of(best_loc) >= float(expanded_search_threshold):
                            break
                if best_loc is None:
                    continue

                marker_cx = (best_loc.marker_left_crop + best_loc.marker_right_crop) / 2.0
                marker_cy = (best_loc.marker_top_crop + best_loc.marker_bottom_crop) / 2.0
                detected_side = "left" if marker_cx < (crop_width / 2.0) else "right"
                lr_label, lr_label_it, lr_binary = _lr_marker_label_from_side(detected_side)
                reasons: List[str] = []
                patch_stats = _lr_marker_patch_stats_from_gray(
                    gray_uint8,
                    (
                        int(best_loc.marker_top_abs),
                        int(best_loc.marker_left_abs),
                        int(best_loc.marker_bottom_abs),
                        int(best_loc.marker_right_abs),
                    ),
                )
                if not image_is_canonical:
                    reasons.append("non_canonical_image_size")
                if float(row_sg.get("confidence", 0.0) or 0.0) < float(min_sugiu_confidence):
                    reasons.append("low_sugiu_conf")
                if float(best_loc.score) < float(min_match_score):
                    reasons.append("low_template_score")
                if bool(patch_stats.get("blank", False)):
                    reasons.append("blank_marker_match")
                if float(patch_stats.get("std", 0.0) or 0.0) < float(LR_MARKER_MIN_TEXTURE_STD):
                    reasons.append("low_texture_marker_match")
                row_lr: Dict[str, object] = {
                        "image_index": int(row_sg.get("image_index", 0) or 0),
                        "image_path": image_path.as_posix(),
                        "image_width": int(width),
                        "image_height": int(height),
                        "canonical_image_size": int(1 if image_is_canonical else 0),
                        "vendor": vendor_name,
                        "status": "",
                        "review_reason": "",
                        "lr_label": lr_label,
                        "lr_label_it": lr_label_it,
                        "lr_binary": int(lr_binary),
                        "detected_marker_side": detected_side,
                        "match_score": float(best_loc.score),
                        "match_patch_is_blank": int(1 if bool(patch_stats.get("blank", False)) else 0),
                        "match_patch_dark_pct": float(patch_stats.get("dark_pct", 0.0) or 0.0),
                        "match_patch_max_value": int(patch_stats.get("max_value", 0) or 0),
                        "match_patch_std": float(patch_stats.get("std", 0.0) or 0.0),
                        "initial_match_score": "" if initial_loc is None else float(initial_loc.score),
                        "full_crop_fallback_score": "" if full_crop_loc is None else float(full_crop_loc.score),
                        "search_strategy": best_loc.search_strategy,
                        "search_scope": best_loc.search_scope,
                        "search_margin_px": int(best_loc.search_margin_px),
                        "template_policy": (
                            folder_template_policy
                            or ("fixed_folder_template" if selected_template_path else "candidate_best_per_frame")
                        ),
                        "folder_fixed_template_path": selected_template_path,
                        "folder_fixed_template_selection_score": float(selected_template_score),
                        "folder_template_seed_image_path": folder_template_seed_image_path,
                        "folder_template_seed_darkness_pct": float(folder_template_seed_darkness_pct),
                        "folder_template_seed_db_template_path": folder_template_seed_db_template_path,
                        "folder_template_rank_json": folder_template_rank_json,
                        "template_path": best_loc.template.path.as_posix(),
                        "template_width": int(best_loc.template.width),
                        "template_height": int(best_loc.template.height),
                        "su_giu_model_pred": sugiu_model_pred,
                        "su_giu_pred": sugiu_pred,
                        "su_giu_conf": float(row_sg.get("confidence", 0.0) or 0.0),
                        "prob_su": float(row_sg.get("prob_su", 0.0) or 0.0),
                        "prob_giu": float(row_sg.get("prob_giu", 0.0) or 0.0),
                        "dual_sugiu_search_used": int(1 if dual_sugiu_search_used else 0),
                        "opposite_sugiu_match_score": opposite_sugiu_match_score,
                        "echo_rect_top_abs": top,
                        "echo_rect_left_abs": left,
                        "echo_rect_bottom_abs": bottom,
                        "echo_rect_right_abs": right,
                        "marker_top_crop": int(best_loc.marker_top_crop),
                        "marker_left_crop": int(best_loc.marker_left_crop),
                        "marker_bottom_crop": int(best_loc.marker_bottom_crop),
                        "marker_right_crop": int(best_loc.marker_right_crop),
                        "marker_top_abs": int(best_loc.marker_top_abs),
                        "marker_left_abs": int(best_loc.marker_left_abs),
                        "marker_bottom_abs": int(best_loc.marker_bottom_abs),
                        "marker_right_abs": int(best_loc.marker_right_abs),
                        "marker_cx_crop_norm": float(marker_cx / max(1, crop_width)),
                        "marker_cy_crop_norm": float(marker_cy / max(1, crop_height)),
                    }
                quadrant_fields = _lr_marker_quadrant_validation_fields(row_lr)
                row_lr.update(quadrant_fields)
                if str(quadrant_fields.get("quadrant_status", "") or "") == "invalid":
                    reasons.append("quadrant_logic_violation")
                    reasons.extend(
                        part
                        for part in str(quadrant_fields.get("quadrant_reason", "") or "").split(";")
                        if part
                    )
                row_lr["status"] = "ok" if not reasons else "review"
                row_lr["review_reason"] = ";".join(dict.fromkeys(reasons))
                out_rows.append(row_lr)
        except Exception:
            continue
    if selected_template_path and len(templates) == 1:
        out_rows = _stabilize_lr_marker_rows_with_spatial_consensus(
            out_rows,
            templates=templates,
            device=device,
            min_match_score=float(min_match_score),
            target_image_width=int(target_image_width or 0),
            target_image_height=int(target_image_height or 0),
        )
    out_rows.sort(key=lambda row: int(row.get("image_index", 0)))
    return out_rows


def _ensure_orientation_marker_bundle_import(bundle_dir: Path, bundle_zip: Path) -> Optional[Path]:
    """Make the portable bundle importable without changing pipeline outputs."""
    root = bundle_dir.expanduser().resolve()
    marker = root / "orientation_marker_detector" / "detector.py"
    if not marker.is_file():
        zip_path = bundle_zip.expanduser().resolve()
        if not zip_path.is_file():
            return None
        root.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(zip_path) as zf:
            zf.extractall(root.parent)
        extracted = root.parent / "41_orientation_marker_detector_bundle"
        if extracted.is_dir():
            root = extracted.resolve()
            marker = root / "orientation_marker_detector" / "detector.py"
    if not marker.is_file():
        return None
    if root.as_posix() not in sys.path:
        sys.path.insert(0, root.as_posix())
    return root


def _install_cv2_fallback_if_needed() -> None:
    """Provide the small cv2 subset used by the bundle when OpenCV is absent."""
    try:
        import cv2  # noqa: F401

        return
    except Exception:
        pass

    cv2_stub = types.ModuleType("cv2")
    cv2_stub.TM_CCOEFF_NORMED = 5
    cv2_stub.INTER_AREA = 3

    def _resize(image: np.ndarray, size: Tuple[int, int], interpolation: int = 3) -> np.ndarray:
        width, height = int(size[0]), int(size[1])
        pil = Image.fromarray(np.asarray(image).astype(np.uint8))
        return np.asarray(pil.resize((width, height), Image.Resampling.BOX), dtype=np.uint8)

    def _match_template(search: np.ndarray, template: np.ndarray, method: int) -> np.ndarray:
        if method != cv2_stub.TM_CCOEFF_NORMED:
            raise ValueError("cv2 fallback supports only TM_CCOEFF_NORMED")
        src = np.asarray(search, dtype=np.float32)
        tpl = np.asarray(template, dtype=np.float32)
        if src.ndim != 2 or tpl.ndim != 2:
            return np.zeros((0, 0), dtype=np.float32)
        th, tw = tpl.shape[:2]
        out_h = int(src.shape[0] - th + 1)
        out_w = int(src.shape[1] - tw + 1)
        if out_h <= 0 or out_w <= 0:
            return np.zeros((0, 0), dtype=np.float32)

        tpl_z = tpl - float(tpl.mean())
        tpl_norm = float(np.sqrt(np.sum(tpl_z * tpl_z)))
        if tpl_norm <= 1e-8:
            return np.zeros((out_h, out_w), dtype=np.float32)

        result = np.empty((out_h, out_w), dtype=np.float32)
        # Chunk rows to keep the sliding-window tensor bounded in memory.
        max_window_elems = 16_000_000
        chunk_rows = max(1, min(out_h, max_window_elems // max(1, out_w * th * tw)))
        for y0 in range(0, out_h, chunk_rows):
            y1 = min(out_h, y0 + chunk_rows)
            src_chunk = src[y0 : y1 + th - 1, :]
            windows = np.lib.stride_tricks.sliding_window_view(src_chunk, (th, tw))
            # windows shape: chunk_rows x out_w x th x tw
            sums = windows.sum(axis=(-1, -2), dtype=np.float32)
            sums_sq = np.square(windows, dtype=np.float32).sum(axis=(-1, -2), dtype=np.float32)
            area = float(th * tw)
            means = sums / area
            var_sum = np.maximum(sums_sq - area * means * means, 1e-8)
            numerator = np.tensordot(windows, tpl_z, axes=((2, 3), (0, 1)))
            denom = np.sqrt(var_sum) * tpl_norm
            result[y0:y1, :] = (numerator / np.maximum(denom, 1e-8)).astype(np.float32)
        return result

    def _min_max_loc(result: np.ndarray):
        arr = np.asarray(result)
        if arr.size == 0:
            return 0.0, 0.0, (0, 0), (0, 0)
        min_flat = int(np.argmin(arr))
        max_flat = int(np.argmax(arr))
        min_y, min_x = np.unravel_index(min_flat, arr.shape)
        max_y, max_x = np.unravel_index(max_flat, arr.shape)
        return float(arr[min_y, min_x]), float(arr[max_y, max_x]), (int(min_x), int(min_y)), (int(max_x), int(max_y))

    cv2_stub.resize = _resize
    cv2_stub.matchTemplate = _match_template
    cv2_stub.minMaxLoc = _min_max_loc
    sys.modules["cv2"] = cv2_stub


def _bundle_crop_rect_from_sugiu_row(row_sg: Dict[str, object]) -> Optional[Tuple[int, int, int, int]]:
    try:
        top = int(float(row_sg.get("crop_top", 0) or 0))
        left = int(float(row_sg.get("crop_left", 0) or 0))
        bottom = int(float(row_sg.get("crop_bottom", 0) or 0)) - 1
        right = int(float(row_sg.get("crop_right", 0) or 0)) - 1
    except (TypeError, ValueError):
        return None
    if bottom <= top or right <= left:
        return None
    return top, left, bottom, right


def _bundle_group_to_lr(group: str) -> Tuple[str, str, int]:
    group = str(group or "").strip().upper()
    if group in {"LR", "LRUD"}:
        return "lr_flipped", "specchiata_a_destra", 1
    if group in {"NF", "UD"}:
        return "not_lr_flipped", "normale", 0
    return "", "", -1


def _lr_marker_echo_rect_penalty_fields(
    *,
    marker_box_abs: Tuple[int, int, int, int],
    echo_rect_abs: Tuple[int, int, int, int],
    raw_score: float,
) -> Dict[str, object]:
    marker_top, marker_left, marker_bottom, marker_right = [int(v) for v in marker_box_abs]
    echo_top, echo_left, echo_bottom, echo_right = [int(v) for v in echo_rect_abs]
    marker_w = max(1, marker_right - marker_left + 1)
    marker_h = max(1, marker_bottom - marker_top + 1)
    marker_area = float(max(1, marker_w * marker_h))
    crop_w = max(1, echo_right - echo_left + 1)
    crop_h = max(1, echo_bottom - echo_top + 1)

    overlap_left = max(marker_left, echo_left)
    overlap_right = min(marker_right, echo_right)
    overlap_top = max(marker_top, echo_top)
    overlap_bottom = min(marker_bottom, echo_bottom)
    overlap_w = max(0, overlap_right - overlap_left + 1)
    overlap_h = max(0, overlap_bottom - overlap_top + 1)
    overlap_ratio = float((overlap_w * overlap_h) / marker_area)

    outside_top_px = max(0, echo_top - marker_top)
    outside_left_px = max(0, echo_left - marker_left)
    outside_bottom_px = max(0, marker_bottom - echo_bottom)
    outside_right_px = max(0, marker_right - echo_right)
    max_outside_px = max(outside_top_px, outside_left_px, outside_bottom_px, outside_right_px)

    fully_above_gap_px = max(0, echo_top - marker_bottom)
    fully_left_gap_px = max(0, echo_left - marker_right)
    fully_below_gap_px = max(0, marker_top - echo_bottom)
    fully_right_gap_px = max(0, marker_left - echo_right)
    fully_outside = overlap_ratio <= 0.0
    top_left_outside = fully_above_gap_px > 0 and fully_left_gap_px > 0

    # The marker may legitimately fall a few pixels outside the rect. The
    # penalty grows with distance; far top-left matches are usually vendor
    # logos, so only those get a much harsher suppression.
    near_margin_px = max(12.0, min(crop_w, crop_h) * 0.025, min(marker_w, marker_h) * 1.5)
    penalty_factor = 1.0
    reasons: List[str] = []
    gap_px = max(fully_above_gap_px, fully_left_gap_px, fully_below_gap_px, fully_right_gap_px)
    outside_distance_px = gap_px if gap_px > 0 else max_outside_px
    outside_severity = float(outside_distance_px) / max(1.0, float(near_margin_px))
    if max_outside_px > 0:
        if outside_severity <= 1.0:
            penalty_factor = 1.0 - 0.08 * outside_severity
        elif outside_severity <= 2.0:
            penalty_factor = 0.92 - 0.17 * (outside_severity - 1.0)
        elif outside_severity <= 4.0:
            penalty_factor = 0.75 - 0.30 * ((outside_severity - 2.0) / 2.0)
        elif outside_severity <= 6.0:
            penalty_factor = 0.45 - 0.20 * ((outside_severity - 4.0) / 2.0)
        else:
            penalty_factor = max(0.08, 0.25 - 0.17 * min((outside_severity - 6.0) / 4.0, 1.0))

        if outside_severity <= 1.0:
            reasons.append("marker_near_outside_echo_rect" if fully_outside else "marker_slightly_outside_echo_rect")
        elif outside_severity <= 2.0:
            reasons.append("marker_near_outside_echo_rect")
        elif outside_severity <= 4.0:
            reasons.append("marker_partly_outside_echo_rect" if overlap_ratio > 0.0 else "marker_fully_outside_echo_rect")
        else:
            reasons.append("marker_mostly_outside_echo_rect" if overlap_ratio > 0.0 else "marker_fully_outside_echo_rect")
        if outside_severity > 6.0:
            reasons.append("marker_far_outside_echo_rect")

        if top_left_outside:
            reasons.append("top_left_vendor_logo_candidate")
            diagonal_severity = max(fully_above_gap_px, fully_left_gap_px) / max(1.0, float(near_margin_px))
            if diagonal_severity <= 1.0:
                top_left_factor = 0.95 - 0.15 * diagonal_severity
            elif diagonal_severity <= 3.0:
                top_left_factor = 0.80 - 0.35 * ((diagonal_severity - 1.0) / 2.0)
            else:
                top_left_factor = max(0.08, 0.45 - 0.37 * min((diagonal_severity - 3.0) / 3.0, 1.0))
            penalty_factor = min(penalty_factor, top_left_factor)

    adjusted_score = max(0.0, min(1.0, float(raw_score) * float(penalty_factor)))
    return {
        "match_score_raw": float(raw_score),
        "match_score_penalty_factor": float(penalty_factor),
        "match_score_penalty_reason": ";".join(dict.fromkeys(reasons)),
        "match_score_adjusted": float(adjusted_score),
        "marker_echo_overlap_ratio": float(overlap_ratio),
        "marker_outside_echo_rect": int(1 if max_outside_px > 0 else 0),
        "marker_outside_echo_top_px": int(outside_top_px),
        "marker_outside_echo_left_px": int(outside_left_px),
        "marker_outside_echo_bottom_px": int(outside_bottom_px),
        "marker_outside_echo_right_px": int(outside_right_px),
        "marker_outside_echo_max_px": int(max_outside_px),
        "marker_outside_echo_gap_px": int(gap_px),
        "marker_outside_echo_severity": float(outside_severity),
        "marker_top_left_outside_echo_rect": int(1 if top_left_outside else 0),
        "marker_echo_near_margin_px": float(near_margin_px),
    }


def _bundle_predict_lr_marker_on_su_giu_rows(
    *,
    su_giu_rows: Sequence[Dict[str, object]],
    vendor_name: str,
    library_root: Path,
    bundle_dir: Path,
    bundle_zip: Path,
    min_match_score: float,
    vertical_delta: float,
    full_crop_fallback_threshold: float,
    expanded_search_threshold: float,
    expanded_search_steps: Sequence[float],
    match_max_side: int,
    selection_images: int,
    target_image_width: int = 0,
    target_image_height: int = 0,
    expected_groups_by_image_path: Optional[Dict[str, str]] = None,
    expected_groups_by_image_index: Optional[Dict[str, str]] = None,
    template_policy: str = "bundle_historical_best",
    scales: Sequence[float] = (1.0,),
    pinned_templates: Optional[Sequence[str]] = None,
) -> Tuple[List[Dict[str, object]], Dict[str, object]]:
    bundle_root = _ensure_orientation_marker_bundle_import(bundle_dir=bundle_dir, bundle_zip=bundle_zip)
    if bundle_root is None:
        return [], {"available": False, "error": "bundle_not_found"}
    _install_cv2_fallback_if_needed()
    try:
        from orientation_marker_detector.detector import (  # type: ignore
            DetectionParams,
            ImageInput,
            analyze_images,
        )
    except Exception as exc:  # pragma: no cover - optional bundle feature
        return [], {"available": False, "error": f"bundle_import_failed:{exc}"}

    params = DetectionParams(
        min_match_score=float(min_match_score),
        vertical_delta=float(vertical_delta),
        fallback_threshold=float(full_crop_fallback_threshold),
        expanded_threshold=float(expanded_search_threshold),
        expanded_steps=tuple(float(x) for x in expanded_search_steps),
        match_max_side=int(match_max_side),
    )
    bundle_library_root = bundle_root / "orientation_marker_detector" / "templates"
    requested_library_root = library_root.expanduser().resolve()
    vendor_dir_requested = next(
        (
            p for p in [requested_library_root / str(vendor_name), requested_library_root / str(vendor_name).capitalize()]
            if p.is_dir()
        ),
        None,
    )
    effective_library_root = requested_library_root if vendor_dir_requested is not None else bundle_library_root
    images: List[object] = []
    source_rows: List[Dict[str, object]] = []
    for row_sg in su_giu_rows:
        image_path_s = str(row_sg.get("image_path", "") or "").strip()
        if not image_path_s:
            continue
        image_path = Path(image_path_s).expanduser().resolve()
        if not image_path.is_file():
            continue
        crop_rect = _bundle_crop_rect_from_sugiu_row(dict(row_sg))
        if crop_rect is None:
            continue
        sugiu_model_pred = str(row_sg.get("pred_label", "") or "").strip().lower()
        sugiu_pred = sugiu_model_pred if sugiu_model_pred in {"su", "giu"} else ""
        try:
            conf = float(row_sg.get("confidence", 0.0) or 0.0)
        except (TypeError, ValueError):
            conf = 0.0
        image_index_s = str(row_sg.get("image_index", "") or "").strip()
        expected_group = ""
        if expected_groups_by_image_path:
            expected_group = str(expected_groups_by_image_path.get(image_path.as_posix(), "") or "").strip().upper()
        if not expected_group and expected_groups_by_image_index and image_index_s:
            expected_group = str(expected_groups_by_image_index.get(image_index_s, "") or "").strip().upper()
        if expected_group not in {"NF", "LR", "UD", "LRUD"}:
            expected_group = ""
        images.append(
            ImageInput(
                image_path=image_path,
                image_id=image_index_s or image_path.name,
                crop_rect=crop_rect,
                crop_source=str(row_sg.get("crop_source", "")),
                sugiu_pred=sugiu_pred,
                sugiu_conf=float(conf),
                expected_group=expected_group,
            )
        )
        source_rows.append(dict(row_sg))

    if not images:
        return [], {"available": True, "error": "no_valid_bundle_inputs", "bundle_root": bundle_root.as_posix()}

    try:
        analysis = analyze_images(
            images,
            vendor=str(vendor_name),
            library_root=effective_library_root,
            params=params,
            selection_images=max(1, int(selection_images)),
            scales=tuple(float(s) for s in scales) or (1.0,),
            pinned_templates=list(pinned_templates or []),
        )
    except Exception as exc:
        return [], {"available": True, "error": f"bundle_analysis_failed:{exc}", "bundle_root": bundle_root.as_posix()}

    out_rows: List[Dict[str, object]] = []
    for result, row_sg in zip(analysis.rows, source_rows):
        result_d = result.to_dict()
        marker_box = result_d.get("marker_box_abs")
        crop_rect = result_d.get("crop_rect_abs")
        image_path_s = str(result_d.get("image_path", row_sg.get("image_path", "")) or "")
        if not marker_box or not crop_rect or not image_path_s:
            row_out: Dict[str, object] = {
                "image_index": int(row_sg.get("image_index", 0) or 0),
                "image_path": image_path_s,
                "vendor": vendor_name,
                "status": "review",
                "review_reason": str(result_d.get("review_reason", "") or "bundle_no_match"),
                "lr_marker_method": "bundle",
                "match_score": float(result_d.get("match_score", 0.0) or 0.0),
                "template_path": str(result_d.get("template_path", "") or ""),
                "template_width": int(result_d.get("template_width", 0) or 0),
                "template_height": int(result_d.get("template_height", 0) or 0),
                "su_giu_model_pred": str(row_sg.get("pred_label", "") or ""),
                "su_giu_pred": str(result_d.get("vertical_final", "") or row_sg.get("pred_label", "") or ""),
                "su_giu_conf": float(row_sg.get("confidence", 0.0) or 0.0),
                "prob_su": float(row_sg.get("prob_su", 0.0) or 0.0),
                "prob_giu": float(row_sg.get("prob_giu", 0.0) or 0.0),
            }
            out_rows.append(row_out)
            continue
        marker_top, marker_left, marker_bottom, marker_right = [int(v) for v in marker_box]
        echo_top, echo_left, echo_bottom, echo_right = [int(v) for v in crop_rect]
        crop_width = max(1, echo_right - echo_left + 1)
        crop_height = max(1, echo_bottom - echo_top + 1)
        marker_cx_abs = (marker_left + marker_right) / 2.0
        marker_cy_abs = (marker_top + marker_bottom) / 2.0
        detected_side = "left" if marker_cx_abs < ((echo_left + echo_right) / 2.0) else "right"
        vertical_final = str(result_d.get("vertical_final", "") or "").strip().lower()
        orientation_group = str(result_d.get("orientation_group", "") or "").strip().upper()
        lr_label, lr_label_it, lr_binary = _bundle_group_to_lr(orientation_group)
        if not lr_label:
            if detected_side == "right":
                lr_label, lr_label_it, lr_binary = "lr_flipped", "specchiata_a_destra", 1
            elif detected_side == "left":
                lr_label, lr_label_it, lr_binary = "not_lr_flipped", "normale", 0
            else:
                lr_label, lr_label_it, lr_binary = "", "", -1
        try:
            with Image.open(Path(image_path_s).expanduser().resolve()) as img:
                width, height = img.size
                gray_uint8 = np.asarray(img.convert("L"), dtype=np.uint8)
        except Exception:
            width = int(target_image_width or 0)
            height = int(target_image_height or 0)
            gray_uint8 = np.zeros((max(1, height), max(1, width)), dtype=np.uint8)
        image_is_canonical = (
            int(target_image_width or 0) <= 0
            or int(target_image_height or 0) <= 0
            or (int(width) == int(target_image_width) and int(height) == int(target_image_height))
        )
        patch_stats = _lr_marker_patch_stats_from_gray(
            gray_uint8,
            (int(marker_top), int(marker_left), int(marker_bottom), int(marker_right)),
        )
        raw_match_score = float(result_d.get("match_score", 0.0) or 0.0)
        penalty_fields = _lr_marker_echo_rect_penalty_fields(
            marker_box_abs=(int(marker_top), int(marker_left), int(marker_bottom), int(marker_right)),
            echo_rect_abs=(int(echo_top), int(echo_left), int(echo_bottom), int(echo_right)),
            raw_score=raw_match_score,
        )
        adjusted_match_score = float(penalty_fields.get("match_score_adjusted", raw_match_score) or 0.0)
        reasons = [
            part
            for part in str(result_d.get("review_reason", "") or "").split(";")
            if part
        ]
        penalty_reason = str(penalty_fields.get("match_score_penalty_reason", "") or "").strip()
        if penalty_reason:
            reasons.extend(part for part in penalty_reason.split(";") if part)
            if adjusted_match_score < raw_match_score:
                reasons.append("outside_echo_rect_score_penalty")
            if raw_match_score >= float(min_match_score) and adjusted_match_score < float(min_match_score):
                reasons.append("outside_echo_rect_penalty_below_threshold")
        if not image_is_canonical:
            reasons.append("non_canonical_image_size")
        if bool(patch_stats.get("blank", False)):
            reasons.append("blank_marker_match")
        if float(patch_stats.get("std", 0.0) or 0.0) < float(LR_MARKER_MIN_TEXTURE_STD):
            reasons.append("low_texture_marker_match")
        row_out = {
            "image_index": int(row_sg.get("image_index", 0) or 0),
            "image_path": image_path_s,
            "image_width": int(width),
            "image_height": int(height),
            "canonical_image_size": int(1 if image_is_canonical else 0),
            "vendor": vendor_name,
            "status": str(result_d.get("status", "") or "review"),
            "review_reason": "",
            "lr_label": lr_label,
            "lr_label_it": lr_label_it,
            "lr_binary": int(lr_binary),
            "detected_marker_side": detected_side,
            "match_score": adjusted_match_score,
            "match_patch_is_blank": int(1 if bool(patch_stats.get("blank", False)) else 0),
            "match_patch_dark_pct": float(patch_stats.get("dark_pct", 0.0) or 0.0),
            "match_patch_max_value": int(patch_stats.get("max_value", 0) or 0),
            "match_patch_std": float(patch_stats.get("std", 0.0) or 0.0),
            "initial_match_score": "",
            "top_half_score": "" if result_d.get("top_half_score") is None else float(result_d.get("top_half_score", 0.0) or 0.0),
            "bottom_half_score": "" if result_d.get("bottom_half_score") is None else float(result_d.get("bottom_half_score", 0.0) or 0.0),
            "full_crop_fallback_score": "" if result_d.get("full_crop_score") is None else float(result_d.get("full_crop_score", 0.0) or 0.0),
            "expanded_score": "" if result_d.get("expanded_score") is None else float(result_d.get("expanded_score", 0.0) or 0.0),
            "search_strategy": str(result_d.get("search_scope", "") or ""),
            "search_scope": str(result_d.get("search_scope", "") or ""),
            "search_margin_px": 0,
            "template_policy": str(template_policy or "bundle_historical_best"),
            "folder_fixed_template_path": str(analysis.template.path),
            "folder_fixed_template_selection_score": raw_match_score,
            "folder_template_seed_image_path": "",
            "folder_template_seed_darkness_pct": -1.0,
            "folder_template_seed_db_template_path": "",
            "folder_template_rank_json": json.dumps(analysis.historical_rank[:25], ensure_ascii=False, separators=(",", ":")),
            "bundle_template_name": str(result_d.get("template_name", "") or ""),
            "template_path": str(result_d.get("template_path", "") or ""),
            "template_width": int(result_d.get("template_width", 0) or 0),
            "template_height": int(result_d.get("template_height", 0) or 0),
            "su_giu_model_pred": str(row_sg.get("pred_label", "") or ""),
            "su_giu_pred": vertical_final or str(row_sg.get("pred_label", "") or ""),
            "su_giu_conf": float(row_sg.get("confidence", 0.0) or 0.0),
            "prob_su": float(row_sg.get("prob_su", 0.0) or 0.0),
            "prob_giu": float(row_sg.get("prob_giu", 0.0) or 0.0),
            "dual_sugiu_search_used": int(1 if str(result_d.get("vertical_correction", "") or "") == "corrected" else 0),
            "opposite_sugiu_match_score": "",
            "orientation_group": orientation_group,
            "bundle_vertical_center": str(result_d.get("vertical_center", "") or ""),
            "bundle_vertical_source": str(result_d.get("vertical_source", "") or ""),
            "bundle_vertical_correction": str(result_d.get("vertical_correction", "") or ""),
            "lr_marker_method": "bundle",
            "echo_rect_top_abs": int(echo_top),
            "echo_rect_left_abs": int(echo_left),
            "echo_rect_bottom_abs": int(echo_bottom),
            "echo_rect_right_abs": int(echo_right),
            "marker_top_crop": int(marker_top - echo_top),
            "marker_left_crop": int(marker_left - echo_left),
            "marker_bottom_crop": int(marker_bottom - echo_top),
            "marker_right_crop": int(marker_right - echo_left),
            "marker_top_abs": int(marker_top),
            "marker_left_abs": int(marker_left),
            "marker_bottom_abs": int(marker_bottom),
            "marker_right_abs": int(marker_right),
            "marker_cx_crop_norm": float((marker_cx_abs - echo_left) / max(1, crop_width)),
            "marker_cy_crop_norm": float((marker_cy_abs - echo_top) / max(1, crop_height)),
            **penalty_fields,
        }
        if adjusted_match_score < float(min_match_score) and raw_match_score >= float(min_match_score):
            row_out["status"] = "review"
        quadrant_fields = _lr_marker_quadrant_validation_fields(row_out)
        row_out.update(quadrant_fields)
        if str(quadrant_fields.get("quadrant_status", "") or "") == "invalid":
            reasons.append("quadrant_logic_violation")
            reasons.extend(
                part
                for part in str(quadrant_fields.get("quadrant_reason", "") or "").split(";")
                if part
            )
            row_out["status"] = "review"
        row_out["review_reason"] = ";".join(dict.fromkeys(reasons))
        out_rows.append(row_out)

    out_rows.sort(key=lambda row: int(row.get("image_index", 0) or 0))
    return out_rows, {
        "available": True,
        "bundle_root": bundle_root.as_posix(),
        "requested_library_root": requested_library_root.as_posix(),
        "effective_library_root": effective_library_root.as_posix(),
        "template_name": analysis.template.name,
        "template_path": str(analysis.template.path),
        "template_policy": str(template_policy or "bundle_historical_best"),
        "historical_rank": analysis.historical_rank[:25],
        "rows": len(out_rows),
    }


def _infer_lr_marker_orientation_group(image_path_text: str) -> str:
    path = Path(str(image_path_text or "").strip())
    normalized = re.sub(r"[^A-Z0-9]+", "_", " ".join([path.name, path.parent.name, str(path)]).upper())
    padded = f"_{normalized.strip('_')}_"
    if re.search(r"_FLIPLR_UD_", padded) or re.search(r"_FLIPUD_LR_", padded):
        return "LRUD"
    if re.search(r"_FLIPLRUD_", padded) or re.search(r"_FLIPUDLR_", padded):
        return "LRUD"
    if re.search(r"(^|_)LRUD(_|$)", normalized):
        return "LRUD"
    if re.search(r"_NOFLIP_", padded):
        return "NF"
    if re.search(r"_FLIPLR_", padded):
        return "LR"
    if re.search(r"_FLIPUD_", padded):
        return "UD"
    if re.search(r"(^|_)UD(_|$)", normalized):
        return "UD"
    if re.search(r"(^|_)LR(_|$)", normalized):
        return "LR"
    if re.search(r"(^|_)NF(_|$)", normalized):
        return "NF"
    return ""


def _infer_lr_marker_orientation_group_from_row(row: Dict[str, object]) -> str:
    exported_group = str(row.get("orientation_group", "") or "").strip().upper()
    quadrant_fields = _lr_marker_quadrant_validation_fields(row)
    if str(row.get("lr_marker_method", "") or "") == "bundle":
        center_group = str(quadrant_fields.get("quadrant_center_group", "") or "").strip().upper()
        if center_group in {"NF", "LR", "UD", "LRUD"}:
            return center_group
        if exported_group in {"NF", "LR", "UD", "LRUD"}:
            return exported_group
    if str(quadrant_fields.get("quadrant_status", "") or "") == "ok":
        quadrant_group = str(quadrant_fields.get("quadrant_group", "") or "").strip().upper()
        if quadrant_group in {"NF", "LR", "UD", "LRUD"}:
            return quadrant_group
    side = str(row.get("detected_marker_side", "") or "").strip().lower()
    sugiu = str(row.get("su_giu_pred", "") or "").strip().lower()
    quadrant_group = _lr_marker_orientation_group_from_quadrant(side, sugiu)
    if quadrant_group and _lr_marker_row_quadrant_valid_or_unknown(dict(row)):
        return quadrant_group
    group = _infer_lr_marker_orientation_group(str(row.get("image_path", "") or ""))
    if group:
        return group
    try:
        cy = float(row.get("marker_cy_crop_norm", 0.0) or 0.0)
    except (TypeError, ValueError):
        return ""
    if side == "left" and cy < 0.5:
        return "NF"
    if side == "right" and cy < 0.5:
        return "LR"
    if side == "left" and cy >= 0.5:
        return "UD"
    if side == "right" and cy >= 0.5:
        return "LRUD"
    return ""


def _lr_marker_group_from_detected_position(row: Dict[str, object]) -> str:
    side = str(row.get("detected_marker_side", "") or "").strip().lower()
    try:
        cy = float(row.get("marker_cy_crop_norm", 0.0) or 0.0)
    except (TypeError, ValueError):
        return ""
    if side == "left" and cy < 0.5:
        return "NF"
    if side == "right" and cy < 0.5:
        return "LR"
    if side == "left" and cy >= 0.5:
        return "UD"
    if side == "right" and cy >= 0.5:
        return "LRUD"
    return ""


def _stabilize_lr_marker_rows_with_spatial_consensus(
    rows: Sequence[Dict[str, object]],
    *,
    templates: Sequence[object],
    device: torch.device,
    min_match_score: float,
    target_image_width: int = 0,
    target_image_height: int = 0,
) -> List[Dict[str, object]]:
    if len(rows) < 8 or len(templates) != 1:
        return [dict(row) for row in rows]
    if any(_infer_lr_marker_orientation_group(str(row.get("image_path", "") or "")) for row in rows):
        return [dict(row) for row in rows]
    assert _lr_marker_locate_match_in_abs_rect is not None
    assert _lr_marker_label_from_side is not None

    out = [dict(row) for row in rows]
    consensus_rows = [
        row for row in out
        if _lr_marker_row_matches_target_size(
            row,
            target_image_width=int(target_image_width or 0),
            target_image_height=int(target_image_height or 0),
        )
        and not _lr_marker_truthy_flag(row.get("match_patch_is_blank", False))
        and _lr_marker_row_quadrant_valid_or_unknown(dict(row))
    ]
    if len(consensus_rows) < 8:
        return out
    group_counts: Counter[str] = Counter(_lr_marker_group_from_detected_position(row) for row in consensus_rows)
    group_counts.pop("", None)
    if not group_counts:
        return out
    majority_group, majority_count = group_counts.most_common(1)[0]
    if majority_count < max(5, int(round(len(consensus_rows) * 0.72))):
        return out

    majority_rows = [row for row in consensus_rows if _lr_marker_group_from_detected_position(row) == majority_group]
    def _median(values: Sequence[float]) -> float:
        return float(statistics.median([float(v) for v in values]))

    median_cx_norm = _median([float(row.get("marker_cx_crop_norm", 0.0) or 0.0) for row in majority_rows])
    median_cy_norm = _median([float(row.get("marker_cy_crop_norm", 0.0) or 0.0) for row in majority_rows])
    median_w = _median([
        float(int(row.get("marker_right_abs", 0) or 0) - int(row.get("marker_left_abs", 0) or 0) + 1)
        for row in majority_rows
    ])
    median_h = _median([
        float(int(row.get("marker_bottom_abs", 0) or 0) - int(row.get("marker_top_abs", 0) or 0) + 1)
        for row in majority_rows
    ])
    tolerance = max(36.0, max(median_w, median_h) * 4.0)
    search_margin = max(28, int(round(max(median_w, median_h) * 4.5)))

    for row in out:
        if not _lr_marker_row_matches_target_size(
            row,
            target_image_width=int(target_image_width or 0),
            target_image_height=int(target_image_height or 0),
        ):
            continue
        if _lr_marker_truthy_flag(row.get("match_patch_is_blank", False)):
            continue
        try:
            echo_top = int(row.get("echo_rect_top_abs", 0) or 0)
            echo_left = int(row.get("echo_rect_left_abs", 0) or 0)
            echo_bottom = int(row.get("echo_rect_bottom_abs", 0) or 0)
            echo_right = int(row.get("echo_rect_right_abs", 0) or 0)
            marker_cx = (int(row.get("marker_left_abs", 0) or 0) + int(row.get("marker_right_abs", 0) or 0)) * 0.5
            marker_cy = (int(row.get("marker_top_abs", 0) or 0) + int(row.get("marker_bottom_abs", 0) or 0)) * 0.5
        except Exception:
            continue
        crop_w = max(1, echo_right - echo_left + 1)
        crop_h = max(1, echo_bottom - echo_top + 1)
        expected_cx = echo_left + median_cx_norm * crop_w
        expected_cy = echo_top + median_cy_norm * crop_h
        group = _lr_marker_group_from_detected_position(row)
        is_outlier = group != majority_group or abs(marker_cx - expected_cx) > tolerance or abs(marker_cy - expected_cy) > tolerance
        if not is_outlier:
            continue

        image_path_s = str(row.get("image_path", "") or "").strip()
        if not image_path_s:
            continue
        try:
            image_path = Path(image_path_s).expanduser().resolve()
            with Image.open(image_path) as img:
                gray_uint8 = np.asarray(img.convert("L"), dtype=np.uint8)
                gray_full = torch.from_numpy(gray_uint8.astype(np.float32) / 255.0).to(
                    device=device,
                    dtype=torch.float32,
                )
                width, height = img.size
            search_rect_abs = (
                max(0, int(round(expected_cy)) - search_margin),
                max(0, int(round(expected_cx)) - search_margin),
                min(height - 1, int(round(expected_cy)) + search_margin),
                min(width - 1, int(round(expected_cx)) + search_margin),
            )
            echo_rect_abs = (echo_top, echo_left, echo_bottom, echo_right)
            loc = _lr_marker_locate_match_in_abs_rect(
                gray_full=gray_full,
                search_rect_abs=search_rect_abs,
                echo_rect_abs=echo_rect_abs,
                templates=templates,
                search_strategy="spatial_consensus_refine",
                search_scope="majority_marker_cluster",
                search_margin_px=int(search_margin),
            )
            consensus_min_score = max(0.25, float(min_match_score) * 0.65)
            forced_consensus = False
            if loc is not None and float(loc.score) >= consensus_min_score:
                marker_left_crop = int(loc.marker_left_crop)
                marker_top_crop = int(loc.marker_top_crop)
                marker_right_crop = int(loc.marker_right_crop)
                marker_bottom_crop = int(loc.marker_bottom_crop)
                marker_left_abs = int(loc.marker_left_abs)
                marker_top_abs = int(loc.marker_top_abs)
                marker_right_abs = int(loc.marker_right_abs)
                marker_bottom_abs = int(loc.marker_bottom_abs)
                score_new = float(loc.score)
                strategy_new = loc.search_strategy
                scope_new = loc.search_scope
                margin_new = int(loc.search_margin_px)
            elif majority_count >= max(5, int(round(len(consensus_rows) * 0.90))):
                forced_consensus = True
                box_w = max(1, int(round(median_w)))
                box_h = max(1, int(round(median_h)))
                marker_left_abs = int(round(expected_cx - box_w * 0.5))
                marker_top_abs = int(round(expected_cy - box_h * 0.5))
                marker_right_abs = marker_left_abs + box_w - 1
                marker_bottom_abs = marker_top_abs + box_h - 1
                marker_left_crop = marker_left_abs - echo_left
                marker_top_crop = marker_top_abs - echo_top
                marker_right_crop = marker_right_abs - echo_left
                marker_bottom_crop = marker_bottom_abs - echo_top
                score_new = min(float(row.get("match_score", 0.0) or 0.0), float(min_match_score) - 1e-6)
                strategy_new = "spatial_consensus_forced"
                scope_new = "majority_marker_cluster"
                margin_new = int(search_margin)
            else:
                continue
            marker_cx_new = (marker_left_crop + marker_right_crop) / 2.0
            marker_cy_new = (marker_top_crop + marker_bottom_crop) / 2.0
            detected_side = "left" if marker_cx_new < (crop_w / 2.0) else "right"
            lr_label, lr_label_it, lr_binary = _lr_marker_label_from_side(detected_side)
            patch_stats = _lr_marker_patch_stats_from_gray(
                gray_uint8,
                (int(marker_top_abs), int(marker_left_abs), int(marker_bottom_abs), int(marker_right_abs)),
            )
            row["detected_marker_side"] = detected_side
            row["lr_label"] = lr_label
            row["lr_label_it"] = lr_label_it
            row["lr_binary"] = int(lr_binary)
            row["match_score"] = float(score_new)
            row["match_patch_is_blank"] = int(1 if bool(patch_stats.get("blank", False)) else 0)
            row["match_patch_dark_pct"] = float(patch_stats.get("dark_pct", 0.0) or 0.0)
            row["match_patch_max_value"] = int(patch_stats.get("max_value", 0) or 0)
            row["match_patch_std"] = float(patch_stats.get("std", 0.0) or 0.0)
            row["search_strategy"] = strategy_new
            row["search_scope"] = scope_new
            row["search_margin_px"] = int(margin_new)
            row["marker_top_crop"] = int(marker_top_crop)
            row["marker_left_crop"] = int(marker_left_crop)
            row["marker_bottom_crop"] = int(marker_bottom_crop)
            row["marker_right_crop"] = int(marker_right_crop)
            row["marker_top_abs"] = int(marker_top_abs)
            row["marker_left_abs"] = int(marker_left_abs)
            row["marker_bottom_abs"] = int(marker_bottom_abs)
            row["marker_right_abs"] = int(marker_right_abs)
            row["marker_cx_crop_norm"] = float(marker_cx_new / max(1, crop_w))
            row["marker_cy_crop_norm"] = float(marker_cy_new / max(1, crop_h))
            quadrant_fields = _lr_marker_quadrant_validation_fields(row)
            row.update(quadrant_fields)
            drop_reasons = {
                "low_template_score",
                "quadrant_logic_violation",
                "marker_crosses_vertical_median",
                "marker_crosses_horizontal_median",
                "missing_marker_side",
                "missing_sugiu_pred",
                "invalid_quadrant_coordinates",
                "invalid_quadrant_rect",
            }
            reasons = [
                part for part in str(row.get("review_reason", "") or "").split(";")
                if part and part not in drop_reasons
            ]
            if bool(patch_stats.get("blank", False)):
                reasons.append("blank_marker_match")
            if str(quadrant_fields.get("quadrant_status", "") or "") == "invalid":
                reasons.append("quadrant_logic_violation")
                reasons.extend(
                    part
                    for part in str(quadrant_fields.get("quadrant_reason", "") or "").split(";")
                    if part
                )
            reasons.append("spatial_consensus_forced" if forced_consensus else "spatial_consensus_refined")
            row["review_reason"] = ";".join(dict.fromkeys(reasons))
            blocking_reasons = {
                "low_sugiu_conf",
                "blank_marker_match",
                "non_canonical_image_size",
                "quadrant_logic_violation",
            }
            if (
                not forced_consensus
                and float(score_new) >= float(min_match_score)
                and not any(part in blocking_reasons for part in reasons)
            ):
                row["status"] = "ok"
            elif forced_consensus:
                row["status"] = "review"
        except Exception:
            continue
    return out


def _build_line16_rect_orientation_from_lr_marker_rows(
    rows: Sequence[Dict[str, object]],
    *,
    min_match_score: float = LR_MARKER_RELIABLE_MATCH_SCORE,
    target_image_width: int = 0,
    target_image_height: int = 0,
) -> Tuple[str, str, int, str]:
    """Build line #16 from reliable marker envelopes while thresholds are pending."""
    order = ("NF", "LR", "UD", "LRUD")
    groups: Dict[str, Dict[str, object]] = {
        key: {"top": None, "left": None, "bottom": None, "right": None, "boxes": 0, "widths": [], "heights": []}
        for key in order
    }
    bundle_method = any(str(row.get("lr_marker_method", "") or "") == "bundle" for row in rows)
    boxes_count = 0
    for row in rows:
        try:
            score = float(row.get("match_score", 0.0) or 0.0)
        except (TypeError, ValueError):
            score = 0.0
        if score < float(min_match_score):
            continue
        if str(row.get("search_strategy", "") or "") == "spatial_consensus_forced":
            continue
        if not _lr_marker_row_matches_target_size(
            dict(row),
            target_image_width=int(target_image_width or 0),
            target_image_height=int(target_image_height or 0),
        ):
            continue
        if str(row.get("lr_marker_method", "") or "") == "bundle":
            q_fields = _lr_marker_quadrant_validation_fields(dict(row))
            center_group = str(q_fields.get("quadrant_center_group", "") or "").strip().upper()
            if center_group not in groups:
                continue
        elif not _lr_marker_row_quadrant_valid_or_unknown(dict(row)):
            continue
        review_reason = str(row.get("review_reason", "") or "")
        if _lr_marker_truthy_flag(row.get("match_patch_is_blank", False)) or "blank_marker_match" in review_reason.split(";"):
            continue
        if not _lr_marker_row_has_textured_match(dict(row)) or "low_texture_marker_match" in review_reason.split(";"):
            continue
        group_key = _infer_lr_marker_orientation_group_from_row(dict(row))
        if group_key not in groups:
            continue
        try:
            top = int(row.get("marker_top_abs", 0) or 0)
            left = int(row.get("marker_left_abs", 0) or 0)
            bottom = int(row.get("marker_bottom_abs", 0) or 0)
            right = int(row.get("marker_right_abs", 0) or 0)
        except (TypeError, ValueError):
            continue
        if bottom <= top or right <= left:
            continue
        env = groups[group_key]
        env["top"] = top if env["top"] is None else min(int(env["top"]), top)
        env["left"] = left if env["left"] is None else min(int(env["left"]), left)
        env["bottom"] = bottom if env["bottom"] is None else max(int(env["bottom"]), bottom)
        env["right"] = right if env["right"] is None else max(int(env["right"]), right)
        env["boxes"] = int(env["boxes"]) + 1
        widths = env.get("widths")
        heights = env.get("heights")
        if isinstance(widths, list):
            widths.append(int(right - left + 1))
        if isinstance(heights, list):
            heights.append(int(bottom - top + 1))
        boxes_count += 1

    unstable_groups: List[str] = []
    for key in order:
        env = groups[key]
        if int(env["boxes"]) <= 0:
            continue
        widths = [int(v) for v in env.get("widths", [])] if isinstance(env.get("widths"), list) else []
        heights = [int(v) for v in env.get("heights", [])] if isinstance(env.get("heights"), list) else []
        median_w = float(statistics.median(widths)) if widths else 0.0
        median_h = float(statistics.median(heights)) if heights else 0.0
        envelope_w = int(env["right"]) - int(env["left"]) + 1
        envelope_h = int(env["bottom"]) - int(env["top"]) + 1
        env["median_width"] = float(median_w)
        env["median_height"] = float(median_h)
        env["envelope_width"] = int(envelope_w)
        env["envelope_height"] = int(envelope_h)
        if (
            not bundle_method
            and (envelope_w > max(80.0, median_w * 6.0) or envelope_h > max(80.0, median_h * 6.0))
        ):
            unstable_groups.append(key)
    if unstable_groups:
        for key in unstable_groups:
            groups[key]["unstable"] = 1

    groups_json = json.dumps(
        {
            "min_match_score": float(min_match_score),
            "target_image_width": int(target_image_width or 0),
            "target_image_height": int(target_image_height or 0),
            "grouping": "bundle_quadrant_marker_official_rect_axes" if bundle_method else "quadrant_marker_sugiu",
            "unstable_groups": unstable_groups,
            "groups": groups,
        },
        ensure_ascii=False,
        separators=(",", ":"),
    )
    if any(int(groups[key]["boxes"]) <= 0 for key in order):
        return "", "lr_marker_envelope_missing_groups", boxes_count, groups_json
    if unstable_groups:
        return "", "lr_marker_envelope_unstable_groups", boxes_count, groups_json

    checks = []
    for key in order:
        env = groups[key]
        checks.append(
            f"{int(env['top'])}|{int(env['left'])}|{int(env['bottom'])}|{int(env['right'])}|"
            f"{LINE16_PENDING_MATCH_TAIL};"
        )
    source = (
        "bundle_marker_orientation_envelopes_pending_thresholds"
        if bundle_method
        else "lr_marker_quadrant_orientation_envelopes_pending_thresholds"
    )
    return "".join(checks), source, boxes_count, groups_json


def _predict_lt_on_rect_crops(
    model: torch.nn.Module,
    class_names: Sequence[str],
    image_paths: Sequence[Path],
    rect_boxes_abs: Sequence[Tuple[float, float, float, float]],
    fallback_rect_box_abs: Optional[Tuple[float, float, float, float]],
    image_size: int,
    batch_size: int,
    device: torch.device,
    rotate_deg_clockwise: int = 0,
) -> List[Dict[str, object]]:
    if not image_paths:
        return []
    if len(class_names) < 2:
        return []

    class_names_up = [str(name).strip().upper() for name in class_names]
    l_idx = class_names_up.index("L") if "L" in class_names_up else 0
    t_idx = class_names_up.index("T") if "T" in class_names_up else min(1, len(class_names_up) - 1)

    jobs: List[Tuple[int, Path, Tuple[float, float, float, float], str]] = []
    for idx, path in enumerate(image_paths):
        if idx < len(rect_boxes_abs):
            jobs.append((idx, path, rect_boxes_abs[idx], "per_image_rect"))
        elif fallback_rect_box_abs is not None:
            jobs.append((idx, path, fallback_rect_box_abs, "global_rect_fallback"))

    if not jobs:
        return []

    out_rows: List[Dict[str, object]] = []
    model.eval()
    with torch.inference_mode():
        for start in range(0, len(jobs), batch_size):
            batch_jobs = jobs[start : start + batch_size]
            tensors: List[torch.Tensor] = []
            metas: List[Tuple[int, Path, Tuple[int, int, int, int], str]] = []
            for image_idx, image_path, crop_box, crop_source in batch_jobs:
                try:
                    tensor, crop_tlbr = _load_rect_crop_classification_tensor(
                        path=image_path,
                        crop_box=crop_box,
                        image_size=image_size,
                        rotate_deg_clockwise=rotate_deg_clockwise,
                    )
                except Exception:
                    continue
                tensors.append(tensor)
                metas.append((image_idx, image_path, crop_tlbr, crop_source))
            if not tensors:
                continue

            images = torch.stack(tensors, dim=0).to(device, non_blocking=True)
            probs = torch.softmax(model(images), dim=1).detach().cpu().numpy()
            pred_idx = np.argmax(probs, axis=1).tolist()

            for i, (image_idx, image_path, crop_tlbr, crop_source) in enumerate(metas):
                pred_i = int(pred_idx[i])
                pred_label = (
                    str(class_names[pred_i]).strip().upper()
                    if 0 <= pred_i < len(class_names)
                    else str(pred_i)
                )
                p_l = float(probs[i][l_idx]) if 0 <= l_idx < probs.shape[1] else 0.0
                p_t = float(probs[i][t_idx]) if 0 <= t_idx < probs.shape[1] else 0.0
                conf = float(probs[i][pred_i]) if 0 <= pred_i < probs.shape[1] else 0.0
                top, left, bottom, right = crop_tlbr
                out_rows.append(
                    {
                        "image_index": int(image_idx),
                        "image_path": image_path.as_posix(),
                        "pred_label": pred_label,
                        "pred_idx": pred_i,
                        "confidence": conf,
                        "prob_l": p_l,
                        "prob_t": p_t,
                        "crop_top": int(top),
                        "crop_left": int(left),
                        "crop_bottom": int(bottom),
                        "crop_right": int(right),
                        "crop_source": crop_source,
                    }
                )

    out_rows.sort(key=lambda row: int(row.get("image_index", 0)))
    return out_rows


def _topk_from_probs(
    probs: np.ndarray,
    class_names: Sequence[str],
    k: int,
) -> List[Tuple[str, float]]:
    if probs.size == 0 or not class_names:
        return []
    limit = max(1, min(int(k), len(class_names)))
    sorted_idx = np.argsort(probs)[::-1][:limit]
    out: List[Tuple[str, float]] = []
    for idx in sorted_idx:
        class_idx = int(idx)
        if class_idx < 0 or class_idx >= len(class_names):
            continue
        out.append((class_names[class_idx], float(probs[class_idx])))
    return out


def _topk_json(items: Sequence[Tuple[str, float]]) -> str:
    payload = [{"label": str(label), "confidence": float(conf)} for label, conf in items]
    return json.dumps(payload, ensure_ascii=False, separators=(",", ":"))


def _prompt_user_topk_or_custom(
    *,
    field_name: str,
    folder_name: str,
    options: Sequence[Tuple[str, float]],
) -> Tuple[str, str]:
    if not sys.stdin.isatty():
        raise RuntimeError(
            f"Low confidence su {field_name} per '{folder_name}', ma stdin non interattivo. "
            "Rilancia da terminale interattivo o usa --low-confidence-policy error/review."
        )

    print(
        f"[USER-REVIEW] {folder_name} | {field_name} bassa confidenza",
        flush=True,
    )
    for i, (label, conf) in enumerate(options, start=1):
        print(f"  {i}) {label} (conf={conf:.4f})", flush=True)
    print("  0) Altro (inserimento manuale)", flush=True)

    while True:
        raw = input("Seleziona opzione [0-{}]: ".format(len(options))).strip()
        if raw.isdigit():
            idx = int(raw)
            if idx == 0:
                while True:
                    custom = input(f"Inserisci valore manuale per {field_name}: ").strip()
                    if custom:
                        return custom, "user_custom"
                    print("Valore vuoto non valido.", flush=True)
            if 1 <= idx <= len(options):
                return options[idx - 1][0], "user_topk"
        print("Scelta non valida.", flush=True)


def _append_encoding_struct_entry(
    *,
    path: Path,
    sheet_name: str,
    entry: Dict[str, str],
) -> None:
    suffix = path.suffix.lower()
    path.parent.mkdir(parents=True, exist_ok=True)

    if suffix in {".csv", ".txt"}:
        exists = path.exists()
        fieldnames = list(entry.keys())
        with path.open("a", encoding="utf-8", newline="") as fh:
            writer = csv.DictWriter(fh, fieldnames=fieldnames)
            if not exists:
                writer.writeheader()
            writer.writerow(entry)
        return

    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        try:
            from openpyxl import Workbook, load_workbook
        except Exception as exc:  # pragma: no cover - optional dependency
            raise RuntimeError(
                "Per aggiornare un file Excel serve openpyxl installato "
                "(pip install openpyxl)."
            ) from exc

        if path.exists():
            wb = load_workbook(path)
        else:
            wb = Workbook()

        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.create_sheet(sheet_name)

        expected_headers = list(entry.keys())
        existing_headers: List[str] = []
        col = 1
        while True:
            value = ws.cell(row=1, column=col).value
            if value is None or str(value).strip() == "":
                break
            existing_headers.append(str(value))
            col += 1

        if not existing_headers:
            for i, header in enumerate(expected_headers, start=1):
                ws.cell(row=1, column=i).value = header
            headers = expected_headers
        else:
            headers = list(existing_headers)
            for header in expected_headers:
                if header not in headers:
                    headers.append(header)
                    ws.cell(row=1, column=len(headers)).value = header

        ws.append([entry.get(h, "") for h in headers])
        wb.save(path)
        return

    raise RuntimeError(
        f"Formato encoding-struct non supportato: {path}. Usa .csv/.txt o .xlsx/.xlsm."
    )


def _load_classification_image_tensor(
    path: Path,
    image_size: int,
    rotate_deg_clockwise: int = 0,
) -> torch.Tensor:
    with Image.open(path) as img:
        image = img.convert("RGB")
        rotate = _normalize_rotation_deg_clockwise(rotate_deg_clockwise)
        if rotate:
            image = image.rotate(-rotate, expand=True)
        image = TF.resize(
            image,
            size=[image_size, image_size],
            interpolation=InterpolationMode.BILINEAR,
            antialias=True,
        )
        tensor = TF.to_tensor(image)
        tensor = TF.normalize(
            tensor,
            mean=(0.485, 0.456, 0.406),
            std=(0.229, 0.224, 0.225),
        )
    return tensor


def _load_rect_image_tensor(
    path: Path,
    image_size: int,
    rotate_deg_clockwise: int = 0,
) -> Tuple[torch.Tensor, int, int, int, int]:
    with Image.open(path) as img:
        image = img.convert("RGB")
        original_width, original_height = image.size
        rotate = _normalize_rotation_deg_clockwise(rotate_deg_clockwise)
        if rotate:
            image = image.rotate(-rotate, expand=True)
        width, height = image.size
        image = TF.resize(
            image,
            size=[image_size, image_size],
            interpolation=InterpolationMode.BILINEAR,
            antialias=True,
        )
        tensor = TF.to_tensor(image)
        tensor = TF.normalize(
            tensor,
            mean=(0.485, 0.456, 0.406),
            std=(0.229, 0.224, 0.225),
        )
    return tensor, width, height, original_width, original_height


def _predict_mean_probs(
    model: torch.nn.Module,
    image_paths: Sequence[Path],
    image_size: int,
    batch_size: int,
    device: torch.device,
    rotate_deg_clockwise: int = 0,
) -> np.ndarray:
    if not image_paths:
        return np.zeros((0,), dtype=np.float32)

    all_probs: List[np.ndarray] = []
    model.eval()
    with torch.inference_mode():
        for start in range(0, len(image_paths), batch_size):
            batch_paths = image_paths[start : start + batch_size]
            batch_tensors: List[torch.Tensor] = []
            for path in batch_paths:
                try:
                    batch_tensors.append(
                        _load_classification_image_tensor(
                            path,
                            image_size=image_size,
                            rotate_deg_clockwise=rotate_deg_clockwise,
                        )
                    )
                except Exception:
                    continue
            if not batch_tensors:
                continue
            images = torch.stack(batch_tensors, dim=0).to(device, non_blocking=True)
            logits = model(images)
            probs = torch.softmax(logits, dim=1).detach().cpu().numpy()
            all_probs.append(probs.astype(np.float32))

    if not all_probs:
        return np.zeros((0,), dtype=np.float32)
    probs_matrix = np.concatenate(all_probs, axis=0)
    return probs_matrix.mean(axis=0)


def _predict_rect_boxes_abs(
    model: torch.nn.Module,
    image_paths: Sequence[Path],
    image_size: int,
    batch_size: int,
    device: torch.device,
    rotate_deg_clockwise: int = 0,
) -> Tuple[List[Tuple[float, float, float, float]], List[Tuple[int, int]]]:
    boxes_abs: List[Tuple[float, float, float, float]] = []
    image_sizes: List[Tuple[int, int]] = []

    if not image_paths:
        return boxes_abs, image_sizes

    model.eval()
    with torch.inference_mode():
        for start in range(0, len(image_paths), batch_size):
            batch_paths = image_paths[start : start + batch_size]
            tensors: List[torch.Tensor] = []
            sizes: List[Tuple[int, int, int, int]] = []
            for path in batch_paths:
                try:
                    tensor, width, height, original_width, original_height = _load_rect_image_tensor(
                        path,
                        image_size=image_size,
                        rotate_deg_clockwise=rotate_deg_clockwise,
                    )
                except Exception:
                    continue
                tensors.append(tensor)
                sizes.append((width, height, original_width, original_height))
            if not tensors:
                continue

            images = torch.stack(tensors, dim=0).to(device, non_blocking=True)
            preds = normalize_box_order(model(images)).detach().cpu()

            for i, pred in enumerate(preds):
                width, height, original_width, original_height = sizes[i]
                x1 = float(pred[0]) * width
                y1 = float(pred[1]) * height
                x2 = float(pred[2]) * width
                y2 = float(pred[3]) * height
                x1 = max(0.0, min(x1, width - 1.0))
                y1 = max(0.0, min(y1, height - 1.0))
                x2 = max(x1 + 1.0, min(x2, float(width)))
                y2 = max(y1 + 1.0, min(y2, float(height)))
                box_abs = _rotate_box_back_to_original_coords(
                    box_rotated=(x1, y1, x2, y2),
                    original_width=original_width,
                    original_height=original_height,
                    rotate_deg_clockwise=rotate_deg_clockwise,
                )
                boxes_abs.append(box_abs)
                image_sizes.append((original_width, original_height))

    return boxes_abs, image_sizes


def _median_box(boxes_abs: Sequence[Tuple[float, float, float, float]]) -> Optional[Tuple[float, float, float, float]]:
    if not boxes_abs:
        return None
    x1 = statistics.median([b[0] for b in boxes_abs])
    y1 = statistics.median([b[1] for b in boxes_abs])
    x2 = statistics.median([b[2] for b in boxes_abs])
    y2 = statistics.median([b[3] for b in boxes_abs])
    if x2 <= x1:
        x2 = x1 + 1.0
    if y2 <= y1:
        y2 = y1 + 1.0
    return x1, y1, x2, y2


def _box_iou_abs(
    box_a: Tuple[float, float, float, float],
    box_b: Tuple[float, float, float, float],
) -> float:
    ax1, ay1, ax2, ay2 = box_a
    bx1, by1, bx2, by2 = box_b
    ix1 = max(ax1, bx1)
    iy1 = max(ay1, by1)
    ix2 = min(ax2, bx2)
    iy2 = min(ay2, by2)
    iw = max(0.0, ix2 - ix1)
    ih = max(0.0, iy2 - iy1)
    inter = iw * ih
    area_a = max(0.0, ax2 - ax1) * max(0.0, ay2 - ay1)
    area_b = max(0.0, bx2 - bx1) * max(0.0, by2 - by1)
    union = area_a + area_b - inter
    if union <= 0.0:
        return 0.0
    return float(inter / union)


def _clamp_box_abs_to_image(
    box: Tuple[float, float, float, float],
    width: int,
    height: int,
) -> Tuple[float, float, float, float]:
    x1, y1, x2, y2 = box
    w = max(2, int(width))
    h = max(2, int(height))
    x1 = max(0.0, min(float(x1), float(w - 1)))
    y1 = max(0.0, min(float(y1), float(h - 1)))
    x2 = max(x1 + 1.0, min(float(x2), float(w)))
    y2 = max(y1 + 1.0, min(float(y2), float(h)))
    return x1, y1, x2, y2


def _postprocess_line13_box(
    boxes_abs: Sequence[Tuple[float, float, float, float]],
    image_sizes: Sequence[Tuple[int, int]],
    iou_threshold: float,
    min_keep: int,
) -> Tuple[Optional[Tuple[float, float, float, float]], Dict[str, object]]:
    total = int(len(boxes_abs))
    diag: Dict[str, object] = {
        "enabled": True,
        "mode": "no_boxes",
        "boxes_total": total,
        "boxes_kept": 0,
        "boxes_dropped": total,
        "iou_threshold": float(iou_threshold),
        "min_keep": int(min_keep),
        "raw_iou_median": 0.0,
        "kept_iou_median": 0.0,
    }
    if total <= 0:
        return None, diag

    raw_box = _median_box(boxes_abs)
    if raw_box is None:
        diag["mode"] = "raw_median_none"
        return None, diag

    ious_to_raw = [_box_iou_abs(b, raw_box) for b in boxes_abs]
    diag["raw_iou_median"] = float(statistics.median(ious_to_raw)) if ious_to_raw else 0.0

    keep_idx = [i for i, v in enumerate(ious_to_raw) if float(v) >= float(iou_threshold)]
    min_keep_eff = max(1, min(int(min_keep), total))
    mode = "iou_threshold"
    if len(keep_idx) < min_keep_eff:
        mode = "topk_by_iou"
        keep_idx = sorted(range(total), key=lambda i: ious_to_raw[i], reverse=True)[:min_keep_eff]
    kept_boxes = [boxes_abs[i] for i in keep_idx]
    refined_box = _median_box(kept_boxes) if kept_boxes else raw_box
    if refined_box is None:
        refined_box = raw_box

    widths = [int(w) for (w, _h) in image_sizes if int(w) > 1]
    heights = [int(h) for (_w, h) in image_sizes if int(h) > 1]
    if widths and heights:
        med_w = max(2, int(round(statistics.median(widths))))
        med_h = max(2, int(round(statistics.median(heights))))
        refined_box = _clamp_box_abs_to_image(refined_box, med_w, med_h)

    kept_ious_to_refined = [_box_iou_abs(b, refined_box) for b in kept_boxes] if kept_boxes else []
    diag.update(
        {
            "mode": mode,
            "boxes_kept": int(len(kept_boxes)),
            "boxes_dropped": int(total - len(kept_boxes)),
            "kept_iou_median": float(statistics.median(kept_ious_to_refined)) if kept_ious_to_refined else 0.0,
            "raw_box": [float(raw_box[0]), float(raw_box[1]), float(raw_box[2]), float(raw_box[3])],
            "refined_box": [float(refined_box[0]), float(refined_box[1]), float(refined_box[2]), float(refined_box[3])],
        }
    )
    return refined_box, diag


def _find_line13_template_path(folder: Path) -> Optional[Path]:
    db_echo_dir: Optional[Path] = None
    try:
        for child in sorted(folder.iterdir(), key=lambda p: p.name.lower()):
            if child.is_dir() and child.name.lower() == "db_echo":
                db_echo_dir = child
                break
    except Exception:
        return None
    if db_echo_dir is None:
        return None

    exact_candidates: List[Path] = []
    fallback_candidates: List[Path] = []
    try:
        for path in sorted(db_echo_dir.rglob("*"), key=lambda p: p.as_posix().lower()):
            if not path.is_file():
                continue
            low = path.name.lower()
            if low == "echo_name.png":
                exact_candidates.append(path)
            elif path.suffix.lower() in IMAGE_SUFFIXES and "echo_name" in low:
                fallback_candidates.append(path)
    except Exception:
        return None

    if exact_candidates:
        return exact_candidates[0]
    if fallback_candidates:
        return fallback_candidates[0]
    return None


def _load_gray_u8(path: Path) -> np.ndarray:
    with Image.open(path) as img:
        arr = np.asarray(img.convert("L"), dtype=np.uint8)
    if arr.ndim != 2 or arr.shape[0] < 2 or arr.shape[1] < 2:
        raise ValueError(f"Immagine grayscale non valida: {path}")
    return arr


def _build_search_roi_from_prior(
    prior_box: Tuple[float, float, float, float],
    *,
    image_width: int,
    image_height: int,
    margin_pct: float,
    min_width: int,
    min_height: int,
) -> Tuple[int, int, int, int]:
    x1, y1, x2, y2 = prior_box
    bw = max(1.0, float(x2) - float(x1))
    bh = max(1.0, float(y2) - float(y1))
    mx = bw * max(0.0, float(margin_pct)) / 100.0
    my = bh * max(0.0, float(margin_pct)) / 100.0

    rx1 = int(max(0, round(float(x1) - mx)))
    ry1 = int(max(0, round(float(y1) - my)))
    rx2 = int(min(int(image_width), round(float(x2) + mx)))
    ry2 = int(min(int(image_height), round(float(y2) + my)))

    if (rx2 - rx1) < int(min_width) or (ry2 - ry1) < int(min_height):
        return (0, 0, int(image_width), int(image_height))
    return (rx1, ry1, rx2, ry2)


def _ncc_best_match(
    image_gray: np.ndarray,
    template_gray: np.ndarray,
) -> Optional[Tuple[int, int, float]]:
    ih, iw = int(image_gray.shape[0]), int(image_gray.shape[1])
    th, tw = int(template_gray.shape[0]), int(template_gray.shape[1])
    if ih < th or iw < tw:
        return None

    img_t = torch.from_numpy(image_gray.astype(np.float64, copy=False)).unsqueeze(0).unsqueeze(0)
    tpl_t = torch.from_numpy(template_gray.astype(np.float64, copy=False)).unsqueeze(0).unsqueeze(0)
    tpl_mean = tpl_t.mean()
    tpl_zm = tpl_t - tpl_mean
    tpl_energy = float(torch.sum(tpl_zm * tpl_zm).item())
    if tpl_energy <= 1e-12:
        return None

    ones = torch.ones((1, 1, th, tw), dtype=torch.float64)
    num = F.conv2d(img_t, tpl_zm)
    sum_i = F.conv2d(img_t, ones)
    sum_i2 = F.conv2d(img_t * img_t, ones)
    n = float(th * tw)
    var_i = (sum_i2 - (sum_i * sum_i) / n).clamp_min(1e-12)
    den = torch.sqrt(var_i * tpl_energy)
    ncc = num / den

    flat = ncc.reshape(-1)
    best_score, best_idx = torch.max(flat, dim=0)
    if not bool(torch.isfinite(best_score).item()):
        return None

    out_w = int(ncc.shape[-1])
    idx = int(best_idx.item())
    y = idx // out_w
    x = idx % out_w
    score = float(best_score.item())
    if score < -1.0:
        score = -1.0
    if score > 1.0:
        score = 1.0
    return (x, y, score)


def _median_pairwise_iou(boxes: Sequence[Tuple[float, float, float, float]]) -> float:
    if len(boxes) < 2:
        return 1.0 if boxes else 0.0
    vals: List[float] = []
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            vals.append(_box_iou_abs(boxes[i], boxes[j]))
    if not vals:
        return 0.0
    return float(statistics.median(vals))


def _postprocess_line13_box_with_template(
    *,
    folder: Path,
    image_paths: Sequence[Path],
    model_box: Optional[Tuple[float, float, float, float]],
    search_margin_pct: float,
    max_images: int,
    min_template_std: float,
    min_template_nonblack_ratio: float,
    min_score: float,
    min_consensus_iou: float,
    min_iou_with_model: float,
) -> Tuple[Optional[Tuple[float, float, float, float]], Dict[str, object]]:
    diag: Dict[str, object] = {
        "enabled": True,
        "attempted": False,
        "applied": False,
        "mode": "init",
        "template_path": "",
        "template_std": 0.0,
        "template_nonblack_ratio": 0.0,
        "matches_used": 0,
        "score_median": 0.0,
        "score_min": 0.0,
        "consensus_iou_median": 0.0,
        "iou_with_model": 0.0,
    }
    if model_box is None:
        diag["mode"] = "skip_no_model_box"
        return None, diag

    template_path = _find_line13_template_path(folder)
    if template_path is None:
        diag["mode"] = "skip_template_missing"
        return model_box, diag
    diag["template_path"] = template_path.as_posix()

    try:
        template_gray_u8 = _load_gray_u8(template_path)
    except Exception:
        diag["mode"] = "skip_template_load_error"
        return model_box, diag

    template_gray = template_gray_u8.astype(np.float64, copy=False)
    tpl_std = float(np.std(template_gray))
    nonblack_ratio = float(np.mean(template_gray_u8 > 10))
    diag["template_std"] = tpl_std
    diag["template_nonblack_ratio"] = nonblack_ratio
    if tpl_std < float(min_template_std):
        diag["mode"] = "skip_template_low_std"
        return model_box, diag
    if nonblack_ratio < float(min_template_nonblack_ratio):
        diag["mode"] = "skip_template_low_nonblack"
        return model_box, diag

    th, tw = int(template_gray.shape[0]), int(template_gray.shape[1])
    if th < 2 or tw < 2:
        diag["mode"] = "skip_template_too_small"
        return model_box, diag

    max_images_eff = max(1, int(max_images))
    match_boxes: List[Tuple[float, float, float, float]] = []
    match_scores: List[float] = []
    match_sizes: List[Tuple[int, int]] = []
    diag["attempted"] = True

    for image_path in list(image_paths)[:max_images_eff]:
        try:
            image_gray_u8 = _load_gray_u8(image_path)
        except Exception:
            continue
        image_gray = image_gray_u8.astype(np.float64, copy=False)
        ih, iw = int(image_gray.shape[0]), int(image_gray.shape[1])
        rx1, ry1, rx2, ry2 = _build_search_roi_from_prior(
            model_box,
            image_width=iw,
            image_height=ih,
            margin_pct=float(search_margin_pct),
            min_width=tw,
            min_height=th,
        )
        roi = image_gray[ry1:ry2, rx1:rx2]
        match = _ncc_best_match(roi, template_gray)
        if match is None and (rx1 != 0 or ry1 != 0 or rx2 != iw or ry2 != ih):
            rx1, ry1, rx2, ry2 = 0, 0, iw, ih
            roi = image_gray
            match = _ncc_best_match(roi, template_gray)
        if match is None:
            continue

        mx, my, score = match
        x1 = float(rx1 + mx)
        y1 = float(ry1 + my)
        x2 = float(x1 + tw)
        y2 = float(y1 + th)
        box_abs = _clamp_box_abs_to_image((x1, y1, x2, y2), iw, ih)
        match_boxes.append(box_abs)
        match_scores.append(float(score))
        match_sizes.append((iw, ih))

    if not match_boxes:
        diag["mode"] = "no_match"
        return model_box, diag

    candidate_box = _median_box(match_boxes)
    if candidate_box is None:
        diag["mode"] = "match_no_candidate"
        return model_box, diag

    med_w = max(2, int(round(statistics.median([w for (w, _h) in match_sizes]))))
    med_h = max(2, int(round(statistics.median([h for (_w, h) in match_sizes]))))
    candidate_box = _clamp_box_abs_to_image(candidate_box, med_w, med_h)

    score_median = float(statistics.median(match_scores))
    score_min = float(min(match_scores))
    consensus_iou = float(_median_pairwise_iou(match_boxes))
    iou_with_model = float(_box_iou_abs(candidate_box, model_box))

    diag.update(
        {
            "matches_used": int(len(match_boxes)),
            "score_median": score_median,
            "score_min": score_min,
            "consensus_iou_median": consensus_iou,
            "iou_with_model": iou_with_model,
            "candidate_box": [
                float(candidate_box[0]),
                float(candidate_box[1]),
                float(candidate_box[2]),
                float(candidate_box[3]),
            ],
        }
    )

    gate_ok = (
        (score_median >= float(min_score))
        and (consensus_iou >= float(min_consensus_iou))
        and (iou_with_model >= float(min_iou_with_model))
    )
    if gate_ok:
        diag["mode"] = "template_applied"
        diag["applied"] = True
        return candidate_box, diag

    diag["mode"] = "template_rejected_gate"
    return model_box, diag


def _is_dark_uniform_strip(values: np.ndarray, *, dark_threshold: float, max_std: float) -> bool:
    arr = np.asarray(values, dtype=np.float32).reshape(-1)
    if arr.size < 2:
        return False
    return float(np.mean(arr)) <= float(dark_threshold) and float(np.std(arr)) <= float(max_std)


def _trim_dark_uniform_borders_local(
    crop_gray_u8: np.ndarray,
    *,
    dark_threshold: float,
    max_std: float,
    max_trim_pct: float,
    min_keep_ratio: float,
) -> Tuple[Tuple[int, int, int, int], Dict[str, object]]:
    h, w = int(crop_gray_u8.shape[0]), int(crop_gray_u8.shape[1])
    diag: Dict[str, object] = {
        "height": h,
        "width": w,
        "top_trim": 0,
        "left_trim": 0,
        "bottom_trim": 0,
        "right_trim": 0,
        "applied": False,
    }
    if h < 2 or w < 2:
        diag["mode"] = "crop_too_small"
        return (0, 0, 0, 0), diag

    min_keep_h = max(2, int(round(float(h) * float(min_keep_ratio))))
    min_keep_w = max(2, int(round(float(w) * float(min_keep_ratio))))
    max_trim_rows = max(0, min(h - min_keep_h, int(round(float(h) * float(max_trim_pct) / 100.0))))
    max_trim_cols = max(0, min(w - min_keep_w, int(round(float(w) * float(max_trim_pct) / 100.0))))

    top = 0
    left = 0
    bottom = h
    right = w
    top_trim = left_trim = bottom_trim = right_trim = 0

    changed = True
    while changed:
        changed = False
        while top_trim < max_trim_rows and (bottom - (top + 1)) >= min_keep_h:
            strip = crop_gray_u8[top, left:right]
            if not _is_dark_uniform_strip(strip, dark_threshold=dark_threshold, max_std=max_std):
                break
            top += 1
            top_trim += 1
            changed = True
        while bottom_trim < max_trim_rows and ((bottom - 1) - top) >= min_keep_h:
            strip = crop_gray_u8[bottom - 1, left:right]
            if not _is_dark_uniform_strip(strip, dark_threshold=dark_threshold, max_std=max_std):
                break
            bottom -= 1
            bottom_trim += 1
            changed = True
        while left_trim < max_trim_cols and (right - (left + 1)) >= min_keep_w:
            strip = crop_gray_u8[top:bottom, left]
            if not _is_dark_uniform_strip(strip, dark_threshold=dark_threshold, max_std=max_std):
                break
            left += 1
            left_trim += 1
            changed = True
        while right_trim < max_trim_cols and ((right - 1) - left) >= min_keep_w:
            strip = crop_gray_u8[top:bottom, right - 1]
            if not _is_dark_uniform_strip(strip, dark_threshold=dark_threshold, max_std=max_std):
                break
            right -= 1
            right_trim += 1
            changed = True

    diag.update(
        {
            "mode": "trim_applied" if (top_trim or left_trim or bottom_trim or right_trim) else "no_trim_signal",
            "top_trim": int(top_trim),
            "left_trim": int(left_trim),
            "bottom_trim": int(bottom_trim),
            "right_trim": int(right_trim),
            "applied": bool(top_trim or left_trim or bottom_trim or right_trim),
        }
    )
    return (top_trim, left_trim, bottom_trim, right_trim), diag


def _postprocess_line13_box_with_dark_trim(
    *,
    image_paths: Sequence[Path],
    input_box: Optional[Tuple[float, float, float, float]],
    max_images: int,
    dark_threshold: float,
    max_std: float,
    max_trim_pct: float,
    min_keep_ratio: float,
) -> Tuple[Optional[Tuple[float, float, float, float]], Dict[str, object]]:
    diag: Dict[str, object] = {
        "enabled": True,
        "attempted": False,
        "applied": False,
        "mode": "init",
        "images_used": 0,
        "trim_signal_images": 0,
        "median_top_trim_frac": 0.0,
        "median_left_trim_frac": 0.0,
        "median_bottom_trim_frac": 0.0,
        "median_right_trim_frac": 0.0,
        "trimmed_box_area_ratio": 1.0,
    }
    if input_box is None:
        diag["mode"] = "skip_no_input_box"
        return None, diag

    max_images_eff = max(1, int(max_images))
    trim_top_fracs: List[float] = []
    trim_left_fracs: List[float] = []
    trim_bottom_fracs: List[float] = []
    trim_right_fracs: List[float] = []
    trim_signal_images = 0
    diag["attempted"] = True

    for image_path in list(image_paths)[:max_images_eff]:
        try:
            image_gray_u8 = _load_gray_u8(image_path)
        except Exception:
            continue
        ih, iw = int(image_gray_u8.shape[0]), int(image_gray_u8.shape[1])
        bx1, by1, bx2, by2 = _clamp_box_abs_to_image(input_box, iw, ih)
        ix1 = max(0, min(iw - 1, int(np.floor(bx1))))
        iy1 = max(0, min(ih - 1, int(np.floor(by1))))
        ix2 = max(ix1 + 1, min(iw, int(np.ceil(bx2))))
        iy2 = max(iy1 + 1, min(ih, int(np.ceil(by2))))
        crop_gray_u8 = image_gray_u8[iy1:iy2, ix1:ix2]
        if crop_gray_u8.size == 0:
            continue
        trims, trim_diag = _trim_dark_uniform_borders_local(
            crop_gray_u8,
            dark_threshold=dark_threshold,
            max_std=max_std,
            max_trim_pct=max_trim_pct,
            min_keep_ratio=min_keep_ratio,
        )
        top_trim, left_trim, bottom_trim, right_trim = trims
        ch = max(1, int(crop_gray_u8.shape[0]))
        cw = max(1, int(crop_gray_u8.shape[1]))
        trim_top_fracs.append(float(top_trim) / float(ch))
        trim_left_fracs.append(float(left_trim) / float(cw))
        trim_bottom_fracs.append(float(bottom_trim) / float(ch))
        trim_right_fracs.append(float(right_trim) / float(cw))
        if bool(trim_diag.get("applied", False)):
            trim_signal_images += 1

    images_used = len(trim_top_fracs)
    diag["images_used"] = int(images_used)
    diag["trim_signal_images"] = int(trim_signal_images)
    if images_used <= 0:
        diag["mode"] = "no_images_loaded"
        return input_box, diag

    top_frac = float(statistics.median(trim_top_fracs))
    left_frac = float(statistics.median(trim_left_fracs))
    bottom_frac = float(statistics.median(trim_bottom_fracs))
    right_frac = float(statistics.median(trim_right_fracs))
    diag["median_top_trim_frac"] = top_frac
    diag["median_left_trim_frac"] = left_frac
    diag["median_bottom_trim_frac"] = bottom_frac
    diag["median_right_trim_frac"] = right_frac

    x1, y1, x2, y2 = input_box
    bw = max(1.0, float(x2) - float(x1))
    bh = max(1.0, float(y2) - float(y1))
    top_trim_px = int(round(bh * top_frac))
    left_trim_px = int(round(bw * left_frac))
    bottom_trim_px = int(round(bh * bottom_frac))
    right_trim_px = int(round(bw * right_frac))
    diag["top_trim_px"] = int(top_trim_px)
    diag["left_trim_px"] = int(left_trim_px)
    diag["bottom_trim_px"] = int(bottom_trim_px)
    diag["right_trim_px"] = int(right_trim_px)

    if (top_trim_px + left_trim_px + bottom_trim_px + right_trim_px) <= 0:
        diag["mode"] = "no_trim_signal"
        return input_box, diag

    candidate_box = (
        float(x1) + float(left_trim_px),
        float(y1) + float(top_trim_px),
        float(x2) - float(right_trim_px),
        float(y2) - float(bottom_trim_px),
    )
    candidate_w = max(0.0, candidate_box[2] - candidate_box[0])
    candidate_h = max(0.0, candidate_box[3] - candidate_box[1])
    min_keep_w = float(bw) * float(min_keep_ratio)
    min_keep_h = float(bh) * float(min_keep_ratio)
    if candidate_w < max(2.0, min_keep_w) or candidate_h < max(2.0, min_keep_h):
        diag["mode"] = "trim_rejected_min_keep"
        return input_box, diag

    input_area = max(1.0, float(bw) * float(bh))
    candidate_area = max(1.0, float(candidate_w) * float(candidate_h))
    diag["trimmed_box_area_ratio"] = float(candidate_area / input_area)
    diag["mode"] = "dark_trim_applied"
    diag["applied"] = True
    return candidate_box, diag


class EchoIdResolver:
    """Resolve line #02 ID_ECHO from historical manifest statistics."""

    def __init__(self, manifest_path: Path) -> None:
        self.by_vendor_probe_video: Dict[Tuple[str, str, int, int], Counter[str]] = defaultdict(Counter)
        self.by_vendor_probe: Dict[Tuple[str, str], Counter[str]] = defaultdict(Counter)
        self.by_vendor_video: Dict[Tuple[str, int, int], Counter[str]] = defaultdict(Counter)
        self.by_vendor: Dict[str, Counter[str]] = defaultdict(Counter)
        self._load_manifest(manifest_path)

    def _load_manifest(self, manifest_path: Path) -> None:
        if not manifest_path.exists():
            raise RuntimeError(f"Manifest storico non trovato: {manifest_path}")

        folder_rows: Dict[str, Dict[str, str]] = {}
        with manifest_path.open("r", encoding="utf-8", newline="") as fh:
            reader = csv.DictReader(fh)
            for row in reader:
                folder = (row.get("dataset_folder") or row.get("model_name") or "").strip()
                if not folder:
                    continue
                if folder not in folder_rows:
                    folder_rows[folder] = row

        for row in folder_rows.values():
            vendor = (row.get("manufacturer") or "").strip()
            id_echo = (row.get("fss_id_echo") or "").strip()
            probe = (row.get("fss_id_probe") or "").strip()
            vx = (row.get("fss_video_x") or "").strip()
            vy = (row.get("fss_video_y") or "").strip()
            if not vendor or not id_echo:
                continue
            self.by_vendor[vendor][id_echo] += 1
            if probe:
                self.by_vendor_probe[(vendor, probe)][id_echo] += 1
            if vx.isdigit() and vy.isdigit():
                vx_i = int(vx)
                vy_i = int(vy)
                self.by_vendor_video[(vendor, vx_i, vy_i)][id_echo] += 1
                if probe:
                    self.by_vendor_probe_video[(vendor, probe, vx_i, vy_i)][id_echo] += 1

    @staticmethod
    def _pick(counter: Counter[str]) -> Tuple[str, float]:
        winner, count = counter.most_common(1)[0]
        total = sum(counter.values())
        support = float(count / max(1, total))
        return winner, support

    @staticmethod
    def _topk(counter: Counter[str], k: int) -> List[Tuple[str, float]]:
        total = sum(counter.values())
        if total <= 0:
            return []
        out: List[Tuple[str, float]] = []
        for value, count in counter.most_common(max(1, int(k))):
            out.append((value, float(count / total)))
        return out

    def candidates(
        self,
        vendor: str,
        probe_id: Optional[str],
        video_x: Optional[int],
        video_y: Optional[int],
        top_k: int = 3,
    ) -> List[Tuple[str, float, str]]:
        if vendor and probe_id and video_x is not None and video_y is not None:
            ctr = self.by_vendor_probe_video.get((vendor, probe_id, video_x, video_y))
            if ctr:
                return [(value, support, "vendor_probe_video") for value, support in self._topk(ctr, top_k)]

        if vendor and probe_id:
            ctr = self.by_vendor_probe.get((vendor, probe_id))
            if ctr:
                return [(value, support, "vendor_probe") for value, support in self._topk(ctr, top_k)]

        if vendor and video_x is not None and video_y is not None:
            ctr = self.by_vendor_video.get((vendor, video_x, video_y))
            if ctr:
                return [(value, support, "vendor_video") for value, support in self._topk(ctr, top_k)]

        if vendor:
            ctr = self.by_vendor.get(vendor)
            if ctr:
                return [(value, support, "vendor_only") for value, support in self._topk(ctr, top_k)]

        return []

    def resolve(
        self,
        vendor: str,
        probe_id: Optional[str],
        video_x: Optional[int],
        video_y: Optional[int],
    ) -> Tuple[str, str, float]:
        if vendor and probe_id and video_x is not None and video_y is not None:
            ctr = self.by_vendor_probe_video.get((vendor, probe_id, video_x, video_y))
            if ctr:
                value, support = self._pick(ctr)
                return value, "vendor_probe_video", support

        if vendor and probe_id:
            ctr = self.by_vendor_probe.get((vendor, probe_id))
            if ctr:
                value, support = self._pick(ctr)
                return value, "vendor_probe", support

        if vendor and video_x is not None and video_y is not None:
            ctr = self.by_vendor_video.get((vendor, video_x, video_y))
            if ctr:
                value, support = self._pick(ctr)
                return value, "vendor_video", support

        if vendor:
            ctr = self.by_vendor.get(vendor)
            if ctr:
                value, support = self._pick(ctr)
                return value, "vendor_only", support

        return "", "unresolved", 0.0


class RectNameEchoResolver:
    """Resolve line #13 RECT_NAME_ECHO from historical .fss examples."""

    def __init__(self, manifest_path: Path) -> None:
        self.by_vendor_video: Dict[Tuple[str, int, int], Counter[str]] = defaultdict(Counter)
        self.by_vendor: Dict[str, Counter[str]] = defaultdict(Counter)
        self.by_video: Dict[Tuple[int, int], Counter[str]] = defaultdict(Counter)
        self.by_global: Counter[str] = Counter()
        self._load_manifest(manifest_path)

    def _load_manifest(self, manifest_path: Path) -> None:
        if not manifest_path.exists():
            raise RuntimeError(f"Manifest storico non trovato: {manifest_path}")

        folder_rows: Dict[str, Dict[str, str]] = {}
        with manifest_path.open("r", encoding="utf-8", newline="") as fh:
            reader = csv.DictReader(fh)
            for row in reader:
                folder = (row.get("dataset_folder") or row.get("model_name") or "").strip()
                if not folder:
                    continue
                if folder not in folder_rows:
                    folder_rows[folder] = row

        for row in folder_rows.values():
            fss_path_s = (row.get("fss_path") or "").strip()
            if not fss_path_s:
                continue
            line_13 = _extract_fss_line_with_rect_offset(Path(fss_path_s), target_line_1based=13)
            if not line_13:
                continue
            if _parse_rect_name_echo_coords(line_13) is None:
                continue

            vendor = (row.get("manufacturer") or "").strip()
            vx = (row.get("fss_video_x") or "").strip()
            vy = (row.get("fss_video_y") or "").strip()
            self.by_global[line_13] += 1
            if vx.isdigit() and vy.isdigit():
                vx_i = int(vx)
                vy_i = int(vy)
                self.by_video[(vx_i, vy_i)][line_13] += 1
                if vendor:
                    self.by_vendor_video[(vendor, vx_i, vy_i)][line_13] += 1
            if vendor:
                self.by_vendor[vendor][line_13] += 1

    @staticmethod
    def _pick(counter: Counter[str]) -> Tuple[str, float]:
        winner, count = counter.most_common(1)[0]
        total = sum(counter.values())
        support = float(count / max(1, total))
        return winner, support

    @staticmethod
    def _topk(counter: Counter[str], k: int) -> List[Tuple[str, float]]:
        total = sum(counter.values())
        if total <= 0:
            return []
        out: List[Tuple[str, float]] = []
        for value, count in counter.most_common(max(1, int(k))):
            out.append((value, float(count / total)))
        return out

    def has_vendor(self, vendor: str) -> bool:
        return bool(vendor and vendor in self.by_vendor and self.by_vendor[vendor])

    def candidates(
        self,
        vendor: str,
        video_x: Optional[int],
        video_y: Optional[int],
        top_k: int = 3,
    ) -> List[Tuple[str, float, str]]:
        if vendor and video_x is not None and video_y is not None:
            ctr = self.by_vendor_video.get((vendor, video_x, video_y))
            if ctr:
                return [(value, support, "vendor_video") for value, support in self._topk(ctr, top_k)]

        if vendor:
            ctr = self.by_vendor.get(vendor)
            if ctr:
                return [(value, support, "vendor_only") for value, support in self._topk(ctr, top_k)]

        if video_x is not None and video_y is not None:
            ctr = self.by_video.get((video_x, video_y))
            if ctr:
                return [(value, support, "video_only") for value, support in self._topk(ctr, top_k)]

        if self.by_global:
            return [(value, support, "global") for value, support in self._topk(self.by_global, top_k)]

        return []

    def resolve(
        self,
        vendor: str,
        video_x: Optional[int],
        video_y: Optional[int],
    ) -> Tuple[str, str, float]:
        candidates = self.candidates(vendor=vendor, video_x=video_x, video_y=video_y, top_k=1)
        if not candidates:
            return "", "unresolved", 0.0
        value, support, source = candidates[0]
        return value, source, support


class RectNameProbeResolver:
    """Resolve line #14 RECT_NAME_PROBE from historical .fss examples."""

    def __init__(self, manifest_path: Path) -> None:
        self.by_vendor_probe_video: Dict[Tuple[str, str, int, int], Counter[str]] = defaultdict(Counter)
        self.by_vendor_probe: Dict[Tuple[str, str], Counter[str]] = defaultdict(Counter)
        self.by_probe_video: Dict[Tuple[str, int, int], Counter[str]] = defaultdict(Counter)
        self.by_probe: Dict[str, Counter[str]] = defaultdict(Counter)
        self.by_vendor_video: Dict[Tuple[str, int, int], Counter[str]] = defaultdict(Counter)
        self.by_vendor: Dict[str, Counter[str]] = defaultdict(Counter)
        self.by_video: Dict[Tuple[int, int], Counter[str]] = defaultdict(Counter)
        self.by_global: Counter[str] = Counter()
        self._load_manifest(manifest_path)

    def _load_manifest(self, manifest_path: Path) -> None:
        if not manifest_path.exists():
            raise RuntimeError(f"Manifest storico non trovato: {manifest_path}")

        folder_rows: Dict[str, Dict[str, str]] = {}
        with manifest_path.open("r", encoding="utf-8", newline="") as fh:
            reader = csv.DictReader(fh)
            for row in reader:
                folder = (row.get("dataset_folder") or row.get("model_name") or "").strip()
                if not folder:
                    continue
                if folder not in folder_rows:
                    folder_rows[folder] = row

        for row in folder_rows.values():
            fss_path_s = (row.get("fss_path") or "").strip()
            if not fss_path_s:
                continue
            line_14 = _extract_fss_line_with_rect_offset(Path(fss_path_s), target_line_1based=14)
            if not line_14:
                continue
            if _parse_rect_name_echo_coords(line_14) is None:
                continue

            vendor = (row.get("manufacturer") or "").strip()
            probe = (row.get("fss_id_probe") or "").strip()
            vx = (row.get("fss_video_x") or "").strip()
            vy = (row.get("fss_video_y") or "").strip()

            self.by_global[line_14] += 1
            if vendor:
                self.by_vendor[vendor][line_14] += 1
            if probe:
                self.by_probe[probe][line_14] += 1
                if vendor:
                    self.by_vendor_probe[(vendor, probe)][line_14] += 1

            if vx.isdigit() and vy.isdigit():
                vx_i = int(vx)
                vy_i = int(vy)
                self.by_video[(vx_i, vy_i)][line_14] += 1
                if vendor:
                    self.by_vendor_video[(vendor, vx_i, vy_i)][line_14] += 1
                if probe:
                    self.by_probe_video[(probe, vx_i, vy_i)][line_14] += 1
                    if vendor:
                        self.by_vendor_probe_video[(vendor, probe, vx_i, vy_i)][line_14] += 1

    @staticmethod
    def _pick(counter: Counter[str]) -> Tuple[str, float]:
        winner, count = counter.most_common(1)[0]
        total = sum(counter.values())
        support = float(count / max(1, total))
        return winner, support

    @staticmethod
    def _topk(counter: Counter[str], k: int) -> List[Tuple[str, float]]:
        total = sum(counter.values())
        if total <= 0:
            return []
        out: List[Tuple[str, float]] = []
        for value, count in counter.most_common(max(1, int(k))):
            out.append((value, float(count / total)))
        return out

    def has_probe(self, probe_id: str) -> bool:
        return bool(probe_id and probe_id in self.by_probe and self.by_probe[probe_id])

    def candidates(
        self,
        vendor: str,
        probe_id: Optional[str],
        video_x: Optional[int],
        video_y: Optional[int],
        top_k: int = 3,
    ) -> List[Tuple[str, float, str]]:
        if vendor and probe_id and video_x is not None and video_y is not None:
            ctr = self.by_vendor_probe_video.get((vendor, probe_id, video_x, video_y))
            if ctr:
                return [(value, support, "vendor_probe_video") for value, support in self._topk(ctr, top_k)]

        if vendor and probe_id:
            ctr = self.by_vendor_probe.get((vendor, probe_id))
            if ctr:
                return [(value, support, "vendor_probe") for value, support in self._topk(ctr, top_k)]

        if probe_id and video_x is not None and video_y is not None:
            ctr = self.by_probe_video.get((probe_id, video_x, video_y))
            if ctr:
                return [(value, support, "probe_video") for value, support in self._topk(ctr, top_k)]

        if probe_id:
            ctr = self.by_probe.get(probe_id)
            if ctr:
                return [(value, support, "probe_only") for value, support in self._topk(ctr, top_k)]

        if vendor and video_x is not None and video_y is not None:
            ctr = self.by_vendor_video.get((vendor, video_x, video_y))
            if ctr:
                return [(value, support, "vendor_video") for value, support in self._topk(ctr, top_k)]

        if vendor:
            ctr = self.by_vendor.get(vendor)
            if ctr:
                return [(value, support, "vendor_only") for value, support in self._topk(ctr, top_k)]

        if video_x is not None and video_y is not None:
            ctr = self.by_video.get((video_x, video_y))
            if ctr:
                return [(value, support, "video_only") for value, support in self._topk(ctr, top_k)]

        if self.by_global:
            return [(value, support, "global") for value, support in self._topk(self.by_global, top_k)]

        return []

    def resolve(
        self,
        vendor: str,
        probe_id: Optional[str],
        video_x: Optional[int],
        video_y: Optional[int],
    ) -> Tuple[str, str, float]:
        candidates = self.candidates(
            vendor=vendor,
            probe_id=probe_id,
            video_x=video_x,
            video_y=video_y,
            top_k=1,
        )
        if not candidates:
            return "", "unresolved", 0.0
        value, support, source = candidates[0]
        return value, source, support


# NOTE: GroupOrientationResolver (resolver storico linea #12, 4=symbol/5=depth) e il suo
# helper _parse_group_orientation_from_fss sono stati rimossi come dead code il 2026-07-09:
# la pipeline forza #12=4 (symbol). Recuperabili dal tag git v0-checkpoint-pre-claude.


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Predice righe chiave del file .fss partendo da cartelle acquisizioni raw "
            "(timestamp_<vga|hdmi>_<WxH>...)."
        )
    )
    parser.add_argument("--dataset-root", type=Path, default=Path("Dataset L_T"))
    parser.add_argument(
        "--vendor-checkpoint",
        type=Path,
        default=Path("artifacts/10_active_pipeline/pipeline_fss_head/models/vendor_training_no_negative_v2_power/best_model.pt"),
    )
    parser.add_argument(
        "--probe-checkpoint",
        type=Path,
        default=Path("artifacts/10_active_pipeline/pipeline_fss_head/models/probe_training_no_negative_v1/best_model.pt"),
    )
    parser.add_argument(
        "--rect-checkpoint",
        type=Path,
        default=Path("artifacts/10_active_pipeline/pipeline_fss_head/models/rect_training_e40_run2/best_model.pt"),
    )
    parser.add_argument(
        "--rect-vendor-map",
        type=Path,
        default=Path("artifacts/10_active_pipeline/pipeline_fss_head/maps/vendor_rect_map_bk_only.json"),
        help=(
            "JSON vendor->checkpoint rect specializzato (es. BK). "
            "Se presente, viene usato con routing su confidenza vendor."
        ),
    )
    parser.add_argument(
        "--disable-rect-vendor-routing",
        action="store_true",
        help="Disattiva il routing rect vendor-specific e usa sempre il checkpoint rect globale.",
    )
    parser.add_argument(
        "--rect-vendor-min-confidence",
        type=float,
        default=0.70,
        help="Soglia confidenza vendor minima per attivare il routing rect vendor-specific.",
    )
    parser.add_argument(
        "--line13-checkpoint",
        type=Path,
        default=Path("artifacts/10_active_pipeline/pipeline_fss_head/models/rect_training_e40_run2/best_model.pt"),
        help="Checkpoint globale regressore bbox per linea #13 (rect_name_echo).",
    )
    parser.add_argument(
        "--line13-image-size",
        type=int,
        default=0,
        help="Image size rete line13 (0=from checkpoint).",
    )
    parser.add_argument(
        "--line13-vendor-map",
        type=Path,
        default=Path("artifacts/10_active_pipeline/pipeline_fss_head/maps/vendor_line13_template_map.json"),
        help="JSON vendor->checkpoint line13 specializzato.",
    )
    parser.add_argument(
        "--disable-line13-vendor-models",
        action="store_true",
        help="Disattiva il calcolo di #13 tramite reti (mantiene resolver storico).",
    )
    parser.add_argument(
        "--disable-line13-vendor-routing",
        action="store_true",
        help=(
            "Con reti line13 attive, disattiva il vincolo vendor-specific e usa esplicitamente "
            "il checkpoint globale #13."
        ),
    )
    parser.add_argument(
        "--line13-vendor-min-confidence",
        type=float,
        default=0.70,
        help=(
            "Parametro legacy mantenuto per compatibilita': con routing vendor-specific #13 attivo "
            "non viene piu' usato come soglia di fallback."
        ),
    )
    parser.add_argument(
        "--disable-line13-postprocess",
        action="store_true",
        help="Disattiva il post-processing bbox su #13 (usa mediana raw).",
    )
    parser.add_argument(
        "--line13-postprocess-iou-threshold",
        type=float,
        default=0.35,
        help="Soglia IoU per filtro consenso bbox su #13.",
    )
    parser.add_argument(
        "--line13-postprocess-min-keep",
        type=int,
        default=3,
        help="Numero minimo di bbox da mantenere nel post-processing #13.",
    )
    parser.add_argument(
        "--disable-line13-template-postprocess",
        action="store_true",
        help=(
            "Disattiva il post-processing template-matching su #13 "
            "(DB_echo/*/echo_name.png con gate di affidabilita')."
        ),
    )
    parser.add_argument(
        "--line13-template-max-images",
        type=int,
        default=3,
        help="Numero massimo di frame per cartella usati nel template-matching #13.",
    )
    parser.add_argument(
        "--line13-template-search-margin-pct",
        type=float,
        default=80.0,
        help="Margine %% del search ROI intorno al box #13 predetto dalla rete.",
    )
    parser.add_argument(
        "--line13-template-min-std",
        type=float,
        default=60.0,
        help="Soglia minima std (0..255) del template per abilitarne l'uso.",
    )
    parser.add_argument(
        "--line13-template-min-nonblack-ratio",
        type=float,
        default=0.20,
        help="Soglia minima frazione pixel non-neri del template (0..1).",
    )
    parser.add_argument(
        "--line13-template-min-score",
        type=float,
        default=0.60,
        help="Soglia minima score NCC mediano per applicare il box template.",
    )
    parser.add_argument(
        "--line13-template-min-consensus-iou",
        type=float,
        default=0.35,
        help="Soglia minima IoU mediana tra match su frame diversi.",
    )
    parser.add_argument(
        "--line13-template-min-iou-with-model",
        type=float,
        default=0.05,
        help="Soglia minima IoU tra box template e box rete #13.",
    )
    parser.add_argument(
        "--disable-line13-dark-trim-postprocess",
        action="store_true",
        help=(
            "Disattiva il fine-tuning finale di #13 che stringe il box "
            "rimuovendo bordi scuri e uniformi del crop."
        ),
    )
    parser.add_argument(
        "--line13-dark-trim-threshold",
        type=float,
        default=70.0,
        help="Soglia media massima (0..255) per considerare un bordo #13 scuro.",
    )
    parser.add_argument(
        "--line13-dark-trim-max-std",
        type=float,
        default=18.0,
        help="Deviazione standard massima per considerare uniforme un bordo #13.",
    )
    parser.add_argument(
        "--line13-dark-trim-max-trim-pct",
        type=float,
        default=35.0,
        help="Percentuale massima di crop rimovibile per lato nel fine-tuning scuro di #13.",
    )
    parser.add_argument(
        "--line13-dark-trim-min-keep-ratio",
        type=float,
        default=0.45,
        help="Frazione minima del box #13 da mantenere dopo il fine-tuning scuro.",
    )
    parser.add_argument(
        "--rect-red-margin-pct",
        type=float,
        default=5.0,
        help="Margine orizzontale (%%) applicato al rect rosso dal segmento top.",
    )
    parser.add_argument(
        "--rect-red-bright-thr",
        type=float,
        default=70.0,
        help="Soglia luminosita' usata nel rilevamento segmento top per il rect rosso.",
    )
    parser.add_argument(
        "--disable-rect-red-line11",
        action="store_true",
        help="Disattiva rect rosso su #11 e mantiene il fallback mediana rettangoli per-image.",
    )
    parser.add_argument(
        "--reference-manifest",
        type=Path,
        default=Path("artifacts/10_active_pipeline/pipeline_fss_head/references/rect_dataset_no_negative_v2/manifest_rect_echo.csv"),
        help="Manifest storico con fss_id_echo/fss_id_probe/fss_video_x/fss_video_y.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("artifacts/10_active_pipeline/pipeline_fss_head/runs/fss_head_from_acquisitions"),
    )
    parser.add_argument("--batch-size", type=int, default=48)
    parser.add_argument(
        "--sample-per-folder",
        type=int,
        default=80,
        help="Numero frame per vendor/probe (0=tutti).",
    )
    parser.add_argument(
        "--rect-sample-per-folder",
        type=int,
        default=0,
        help=(
            "Numero frame per il rect (#11) con campionamento uniforme. "
            "0 = tutti i frame della cartella (default produzione)."
        ),
    )
    parser.add_argument(
        "--disable-rotation-normalization",
        action="store_true",
        help=(
            "Disattiva la normalizzazione rotazione (0/90/180/270) post-dedup. "
            "Default: attiva con OCR-OSD su campione dinamico."
        ),
    )
    parser.add_argument(
        "--rotation-max-samples",
        type=int,
        default=24,
        help="Numero massimo frame campionati per stima rotazione OCR-OSD (0=tutti i frame unici).",
    )
    parser.add_argument(
        "--rotation-osd-timeout-sec",
        type=float,
        default=2.5,
        help="Timeout per frame OCR-OSD (secondi).",
    )
    parser.add_argument(
        "--rotation-osd-min-confidence",
        type=float,
        default=0.0,
        help="Soglia minima Orientation confidence OSD per contare un voto.",
    )
    parser.add_argument(
        "--rotation-min-votes",
        type=int,
        default=2,
        help="Numero minimo voti OSD per accettare la rotazione di cartella.",
    )
    parser.add_argument(
        "--rotation-min-ratio",
        type=float,
        default=0.60,
        help="Rapporto minimo voti best/total_votes OSD per accettare la rotazione.",
    )
    parser.add_argument(
        "--disable-rotation-ocr-validation",
        action="store_true",
        help=(
            "Disattiva la validazione OCR della rotazione OSD. "
            "Default: attiva, applicata quando OSD propone una rotazione != 0."
        ),
    )
    parser.add_argument(
        "--rotation-ocr-samples",
        type=int,
        default=2,
        help="Numero massimo frame usati per validazione OCR della rotazione.",
    )
    parser.add_argument(
        "--rotation-ocr-timeout-sec",
        type=float,
        default=3.0,
        help="Timeout OCR per frame durante validazione rotazione.",
    )
    parser.add_argument(
        "--rotation-ocr-lang",
        type=str,
        default="eng",
        help="Lingua OCR per validazione rotazione.",
    )
    parser.add_argument(
        "--rotation-ocr-psm",
        type=int,
        default=6,
        help="PSM Tesseract per validazione OCR rotazione.",
    )
    parser.add_argument(
        "--rotation-ocr-min-word-conf",
        type=float,
        default=35.0,
        help="Conf minima parola OCR per contribuire allo score rotazione.",
    )
    parser.add_argument(
        "--rotation-ocr-score-word-bonus",
        type=float,
        default=4.0,
        help="Bonus per parola OCR valida nello score rotazione.",
    )
    parser.add_argument(
        "--rotation-ocr-min-words",
        type=int,
        default=2,
        help="Numero minimo parole OCR sul best angle per permettere override.",
    )
    parser.add_argument(
        "--rotation-ocr-min-score-delta",
        type=float,
        default=12.0,
        help="Delta minimo score OCR (best-candidate) per override rotazione.",
    )
    parser.add_argument("--vendor-image-size", type=int, default=0, help="0=from checkpoint")
    parser.add_argument("--probe-image-size", type=int, default=0, help="0=from checkpoint")
    parser.add_argument("--rect-image-size", type=int, default=0, help="0=from checkpoint")
    parser.add_argument(
        "--su-giu-rect-checkpoint",
        type=Path,
        default=Path("artifacts/20_datasets/us_orientation_binary_20260320/rect_pipeline/model_su_giu_rect/best_model.pt"),
        help="Checkpoint classificatore binario su/giu applicato sui crop del rect.",
    )
    parser.add_argument(
        "--su-giu-rect-image-size",
        type=int,
        default=0,
        help="Image size classificatore su/giu (0=from checkpoint).",
    )
    parser.add_argument(
        "--su-giu-rect-batch-size",
        type=int,
        default=0,
        help="Batch size classificatore su/giu (0=usa --batch-size).",
    )
    parser.add_argument(
        "--disable-su-giu-rect-classifier",
        action="store_true",
        help="Disattiva classificazione su/giu sui crop del rect.",
    )
    parser.add_argument(
        "--lr-marker-template-roots",
        type=Path,
        nargs="+",
        default=[DEFAULT_LR_MARKER_VENDOR_LIBRARY, Path("/Volumes/SSD_esi1_n1")],
        help="Root libreria vendor e/o root da scandire per i template DB_echo/orientation_0 usati nello step LR classico.",
    )
    parser.add_argument(
        "--disable-lr-marker-classical",
        action="store_true",
        help="Disattiva inferenza LR classica basata sui marker orientation.",
    )
    parser.add_argument(
        "--lr-marker-method",
        type=str,
        choices=("bundle", "classical"),
        default="bundle",
        help=(
            "Metodo ufficiale per marker orientamento: bundle usa la regola del bundle "
            "con assi mediani del rettangolo ecografico ufficiale; classical mantiene il vecchio blocco."
        ),
    )
    parser.add_argument(
        "--lr-marker-bundle-dir",
        type=Path,
        default=DEFAULT_ORIENTATION_MARKER_BUNDLE_DIR,
        help="Cartella unpacked del bundle orientation_marker_detector.",
    )
    parser.add_argument(
        "--lr-marker-bundle-zip",
        type=Path,
        default=DEFAULT_ORIENTATION_MARKER_BUNDLE_ZIP,
        help="Zip del bundle da usare se --lr-marker-bundle-dir non e' gia presente.",
    )
    parser.add_argument(
        "--lr-marker-bundle-library-root",
        type=Path,
        default=DEFAULT_ORIENTATION_MARKER_BUNDLE_LIBRARY_ROOT,
        help=(
            "Libreria marker usata dal metodo bundle. Default: templates interni al bundle. "
            "Se la libreria indicata non contiene il vendor, la pipeline ricade automaticamente "
            "su orientation_marker_detector/templates del bundle."
        ),
    )
    parser.add_argument(
        "--lr-marker-bundle-vertical-delta",
        type=float,
        default=0.05,
        help="Delta score minimo per correggere SU/GIU quando l'altra meta' del rect e' migliore.",
    )
    parser.add_argument(
        "--lr-marker-bundle-match-max-side",
        type=int,
        default=720,
        help="Lato massimo usato dal bundle per accelerare matchTemplate.",
    )
    parser.add_argument(
        "--lr-marker-bundle-selection-images",
        type=int,
        default=18,
        help="Numero immagini usate dal bundle per scegliere il miglior template storico del vendor.",
    )
    parser.add_argument(
        "--lr-marker-scales",
        type=str,
        default="1.0",
        help="Scale del template provate in selezione e per-immagine (la dimensione del marker "
             "varia tra ecografi/frame). Es. '0.75,1.0,1.3,1.7,2.2'. Default '1.0' = comportamento storico.",
    )
    parser.add_argument(
        "--lr-marker-pinned-templates",
        type=Path,
        default=None,
        help="JSON: nome-cartella -> [vendor/marker_NNN.png, ...] template verificati dalla review umana. "
             "Provati per-immagine accanto al template automatico; vincono solo con score alto e margine.",
    )
    parser.add_argument("--lr-marker-max-config-depth", type=int, default=4)
    parser.add_argument("--lr-marker-min-match-score", type=float, default=LR_MARKER_RELIABLE_MATCH_SCORE)
    parser.add_argument("--lr-marker-full-crop-fallback-threshold", type=float, default=0.55)
    parser.add_argument("--lr-marker-expanded-search-threshold", type=float, default=0.66)
    parser.add_argument("--lr-marker-expanded-search-steps", type=str, default="0.03,0.06,0.10,0.15,0.20")
    parser.add_argument("--lr-marker-blank-template-max-value", type=int, default=3)
    parser.add_argument("--lr-marker-min-sugiu-confidence", type=float, default=0.80)
    parser.add_argument(
        "--lr-marker-template-policy",
        type=str,
        choices=("historical_best_then_derived", "historical_best", "derived_folder"),
        default="historical_best_then_derived",
        help=(
            "Policy template LR marker per cartella: historical_best_then_derived usa prima "
            "il template storico best-match e ricorre al template ad hoc solo se il match "
            "storico e' debole; historical_best usa sempre e solo lo storico; derived_folder "
            "crea subito un template ad hoc dal miglior match storico."
        ),
    )
    parser.add_argument(
        "--lr-marker-manual-seeds-file",
        type=Path,
        default=None,
        help=(
            "JSON opzionale con rettangoli marker corretti manualmente. Se presente, "
            "la pipeline crea template manuali di cartella e li prova prima dello storico."
        ),
    )
    parser.add_argument(
        "--lr-marker-review-file",
        type=Path,
        default=DEFAULT_LR_MARKER_REVIEW_FILE,
        help="JSON opzionale con template LR marker storici accettati/scartati per vendor.",
    )
    parser.add_argument(
        "--lr-marker-exclude-vendors",
        nargs="*",
        default=["Biopsee"],
        help="Vendor esclusi dallo step LR marker classico.",
    )
    parser.add_argument(
        "--lt-rect-checkpoint",
        type=Path,
        default=Path("artifacts/30_models/lt_training_transrectal_rect_only_v2_trecall/best_model.pt"),
        help="Checkpoint classificatore binario L/T applicato sui crop del rect.",
    )
    parser.add_argument(
        "--lt-rect-image-size",
        type=int,
        default=0,
        help="Image size classificatore L/T (0=from checkpoint).",
    )
    parser.add_argument(
        "--lt-rect-batch-size",
        type=int,
        default=0,
        help="Batch size classificatore L/T (0=usa --batch-size).",
    )
    parser.add_argument(
        "--lt-min-confidence",
        type=float,
        default=0.55,
        help="Soglia media confidenza L/T sotto cui aggiungere review flag.",
    )
    parser.add_argument(
        "--disable-lt-rect-classifier",
        action="store_true",
        help="Disattiva classificazione L/T sui crop del rect.",
    )
    parser.add_argument(
        "--disable-rect-depth-autonomous",
        action="store_true",
        help="Disattiva la tappa ufficiale RECT_DEPTH autonoma OCR/classica/ranker.",
    )
    parser.add_argument(
        "--rect-depth-max-images",
        type=int,
        default=0,
        help="Numero massimo immagini per RECT_DEPTH autonomo (0=tutta la cartella).",
    )
    parser.add_argument(
        "--rect-depth-max-candidates-per-sample",
        type=int,
        default=24,
        help="Numero massimo candidati OCR/classici per immagine nello step RECT_DEPTH.",
    )
    parser.add_argument(
        "--rect-depth-ocr-timeout",
        type=float,
        default=8.0,
        help="Timeout OCR per immagine nello step RECT_DEPTH.",
    )
    parser.add_argument(
        "--rect-depth-subprocess-timeout-sec",
        type=float,
        default=0.0,
        help="Timeout globale subprocess RECT_DEPTH (0=nessun timeout globale).",
    )
    parser.add_argument(
        "--rect-depth-scale-side-preference",
        type=str,
        default="",
        choices=("", "auto", "right", "left", "destra", "sinistra", "dx", "sx"),
        help="Preferenza lato scala per RECT_DEPTH; vuoto=profilo vendor.",
    )
    parser.add_argument(
        "--rect-depth-min-accepted-ratio",
        type=float,
        default=0.80,
        help="Quota minima accepted per considerare ok lo step RECT_DEPTH a livello cartella.",
    )
    parser.add_argument(
        "--disable-scale-stage",
        action="store_true",
        help="Disattiva lo stadio scala (#18-#21). Con questo flag la pipeline si comporta "
             "esattamente come prima della sua introduzione.",
    )
    parser.add_argument(
        "--scale-max-frames",
        type=int,
        default=48,
        help="Frame studiati al massimo dallo stadio scala (l'OCR delle etichette e la parte "
             "costosa). I frame sono scelti fra quelli per cui la depth ha dato un valore.",
    )
    parser.add_argument(
        "--scale-subprocess-timeout-sec",
        type=float,
        default=900.0,
        help="Timeout del subprocess dello stadio scala (0 = nessun timeout).",
    )
    parser.add_argument(
        "--scale-min-accepted-ratio",
        type=float,
        default=0.80,
        help="Quota minima di depth accepted per considerare ok lo stadio scala.",
    )
    parser.add_argument(
        "--scale-corrections",
        type=Path,
        default=None,
        help="JSON delle correzioni scala esportate dalla pagina di studio "
             "(tools/scale/study_scale_folder.py). Da qui viene usata solo la colonna del "
             "righello, che e un fatto di cartella; le correzioni per singolo frame restano "
             "nella pagina di review.",
    )
    parser.add_argument("--device", type=str, default=None)
    parser.add_argument("--exclude-folder-regex", type=str, default=None)
    parser.add_argument("--max-folders", type=int, default=0, help="0=tutte")
    parser.add_argument(
        "--fss-version",
        type=str,
        default="4.0",
        help="Valore linea #01 del file .fss (default: 4.0).",
    )
    parser.add_argument(
        "--vendor-min-confidence",
        type=float,
        default=0.50,
        help="Soglia confidenza vendor sotto cui attivare la policy low-confidence.",
    )
    parser.add_argument(
        "--vendor-min-margin",
        type=float,
        default=0.0,
        help=(
            "Soglia minima sul margine vendor (top1 - top2). Sotto soglia attiva la stessa "
            "policy low-confidence del vendor. Default 0.0 = disattivato (comportamento storico)."
        ),
    )
    parser.add_argument(
        "--disable-vendor-ocr-fallback",
        action="store_true",
        help=(
            "Disattiva fallback OCR sul vendor quando la confidenza CNN e' sotto soglia. "
            "Default: attivo."
        ),
    )
    parser.add_argument(
        "--vendor-ocr-samples",
        type=int,
        default=6,
        help="Numero massimo frame usati per fallback OCR vendor sotto soglia.",
    )
    parser.add_argument(
        "--vendor-ocr-timeout-sec",
        type=float,
        default=3.0,
        help="Timeout OCR per frame durante fallback vendor.",
    )
    parser.add_argument(
        "--vendor-ocr-lang",
        type=str,
        default="eng",
        help="Lingua OCR per fallback vendor.",
    )
    parser.add_argument(
        "--vendor-ocr-psm",
        type=int,
        default=6,
        help="PSM Tesseract per fallback vendor.",
    )
    parser.add_argument(
        "--vendor-ocr-min-word-conf",
        type=float,
        default=35.0,
        help="Conf minima parola OCR per fallback vendor.",
    )
    parser.add_argument(
        "--vendor-ocr-min-score",
        type=float,
        default=2.0,
        help="Score OCR minimo per accettare override vendor.",
    )
    parser.add_argument(
        "--vendor-ocr-min-score-delta",
        type=float,
        default=0.7,
        help="Delta score OCR minimo tra top1 e top2 per override vendor.",
    )
    parser.add_argument(
        "--vendor-ocr-min-hits",
        type=int,
        default=2,
        help="Numero minimo hit OCR sul vendor candidato per override.",
    )
    parser.add_argument(
        "--probe-min-confidence",
        type=float,
        default=0.40,
        help="Soglia confidenza probe sotto cui attivare la policy low-confidence.",
    )
    parser.add_argument(
        "--probe-type-summary-csv",
        type=Path,
        default=Path("artifacts/60_metadata/probe_model_to_type_summary_2026-03-31.csv"),
        help="CSV summary probe->primary_probe_type usato dal router linea #04.",
    )
    parser.add_argument(
        "--probe-type-evidence-csv",
        type=Path,
        default=Path("artifacts/60_metadata/probe_id_probe_type_evidence_from_line16_manifest_2026-03-30.csv"),
        help="CSV evidence probe_id+probe_type_id+share usato dal router linea #04.",
    )
    parser.add_argument(
        "--disable-probe-type-router",
        action="store_true",
        help="Disattiva routing probe type per la linea #04 (fallback: valore vuoto).",
    )
    parser.add_argument(
        "--line13-min-support",
        type=float,
        default=0.00,
        help="Soglia supporto storico per #13 sotto cui attivare la policy low-confidence.",
    )
    parser.add_argument(
        "--line14-min-support",
        type=float,
        default=0.00,
        help="Soglia supporto storico per #14 sotto cui attivare la policy low-confidence.",
    )
    parser.add_argument(
        "--low-confidence-policy",
        type=str,
        choices=["ask_user", "error", "review"],
        default="ask_user",
        help="Gestione bassa confidenza: ask_user (top-k), error (stop), review (flag).",
    )
    parser.add_argument(
        "--interactive-topk",
        type=int,
        default=3,
        help="Numero opzioni da proporre all'utente in modalita ask_user.",
    )
    parser.add_argument(
        "--encoding-struct-path",
        type=Path,
        default=None,
        help=(
            "File encoding struct da aggiornare quando l'utente inserisce un valore manuale. "
            "Supportati .csv/.txt/.xlsx/.xlsm."
        ),
    )
    parser.add_argument(
        "--encoding-struct-sheet",
        type=str,
        default=".fss",
        help="Nome foglio Excel per encoding struct (ignorato per CSV).",
    )
    parser.add_argument("--log-interval", type=int, default=25)
    parser.add_argument(
        "--stage-events",
        action="store_true",
        help="Stampa una riga '##STAGE {json}' alla fine di ogni stadio di ogni cartella "
             "(dedup, rotazione, vendor, probe, rect, orientamento, depth, scala, folder_done). "
             "Serve al tool di revisione per mostrare i risultati mentre la cartella gira; "
             "senza il flag l'output della pipeline e' identico a prima.",
    )
    return parser


def main() -> int:
    args = _build_parser().parse_args()

    global _STAGE_EVENTS_ENABLED
    _STAGE_EVENTS_ENABLED = bool(args.stage_events)

    if args.batch_size <= 0:
        raise ValueError("--batch-size deve essere > 0.")
    if args.sample_per_folder < 0 or args.rect_sample_per_folder < 0 or args.rotation_max_samples < 0:
        raise ValueError(
            "--sample-per-folder, --rect-sample-per-folder e --rotation-max-samples devono essere >= 0."
        )
    if args.interactive_topk <= 0:
        raise ValueError("--interactive-topk deve essere > 0.")
    if args.rotation_osd_timeout_sec <= 0:
        raise ValueError("--rotation-osd-timeout-sec deve essere > 0.")
    if args.rotation_osd_min_confidence < 0:
        raise ValueError("--rotation-osd-min-confidence deve essere >= 0.")
    if args.rotation_min_votes <= 0:
        raise ValueError("--rotation-min-votes deve essere > 0.")
    if not (0.0 <= args.rotation_min_ratio <= 1.0):
        raise ValueError("--rotation-min-ratio deve essere tra 0 e 1.")
    if args.rotation_ocr_samples <= 0:
        raise ValueError("--rotation-ocr-samples deve essere > 0.")
    if args.rotation_ocr_timeout_sec <= 0:
        raise ValueError("--rotation-ocr-timeout-sec deve essere > 0.")
    if args.rotation_ocr_min_word_conf < 0:
        raise ValueError("--rotation-ocr-min-word-conf deve essere >= 0.")
    if args.rotation_ocr_score_word_bonus < 0:
        raise ValueError("--rotation-ocr-score-word-bonus deve essere >= 0.")
    if args.rotation_ocr_min_words < 0:
        raise ValueError("--rotation-ocr-min-words deve essere >= 0.")
    if args.rotation_ocr_min_score_delta < 0:
        raise ValueError("--rotation-ocr-min-score-delta deve essere >= 0.")
    if not (0.0 <= args.vendor_min_confidence <= 1.0):
        raise ValueError("--vendor-min-confidence deve essere tra 0 e 1.")
    if args.vendor_ocr_samples <= 0:
        raise ValueError("--vendor-ocr-samples deve essere > 0.")
    if args.vendor_ocr_timeout_sec <= 0:
        raise ValueError("--vendor-ocr-timeout-sec deve essere > 0.")
    if args.vendor_ocr_min_word_conf < 0:
        raise ValueError("--vendor-ocr-min-word-conf deve essere >= 0.")
    if args.vendor_ocr_min_score < 0:
        raise ValueError("--vendor-ocr-min-score deve essere >= 0.")
    if args.vendor_ocr_min_score_delta < 0:
        raise ValueError("--vendor-ocr-min-score-delta deve essere >= 0.")
    if args.vendor_ocr_min_hits <= 0:
        raise ValueError("--vendor-ocr-min-hits deve essere > 0.")
    if not (0.0 <= args.probe_min_confidence <= 1.0):
        raise ValueError("--probe-min-confidence deve essere tra 0 e 1.")
    if not (0.0 <= args.rect_vendor_min_confidence <= 1.0):
        raise ValueError("--rect-vendor-min-confidence deve essere tra 0 e 1.")
    if not (0.0 <= args.line13_vendor_min_confidence <= 1.0):
        raise ValueError("--line13-vendor-min-confidence deve essere tra 0 e 1.")
    if not (0.0 <= args.line13_postprocess_iou_threshold <= 1.0):
        raise ValueError("--line13-postprocess-iou-threshold deve essere tra 0 e 1.")
    if args.line13_postprocess_min_keep <= 0:
        raise ValueError("--line13-postprocess-min-keep deve essere > 0.")
    if args.line13_template_max_images <= 0:
        raise ValueError("--line13-template-max-images deve essere > 0.")
    if args.line13_template_search_margin_pct < 0.0:
        raise ValueError("--line13-template-search-margin-pct deve essere >= 0.")
    if args.line13_template_min_std < 0.0:
        raise ValueError("--line13-template-min-std deve essere >= 0.")
    if not (0.0 <= args.line13_template_min_nonblack_ratio <= 1.0):
        raise ValueError("--line13-template-min-nonblack-ratio deve essere tra 0 e 1.")
    if not (-1.0 <= args.line13_template_min_score <= 1.0):
        raise ValueError("--line13-template-min-score deve essere tra -1 e 1.")
    if not (0.0 <= args.line13_template_min_consensus_iou <= 1.0):
        raise ValueError("--line13-template-min-consensus-iou deve essere tra 0 e 1.")
    if not (0.0 <= args.line13_template_min_iou_with_model <= 1.0):
        raise ValueError("--line13-template-min-iou-with-model deve essere tra 0 e 1.")
    if not (0.0 <= args.line13_dark_trim_threshold <= 255.0):
        raise ValueError("--line13-dark-trim-threshold deve essere tra 0 e 255.")
    if args.line13_dark_trim_max_std < 0.0:
        raise ValueError("--line13-dark-trim-max-std deve essere >= 0.")
    if not (0.0 <= args.line13_dark_trim_max_trim_pct <= 49.0):
        raise ValueError("--line13-dark-trim-max-trim-pct deve essere tra 0 e 49.")
    if not (0.10 <= args.line13_dark_trim_min_keep_ratio <= 1.0):
        raise ValueError("--line13-dark-trim-min-keep-ratio deve essere tra 0.10 e 1.")
    if not (0.0 <= args.rect_red_margin_pct <= 30.0):
        raise ValueError("--rect-red-margin-pct deve essere tra 0 e 30.")
    if not (0.0 <= args.rect_red_bright_thr <= 255.0):
        raise ValueError("--rect-red-bright-thr deve essere tra 0 e 255.")
    if args.su_giu_rect_batch_size < 0:
        raise ValueError("--su-giu-rect-batch-size deve essere >= 0.")
    if args.su_giu_rect_image_size < 0:
        raise ValueError("--su-giu-rect-image-size deve essere >= 0.")
    if args.lr_marker_max_config_depth < 0:
        raise ValueError("--lr-marker-max-config-depth deve essere >= 0.")
    if not (0.0 <= args.lr_marker_min_match_score <= 1.0):
        raise ValueError("--lr-marker-min-match-score deve essere tra 0 e 1.")
    if not (0.0 <= args.lr_marker_full_crop_fallback_threshold <= 1.0):
        raise ValueError("--lr-marker-full-crop-fallback-threshold deve essere tra 0 e 1.")
    if not (0.0 <= args.lr_marker_expanded_search_threshold <= 1.0):
        raise ValueError("--lr-marker-expanded-search-threshold deve essere tra 0 e 1.")
    if not (0.0 <= args.lr_marker_min_sugiu_confidence <= 1.0):
        raise ValueError("--lr-marker-min-sugiu-confidence deve essere tra 0 e 1.")
    if not (0.0 <= args.lr_marker_bundle_vertical_delta <= 1.0):
        raise ValueError("--lr-marker-bundle-vertical-delta deve essere tra 0 e 1.")
    if args.lr_marker_bundle_match_max_side < 0:
        raise ValueError("--lr-marker-bundle-match-max-side deve essere >= 0.")
    if args.lr_marker_bundle_selection_images <= 0:
        raise ValueError("--lr-marker-bundle-selection-images deve essere > 0.")
    if args.lt_rect_batch_size < 0:
        raise ValueError("--lt-rect-batch-size deve essere >= 0.")
    if args.lt_rect_image_size < 0:
        raise ValueError("--lt-rect-image-size deve essere >= 0.")
    if args.line13_image_size < 0:
        raise ValueError("--line13-image-size deve essere >= 0.")
    if not (0.0 <= args.lt_min_confidence <= 1.0):
        raise ValueError("--lt-min-confidence deve essere tra 0 e 1.")
    if args.rect_depth_max_images < 0:
        raise ValueError("--rect-depth-max-images deve essere >= 0.")
    if args.rect_depth_max_candidates_per_sample <= 0:
        raise ValueError("--rect-depth-max-candidates-per-sample deve essere > 0.")
    if args.rect_depth_ocr_timeout <= 0:
        raise ValueError("--rect-depth-ocr-timeout deve essere > 0.")
    if args.rect_depth_subprocess_timeout_sec < 0:
        raise ValueError("--rect-depth-subprocess-timeout-sec deve essere >= 0.")
    if not (0.0 <= args.rect_depth_min_accepted_ratio <= 1.0):
        raise ValueError("--rect-depth-min-accepted-ratio deve essere tra 0 e 1.")
    if not (0.0 <= args.line13_min_support <= 1.0):
        raise ValueError("--line13-min-support deve essere tra 0 e 1.")
    if not (0.0 <= args.line14_min_support <= 1.0):
        raise ValueError("--line14-min-support deve essere tra 0 e 1.")
    if int(args.rect_sample_per_folder) > 0:
        print(
            f"[rect] --rect-sample-per-folder={int(args.rect_sample_per_folder)}: "
            "campionamento uniforme dei frame per il rect (0=tutti, default produzione).",
            flush=True,
        )
    if float(args.scale_subprocess_timeout_sec) < 0:
        raise ValueError("--scale-subprocess-timeout-sec deve essere >= 0.")
    if not (0.0 <= args.scale_min_accepted_ratio <= 1.0):
        raise ValueError("--scale-min-accepted-ratio deve essere tra 0 e 1.")
    if int(args.scale_max_frames) <= 0:
        raise ValueError("--scale-max-frames deve essere >= 1.")

    dataset_root = args.dataset_root.expanduser().resolve()
    output_dir = args.output_dir.expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    lr_marker_manual_seed_rows = _load_lr_marker_manual_seed_rows(args.lr_marker_manual_seeds_file)
    encoding_struct_path = (
        args.encoding_struct_path.expanduser().resolve()
        if args.encoding_struct_path is not None
        else (output_dir / "encoding_struct_fss_updates.csv")
    )

    vendor_ckpt_path = args.vendor_checkpoint.expanduser().resolve()
    probe_ckpt_path = args.probe_checkpoint.expanduser().resolve()
    rect_ckpt_path = args.rect_checkpoint.expanduser().resolve()
    line13_ckpt_path = args.line13_checkpoint.expanduser().resolve()
    manifest_ref_path = args.reference_manifest.expanduser().resolve()
    probe_type_summary_csv = args.probe_type_summary_csv.expanduser().resolve()
    probe_type_evidence_csv = args.probe_type_evidence_csv.expanduser().resolve()

    device = choose_device(args.device)
    print(f"Device: {device}", flush=True)

    vendor_ckpt = torch.load(vendor_ckpt_path, map_location="cpu", weights_only=False)
    vendor_classes = list(vendor_ckpt.get("class_names") or [])
    if not vendor_classes:
        raise RuntimeError("Checkpoint vendor senza class_names.")
    vendor_image_size = int(args.vendor_image_size) or int(vendor_ckpt.get("args", {}).get("image_size", 320))
    vendor_model = VendorClassifier(num_classes=len(vendor_classes), pretrained=False).to(device)
    vendor_model.load_state_dict(vendor_ckpt["model_state_dict"])
    vendor_model.eval()
    vendor_ocr_alias_map = _build_vendor_ocr_alias_map(vendor_classes)

    probe_ckpt = torch.load(probe_ckpt_path, map_location="cpu", weights_only=False)
    probe_classes = list(probe_ckpt.get("class_names") or [])
    if not probe_classes:
        raise RuntimeError("Checkpoint probe senza class_names.")
    probe_image_size = int(args.probe_image_size) or int(probe_ckpt.get("args", {}).get("image_size", 320))
    probe_model = ProbeClassifier(num_classes=len(probe_classes), pretrained=False).to(device)
    probe_model.load_state_dict(probe_ckpt["model_state_dict"])
    probe_model.eval()

    rect_ckpt = torch.load(rect_ckpt_path, map_location="cpu", weights_only=False)
    rect_image_size = int(args.rect_image_size) or int(rect_ckpt.get("args", {}).get("image_size", 320))
    rect_model = RectRegressor(pretrained=False).to(device)
    rect_model.load_state_dict(rect_ckpt["model_state_dict"])
    rect_model.eval()

    line13_model: Optional[torch.nn.Module] = None
    line13_image_size = 0
    if not args.disable_line13_vendor_models:
        if not line13_ckpt_path.exists():
            raise FileNotFoundError(f"Checkpoint line13 non trovato: {line13_ckpt_path}")
        line13_ckpt = torch.load(line13_ckpt_path, map_location="cpu", weights_only=False)
        line13_image_size = int(args.line13_image_size) or int(
            line13_ckpt.get("args", {}).get("image_size", rect_image_size)
        )
        line13_model = RectRegressor(pretrained=False).to(device)
        line13_model.load_state_dict(line13_ckpt["model_state_dict"])
        line13_model.eval()

    su_giu_model: Optional[torch.nn.Module] = None
    su_giu_ckpt_path = args.su_giu_rect_checkpoint.expanduser().resolve()
    su_giu_image_size = 0
    su_giu_batch_size = max(1, int(args.su_giu_rect_batch_size) if int(args.su_giu_rect_batch_size) > 0 else int(args.batch_size))
    su_giu_class_names: List[str] = list(SU_GIU_LABELS)
    if not args.disable_su_giu_rect_classifier:
        if not su_giu_ckpt_path.exists():
            raise FileNotFoundError(f"Checkpoint su/giu non trovato: {su_giu_ckpt_path}")
        su_giu_ckpt = torch.load(su_giu_ckpt_path, map_location="cpu", weights_only=False)
        su_giu_class_names = [str(x).strip().lower() for x in (su_giu_ckpt.get("class_names") or list(SU_GIU_LABELS))]
        if len(su_giu_class_names) != 2:
            raise RuntimeError(
                "Checkpoint su/giu non valido: attese 2 classi, "
                f"ricevute {len(su_giu_class_names)} ({su_giu_class_names})."
            )
        su_giu_image_size = int(args.su_giu_rect_image_size) or int(su_giu_ckpt.get("args", {}).get("image_size", 256))
        su_giu_model = SuGiuRectClassifier().to(device)
        su_giu_model.load_state_dict(su_giu_ckpt["model_state_dict"])
        su_giu_model.eval()

    lt_model: Optional[torch.nn.Module] = None
    lt_ckpt_path = args.lt_rect_checkpoint.expanduser().resolve()
    lt_image_size = 0
    lt_batch_size = max(1, int(args.lt_rect_batch_size) if int(args.lt_rect_batch_size) > 0 else int(args.batch_size))
    lt_class_names: List[str] = list(LT_LABELS)
    if not args.disable_lt_rect_classifier:
        if not lt_ckpt_path.exists():
            raise FileNotFoundError(f"Checkpoint L/T non trovato: {lt_ckpt_path}")
        lt_ckpt = torch.load(lt_ckpt_path, map_location="cpu", weights_only=False)
        lt_class_names = [str(x).strip().upper() for x in (lt_ckpt.get("class_names") or list(LT_LABELS))]
        if len(lt_class_names) != 2:
            raise RuntimeError(
                "Checkpoint L/T non valido: attese 2 classi, "
                f"ricevute {len(lt_class_names)} ({lt_class_names})."
            )
        lt_image_size = int(args.lt_rect_image_size) or int(lt_ckpt.get("args", {}).get("image_size", 256))
        lt_model = LTRectClassifier().to(device)
        lt_model.load_state_dict(lt_ckpt["model_state_dict"])
        lt_model.eval()

    echo_resolver = EchoIdResolver(manifest_ref_path)
    rect_name_echo_resolver = RectNameEchoResolver(manifest_ref_path)
    rect_name_probe_resolver = RectNameProbeResolver(manifest_ref_path)

    exclude_re = re.compile(args.exclude_folder_regex) if args.exclude_folder_regex else None
    folders = [p for p in sorted(dataset_root.iterdir()) if p.is_dir()]
    if exclude_re:
        folders = [p for p in folders if not exclude_re.search(p.name)]
    if args.max_folders > 0:
        folders = folders[: args.max_folders]

    if not folders:
        raise RuntimeError("Nessuna cartella trovata nel dataset-root.")

    predictions: List[FolderPrediction] = []
    warnings: List[str] = []
    status_counter: Counter[str] = Counter()
    rect_route_counter: Counter[str] = Counter()
    rect_vendor_specialized_usage: Counter[str] = Counter()
    su_giu_source_counter: Counter[str] = Counter()
    su_giu_majority_counter: Counter[str] = Counter()
    lr_marker_source_counter: Counter[str] = Counter()
    lr_marker_majority_counter: Counter[str] = Counter()
    lt_source_counter: Counter[str] = Counter()
    lt_majority_counter: Counter[str] = Counter()
    rotation_source_counter: Counter[str] = Counter()
    rotation_degree_counter: Counter[int] = Counter()
    total_images_raw = 0
    total_images_unique = 0
    total_duplicates_removed = 0
    manual_entries_written = 0
    su_giu_per_image_rows: List[Dict[str, object]] = []
    lr_marker_per_image_rows: List[Dict[str, object]] = []
    lr_marker_manual_seed_template_meta: List[Dict[str, object]] = []
    lt_per_image_rows: List[Dict[str, object]] = []
    rect_depth_per_image_rows: List[Dict[str, object]] = []
    rect_depth_status_counter: Counter[str] = Counter()
    rect_depth_mode_counter: Counter[str] = Counter()
    scale_per_image_rows: List[Dict[str, object]] = []
    scale_status_counter: Counter[str] = Counter()
    scale_source_counter: Counter[str] = Counter()
    lr_marker_templates_cache: Dict[str, Tuple[List[object], Dict[str, object]]] = {}
    lr_marker_template_roots = [p.expanduser().resolve() for p in args.lr_marker_template_roots]
    lr_marker_expanded_search_steps = (
        _lr_marker_parse_float_steps(str(args.lr_marker_expanded_search_steps))
        if _lr_marker_parse_float_steps is not None
        else tuple()
    )
    lr_marker_scales = tuple(
        float(s) for s in str(args.lr_marker_scales).split(",") if s.strip()
    ) or (1.0,)
    lr_marker_pinned_map: Dict[str, List[str]] = {}
    if args.lr_marker_pinned_templates is not None:
        pinned_path = args.lr_marker_pinned_templates.expanduser().resolve()
        if pinned_path.is_file():
            try:
                raw_pinned = json.loads(pinned_path.read_text(encoding="utf-8"))
                lr_marker_pinned_map = {
                    str(k): [str(x) for x in v]
                    for k, v in raw_pinned.items()
                    if isinstance(v, list) and not str(k).startswith("_")
                }
                print(f"[lr-marker] pinned templates: {len(lr_marker_pinned_map)} folders", flush=True)
            except Exception as exc:
                print(f"[lr-marker] WARN cannot read pinned templates {pinned_path}: {exc}", flush=True)
    rect_red_by_folder: Dict[str, Dict[str, object]] = {}
    rect_vendor_map_path: Optional[Path] = None
    rect_vendor_map_loaded: Dict[str, str] = {}
    rect_vendor_models: Dict[str, Tuple[torch.nn.Module, int, Path]] = {}
    line13_vendor_map_path: Optional[Path] = None
    line13_vendor_map_loaded: Dict[str, str] = {}
    line13_vendor_ckpt_paths: Dict[str, Path] = {}
    line13_vendor_model_cache: Dict[str, Tuple[torch.nn.Module, int, Path]] = {}
    line13_route_counter: Counter[str] = Counter()
    line13_vendor_specialized_usage: Counter[str] = Counter()
    line13_postprocess_mode_counter: Counter[str] = Counter()
    line13_postprocess_total_boxes = 0
    line13_postprocess_total_kept = 0
    line13_postprocess_total_dropped = 0
    line13_template_postprocess_mode_counter: Counter[str] = Counter()
    line13_template_postprocess_attempted_total = 0
    line13_template_postprocess_applied_total = 0
    line13_dark_trim_mode_counter: Counter[str] = Counter()
    line13_dark_trim_attempted_total = 0
    line13_dark_trim_applied_total = 0
    probe_type_router: Optional[ProbeTypeRouter] = None
    probe_type_router_entries = 0

    def _register_manual_choice(
        *,
        folder_name: str,
        field_name: str,
        chosen_value: str,
        topk_options: Sequence[Tuple[str, float]],
        threshold: Optional[float],
        confidence: Optional[float],
    ) -> None:
        nonlocal manual_entries_written
        entry = {
            "timestamp_utc": datetime.now(timezone.utc).isoformat(),
            "folder_name": folder_name,
            "field_name": field_name,
            "chosen_value": chosen_value,
            "model_confidence": f"{confidence:.6f}" if confidence is not None else "",
            "threshold": f"{threshold:.6f}" if threshold is not None else "",
            "topk_json": json.dumps(
                [{"value": value, "confidence": conf} for value, conf in topk_options],
                ensure_ascii=False,
            ),
            "note": "manual_user_input_on_low_confidence",
        }
        _append_encoding_struct_entry(
            path=encoding_struct_path,
            sheet_name=args.encoding_struct_sheet,
            entry=entry,
        )
        manual_entries_written += 1

    if args.disable_probe_type_router:
        warnings.append("probe_type_router disattivato da flag --disable-probe-type-router.")
    else:
        if not probe_type_summary_csv.exists():
            warnings.append(f"probe_type_router: summary csv non trovato: {probe_type_summary_csv}")
        if not probe_type_evidence_csv.exists():
            warnings.append(f"probe_type_router: evidence csv non trovato: {probe_type_evidence_csv}")
        try:
            probe_type_router = ProbeTypeRouter.from_csvs(
                probe_summary_csv=probe_type_summary_csv,
                probe_type_evidence_csv=probe_type_evidence_csv,
            )
            probe_type_router_entries = sum(1 for _ in probe_type_router.iter_decisions())
            print(
                f"[probe_type_router] loaded entries={probe_type_router_entries} "
                f"(summary={probe_type_summary_csv.name}, evidence={probe_type_evidence_csv.name})",
                flush=True,
            )
        except Exception as exc:
            warnings.append(f"probe_type_router load failed: {exc}")
            probe_type_router = None

    if not args.disable_rect_vendor_routing:
        rect_vendor_map_path = args.rect_vendor_map.expanduser().resolve()
        if rect_vendor_map_path.exists():
            try:
                raw_vendor_map = json.loads(rect_vendor_map_path.read_text(encoding="utf-8"))
            except Exception as exc:
                raise RuntimeError(f"Impossibile leggere rect vendor map: {rect_vendor_map_path}") from exc

            if not isinstance(raw_vendor_map, dict):
                raise RuntimeError(
                    f"Rect vendor map non valida (atteso JSON object vendor->checkpoint): {rect_vendor_map_path}"
                )

            for vendor_name_raw, vendor_ckpt_raw in raw_vendor_map.items():
                vendor_name = str(vendor_name_raw).strip()
                if not vendor_name:
                    continue

                vendor_ckpt_path_map = Path(str(vendor_ckpt_raw)).expanduser()
                if not vendor_ckpt_path_map.is_absolute():
                    candidate_from_map_dir = (rect_vendor_map_path.parent / vendor_ckpt_path_map).resolve()
                    candidate_from_cwd = (Path.cwd() / vendor_ckpt_path_map).resolve()
                    if candidate_from_map_dir.exists():
                        vendor_ckpt_path_map = candidate_from_map_dir
                    elif candidate_from_cwd.exists():
                        vendor_ckpt_path_map = candidate_from_cwd
                    else:
                        vendor_ckpt_path_map = candidate_from_map_dir
                if not vendor_ckpt_path_map.exists():
                    warnings.append(
                        f"rect vendor map: checkpoint non trovato per vendor '{vendor_name}' -> {vendor_ckpt_path_map}"
                    )
                    continue

                try:
                    vendor_rect_ckpt = torch.load(vendor_ckpt_path_map, map_location="cpu", weights_only=False)
                    vendor_rect_image_size = int(args.rect_image_size) or int(
                        vendor_rect_ckpt.get("args", {}).get("image_size", rect_image_size)
                    )
                    vendor_rect_model = RectRegressor(pretrained=False).to(device)
                    vendor_rect_model.load_state_dict(vendor_rect_ckpt["model_state_dict"])
                    vendor_rect_model.eval()
                except Exception as exc:
                    warnings.append(
                        f"rect vendor map: errore caricamento '{vendor_name}' ({vendor_ckpt_path_map}): {exc}"
                    )
                    continue

                rect_vendor_models[vendor_name] = (
                    vendor_rect_model,
                    vendor_rect_image_size,
                    vendor_ckpt_path_map,
                )
                rect_vendor_map_loaded[vendor_name] = vendor_ckpt_path_map.as_posix()
        else:
            warnings.append(
                f"rect vendor map non trovata: {rect_vendor_map_path}. "
                "Routing rect vendor-specific disattivato (uso modello globale)."
            )

    if line13_model is not None and (not args.disable_line13_vendor_routing):
        line13_vendor_map_path = args.line13_vendor_map.expanduser().resolve()
        if line13_vendor_map_path.exists():
            try:
                raw_line13_map = json.loads(line13_vendor_map_path.read_text(encoding="utf-8"))
            except Exception as exc:
                raise RuntimeError(f"Impossibile leggere line13 vendor map: {line13_vendor_map_path}") from exc

            if not isinstance(raw_line13_map, dict):
                raise RuntimeError(
                    f"Line13 vendor map non valida (atteso JSON object vendor->checkpoint): {line13_vendor_map_path}"
                )

            for vendor_name_raw, vendor_ckpt_raw in raw_line13_map.items():
                vendor_name = str(vendor_name_raw).strip()
                if not vendor_name:
                    continue
                vendor_ckpt_path_map = Path(str(vendor_ckpt_raw)).expanduser()
                if not vendor_ckpt_path_map.is_absolute():
                    candidate_from_map_dir = (line13_vendor_map_path.parent / vendor_ckpt_path_map).resolve()
                    candidate_from_cwd = (Path.cwd() / vendor_ckpt_path_map).resolve()
                    if candidate_from_map_dir.exists():
                        vendor_ckpt_path_map = candidate_from_map_dir
                    elif candidate_from_cwd.exists():
                        vendor_ckpt_path_map = candidate_from_cwd
                    else:
                        vendor_ckpt_path_map = candidate_from_map_dir
                if not vendor_ckpt_path_map.exists():
                    warnings.append(
                        "line13 vendor map: checkpoint non trovato per vendor "
                        f"'{vendor_name}' -> {vendor_ckpt_path_map}"
                    )
                    continue
                line13_vendor_ckpt_paths[vendor_name] = vendor_ckpt_path_map
                line13_vendor_map_loaded[vendor_name] = vendor_ckpt_path_map.as_posix()
        else:
            warnings.append(
                f"line13 vendor map non trovata: {line13_vendor_map_path}. "
                "Con routing vendor-specific attivo, #13 restera' in review per i vendor senza checkpoint dedicato."
            )
        missing_line13_vendors = sorted(set(vendor_classes) - set(line13_vendor_ckpt_paths))
        if missing_line13_vendors:
            warnings.append(
                "line13 vendor map incompleta: mancano checkpoint dedicati per "
                + ", ".join(missing_line13_vendors)
            )

    def _get_line13_vendor_model(vendor_name: str) -> Optional[Tuple[torch.nn.Module, int, Path]]:
        cached = line13_vendor_model_cache.get(vendor_name)
        if cached is not None:
            return cached
        ckpt_path = line13_vendor_ckpt_paths.get(vendor_name)
        if ckpt_path is None:
            return None
        try:
            vendor_line13_ckpt = torch.load(ckpt_path, map_location="cpu", weights_only=False)
            vendor_line13_image_size = int(args.line13_image_size) or int(
                vendor_line13_ckpt.get("args", {}).get("image_size", line13_image_size or rect_image_size)
            )
            vendor_line13_model = RectRegressor(pretrained=False).to(device)
            vendor_line13_model.load_state_dict(vendor_line13_ckpt["model_state_dict"])
            vendor_line13_model.eval()
        except Exception as exc:
            warnings.append(
                f"line13 vendor map: errore caricamento '{vendor_name}' ({ckpt_path}): {exc}"
            )
            return None
        out = (vendor_line13_model, vendor_line13_image_size, ckpt_path)
        line13_vendor_model_cache[vendor_name] = out
        return out

    start_time = time.time()
    for idx, folder in enumerate(folders, start=1):
        raw_images = _collect_acquisition_images(folder)
        if not raw_images:
            warnings.append(f"{folder.name}: nessuna immagine acquisizione trovata.")
            continue

        all_images, duplicates_removed = _deduplicate_exact_images(raw_images)
        if not all_images:
            warnings.append(f"{folder.name}: dedup ha rimosso tutte le immagini.")
            continue
        total_images_raw += len(raw_images)
        total_images_unique += len(all_images)
        total_duplicates_removed += duplicates_removed
        _emit_stage_event(
            "dedup",
            folder.name,
            folder_index=idx,
            folders_total=len(folders),
            folder_path=folder.as_posix(),
            images_raw=len(raw_images),
            images_unique=len(all_images),
            duplicates_removed=duplicates_removed,
        )

        rotation_deg_clockwise = 0
        rotation_source = "disabled"
        rotation_vote_ratio = 1.0 if args.disable_rotation_normalization else 0.0
        rotation_votes_total = 0
        rotation_samples_checked = 0
        rotation_osd_debug: Dict[str, object] = {
            "sample_limit": int(args.rotation_max_samples),
            "min_confidence": float(args.rotation_osd_min_confidence),
            "min_votes": int(args.rotation_min_votes),
            "min_ratio": float(args.rotation_min_ratio),
            "votes": {str(d): 0 for d in ROTATION_VALUES_CLOCKWISE},
            "samples": [],
            "best_rotation": 0,
            "best_count": 0,
            "ratio": 0.0,
        }
        rotation_ocr_debug: Dict[str, object] = {}
        rotation_decision_reason = "Rotazione non normalizzata: step disabilitato."
        if not args.disable_rotation_normalization:
            (
                rotation_deg_clockwise,
                rotation_source,
                rotation_vote_ratio,
                rotation_votes_total,
                rotation_samples_checked,
                tesseract_missing,
                rotation_osd_debug,
            ) = _estimate_folder_rotation_with_osd(
                image_paths=all_images,
                sample_limit=args.rotation_max_samples,
                timeout_sec=args.rotation_osd_timeout_sec,
                min_confidence=args.rotation_osd_min_confidence,
                min_votes=args.rotation_min_votes,
                min_ratio=args.rotation_min_ratio,
            )
            rotation_decision_reason = (
                f"OSD: best={rotation_osd_debug.get('best_rotation', 0)} "
                f"votes={rotation_osd_debug.get('best_count', 0)}/{rotation_votes_total} "
                f"(ratio={rotation_vote_ratio:.3f})."
            )
            if tesseract_missing:
                warnings.append(
                    f"{folder.name}: tesseract non disponibile, rotazione disattivata su questa cartella."
                )
                rotation_decision_reason = "Tesseract OSD non disponibile: rotazione forzata a 0°."
            elif (
                (not args.disable_rotation_ocr_validation)
                and rotation_source == "osd_majority"
                and int(rotation_deg_clockwise) != 0
            ):
                osd_candidate_rotation = int(rotation_deg_clockwise)
                (
                    rotation_deg_clockwise,
                    rotation_source,
                    tesseract_missing_ocr,
                    rotation_ocr_debug,
                ) = _refine_folder_rotation_with_ocr(
                    image_paths=all_images,
                    candidate_rotate_deg=int(rotation_deg_clockwise),
                    base_source=rotation_source,
                    sample_limit=args.rotation_ocr_samples,
                    timeout_sec=args.rotation_ocr_timeout_sec,
                    lang=args.rotation_ocr_lang,
                    psm=args.rotation_ocr_psm,
                    min_word_conf=args.rotation_ocr_min_word_conf,
                    score_word_bonus=args.rotation_ocr_score_word_bonus,
                    min_words=args.rotation_ocr_min_words,
                    min_score_delta=args.rotation_ocr_min_score_delta,
                )
                if tesseract_missing_ocr:
                    warnings.append(
                        f"{folder.name}: tesseract non disponibile durante validazione OCR rotazione."
                    )
                    rotation_decision_reason = (
                        f"OSD candidato={osd_candidate_rotation}°. "
                        "Validazione OCR non disponibile: mantengo candidato OSD."
                    )
                else:
                    ocr_selected = int(rotation_ocr_debug.get("selected_rotation", rotation_deg_clockwise))
                    ocr_delta = float(rotation_ocr_debug.get("delta_score", 0.0))
                    rotation_decision_reason = (
                        f"OSD candidato={osd_candidate_rotation}° -> OCR selezione={ocr_selected}° "
                        f"(delta_score={ocr_delta:.3f}, soglia_delta={args.rotation_ocr_min_score_delta:.3f}). "
                        f"Decision source={rotation_source}."
                    )
            elif rotation_source == "osd_majority":
                rotation_decision_reason = (
                    f"OSD majority accettata: rotazione={int(rotation_deg_clockwise)}° "
                    f"con ratio={rotation_vote_ratio:.3f} (soglia={args.rotation_min_ratio:.3f})."
                )
            elif rotation_source == "osd_low_support":
                rotation_decision_reason = (
                    f"OSD supporto insufficiente (ratio={rotation_vote_ratio:.3f}); "
                    "fallback a 0°."
                )
            elif rotation_source == "osd_no_votes":
                rotation_decision_reason = "OSD senza voti validi; fallback a 0°."
        rotation_source_counter[rotation_source] += 1
        rotation_degree_counter[int(rotation_deg_clockwise)] += 1

        cls_images = _select_uniform_subset(all_images, args.sample_per_folder)
        rect_images = _select_uniform_subset(all_images, args.rect_sample_per_folder)

        capture_input, capture_x, capture_y, capture_support = _infer_capture_metadata_majority(all_images)

        vendor_probs = _predict_mean_probs(
            model=vendor_model,
            image_paths=cls_images,
            image_size=vendor_image_size,
            batch_size=args.batch_size,
            device=device,
            rotate_deg_clockwise=rotation_deg_clockwise,
        )
        if vendor_probs.size == 0:
            raise RuntimeError(
                f"{folder.name}: inferenza vendor fallita (nessuna probabilita valida). "
                "Interrompo come richiesto."
            )
        vendor_sorted_idx = np.argsort(vendor_probs)[::-1]
        vendor_idx = int(vendor_sorted_idx[0])
        vendor_pred = vendor_classes[vendor_idx]
        vendor_conf = float(vendor_probs[vendor_idx])
        vendor_top2 = float(vendor_probs[int(vendor_sorted_idx[1])]) if len(vendor_sorted_idx) > 1 else 0.0
        vendor_margin = vendor_conf - vendor_top2
        vendor_top3 = _topk_from_probs(vendor_probs, vendor_classes, k=3)
        vendor_ocr_used = ""
        vendor_ocr_candidate = ""
        vendor_ocr_candidate_score = ""
        vendor_ocr_candidate_hits = ""
        vendor_ocr_delta_score = ""
        vendor_ocr_text = ""
        vendor_ocr_top3_json = ""
        vendor_ocr_samples_checked = ""
        vendor_ocr_reason = ""
        vendor_manual_resolved = False
        vendor_source = "cnn_classifier"
        vendor_decision_reason = (
            f"CNN top-1 sopra soglia ({vendor_conf:.4f} >= {args.vendor_min_confidence:.4f})."
        )
        vendor_low_margin = vendor_margin < float(args.vendor_min_margin)
        if vendor_conf < args.vendor_min_confidence or vendor_low_margin:
            vendor_topk = _topk_from_probs(vendor_probs, vendor_classes, k=args.interactive_topk)
            if vendor_conf < args.vendor_min_confidence:
                vendor_decision_reason = (
                    f"CNN top-1 sotto soglia ({vendor_conf:.4f} < {args.vendor_min_confidence:.4f})."
                )
            else:
                vendor_decision_reason = (
                    f"CNN margine top1-top2 sotto soglia "
                    f"({vendor_margin:.4f} < {float(args.vendor_min_margin):.4f})."
                )
            if not args.disable_vendor_ocr_fallback:
                ocr_diag = _resolve_vendor_with_ocr(
                    image_paths=all_images,
                    rotate_deg_clockwise=int(rotation_deg_clockwise),
                    vendor_alias_map=vendor_ocr_alias_map,
                    sample_limit=int(args.vendor_ocr_samples),
                    timeout_sec=float(args.vendor_ocr_timeout_sec),
                    lang=str(args.vendor_ocr_lang),
                    psm=int(args.vendor_ocr_psm),
                    min_word_conf=float(args.vendor_ocr_min_word_conf),
                    min_score=float(args.vendor_ocr_min_score),
                    min_score_delta=float(args.vendor_ocr_min_score_delta),
                    min_hits=int(args.vendor_ocr_min_hits),
                    cnn_vendor_hint=str(vendor_pred),
                )
                vendor_ocr_candidate = str(ocr_diag.get("candidate", "") or "")
                vendor_ocr_candidate_score = f"{float(ocr_diag.get('candidate_score', 0.0) or 0.0):.4f}"
                vendor_ocr_candidate_hits = str(int(ocr_diag.get("candidate_hits", 0) or 0))
                vendor_ocr_delta_score = f"{float(ocr_diag.get('delta_score', 0.0) or 0.0):.4f}"
                vendor_ocr_text = str(ocr_diag.get("text_excerpt", "") or "")
                vendor_ocr_top3_json = json.dumps(
                    ocr_diag.get("top3", []),
                    ensure_ascii=False,
                    separators=(",", ":"),
                )
                vendor_ocr_samples_checked = str(int(ocr_diag.get("samples_checked", 0) or 0))
                vendor_ocr_reason = str(ocr_diag.get("reason", "") or "")

                if bool(ocr_diag.get("tesseract_missing", False)):
                    vendor_ocr_used = "0"
                    warnings.append(
                        f"{folder.name}: tesseract non disponibile durante fallback OCR vendor."
                    )
                    vendor_ocr_reason = "tesseract_unavailable"
                elif bool(ocr_diag.get("accepted", False)) and vendor_ocr_candidate:
                    vendor_ocr_used = "1"
                    prev_vendor = str(vendor_pred)
                    prev_conf = float(vendor_conf)
                    vendor_pred = vendor_ocr_candidate
                    vendor_source = "ocr_fallback"
                    vendor_manual_resolved = True
                    ocr_conf_proxy = float(ocr_diag.get("confidence_proxy", 0.0) or 0.0)
                    vendor_conf = max(
                        float(vendor_conf),
                        float(args.vendor_min_confidence) + 0.02,
                        ocr_conf_proxy,
                    )
                    vendor_margin = max(float(vendor_margin), float(ocr_diag.get("delta_score", 0.0) or 0.0))
                    ocr_top3 = ocr_diag.get("top3", [])
                    if isinstance(ocr_top3, list):
                        ocr_top3_clean: List[Tuple[str, float]] = []
                        for item in ocr_top3[:3]:
                            if not isinstance(item, dict):
                                continue
                            lbl = str(item.get("label", "") or "").strip()
                            conf_val = float(item.get("confidence", 0.0) or 0.0)
                            if lbl:
                                ocr_top3_clean.append((lbl, conf_val))
                        if ocr_top3_clean:
                            vendor_top3 = ocr_top3_clean
                    vendor_decision_reason = (
                        f"CNN top-1 sotto soglia ({prev_conf:.4f} < {args.vendor_min_confidence:.4f}). "
                        f"OCR fallback accettato: {prev_vendor} -> {vendor_pred} "
                        f"(score={vendor_ocr_candidate_score}, delta={vendor_ocr_delta_score}, "
                        f"hits={vendor_ocr_candidate_hits})."
                    )
                else:
                    vendor_ocr_used = "0"
                    vendor_decision_reason += (
                        " OCR fallback non conclusivo "
                        f"(reason={vendor_ocr_reason}, candidate={vendor_ocr_candidate or '-'}, "
                        f"score={vendor_ocr_candidate_score or '-'}, delta={vendor_ocr_delta_score or '-'}, "
                        f"hits={vendor_ocr_candidate_hits or '-'})."
                    )

            if not vendor_manual_resolved:
                if args.low_confidence_policy == "error":
                    raise RuntimeError(
                        f"{folder.name}: vendor_conf={vendor_conf:.4f} (soglia {args.vendor_min_confidence:.4f}), "
                        f"vendor_margin={vendor_margin:.4f} (soglia {float(args.vendor_min_margin):.4f}). "
                        f"Top-k: {vendor_topk}"
                    )
                if args.low_confidence_policy == "ask_user":
                    selected_vendor, selection_source = _prompt_user_topk_or_custom(
                        field_name="vendor",
                        folder_name=folder.name,
                        options=vendor_topk,
                    )
                    vendor_pred = selected_vendor
                    if vendor_pred in vendor_classes:
                        vendor_conf = float(vendor_probs[vendor_classes.index(vendor_pred)])
                    else:
                        vendor_conf = 1.0
                    vendor_manual_resolved = True
                    if selection_source == "user_custom":
                        vendor_source = "manual_custom"
                        vendor_decision_reason = "Scelta manuale utente (custom) su vendor."
                    else:
                        vendor_source = "cnn_user_topk"
                        vendor_decision_reason = "Scelta utente da top-k CNN vendor."
                    if selection_source == "user_custom":
                        _register_manual_choice(
                            folder_name=folder.name,
                            field_name="vendor",
                            chosen_value=vendor_pred,
                            topk_options=vendor_topk,
                            threshold=args.vendor_min_confidence,
                            confidence=vendor_conf,
                        )
                # policy=review -> nessuna azione, rimane flag review
                if args.low_confidence_policy == "review":
                    vendor_source = "cnn_classifier"
                    vendor_decision_reason += " Mantengo output CNN e segnalo review."

        _emit_stage_event(
            "rotazione",
            folder.name,
            rotation_deg_clockwise=rotation_deg_clockwise,
            rotation_source=rotation_source,
            vote_ratio=rotation_vote_ratio,
        )
        _emit_stage_event(
            "vendor",
            folder.name,
            vendor=vendor_pred,
            confidence=vendor_conf,
            margin=vendor_margin,
            source=vendor_source,
            threshold=float(args.vendor_min_confidence),
            top3=[{"label": label, "prob": prob} for label, prob in vendor_top3],
        )

        # #13 RECT_NAME_ECHO: routing vendor-specific prima del probe.
        pre_video_x = capture_x
        pre_video_y = capture_y
        if pre_video_x is None or pre_video_y is None:
            for path in all_images:
                try:
                    with Image.open(path) as img:
                        w0, h0 = img.size
                    pre_video_x = pre_video_x if pre_video_x is not None else int(w0)
                    pre_video_y = pre_video_y if pre_video_y is not None else int(h0)
                    break
                except Exception:
                    continue

        vendor_seen_for_line13 = rect_name_echo_resolver.has_vendor(vendor_pred)
        line_13 = ""
        line_13_source = "unresolved"
        line_13_support = 0.0
        line13_topk = [
            (value, support)
            for value, support, _source in rect_name_echo_resolver.candidates(
                vendor=vendor_pred,
                video_x=pre_video_x,
                video_y=pre_video_y,
                top_k=args.interactive_topk,
            )
        ]
        line_13_model_checkpoint = ""
        line_13_model_route = "disabled"
        line_13_pre_dark_trim_rect_name_echo = ""
        line13_manual_resolved = False
        line13_unseen_vendor = False
        line13_missing_vendor_model = False
        line13_vendor_model_no_prediction = False
        line13_force_review_only = False
        line13_post_diag: Dict[str, object] = {
            "enabled": False,
            "mode": "model_disabled",
            "boxes_total": 0,
            "boxes_kept": 0,
            "boxes_dropped": 0,
        }
        line13_template_post_diag: Dict[str, object] = {
            "enabled": not args.disable_line13_template_postprocess,
            "attempted": False,
            "applied": False,
            "mode": "model_disabled",
        }
        line13_dark_trim_diag: Dict[str, object] = {
            "enabled": not args.disable_line13_dark_trim_postprocess,
            "attempted": False,
            "applied": False,
            "mode": "model_disabled",
        }

        if line13_model is None:
            line_13, line_13_source, line_13_support = rect_name_echo_resolver.resolve(
                vendor=vendor_pred,
                video_x=pre_video_x,
                video_y=pre_video_y,
            )
            if not vendor_seen_for_line13:
                # Non applichiamo fallback automatico "cieco" quando il vendor non esiste nello storico #13.
                line_13 = ""
                line_13_source = "unseen_vendor"
                line_13_support = 0.0
            line13_unseen_vendor = not vendor_seen_for_line13
            line13_postprocess_mode_counter["model_disabled"] += 1
            line13_template_postprocess_mode_counter["model_disabled"] += 1
            line13_dark_trim_mode_counter["model_disabled"] += 1
        else:
            line13_model_selected: Optional[torch.nn.Module] = None
            line13_model_image_size_selected = int(line13_image_size) or int(rect_image_size)
            if args.disable_line13_vendor_routing:
                line13_model_selected = line13_model
                line_13_model_checkpoint = line13_ckpt_path.as_posix()
                line_13_model_route = "global_routing_disabled"
            else:
                vendor_line13_entry = _get_line13_vendor_model(vendor_pred) if vendor_pred else None
                if vendor_line13_entry is not None:
                    line13_model_selected, line13_model_image_size_selected, line13_vendor_ckpt_path = vendor_line13_entry
                    line_13_model_checkpoint = line13_vendor_ckpt_path.as_posix()
                    line_13_model_route = "vendor_specialized"
                    if vendor_pred:
                        line13_vendor_specialized_usage[vendor_pred] += 1
                else:
                    line13_missing_vendor_model = True
                    line13_force_review_only = True
                    line_13_model_route = "vendor_model_missing"
                    line_13_source = "vendor_model_missing"
            line13_route_counter[line_13_model_route] += 1

            if line13_model_selected is not None:
                line13_boxes_abs, line13_sizes = _predict_rect_boxes_abs(
                    model=line13_model_selected,
                    image_paths=rect_images,
                    image_size=line13_model_image_size_selected,
                    batch_size=args.batch_size,
                    device=device,
                    rotate_deg_clockwise=rotation_deg_clockwise,
                )
                if args.disable_line13_postprocess:
                    line13_box_model = _median_box(line13_boxes_abs)
                    line13_post_diag = {
                        "enabled": False,
                        "mode": "disabled",
                        "boxes_total": int(len(line13_boxes_abs)),
                        "boxes_kept": int(len(line13_boxes_abs)),
                        "boxes_dropped": 0,
                    }
                else:
                    line13_box_model, line13_post_diag = _postprocess_line13_box(
                        boxes_abs=line13_boxes_abs,
                        image_sizes=line13_sizes,
                        iou_threshold=float(args.line13_postprocess_iou_threshold),
                        min_keep=int(args.line13_postprocess_min_keep),
                    )

                line13_post_mode = str(line13_post_diag.get("mode", "unknown") or "unknown")
                line13_postprocess_mode_counter[line13_post_mode] += 1
                line13_postprocess_total_boxes += int(line13_post_diag.get("boxes_total", 0) or 0)
                line13_postprocess_total_kept += int(line13_post_diag.get("boxes_kept", 0) or 0)
                line13_postprocess_total_dropped += int(line13_post_diag.get("boxes_dropped", 0) or 0)

                line13_box_final = line13_box_model
                if line13_box_model is not None:
                    if args.disable_line13_template_postprocess:
                        line13_template_post_diag = {
                            "enabled": False,
                            "attempted": False,
                            "applied": False,
                            "mode": "disabled",
                        }
                    else:
                        line13_box_final, line13_template_post_diag = _postprocess_line13_box_with_template(
                            folder=folder,
                            image_paths=rect_images,
                            model_box=line13_box_model,
                            search_margin_pct=float(args.line13_template_search_margin_pct),
                            max_images=int(args.line13_template_max_images),
                            min_template_std=float(args.line13_template_min_std),
                            min_template_nonblack_ratio=float(args.line13_template_min_nonblack_ratio),
                            min_score=float(args.line13_template_min_score),
                            min_consensus_iou=float(args.line13_template_min_consensus_iou),
                            min_iou_with_model=float(args.line13_template_min_iou_with_model),
                        )
                else:
                    line13_template_post_diag = {
                        "enabled": not args.disable_line13_template_postprocess,
                        "attempted": False,
                        "applied": False,
                        "mode": "skip_no_model_box",
                    }

                line13_template_mode = str(line13_template_post_diag.get("mode", "unknown") or "unknown")
                line13_template_postprocess_mode_counter[line13_template_mode] += 1
                if bool(line13_template_post_diag.get("attempted", False)):
                    line13_template_postprocess_attempted_total += 1
                if bool(line13_template_post_diag.get("applied", False)):
                    line13_template_postprocess_applied_total += 1

                if line13_box_final is not None:
                    x1_pre, y1_pre, x2_pre, y2_pre = line13_box_final
                    line_13_pre_dark_trim_rect_name_echo = _format_line13_with_disabled_tail(
                        int(round(y1_pre)),
                        int(round(x1_pre)),
                        int(round(y2_pre)),
                        int(round(x2_pre)),
                    )
                    if args.disable_line13_dark_trim_postprocess:
                        line13_dark_trim_diag = {
                            "enabled": False,
                            "attempted": False,
                            "applied": False,
                            "mode": "disabled",
                        }
                    else:
                        line13_box_final, line13_dark_trim_diag = _postprocess_line13_box_with_dark_trim(
                            image_paths=rect_images,
                            input_box=line13_box_final,
                            max_images=int(args.line13_template_max_images),
                            dark_threshold=float(args.line13_dark_trim_threshold),
                            max_std=float(args.line13_dark_trim_max_std),
                            max_trim_pct=float(args.line13_dark_trim_max_trim_pct),
                            min_keep_ratio=float(args.line13_dark_trim_min_keep_ratio),
                        )
                else:
                    line13_dark_trim_diag = {
                        "enabled": not args.disable_line13_dark_trim_postprocess,
                        "attempted": False,
                        "applied": False,
                        "mode": "skip_no_input_box",
                    }

                line13_dark_trim_mode = str(line13_dark_trim_diag.get("mode", "unknown") or "unknown")
                line13_dark_trim_mode_counter[line13_dark_trim_mode] += 1
                if bool(line13_dark_trim_diag.get("attempted", False)):
                    line13_dark_trim_attempted_total += 1
                if bool(line13_dark_trim_diag.get("applied", False)):
                    line13_dark_trim_applied_total += 1

                if line13_box_final is not None:
                    x1_m, y1_m, x2_m, y2_m = line13_box_final
                    top_m = str(int(round(y1_m)))
                    left_m = str(int(round(x1_m)))
                    bottom_m = str(int(round(y2_m)))
                    right_m = str(int(round(x2_m)))
                    line_13 = _format_line13_with_disabled_tail(
                        int(top_m),
                        int(left_m),
                        int(bottom_m),
                        int(right_m),
                    )
                    line_13_source = (
                        f"vendor_template_model_{line_13_model_route}"
                        f"_postproc_{line13_post_mode}"
                        f"_template_{line13_template_mode}"
                        f"_darktrim_{line13_dark_trim_mode}"
                    )
                    line_13_support = 1.0
                else:
                    line13_vendor_model_no_prediction = True
                    if not args.disable_line13_vendor_routing:
                        line13_force_review_only = True
                    line_13 = ""
                    line_13_source = f"vendor_template_model_{line_13_model_route}_no_prediction"
                    line_13_support = 0.0
                    line_13_model_route = f"{line_13_model_route}_no_prediction"
                    line13_route_counter[line_13_model_route] += 1
            else:
                line13_postprocess_mode_counter["model_not_run"] += 1
                line13_template_postprocess_mode_counter["model_not_run"] += 1
                line13_dark_trim_mode_counter["model_not_run"] += 1

        if (not line_13) or (line_13_support < args.line13_min_support):
            if line13_missing_vendor_model:
                reason = "manca modello vendor-specific per #13"
            elif line13_vendor_model_no_prediction:
                reason = "rete vendor-specific #13 senza predizione valida"
            else:
                reason = "vendor mai visto in storico #13" if line13_unseen_vendor else "support basso/non risolto"
            if line13_force_review_only:
                pass
            elif args.low_confidence_policy == "error":
                raise RuntimeError(
                    f"{folder.name}: #13 RECT_NAME_ECHO non affidabile ({reason}) "
                    f"(support={line_13_support:.4f}, soglia={args.line13_min_support:.4f}). "
                    f"Top-k: {line13_topk}"
                )
            elif args.low_confidence_policy == "ask_user":
                selected_line13, selection_source = _prompt_user_topk_or_custom(
                    field_name="line_13_rect_name_echo",
                    folder_name=folder.name,
                    options=line13_topk,
                )
                fallback_template = line13_topk[0][0] if line13_topk else ""
                selected_line13 = _normalize_rect_name_user_value(selected_line13, fallback_template)
                line_13 = selected_line13
                line_13_source = selection_source
                option_map = {value: support for value, support in line13_topk}
                line_13_support = float(option_map.get(line_13, 1.0 if selection_source == "user_custom" else 0.0))
                line13_manual_resolved = True
                if selection_source == "user_custom":
                    _register_manual_choice(
                        folder_name=folder.name,
                        field_name="line_13_rect_name_echo",
                        chosen_value=line_13,
                        topk_options=line13_topk,
                        threshold=args.line13_min_support,
                        confidence=line_13_support,
                    )

        # Per linea #13 usiamo solo le coordinate rettangolo.
        # La parte dopo le coordinate rimane disabilitata finché non avremo un calcolo dedicato.
        line_13 = _force_line13_disabled_tail(line_13)

        line13_top = line13_left = line13_bottom = line13_right = ""
        line13_coords = _parse_rect_name_echo_coords(line_13) if line_13 else None
        if line13_coords is not None:
            line13_top, line13_left, line13_bottom, line13_right = (
                str(line13_coords[0]),
                str(line13_coords[1]),
                str(line13_coords[2]),
                str(line13_coords[3]),
            )

        probe_probs = _predict_mean_probs(
            model=probe_model,
            image_paths=cls_images,
            image_size=probe_image_size,
            batch_size=args.batch_size,
            device=device,
            rotate_deg_clockwise=rotation_deg_clockwise,
        )
        probe_id = ""
        probe_conf = 0.0
        probe_manual_resolved = False
        probe_top3: List[Tuple[str, float]] = []
        probe_source = "cnn_classifier"
        probe_decision_reason = "Classificatore probe senza output valido."
        if probe_probs.size > 0:
            probe_idx = int(np.argmax(probe_probs))
            probe_id = probe_classes[probe_idx]
            probe_conf = float(probe_probs[probe_idx])
            probe_top3 = _topk_from_probs(probe_probs, probe_classes, k=3)
            probe_decision_reason = (
                f"CNN top-1 probe sopra soglia ({probe_conf:.4f} >= {args.probe_min_confidence:.4f})."
            )
        if probe_conf < args.probe_min_confidence:
            probe_topk = _topk_from_probs(probe_probs, probe_classes, k=args.interactive_topk)
            probe_decision_reason = (
                f"CNN top-1 probe sotto soglia ({probe_conf:.4f} < {args.probe_min_confidence:.4f})."
            )
            if args.low_confidence_policy == "error":
                raise RuntimeError(
                    f"{folder.name}: probe_conf={probe_conf:.4f} < {args.probe_min_confidence:.4f}. "
                    f"Top-k: {probe_topk}"
                )
            if args.low_confidence_policy == "ask_user":
                selected_probe, selection_source = _prompt_user_topk_or_custom(
                    field_name="fss_id_probe",
                    folder_name=folder.name,
                    options=probe_topk,
                )
                probe_id = selected_probe
                if probe_id in probe_classes:
                    probe_conf = float(probe_probs[probe_classes.index(probe_id)])
                else:
                    probe_conf = 1.0
                probe_manual_resolved = True
                if selection_source == "user_custom":
                    probe_source = "manual_custom"
                    probe_decision_reason = "Scelta manuale utente (custom) su probe."
                else:
                    probe_source = "cnn_user_topk"
                    probe_decision_reason = "Scelta utente da top-k CNN probe."
                if selection_source == "user_custom":
                    _register_manual_choice(
                        folder_name=folder.name,
                        field_name="fss_id_probe",
                        chosen_value=probe_id,
                        topk_options=probe_topk,
                        threshold=args.probe_min_confidence,
                        confidence=probe_conf,
                    )
            # policy=review -> nessuna azione, rimane flag review
            if args.low_confidence_policy == "review":
                probe_source = "cnn_classifier"
                probe_decision_reason += " Mantengo output CNN e segnalo review."

        _emit_stage_event(
            "probe",
            folder.name,
            probe_id=probe_id,
            confidence=probe_conf,
            source=probe_source,
            threshold=float(args.probe_min_confidence),
            top3=[{"label": label, "prob": prob} for label, prob in probe_top3],
        )

        line_04_probe_type = ""
        line_04_probe_type_source = "router_disabled" if args.disable_probe_type_router else "router_unavailable"
        line_04_probe_type_strategy = "not_computed"
        line_04_probe_type_needs_secondary_model = ""
        line_04_probe_type_candidate_ints = ""
        if probe_type_router is not None:
            decision = probe_type_router.resolve(probe_id)
            line_04_probe_type = str(decision.type_value or "").strip()
            line_04_probe_type_source = str(decision.source or "").strip() or "none"
            line_04_probe_type_strategy = str(decision.strategy or "").strip() or "unknown"
            line_04_probe_type_needs_secondary_model = "1" if bool(decision.needs_secondary_model) else "0"
            line_04_probe_type_candidate_ints = ",".join(str(x) for x in decision.candidate_type_ints)

        # #14 RECT_NAME_PROBE: resolver storico dipendente da probe.
        probe_seen_for_line14 = rect_name_probe_resolver.has_probe(probe_id)
        line_14, line_14_source, line_14_support = rect_name_probe_resolver.resolve(
            vendor=vendor_pred,
            probe_id=probe_id if probe_id else None,
            video_x=pre_video_x,
            video_y=pre_video_y,
        )
        if not probe_seen_for_line14:
            # Non applichiamo fallback automatico "cieco" quando la probe non esiste nello storico #14.
            line_14 = ""
            line_14_source = "unseen_probe"
            line_14_support = 0.0
        line14_topk = [
            (value, support)
            for value, support, _source in rect_name_probe_resolver.candidates(
                vendor=vendor_pred,
                probe_id=probe_id if probe_id else None,
                video_x=pre_video_x,
                video_y=pre_video_y,
                top_k=args.interactive_topk,
            )
        ]
        line14_manual_resolved = False
        line14_unseen_probe = not probe_seen_for_line14
        if (not line_14) or (line_14_support < args.line14_min_support):
            reason = "probe mai vista in storico #14" if line14_unseen_probe else "support basso/non risolto"
            if args.low_confidence_policy == "error":
                raise RuntimeError(
                    f"{folder.name}: #14 RECT_NAME_PROBE non affidabile ({reason}) "
                    f"(support={line_14_support:.4f}, soglia={args.line14_min_support:.4f}). "
                    f"Top-k: {line14_topk}"
                )
            if args.low_confidence_policy == "ask_user":
                selected_line14, selection_source = _prompt_user_topk_or_custom(
                    field_name="line_14_rect_name_probe",
                    folder_name=folder.name,
                    options=line14_topk,
                )
                fallback_template = line14_topk[0][0] if line14_topk else ""
                selected_line14 = _normalize_rect_name_user_value(selected_line14, fallback_template)
                line_14 = selected_line14
                line_14_source = selection_source
                option_map = {value: support for value, support in line14_topk}
                line_14_support = float(option_map.get(line_14, 1.0 if selection_source == "user_custom" else 0.0))
                line14_manual_resolved = True
                if selection_source == "user_custom":
                    _register_manual_choice(
                        folder_name=folder.name,
                        field_name="line_14_rect_name_probe",
                        chosen_value=line_14,
                        topk_options=line14_topk,
                        threshold=args.line14_min_support,
                        confidence=line_14_support,
                    )

        line14_top = line14_left = line14_bottom = line14_right = ""
        line14_coords = _parse_rect_name_echo_coords(line_14) if line_14 else None
        if line14_coords is not None:
            line14_top, line14_left, line14_bottom, line14_right = (
                str(line14_coords[0]),
                str(line14_coords[1]),
                str(line14_coords[2]),
                str(line14_coords[3]),
            )

        rect_model_selected = rect_model
        rect_image_size_selected = rect_image_size
        rect_source = "global"
        rect_model_checkpoint = rect_ckpt_path.as_posix()
        specialized_entry = rect_vendor_models.get(vendor_pred)
        if specialized_entry is not None:
            if vendor_conf >= args.rect_vendor_min_confidence:
                rect_model_selected, rect_image_size_selected, specialized_ckpt = specialized_entry
                rect_source = "vendor_specialized"
                rect_model_checkpoint = specialized_ckpt.as_posix()
                rect_vendor_specialized_usage[vendor_pred] += 1
            else:
                rect_source = "global_low_vendor_conf"
        rect_route_counter[rect_source] += 1

        rect_boxes_abs, rect_sizes = _predict_rect_boxes_abs(
            model=rect_model_selected,
            image_paths=rect_images,
            image_size=rect_image_size_selected,
            batch_size=args.batch_size,
            device=device,
            rotate_deg_clockwise=rotation_deg_clockwise,
        )
        rect_box = _median_box(rect_boxes_abs)

        # Fallback to observed most common image size if filename metadata is unavailable.
        out_video_x = capture_x
        out_video_y = capture_y
        if (out_video_x is None or out_video_y is None) and rect_sizes:
            size_hist = Counter(rect_sizes)
            (w0, h0), _ = size_hist.most_common(1)[0]
            out_video_x = out_video_x if out_video_x is not None else w0
            out_video_y = out_video_y if out_video_y is not None else h0

        rect_top = rect_left = rect_bottom = rect_right = ""
        line_11 = ""
        line_11_method = "median_rect_fallback"
        line_11_rect_red_margin_pct = float(args.rect_red_margin_pct)
        line_11_rect_red_winner_group = ""
        rect_red_payload: Dict[str, object] = {
            "available": False,
            "error": "not_computed",
            "margin_pct": float(args.rect_red_margin_pct),
            "bright_thr": float(args.rect_red_bright_thr),
        }
        if rect_box is not None:
            x1, y1, x2, y2 = rect_box
            rect_left = str(int(round(x1)))
            rect_top = str(int(round(y1)))
            rect_right = str(int(round(x2)))
            rect_bottom = str(int(round(y2)))
            line_11 = f"{rect_top}|{rect_left}|{rect_bottom}|{rect_right}|"
        line_11_median_text = line_11
        line_11_median_top = rect_top
        line_11_median_left = rect_left
        line_11_median_bottom = rect_bottom
        line_11_median_right = rect_right
        _emit_stage_event(
            "rect",
            folder.name,
            provisional=True,
            line_11=line_11,
            method="median_rect_fallback",
            source=rect_source,
            model_checkpoint=rect_model_checkpoint,
            images_used=len(rect_boxes_abs),
        )

        su_giu_images_predicted = 0
        su_giu_majority_label = ""
        su_giu_majority_vote_ratio = 0.0
        su_giu_mean_confidence = 0.0
        su_giu_mean_prob_su = 0.0
        su_giu_mean_prob_giu = 0.0
        su_giu_source = "disabled"
        su_giu_rows: List[Dict[str, object]] = []
        su_giu_checkpoint = su_giu_ckpt_path.as_posix() if su_giu_model is not None else ""
        if su_giu_model is not None:
            su_giu_rows = _predict_su_giu_on_rect_crops(
                model=su_giu_model,
                class_names=su_giu_class_names,
                image_paths=rect_images,
                rect_boxes_abs=rect_boxes_abs,
                fallback_rect_box_abs=rect_box,
                image_size=su_giu_image_size,
                batch_size=su_giu_batch_size,
                device=device,
                rotate_deg_clockwise=rotation_deg_clockwise,
            )
            su_giu_images_predicted = len(su_giu_rows)
            if su_giu_rows:
                pred_counter: Counter[str] = Counter()
                used_global_fallback = False
                conf_values: List[float] = []
                prob_su_values: List[float] = []
                prob_giu_values: List[float] = []
                for row_sg in su_giu_rows:
                    pred_label = str(row_sg.get("pred_label", "")).strip().lower()
                    if pred_label:
                        pred_counter[pred_label] += 1
                    if str(row_sg.get("crop_source", "")) == "global_rect_fallback":
                        used_global_fallback = True
                    conf_values.append(float(row_sg.get("confidence", 0.0) or 0.0))
                    prob_su_values.append(float(row_sg.get("prob_su", 0.0) or 0.0))
                    prob_giu_values.append(float(row_sg.get("prob_giu", 0.0) or 0.0))
                    su_giu_per_image_rows.append(
                        {
                            "folder_path": folder.as_posix(),
                            "folder_name": folder.name,
                            "image_index": int(row_sg.get("image_index", 0) or 0),
                            "image_path": str(row_sg.get("image_path", "") or ""),
                            "pred_label": pred_label,
                            "pred_idx": int(row_sg.get("pred_idx", 0) or 0),
                            "confidence": float(row_sg.get("confidence", 0.0) or 0.0),
                            "prob_su": float(row_sg.get("prob_su", 0.0) or 0.0),
                            "prob_giu": float(row_sg.get("prob_giu", 0.0) or 0.0),
                            "crop_top": int(row_sg.get("crop_top", 0) or 0),
                            "crop_left": int(row_sg.get("crop_left", 0) or 0),
                            "crop_bottom": int(row_sg.get("crop_bottom", 0) or 0),
                            "crop_right": int(row_sg.get("crop_right", 0) or 0),
                            "crop_source": str(row_sg.get("crop_source", "") or ""),
                        }
                    )
                if pred_counter:
                    su_giu_majority_label, majority_count = max(
                        pred_counter.items(),
                        key=lambda item: (item[1], item[0] == "su", item[0] == "giu"),
                    )
                    su_giu_majority_vote_ratio = float(majority_count / max(1, su_giu_images_predicted))
                    su_giu_majority_counter[su_giu_majority_label] += 1
                su_giu_mean_confidence = float(sum(conf_values) / max(1, len(conf_values)))
                su_giu_mean_prob_su = float(sum(prob_su_values) / max(1, len(prob_su_values)))
                su_giu_mean_prob_giu = float(sum(prob_giu_values) / max(1, len(prob_giu_values)))
                su_giu_source = "per_image_rect_with_global_fallback" if used_global_fallback else "per_image_rect"
            else:
                su_giu_source = "no_predictions"
            su_giu_source_counter[su_giu_source] += 1
        else:
            su_giu_source_counter[su_giu_source] += 1

        lr_marker_images_predicted = 0
        lr_marker_majority_label = ""
        lr_marker_majority_vote_ratio = 0.0
        lr_marker_best_label = ""
        lr_marker_best_score = 0.0
        lr_marker_best_template_path = ""
        lr_marker_best_search_strategy = ""
        lr_marker_source = "disabled"
        line_16_rect_orientation = ""
        line_16_source = "lr_marker_unavailable"
        line_16_marker_boxes_count = 0
        line_16_groups_json = "{}"
        lr_marker_rows: List[Dict[str, object]] = []
        if bool(args.disable_lr_marker_classical):
            lr_marker_source = "disabled"
        elif not su_giu_rows:
            lr_marker_source = "missing_su_giu_predictions"
        elif not vendor_pred:
            lr_marker_source = "missing_vendor"
        elif str(args.lr_marker_method) == "bundle":
            def _bundle_rows_choice_score(rows: Sequence[Dict[str, object]]) -> Tuple[int, int, float]:
                reliable = [
                    row for row in _lr_marker_reliable_rows(
                        rows,
                        min_match_score=float(args.lr_marker_min_match_score),
                        target_image_width=int(out_video_x or 0),
                        target_image_height=int(out_video_y or 0),
                    )
                    if str(row.get("status", "") or "") == "ok"
                ]
                groups = {
                    _infer_lr_marker_orientation_group_from_row(dict(row))
                    for row in reliable
                }
                groups.discard("")
                best = max((float(row.get("match_score", 0.0) or 0.0) for row in reliable), default=0.0)
                return (len(groups), len(reliable), best)

            lr_marker_rows, bundle_meta = _bundle_predict_lr_marker_on_su_giu_rows(
                su_giu_rows=su_giu_rows,
                vendor_name=vendor_pred,
                library_root=args.lr_marker_bundle_library_root.expanduser().resolve(),
                bundle_dir=args.lr_marker_bundle_dir.expanduser().resolve(),
                bundle_zip=args.lr_marker_bundle_zip.expanduser().resolve(),
                min_match_score=float(args.lr_marker_min_match_score),
                vertical_delta=float(args.lr_marker_bundle_vertical_delta),
                full_crop_fallback_threshold=float(args.lr_marker_full_crop_fallback_threshold),
                expanded_search_threshold=float(args.lr_marker_expanded_search_threshold),
                expanded_search_steps=lr_marker_expanded_search_steps,
                match_max_side=int(args.lr_marker_bundle_match_max_side),
                selection_images=int(args.lr_marker_bundle_selection_images),
                target_image_width=int(out_video_x or 0),
                target_image_height=int(out_video_y or 0),
                template_policy="bundle_historical_best",
                scales=lr_marker_scales,
                pinned_templates=lr_marker_pinned_map.get(folder.name, []),
            )
            historical_bundle_meta = dict(bundle_meta)
            bundle_policy_effective = "bundle_historical_best"
            manual_bundle_library_root, manual_bundle_meta = _build_lr_marker_bundle_manual_seed_library(
                seeds=lr_marker_manual_seed_rows,
                su_giu_rows=su_giu_rows,
                folder=folder,
                output_dir=output_dir,
                vendor_name=vendor_pred,
                blank_template_max_value=int(args.lr_marker_blank_template_max_value),
            )
            if manual_bundle_meta:
                lr_marker_manual_seed_template_meta.extend(manual_bundle_meta)
            if manual_bundle_library_root is not None:
                expected_by_path = {
                    str(meta.get("image_path", "") or ""): str(meta.get("correction", "") or "").strip().upper()
                    for meta in manual_bundle_meta
                    if str(meta.get("status", "") or "") == "accepted"
                    and str(meta.get("image_path", "") or "")
                    and str(meta.get("correction", "") or "").strip().upper() in {"NF", "LR", "UD", "LRUD"}
                }
                expected_by_index = {
                    str(meta.get("image_index", "") or ""): str(meta.get("correction", "") or "").strip().upper()
                    for meta in manual_bundle_meta
                    if str(meta.get("status", "") or "") == "accepted"
                    and str(meta.get("image_index", "") or "")
                    and str(meta.get("correction", "") or "").strip().upper() in {"NF", "LR", "UD", "LRUD"}
                }
                manual_lr_marker_rows, manual_bundle_run_meta = _bundle_predict_lr_marker_on_su_giu_rows(
                    su_giu_rows=su_giu_rows,
                    vendor_name=vendor_pred,
                    library_root=manual_bundle_library_root,
                    bundle_dir=args.lr_marker_bundle_dir.expanduser().resolve(),
                    bundle_zip=args.lr_marker_bundle_zip.expanduser().resolve(),
                    min_match_score=float(args.lr_marker_min_match_score),
                    vertical_delta=float(args.lr_marker_bundle_vertical_delta),
                    full_crop_fallback_threshold=float(args.lr_marker_full_crop_fallback_threshold),
                    expanded_search_threshold=float(args.lr_marker_expanded_search_threshold),
                    expanded_search_steps=lr_marker_expanded_search_steps,
                    match_max_side=int(args.lr_marker_bundle_match_max_side),
                    selection_images=max(int(args.lr_marker_bundle_selection_images), len(expected_by_path), 1),
                    target_image_width=int(out_video_x or 0),
                    target_image_height=int(out_video_y or 0),
                    expected_groups_by_image_path=expected_by_path,
                    expected_groups_by_image_index=expected_by_index,
                    template_policy="bundle_manual_seed",
                )
                accepted_manual_metas = [
                    meta for meta in manual_bundle_meta
                    if str(meta.get("status", "") or "") == "accepted"
                    and str(meta.get("template_path", "") or "")
                ]
                manual_ensemble_candidates: List[Dict[str, object]] = []
                safe_vendor = re.sub(r"[^A-Za-z0-9_.-]+", "_", vendor_pred.strip() or "vendor").strip("_") or "vendor"
                single_root_base = output_dir / "lr_marker_bundle_manual_seed_single_template_libraries"
                for meta in accepted_manual_metas:
                    try:
                        seed_index = int(meta.get("seed_index", len(manual_ensemble_candidates) + 1) or len(manual_ensemble_candidates) + 1)
                    except (TypeError, ValueError):
                        seed_index = len(manual_ensemble_candidates) + 1
                    template_path = Path(str(meta.get("template_path", "") or "")).expanduser().resolve()
                    if not template_path.is_file():
                        continue
                    single_root = single_root_base / f"seed_{seed_index:03d}"
                    single_vendor_dir = single_root / safe_vendor
                    single_template_path = single_vendor_dir / "marker_001.png"
                    try:
                        single_vendor_dir.mkdir(parents=True, exist_ok=True)
                        single_template_path.write_bytes(template_path.read_bytes())
                        (single_root / "review_decisions.json").write_text(
                            json.dumps(
                                {
                                    "source": "pipeline_single_manual_lr_marker_seed",
                                    "seed_index": seed_index,
                                    "vendors": {
                                        safe_vendor: {
                                            "vendor": safe_vendor,
                                            "accepted": [f"{safe_vendor}/{single_template_path.name}"],
                                            "rejected": [],
                                        }
                                    },
                                },
                                ensure_ascii=False,
                                indent=2,
                            ),
                            encoding="utf-8",
                        )
                    except Exception:
                        continue
                    seed_rows, _seed_bundle_meta = _bundle_predict_lr_marker_on_su_giu_rows(
                        su_giu_rows=su_giu_rows,
                        vendor_name=vendor_pred,
                        library_root=single_root,
                        bundle_dir=args.lr_marker_bundle_dir.expanduser().resolve(),
                        bundle_zip=args.lr_marker_bundle_zip.expanduser().resolve(),
                        min_match_score=float(args.lr_marker_min_match_score),
                        vertical_delta=float(args.lr_marker_bundle_vertical_delta),
                        full_crop_fallback_threshold=float(args.lr_marker_full_crop_fallback_threshold),
                        expanded_search_threshold=float(args.lr_marker_expanded_search_threshold),
                        expanded_search_steps=lr_marker_expanded_search_steps,
                        match_max_side=int(args.lr_marker_bundle_match_max_side),
                        selection_images=1,
                        target_image_width=int(out_video_x or 0),
                        target_image_height=int(out_video_y or 0),
                        expected_groups_by_image_path=expected_by_path,
                        expected_groups_by_image_index=expected_by_index,
                        template_policy="bundle_manual_seed_ensemble",
                    )
                    for row in seed_rows:
                        row["manual_seed_index"] = seed_index
                        row["folder_template_seed_image_path"] = str(meta.get("image_path", "") or "")
                        row["folder_template_seed_db_template_path"] = str(meta.get("template_path", "") or "")
                    manual_ensemble_candidates.extend(seed_rows)
                if manual_ensemble_candidates:
                    def _manual_ensemble_rank(row: Dict[str, object]) -> Tuple[int, int, int, float, float]:
                        image_path_s = str(row.get("image_path", "") or "")
                        image_index_s = str(row.get("image_index", "") or "")
                        expected_group = expected_by_path.get(image_path_s, "") or expected_by_index.get(image_index_s, "")
                        group = str(row.get("orientation_group", "") or "").strip().upper()
                        expected_hit = int(1 if expected_group and group == expected_group else 0)
                        ok = int(1 if str(row.get("status", "") or "") == "ok" else 0)
                        inside = int(0 if _lr_marker_truthy_flag(row.get("marker_outside_echo_rect", 0)) else 1)
                        try:
                            score = float(row.get("match_score", 0.0) or 0.0)
                        except (TypeError, ValueError):
                            score = 0.0
                        try:
                            overlap = float(row.get("marker_echo_overlap_ratio", 0.0) or 0.0)
                        except (TypeError, ValueError):
                            overlap = 0.0
                        return (expected_hit, ok, inside, score, overlap)

                    combined_rows_by_key: Dict[str, Dict[str, object]] = {}
                    for row in manual_ensemble_candidates:
                        key = str(row.get("image_index", "") or row.get("image_path", "") or "")
                        if not key:
                            continue
                        current = combined_rows_by_key.get(key)
                        if current is None or _manual_ensemble_rank(dict(row)) > _manual_ensemble_rank(dict(current)):
                            combined_rows_by_key[key] = dict(row)
                    manual_lr_marker_rows = sorted(
                        combined_rows_by_key.values(),
                        key=lambda row: int(row.get("image_index", 0) or 0),
                    )
                    manual_bundle_run_meta = dict(manual_bundle_run_meta)
                    manual_bundle_run_meta["template_policy"] = "bundle_manual_seed_ensemble"
                    manual_bundle_run_meta["manual_seed_ensemble_templates"] = len(accepted_manual_metas)
                    manual_bundle_run_meta["manual_seed_ensemble_candidates"] = len(manual_ensemble_candidates)
                if manual_lr_marker_rows and _bundle_rows_choice_score(manual_lr_marker_rows) > _bundle_rows_choice_score(lr_marker_rows):
                    chosen_bundle_meta = dict(manual_bundle_run_meta)
                    chosen_bundle_meta["historical_bundle_meta"] = historical_bundle_meta
                    chosen_bundle_meta["manual_seed_template_count"] = len(expected_by_path)
                    bundle_meta = chosen_bundle_meta
                    lr_marker_rows = manual_lr_marker_rows
                    bundle_policy_effective = str(bundle_meta.get("template_policy", "") or "bundle_manual_seed")
            if not lr_marker_rows:
                lr_marker_source = str(bundle_meta.get("error", "") or "bundle_no_marker_predictions")
            else:
                lr_marker_images_predicted = len(lr_marker_rows)
                for row_lr in lr_marker_rows:
                    row_lr["template_policy_requested"] = "bundle"
                    row_lr["template_policy_effective"] = bundle_policy_effective
                    row_lr["template_fallback_reason"] = ""
                    lr_marker_per_image_rows.append(
                        {
                            "folder_path": folder.as_posix(),
                            "folder_name": folder.name,
                            **row_lr,
                        }
                    )
                decision_lr_rows = [
                    row for row in _lr_marker_reliable_rows(
                        lr_marker_rows,
                        min_match_score=float(args.lr_marker_min_match_score),
                        target_image_width=int(out_video_x or 0),
                        target_image_height=int(out_video_y or 0),
                    )
                    if str(row.get("status", "") or "") == "ok"
                ]
                label_counter: Counter[str] = Counter()
                for row_lr in decision_lr_rows:
                    label = str(row_lr.get("lr_label", "") or "").strip()
                    if label:
                        label_counter[label] += 1
                if label_counter:
                    lr_marker_majority_label, majority_count = max(
                        label_counter.items(),
                        key=lambda item: (item[1], item[0] == "not_lr_flipped", item[0]),
                    )
                    lr_marker_majority_vote_ratio = float(majority_count / max(1, len(decision_lr_rows)))
                    lr_marker_majority_counter[lr_marker_majority_label] += 1
                if decision_lr_rows:
                    best_lr = max(decision_lr_rows, key=lambda item: float(item.get("match_score", 0.0) or 0.0))
                    lr_marker_best_label = str(best_lr.get("lr_label", "") or "")
                    lr_marker_best_score = float(best_lr.get("match_score", 0.0) or 0.0)
                    lr_marker_best_template_path = str(best_lr.get("template_path", "") or "")
                    lr_marker_best_search_strategy = str(best_lr.get("search_strategy", "") or "")
                    lr_marker_source = (
                        "bundle_manual_seed_template"
                        if bundle_policy_effective.startswith("bundle_manual_seed")
                        else "bundle_historical_best_template"
                    )
                elif any(not _lr_marker_row_quadrant_valid_or_unknown(dict(row_lr)) for row_lr in lr_marker_rows):
                    lr_marker_source = "bundle_quadrant_logic_failed"
                else:
                    lr_marker_source = "bundle_no_reliable_marker_predictions"
        elif not _lr_marker_tooling_available():
            lr_marker_source = "tooling_unavailable"
        else:
            vendor_key = str(vendor_pred).strip().lower()
            if vendor_key not in lr_marker_templates_cache:
                lr_marker_templates_cache[vendor_key] = _load_lr_marker_templates_by_vendor(
                    roots=lr_marker_template_roots,
                    vendor_name=vendor_pred,
                    device=device,
                    max_depth=int(args.lr_marker_max_config_depth),
                    blank_template_max_value=int(args.lr_marker_blank_template_max_value),
                    exclude_vendors=tuple(args.lr_marker_exclude_vendors),
                    review_file=args.lr_marker_review_file,
                )
            vendor_templates, vendor_template_meta = lr_marker_templates_cache[vendor_key]
            if not vendor_templates:
                lr_marker_source = str(vendor_template_meta.get("error", "") or "no_usable_vendor_templates")
            else:
                lr_marker_template_policy_effective = str(args.lr_marker_template_policy)
                lr_marker_template_fallback_reason = ""

                def _run_lr_marker_with_optional_derived(use_derived_template: bool) -> List[Dict[str, object]]:
                    return _predict_lr_marker_on_su_giu_rows(
                        su_giu_rows=su_giu_rows,
                        templates=vendor_templates,
                        vendor_name=vendor_pred,
                        min_match_score=float(args.lr_marker_min_match_score),
                        full_crop_fallback_threshold=float(args.lr_marker_full_crop_fallback_threshold),
                        expanded_search_threshold=float(args.lr_marker_expanded_search_threshold),
                        expanded_search_steps=lr_marker_expanded_search_steps,
                        min_sugiu_confidence=float(args.lr_marker_min_sugiu_confidence),
                        device=device,
                        derived_template_dir=(
                            output_dir / "lr_marker_folder_templates" if bool(use_derived_template) else None
                        ),
                        blank_template_max_value=int(args.lr_marker_blank_template_max_value),
                        target_image_width=int(out_video_x or 0),
                        target_image_height=int(out_video_y or 0),
                        prefer_large_historical_templates=not bool(
                            vendor_template_meta.get("review_active_for_vendor", False)
                        ),
                    )

                manual_seed_templates, manual_seed_meta = _build_lr_marker_manual_seed_templates(
                    seeds=lr_marker_manual_seed_rows,
                    su_giu_rows=su_giu_rows,
                    folder=folder,
                    output_dir=output_dir,
                    vendor_name=vendor_pred,
                    device=device,
                    blank_template_max_value=int(args.lr_marker_blank_template_max_value),
                )
                if manual_seed_meta:
                    for item in manual_seed_meta:
                        item["folder_path"] = folder.as_posix()
                        item["folder_name"] = folder.name
                        item["vendor"] = vendor_pred
                    lr_marker_manual_seed_template_meta.extend(manual_seed_meta)

                manual_seed_fallback_reason = ""
                if manual_seed_templates:
                    manual_rows = _predict_lr_marker_on_su_giu_rows(
                        su_giu_rows=su_giu_rows,
                        templates=manual_seed_templates,
                        vendor_name=vendor_pred,
                        min_match_score=float(args.lr_marker_min_match_score),
                        full_crop_fallback_threshold=float(args.lr_marker_full_crop_fallback_threshold),
                        expanded_search_threshold=float(args.lr_marker_expanded_search_threshold),
                        expanded_search_steps=lr_marker_expanded_search_steps,
                        min_sugiu_confidence=float(args.lr_marker_min_sugiu_confidence),
                        device=device,
                        force_single_template=True,
                        blank_template_max_value=int(args.lr_marker_blank_template_max_value),
                        folder_template_policy="fixed_manual_seed_template",
                        single_template_policy="fixed_manual_seed_template",
                        target_image_width=int(out_video_x or 0),
                        target_image_height=int(out_video_y or 0),
                        prefer_large_historical_templates=False,
                    )
                    manual_reliable_rows = _lr_marker_reliable_rows(
                        manual_rows,
                        min_match_score=float(args.lr_marker_min_match_score),
                        target_image_width=int(out_video_x or 0),
                        target_image_height=int(out_video_y or 0),
                    )
                    if manual_reliable_rows:
                        lr_marker_rows = manual_rows
                        lr_marker_template_policy_effective = "manual_seed"
                    else:
                        manual_seed_fallback_reason = "manual_seed_no_reliable_rows"

                if not lr_marker_rows:
                    if str(args.lr_marker_template_policy) == "derived_folder":
                        lr_marker_rows = _run_lr_marker_with_optional_derived(True)
                        lr_marker_template_policy_effective = "derived_folder"
                    elif str(args.lr_marker_template_policy) == "historical_best_then_derived":
                        lr_marker_rows = _run_lr_marker_with_optional_derived(False)
                        should_fallback, fallback_reason = _lr_marker_should_use_derived_fallback(
                            lr_marker_rows,
                            min_match_score=float(args.lr_marker_min_match_score),
                            target_image_width=int(out_video_x or 0),
                            target_image_height=int(out_video_y or 0),
                        )
                        lr_marker_template_policy_effective = "historical_best"
                        if should_fallback:
                            fallback_rows = _run_lr_marker_with_optional_derived(True)
                            if fallback_rows:
                                lr_marker_rows = fallback_rows
                                lr_marker_template_policy_effective = "derived_folder"
                                lr_marker_template_fallback_reason = fallback_reason
                            else:
                                lr_marker_rows = []
                                lr_marker_template_policy_effective = "derived_folder"
                                lr_marker_template_fallback_reason = f"{fallback_reason};derived_no_rows"
                    else:
                        lr_marker_rows = _run_lr_marker_with_optional_derived(False)
                        lr_marker_template_policy_effective = "historical_best"

                if manual_seed_fallback_reason and lr_marker_template_policy_effective != "manual_seed":
                    lr_marker_template_fallback_reason = (
                        f"{manual_seed_fallback_reason};{lr_marker_template_fallback_reason}"
                        if lr_marker_template_fallback_reason
                        else manual_seed_fallback_reason
                    )

                if lr_marker_template_fallback_reason:
                    for row_lr in lr_marker_rows:
                        row_lr["template_fallback_reason"] = lr_marker_template_fallback_reason
                        row_lr["template_policy_requested"] = str(args.lr_marker_template_policy)
                        row_lr["template_policy_effective"] = lr_marker_template_policy_effective
                else:
                    for row_lr in lr_marker_rows:
                        row_lr["template_policy_requested"] = str(args.lr_marker_template_policy)
                        row_lr["template_policy_effective"] = lr_marker_template_policy_effective
                lr_marker_images_predicted = len(lr_marker_rows)
                if lr_marker_rows:
                    decision_lr_rows = _lr_marker_reliable_rows(
                        lr_marker_rows,
                        min_match_score=float(args.lr_marker_min_match_score),
                        target_image_width=int(out_video_x or 0),
                        target_image_height=int(out_video_y or 0),
                    )
                    label_counter: Counter[str] = Counter()
                    for row_lr in decision_lr_rows:
                        label = str(row_lr.get("lr_label", "") or "").strip()
                        if label:
                            label_counter[label] += 1
                    for row_lr in lr_marker_rows:
                        lr_marker_per_image_rows.append(
                            {
                                "folder_path": folder.as_posix(),
                                "folder_name": folder.name,
                                **row_lr,
                            }
                        )
                    if label_counter:
                        lr_marker_majority_label, majority_count = max(
                            label_counter.items(),
                            key=lambda item: (item[1], item[0] == "not_lr_flipped", item[0]),
                        )
                        lr_marker_majority_vote_ratio = float(majority_count / max(1, len(decision_lr_rows)))
                        lr_marker_majority_counter[lr_marker_majority_label] += 1
                    if decision_lr_rows:
                        best_lr = max(decision_lr_rows, key=lambda item: float(item.get("match_score", 0.0) or 0.0))
                        lr_marker_best_label = str(best_lr.get("lr_label", "") or "")
                        lr_marker_best_score = float(best_lr.get("match_score", 0.0) or 0.0)
                        lr_marker_best_template_path = str(best_lr.get("template_path", "") or "")
                        lr_marker_best_search_strategy = str(best_lr.get("search_strategy", "") or "")
                        if lr_marker_template_policy_effective == "manual_seed":
                            lr_marker_source = "classical_manual_seed_template"
                        elif lr_marker_template_policy_effective == "derived_folder":
                            lr_marker_source = "classical_vendor_derived_folder_template"
                        else:
                            lr_marker_source = "classical_vendor_historical_best_template"
                        if lr_marker_template_fallback_reason:
                            lr_marker_source = f"{lr_marker_source}_fallback"
                    elif any(not _lr_marker_row_quadrant_valid_or_unknown(dict(row_lr)) for row_lr in lr_marker_rows):
                        lr_marker_source = "lr_marker_quadrant_logic_failed"
                    else:
                        lr_marker_source = "no_reliable_marker_predictions"
                else:
                    lr_marker_source = "no_marker_predictions"
        (
            line_16_rect_orientation,
            line_16_source,
            line_16_marker_boxes_count,
            line_16_groups_json,
        ) = _build_line16_rect_orientation_from_lr_marker_rows(
            lr_marker_rows,
            min_match_score=float(args.lr_marker_min_match_score),
            target_image_width=int(out_video_x or 0),
            target_image_height=int(out_video_y or 0),
        )
        lr_marker_source_counter[lr_marker_source] += 1
        _emit_stage_event(
            "orientamento",
            folder.name,
            su_giu_images=su_giu_images_predicted,
            su_giu_majority=su_giu_majority_label,
            su_giu_source=su_giu_source,
            lr_marker_images=lr_marker_images_predicted,
            lr_marker_majority=lr_marker_majority_label,
            lr_marker_source=lr_marker_source,
            line_16=line_16_rect_orientation,
            line_16_source=line_16_source,
        )

        lt_images_predicted = 0
        lt_majority_label = ""
        lt_majority_vote_ratio = 0.0
        lt_mean_confidence = 0.0
        lt_mean_prob_l = 0.0
        lt_mean_prob_t = 0.0
        lt_source = "disabled"
        lt_rows: List[Dict[str, object]] = []
        lt_checkpoint = lt_ckpt_path.as_posix() if lt_model is not None else ""
        if lt_model is not None:
            lt_rows = _predict_lt_on_rect_crops(
                model=lt_model,
                class_names=lt_class_names,
                image_paths=rect_images,
                rect_boxes_abs=rect_boxes_abs,
                fallback_rect_box_abs=rect_box,
                image_size=lt_image_size,
                batch_size=lt_batch_size,
                device=device,
                rotate_deg_clockwise=rotation_deg_clockwise,
            )
            lt_images_predicted = len(lt_rows)
            if lt_rows:
                pred_counter: Counter[str] = Counter()
                used_global_fallback = False
                conf_values: List[float] = []
                prob_l_values: List[float] = []
                prob_t_values: List[float] = []
                for row_lt in lt_rows:
                    pred_label = str(row_lt.get("pred_label", "")).strip().upper()
                    if pred_label:
                        pred_counter[pred_label] += 1
                    if str(row_lt.get("crop_source", "")) == "global_rect_fallback":
                        used_global_fallback = True
                    conf_values.append(float(row_lt.get("confidence", 0.0) or 0.0))
                    prob_l_values.append(float(row_lt.get("prob_l", 0.0) or 0.0))
                    prob_t_values.append(float(row_lt.get("prob_t", 0.0) or 0.0))
                    lt_per_image_rows.append(
                        {
                            "folder_path": folder.as_posix(),
                            "folder_name": folder.name,
                            "image_index": int(row_lt.get("image_index", 0) or 0),
                            "image_path": str(row_lt.get("image_path", "") or ""),
                            "pred_label": pred_label,
                            "pred_idx": int(row_lt.get("pred_idx", 0) or 0),
                            "confidence": float(row_lt.get("confidence", 0.0) or 0.0),
                            "prob_l": float(row_lt.get("prob_l", 0.0) or 0.0),
                            "prob_t": float(row_lt.get("prob_t", 0.0) or 0.0),
                            "crop_top": int(row_lt.get("crop_top", 0) or 0),
                            "crop_left": int(row_lt.get("crop_left", 0) or 0),
                            "crop_bottom": int(row_lt.get("crop_bottom", 0) or 0),
                            "crop_right": int(row_lt.get("crop_right", 0) or 0),
                            "crop_source": str(row_lt.get("crop_source", "") or ""),
                        }
                    )
                if pred_counter:
                    lt_majority_label, majority_count = max(
                        pred_counter.items(),
                        key=lambda item: (item[1], item[0] == "L", item[0] == "T"),
                    )
                    lt_majority_vote_ratio = float(majority_count / max(1, lt_images_predicted))
                    lt_majority_counter[lt_majority_label] += 1
                lt_mean_confidence = float(sum(conf_values) / max(1, len(conf_values)))
                lt_mean_prob_l = float(sum(prob_l_values) / max(1, len(prob_l_values)))
                lt_mean_prob_t = float(sum(prob_t_values) / max(1, len(prob_t_values)))
                lt_source = "per_image_rect_with_global_fallback" if used_global_fallback else "per_image_rect"
            else:
                lt_source = "no_predictions"
            lt_source_counter[lt_source] += 1
        else:
            lt_source_counter[lt_source] += 1

        # Line #04 split rule requested:
        # if router gives "3-4", use L/T classifier majority on rect crops:
        # L -> 3, T -> 4. Keep "3-4" only when LT is unavailable/unresolved.
        if line_04_probe_type == "3-4":
            if lt_majority_label == "L":
                line_04_probe_type = "3"
                line_04_probe_type_source = f"{line_04_probe_type_source}+lt_rect_majority"
                line_04_probe_type_strategy = f"{line_04_probe_type_strategy}+split_3_from_L"
                line_04_probe_type_needs_secondary_model = "0"
            elif lt_majority_label == "T":
                line_04_probe_type = "4"
                line_04_probe_type_source = f"{line_04_probe_type_source}+lt_rect_majority"
                line_04_probe_type_strategy = f"{line_04_probe_type_strategy}+split_4_from_T"
                line_04_probe_type_needs_secondary_model = "0"
            else:
                line_04_probe_type_source = f"{line_04_probe_type_source}+lt_unresolved"
                line_04_probe_type_strategy = f"{line_04_probe_type_strategy}+split_unresolved"

        if args.disable_rect_red_line11:
            rect_red_payload = {
                "available": False,
                "error": "rect_red_disabled_by_flag",
                "margin_pct": float(args.rect_red_margin_pct),
                "bright_thr": float(args.rect_red_bright_thr),
            }
        else:
            orientation_by_image: Dict[str, str] = {}
            for row_sg in su_giu_rows:
                image_path_s = str(row_sg.get("image_path", "") or "").strip()
                label = str(row_sg.get("pred_label", "") or "").strip().lower()
                if not image_path_s or not label:
                    continue
                orientation_by_image[image_path_s] = label
                try:
                    orientation_by_image[Path(image_path_s).expanduser().resolve().as_posix()] = label
                except Exception:
                    pass

            rect_red_records: List[Dict[str, object]] = []
            paired_n = min(len(rect_images), len(rect_boxes_abs), len(rect_sizes))
            for img_i in range(paired_n):
                image_path_obj = rect_images[img_i]
                x1, y1, x2, y2 = rect_boxes_abs[img_i]
                image_w, image_h = rect_sizes[img_i]
                if image_w <= 1 or image_h <= 1:
                    continue
                x1c = max(0.0, min(float(x1), float(image_w) - 1.0))
                y1c = max(0.0, min(float(y1), float(image_h) - 1.0))
                x2c = max(x1c + 1.0, min(float(x2), float(image_w)))
                y2c = max(y1c + 1.0, min(float(y2), float(image_h)))
                image_path_abs = image_path_obj.as_posix()
                try:
                    image_path_abs = image_path_obj.expanduser().resolve().as_posix()
                except Exception:
                    image_path_abs = image_path_obj.as_posix()
                orientation_label = orientation_by_image.get(image_path_abs) or orientation_by_image.get(
                    image_path_obj.as_posix(),
                    "unknown",
                )
                rect_red_records.append(
                    {
                        "image_id": image_path_obj.name,
                        "image_path": image_path_abs,
                        "image_width": int(image_w),
                        "image_height": int(image_h),
                        "orientation_label": orientation_label,
                        "pred_rect_norm": {
                            "x": float(x1c / float(image_w)),
                            "y": float(y1c / float(image_h)),
                            "w": float((x2c - x1c) / float(image_w)),
                            "h": float((y2c - y1c) / float(image_h)),
                        },
                    }
                )

            if not rect_red_records:
                rect_red_payload = {
                    "available": False,
                    "error": "no_rect_records_for_rect_red",
                    "margin_pct": float(args.rect_red_margin_pct),
                    "bright_thr": float(args.rect_red_bright_thr),
                    "records_total": 0,
                }
            else:
                try:
                    rect_red_payload = compute_rect_red_pipeline(
                        records=rect_red_records,
                        output_width=int(out_video_x or 0),
                        output_height=int(out_video_y or 0),
                        margin_pct=float(args.rect_red_margin_pct),
                        bright_thr=float(args.rect_red_bright_thr),
                        detect_segments=True,
                    )
                except Exception as exc:
                    rect_red_payload = {
                        "available": False,
                        "error": f"rect_red_exception:{exc}",
                        "margin_pct": float(args.rect_red_margin_pct),
                        "bright_thr": float(args.rect_red_bright_thr),
                        "records_total": int(len(rect_red_records)),
                    }

        if bool(rect_red_payload.get("available", False)):
            line11_red_obj = rect_red_payload.get("line11_red", {})
            red_text = ""
            red_tlbr = {}
            if isinstance(line11_red_obj, dict):
                red_text = str(line11_red_obj.get("text", "") or "").strip()
                red_tlbr_obj = line11_red_obj.get("rect_tlbr", {})
                if isinstance(red_tlbr_obj, dict):
                    red_tlbr = red_tlbr_obj
            red_accepted = bool(rect_red_payload.get("line11_red_accepted", True))
            if red_text and red_accepted:
                line_11 = red_text
                try:
                    rect_top = str(int(red_tlbr.get("top", 0)))
                    rect_left = str(int(red_tlbr.get("left", 0)))
                    rect_bottom = str(int(red_tlbr.get("bottom", 0)))
                    rect_right = str(int(red_tlbr.get("right", 0)))
                except Exception:
                    pass
                line_11_method = "segment_top_red_rect"
                line_11_rect_red_winner_group = str(rect_red_payload.get("winner_group", "") or "")

        if not line_11:
            line_11 = line_11_median_text
            rect_top = line_11_median_top
            rect_left = line_11_median_left
            rect_bottom = line_11_median_bottom
            rect_right = line_11_median_right
            line_11_method = "median_rect_fallback"
            line_11_rect_red_winner_group = ""

        _emit_stage_event(
            "rect",
            folder.name,
            provisional=False,
            line_11=line_11,
            method=line_11_method,
            source=rect_source,
            model_checkpoint=rect_model_checkpoint,
            median_fallback=line_11_median_text,
            winner_group=line_11_rect_red_winner_group,
            images_used=len(rect_boxes_abs),
        )

        rect_red_payload_out = (
            dict(rect_red_payload)
            if isinstance(rect_red_payload, dict)
            else {"available": False, "error": "invalid_rect_red_payload_type"}
        )
        rect_red_payload_out["folder_path"] = folder.as_posix()
        rect_red_payload_out["folder_name"] = folder.name
        rect_red_payload_out["line11_median_fallback"] = line_11_median_text
        rect_red_payload_out["line11_applied"] = line_11
        rect_red_payload_out["line11_method"] = line_11_method
        rect_red_payload_out["line11_source"] = rect_source
        rect_red_payload_out["line11_model_checkpoint"] = rect_model_checkpoint
        rect_red_by_folder[folder.as_posix()] = rect_red_payload_out

        rect_depth_status = "ok" if args.disable_rect_depth_autonomous else "review"
        rect_depth_images_predicted = 0
        rect_depth_accepted_count = 0
        rect_depth_review_count = 0
        rect_depth_reject_count = 0
        rect_depth_missing_count = 0
        rect_depth_acceptance_ratio = 0.0
        rect_depth_majority_mode = ""
        rect_depth_unique_depths_json = "[]"
        rect_depth_source = "disabled" if args.disable_rect_depth_autonomous else "not_run"
        rect_depth_output_dir = ""
        rect_depth_predictions_csv = ""
        rect_depth_summary_json = ""
        rect_depth_rows: List[Dict[str, object]] = []
        if not args.disable_rect_depth_autonomous:
            rect_depth_summary, rect_depth_rows = _run_rect_depth_autonomous_stage(
                folder=folder,
                output_dir=output_dir,
                folder_index=idx,
                python_bin=sys.executable,
                vendor_pred=vendor_pred,
                probe_id=probe_id,
                line_11_rect_echo=line_11,
                video_x=out_video_x,
                video_y=out_video_y,
                rotation_deg_clockwise=int(rotation_deg_clockwise),
                max_images=int(args.rect_depth_max_images),
                max_candidates_per_sample=int(args.rect_depth_max_candidates_per_sample),
                ocr_timeout=float(args.rect_depth_ocr_timeout),
                scale_side_preference=str(args.rect_depth_scale_side_preference),
                subprocess_timeout_sec=float(args.rect_depth_subprocess_timeout_sec),
                min_accepted_ratio=float(args.rect_depth_min_accepted_ratio),
            )
            rect_depth_per_image_rows.extend(rect_depth_rows)
            rect_depth_status = str(rect_depth_summary.get("status", "review") or "review")
            rect_depth_images_predicted = int(rect_depth_summary.get("images_predicted", 0) or 0)
            rect_depth_accepted_count = int(rect_depth_summary.get("accepted_count", 0) or 0)
            rect_depth_review_count = int(rect_depth_summary.get("review_count", 0) or 0)
            rect_depth_reject_count = int(rect_depth_summary.get("reject_count", 0) or 0)
            rect_depth_missing_count = int(rect_depth_summary.get("missing_count", 0) or 0)
            rect_depth_acceptance_ratio = float(rect_depth_summary.get("acceptance_ratio", 0.0) or 0.0)
            rect_depth_majority_mode = str(rect_depth_summary.get("majority_mode", "") or "")
            rect_depth_unique_depths_json = str(rect_depth_summary.get("unique_depths_json", "[]") or "[]")
            rect_depth_source = str(rect_depth_summary.get("source", "") or "")
            rect_depth_output_dir = str(rect_depth_summary.get("output_dir", "") or "")
            rect_depth_predictions_csv = str(rect_depth_summary.get("predictions_csv", "") or "")
            rect_depth_summary_json = str(rect_depth_summary.get("summary_json", "") or "")
            rect_depth_status_counter[rect_depth_status] += 1
            if rect_depth_majority_mode:
                rect_depth_mode_counter[rect_depth_majority_mode] += 1
            rect_depth_error = str(rect_depth_summary.get("error", "") or "").strip()
            if rect_depth_error:
                warnings.append(f"{folder.name}: rect_depth_autonomous {rect_depth_status}: {rect_depth_error}")
        else:
            rect_depth_status_counter[rect_depth_status] += 1
        _emit_stage_event(
            "depth",
            folder.name,
            status=rect_depth_status,
            images_predicted=rect_depth_images_predicted,
            accepted=rect_depth_accepted_count,
            review=rect_depth_review_count,
            reject=rect_depth_reject_count,
            missing=rect_depth_missing_count,
            acceptance_ratio=rect_depth_acceptance_ratio,
            majority_mode=rect_depth_majority_mode,
            unique_depths_json=rect_depth_unique_depths_json,
            source=rect_depth_source,
            output_dir=rect_depth_output_dir,
            predictions_csv=rect_depth_predictions_csv,
        )

        # ---- scala #18-#21 -------------------------------------------------------------
        # Runs after the depth because line #21 is *per depth*: #17 defines the groups and
        # #18-#22 carry one entry each, in that order. It also consumes the up/down the
        # marker already settled and the per-image rect, so nothing here is recomputed.
        line_18_vect_depth = ""
        line_19_pixel_ratio_x = ""
        line_20_pixel_ratio_y = ""
        line_21_scale_line = ""
        scale_status = "ok" if args.disable_scale_stage else "review"
        scale_source = "disabled" if args.disable_scale_stage else "not_run"
        scale_profile = ""
        scale_frames_studied = 0
        scale_depths_total = 0
        scale_depths_accepted = 0
        scale_depths_review = 0
        scale_depths_reject = 0
        scale_depths_interpolated = 0
        scale_acceptance_ratio = 0.0
        scale_ruler_x = ""
        scale_output_dir = ""
        scale_per_image_csv = ""
        scale_per_depth_csv = ""
        scale_review_reasons: List[str] = []
        if not args.disable_scale_stage:
            scale_frames = _build_scale_frames_context(
                images=all_images,
                su_giu_rows=su_giu_rows,
                lr_marker_rows=lr_marker_rows,
                rect_depth_rows=rect_depth_rows,
                line_11_rect_echo=line_11,
            )
            scale_summary, scale_rows = _run_scale_stage(
                folder=folder,
                output_dir=output_dir,
                folder_index=idx,
                python_bin=sys.executable,
                vendor_pred=vendor_pred,
                vendor_conf=float(vendor_conf),
                probe_id=probe_id,
                line_11_rect_echo=line_11,
                video_x=out_video_x,
                video_y=out_video_y,
                rotation_deg_clockwise=int(rotation_deg_clockwise),
                frames=scale_frames,
                max_frames=int(args.scale_max_frames),
                subprocess_timeout_sec=float(args.scale_subprocess_timeout_sec),
                min_accepted_ratio=float(args.scale_min_accepted_ratio),
                corrections_path=args.scale_corrections,
            )
            scale_per_image_rows.extend(scale_rows)
            scale_status = str(scale_summary.get("status", "review") or "review")
            scale_source = str(scale_summary.get("source", "") or "")
            scale_profile = str(scale_summary.get("profile", "") or "")
            scale_frames_studied = int(scale_summary.get("frames_studied", 0) or 0)
            scale_depths_total = int(scale_summary.get("depths_total", 0) or 0)
            scale_depths_accepted = int(scale_summary.get("depths_accepted", 0) or 0)
            scale_depths_review = int(scale_summary.get("depths_review", 0) or 0)
            scale_depths_reject = int(scale_summary.get("depths_reject", 0) or 0)
            scale_depths_interpolated = int(scale_summary.get("depths_interpolated", 0) or 0)
            scale_acceptance_ratio = float(scale_summary.get("acceptance_ratio", 0.0) or 0.0)
            scale_ruler_x = str(scale_summary.get("ruler_x", "") or "")
            scale_output_dir = str(scale_summary.get("output_dir", "") or "")
            scale_per_image_csv = str(scale_summary.get("per_image_csv", "") or "")
            scale_per_depth_csv = str(scale_summary.get("per_depth_csv", "") or "")
            line_18_vect_depth = str(scale_summary.get("line_18_vect_depth", "") or "")
            line_19_pixel_ratio_x = str(scale_summary.get("line_19_pixel_ratio_x", "") or "")
            line_20_pixel_ratio_y = str(scale_summary.get("line_20_pixel_ratio_y", "") or "")
            line_21_scale_line = str(scale_summary.get("line_21_scale_line", "") or "")
            scale_review_reasons = [
                str(r) for r in (scale_summary.get("review_reasons") or []) if str(r).strip()
            ]
            scale_status_counter[scale_status] += 1
            if scale_source:
                scale_source_counter[scale_source] += 1
            scale_error = str(scale_summary.get("error", "") or "").strip()
            if scale_error:
                warnings.append(f"{folder.name}: scala {scale_status}: {scale_error}")
        else:
            scale_status_counter[scale_status] += 1
        _emit_stage_event(
            "scala",
            folder.name,
            status=scale_status,
            source=scale_source,
            profile=scale_profile,
            frames_studied=scale_frames_studied,
            depths_total=scale_depths_total,
            depths_accepted=scale_depths_accepted,
            depths_review=scale_depths_review,
            depths_reject=scale_depths_reject,
            depths_interpolated=scale_depths_interpolated,
            acceptance_ratio=scale_acceptance_ratio,
            ruler_x=scale_ruler_x,
            line_18=line_18_vect_depth,
            line_21=line_21_scale_line,
            review_reasons=scale_review_reasons,
            output_dir=scale_output_dir,
            per_image_csv=scale_per_image_csv,
            per_depth_csv=scale_per_depth_csv,
        )

        id_echo, id_echo_source, id_echo_support = echo_resolver.resolve(
            vendor=vendor_pred,
            probe_id=probe_id if probe_id else None,
            video_x=out_video_x,
            video_y=out_video_y,
        )
        if not id_echo:
            id_echo_topk = [
                (value, support)
                for value, support, _source in echo_resolver.candidates(
                    vendor=vendor_pred,
                    probe_id=probe_id if probe_id else None,
                    video_x=out_video_x,
                    video_y=out_video_y,
                    top_k=args.interactive_topk,
                )
            ]
            if args.low_confidence_policy == "error":
                raise RuntimeError(
                    f"{folder.name}: ID_ECHO non risolto da mapping storico. "
                    f"Candidate top-k: {id_echo_topk}"
                )
            if args.low_confidence_policy == "ask_user":
                selected_echo, selection_source = _prompt_user_topk_or_custom(
                    field_name="fss_id_echo",
                    folder_name=folder.name,
                    options=id_echo_topk,
                )
                id_echo = selected_echo
                id_echo_source = selection_source
                id_echo_support = 1.0
                if selection_source == "user_custom":
                    _register_manual_choice(
                        folder_name=folder.name,
                        field_name="fss_id_echo",
                        chosen_value=id_echo,
                        topk_options=id_echo_topk,
                        threshold=None,
                        confidence=None,
                    )
        line_12_group_orientation = str(GROUP_ORIENTATION_SYMBOL)
        line_12_source = "forced_constant"
        line_12_support = 1.0
        line_12_label = GROUP_ORIENTATION_LABEL[GROUP_ORIENTATION_SYMBOL]

        review_reasons: List[str] = []
        if not id_echo:
            review_reasons.append("missing_id_echo")
        if not probe_id:
            review_reasons.append("missing_probe_id")
        if not line_13:
            review_reasons.append("missing_rect_name_echo")
        if line13_missing_vendor_model:
            review_reasons.append("missing_vendor_line13_model")
        if line13_vendor_model_no_prediction:
            review_reasons.append("line13_vendor_model_no_prediction")
        if line13_unseen_vendor and not line13_manual_resolved:
            review_reasons.append("unseen_vendor_line13")
        if line_13 and line13_coords is None:
            review_reasons.append("invalid_rect_name_echo")
        if line_13_support < args.line13_min_support and not line13_manual_resolved:
            review_reasons.append("low_rect_name_echo_support")
        if not line_14:
            review_reasons.append("missing_rect_name_probe")
        if line14_unseen_probe and not line14_manual_resolved:
            review_reasons.append("unseen_probe_line14")
        if line_14 and line14_coords is None:
            review_reasons.append("invalid_rect_name_probe")
        if line_14_support < args.line14_min_support and not line14_manual_resolved:
            review_reasons.append("low_rect_name_probe_support")
        if not line_12_group_orientation:
            review_reasons.append("missing_group_orientation")
        if not line_11:
            review_reasons.append("missing_rect_echo")
        if capture_input is None:
            review_reasons.append("missing_video_input")
        if out_video_x is None or out_video_y is None:
            review_reasons.append("missing_video_size")
        if vendor_conf < args.vendor_min_confidence and not vendor_manual_resolved:
            review_reasons.append("low_vendor_conf")
        if vendor_margin < float(args.vendor_min_margin) and not vendor_manual_resolved:
            review_reasons.append("low_vendor_margin")
        if probe_conf < args.probe_min_confidence and not probe_manual_resolved:
            review_reasons.append("low_probe_conf")
        if (
            probe_type_router is not None
            and (not line_04_probe_type or line_04_probe_type.upper() == "UNKNOWN")
        ):
            review_reasons.append("missing_probe_type_line4")
        if line_04_probe_type == "3-4":
            review_reasons.append("probe_type_line4_unresolved_3_4")
        if (not args.disable_rotation_normalization) and rotation_source in {
            "osd_unavailable",
            "osd_no_votes",
            "osd_low_support",
        }:
            review_reasons.append("rotation_not_reliable")
        # #12 e' forzata a 4 (symbol) con support 1.0: il check di support e' inutile
        # finche' non esiste un resolver reale per la linea #12.
        if su_giu_model is not None and su_giu_images_predicted <= 0:
            review_reasons.append("missing_su_giu_predictions")
        if (
            not bool(args.disable_lr_marker_classical)
            and lr_marker_images_predicted <= 0
            and vendor_pred
            and str(vendor_pred).strip().lower()
            not in {str(v).strip().lower() for v in args.lr_marker_exclude_vendors if str(v).strip()}
        ):
            review_reasons.append("missing_lr_marker_predictions")
        if (
            not bool(args.disable_lr_marker_classical)
            and lr_marker_images_predicted > 0
            and vendor_pred
            and str(vendor_pred).strip().lower()
            not in {str(v).strip().lower() for v in args.lr_marker_exclude_vendors if str(v).strip()}
            and not str(line_16_rect_orientation or "").strip()
        ):
            review_reasons.append(str(line_16_source or "missing_line16_rect_orientation"))
        if lt_model is not None and lt_images_predicted <= 0:
            review_reasons.append("missing_lt_predictions")
        if lt_model is not None and lt_images_predicted > 0 and lt_mean_confidence < float(args.lt_min_confidence):
            review_reasons.append("low_lt_conf")
        if not args.disable_rect_depth_autonomous:
            if rect_depth_images_predicted <= 0:
                review_reasons.append("missing_rect_depth_predictions")
            elif rect_depth_status != "ok":
                review_reasons.append("rect_depth_autonomous_review")
        if not args.disable_scale_stage:
            # The stage's own reasons are kept verbatim (they say *why* the scale is unsure),
            # and the folder is marked review whenever the scale is not ok.
            review_reasons.extend(scale_review_reasons)
            if scale_depths_total <= 0:
                review_reasons.append("missing_scale_predictions")
            elif scale_status != "ok" and "scale_review" not in review_reasons:
                review_reasons.append("scale_review")

        status = "ok" if not review_reasons else "review"
        status_counter[status] += 1

        pred = FolderPrediction(
            folder_path=folder,
            folder_name=folder.name,
            images_total=len(all_images),
            images_total_raw=len(raw_images),
            images_duplicates_removed=duplicates_removed,
            images_used_cls=len(cls_images),
            images_used_rect=len(rect_images),
            line_01_version=args.fss_version,
            line_02_id_echo=id_echo,
            line_02_source=id_echo_source,
            line_02_support=id_echo_support,
            line_03_id_probe=probe_id,
            line_03_probe_conf=probe_conf,
            line_04_probe_type=line_04_probe_type,
            line_04_probe_type_source=line_04_probe_type_source,
            line_04_probe_type_strategy=line_04_probe_type_strategy,
            line_04_probe_type_needs_secondary_model=line_04_probe_type_needs_secondary_model,
            line_04_probe_type_candidate_ints=line_04_probe_type_candidate_ints,
            line_06_video_input=(capture_input or ""),
            line_06_video_input_code=(
                str(_video_input_code_0hdmi_1vga(capture_input))
                if _video_input_code_0hdmi_1vga(capture_input) is not None
                else ""
            ),
            line_07_video_input_size_x=str(out_video_x) if out_video_x is not None else "",
            line_08_video_input_size_y=str(out_video_y) if out_video_y is not None else "",
            line_09_video_x_size=str(out_video_x) if out_video_x is not None else "",
            line_10_video_y_size=str(out_video_y) if out_video_y is not None else "",
            line_11_rect_echo=line_11,
            line_11_top=rect_top,
            line_11_left=rect_left,
            line_11_bottom=rect_bottom,
            line_11_right=rect_right,
            line_11_source=rect_source,
            line_11_model_checkpoint=rect_model_checkpoint,
            line_11_method=line_11_method,
            line_11_rect_red_margin_pct=float(line_11_rect_red_margin_pct),
            line_11_rect_red_winner_group=line_11_rect_red_winner_group,
            line_12_group_orientation=line_12_group_orientation,
            line_12_group_orientation_label=line_12_label,
            line_12_source=line_12_source,
            line_12_support=line_12_support,
            su_giu_images_predicted=su_giu_images_predicted,
            su_giu_majority_label=su_giu_majority_label,
            su_giu_majority_vote_ratio=su_giu_majority_vote_ratio,
            su_giu_mean_confidence=su_giu_mean_confidence,
            su_giu_mean_prob_su=su_giu_mean_prob_su,
            su_giu_mean_prob_giu=su_giu_mean_prob_giu,
            su_giu_source=su_giu_source,
            su_giu_checkpoint=su_giu_checkpoint,
            lr_marker_images_predicted=lr_marker_images_predicted,
            lr_marker_majority_label=lr_marker_majority_label,
            lr_marker_majority_vote_ratio=lr_marker_majority_vote_ratio,
            lr_marker_best_label=lr_marker_best_label,
            lr_marker_best_score=lr_marker_best_score,
            lr_marker_best_template_path=lr_marker_best_template_path,
            lr_marker_best_search_strategy=lr_marker_best_search_strategy,
            lr_marker_source=lr_marker_source,
            lr_marker_method=str(args.lr_marker_method),
            line_16_rect_orientation=line_16_rect_orientation,
            line_16_source=line_16_source,
            line_16_marker_boxes_count=line_16_marker_boxes_count,
            line_16_groups_json=line_16_groups_json,
            rect_depth_status=rect_depth_status,
            rect_depth_images_predicted=rect_depth_images_predicted,
            rect_depth_accepted_count=rect_depth_accepted_count,
            rect_depth_review_count=rect_depth_review_count,
            rect_depth_reject_count=rect_depth_reject_count,
            rect_depth_missing_count=rect_depth_missing_count,
            rect_depth_acceptance_ratio=rect_depth_acceptance_ratio,
            rect_depth_majority_mode=rect_depth_majority_mode,
            rect_depth_unique_depths_json=rect_depth_unique_depths_json,
            rect_depth_source=rect_depth_source,
            rect_depth_output_dir=rect_depth_output_dir,
            rect_depth_predictions_csv=rect_depth_predictions_csv,
            rect_depth_summary_json=rect_depth_summary_json,
            line_18_vect_depth=line_18_vect_depth,
            line_19_pixel_ratio_x=line_19_pixel_ratio_x,
            line_20_pixel_ratio_y=line_20_pixel_ratio_y,
            line_21_scale_line=line_21_scale_line,
            scale_status=scale_status,
            scale_source=scale_source,
            scale_profile=scale_profile,
            scale_frames_studied=scale_frames_studied,
            scale_depths_total=scale_depths_total,
            scale_depths_accepted=scale_depths_accepted,
            scale_depths_review=scale_depths_review,
            scale_depths_reject=scale_depths_reject,
            scale_depths_interpolated=scale_depths_interpolated,
            scale_acceptance_ratio=scale_acceptance_ratio,
            scale_ruler_x=scale_ruler_x,
            scale_output_dir=scale_output_dir,
            scale_per_image_csv=scale_per_image_csv,
            scale_per_depth_csv=scale_per_depth_csv,
            lt_images_predicted=lt_images_predicted,
            lt_majority_label=lt_majority_label,
            lt_majority_vote_ratio=lt_majority_vote_ratio,
            lt_mean_confidence=lt_mean_confidence,
            lt_mean_prob_l=lt_mean_prob_l,
            lt_mean_prob_t=lt_mean_prob_t,
            lt_source=lt_source,
            lt_checkpoint=lt_checkpoint,
            line_13_rect_name_echo=line_13,
            line_13_top=line13_top,
            line_13_left=line13_left,
            line_13_bottom=line13_bottom,
            line_13_right=line13_right,
            line_13_source=line_13_source,
            line_13_support=line_13_support,
            line_13_model_checkpoint=line_13_model_checkpoint,
            line_13_model_route=line_13_model_route,
            line_13_pre_dark_trim_rect_name_echo=line_13_pre_dark_trim_rect_name_echo,
            line_14_rect_name_probe=line_14,
            line_14_top=line14_top,
            line_14_left=line14_left,
            line_14_bottom=line14_bottom,
            line_14_right=line14_right,
            line_14_source=line_14_source,
            line_14_support=line_14_support,
            vendor_pred=vendor_pred,
            vendor_conf=vendor_conf,
            vendor_margin=vendor_margin,
            vendor_source=vendor_source,
            vendor_decision_reason=vendor_decision_reason,
            vendor_top3_json=_topk_json(vendor_top3),
            vendor_threshold=float(args.vendor_min_confidence),
            vendor_ocr_used=vendor_ocr_used,
            vendor_ocr_candidate=vendor_ocr_candidate,
            vendor_ocr_candidate_score=vendor_ocr_candidate_score,
            vendor_ocr_candidate_hits=vendor_ocr_candidate_hits,
            vendor_ocr_delta_score=vendor_ocr_delta_score,
            vendor_ocr_text=vendor_ocr_text,
            vendor_ocr_top3_json=vendor_ocr_top3_json,
            vendor_ocr_samples_checked=vendor_ocr_samples_checked,
            vendor_ocr_reason=vendor_ocr_reason,
            probe_source=probe_source,
            probe_decision_reason=probe_decision_reason,
            probe_top3_json=_topk_json(probe_top3),
            probe_threshold=float(args.probe_min_confidence),
            capture_metadata_support=capture_support,
            rotation_deg_clockwise=int(rotation_deg_clockwise),
            rotation_source=rotation_source,
            rotation_vote_ratio=float(rotation_vote_ratio),
            rotation_votes_total=int(rotation_votes_total),
            rotation_samples_checked=int(rotation_samples_checked),
            rotation_decision_reason=rotation_decision_reason,
            rotation_osd_debug_json=json.dumps(rotation_osd_debug, ensure_ascii=False, separators=(",", ":")),
            rotation_ocr_debug_json=json.dumps(rotation_ocr_debug, ensure_ascii=False, separators=(",", ":")),
            status=status,
            review_reasons=",".join(review_reasons),
        )
        predictions.append(pred)
        if _STAGE_EVENTS_ENABLED:
            # The whole prediction, so the review tool has every line and every source without
            # waiting for the CSV (written only when the last folder is done).
            _emit_stage_event(
                "folder_done",
                folder.name,
                folder_index=idx,
                folders_total=len(folders),
                status=status,
                review_reasons=review_reasons,
                prediction={k: str(v) for k, v in asdict(pred).items()},
            )

        if args.log_interval > 0 and (idx % args.log_interval == 0 or idx == len(folders)):
            elapsed = time.time() - start_time
            print(
                f"[{idx:04d}/{len(folders):04d}] folders processed | "
                f"ok={status_counter['ok']} review={status_counter['review']} "
                f"| elapsed {elapsed:.1f}s",
                flush=True,
            )

    csv_path = output_dir / "folder_fss_head_predictions.csv"
    with csv_path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(
            fh,
                fieldnames=[
                    "folder_path",
                    "folder_name",
                    "images_total",
                    "images_total_raw",
                    "images_duplicates_removed",
                    "images_used_cls",
                    "images_used_rect",
                    "rotation_deg_clockwise",
                    "rotation_source",
                    "rotation_vote_ratio",
                    "rotation_votes_total",
                    "rotation_samples_checked",
                    "rotation_decision_reason",
                    "rotation_osd_debug_json",
                    "rotation_ocr_debug_json",
                "line_01_version",
                "line_02_id_echo",
                "line_02_id_echo_source",
                "line_02_id_echo_support",
                "line_03_id_probe",
                "line_03_probe_confidence",
                "line_04_probe_type",
                "line_04_probe_type_source",
                "line_04_probe_type_strategy",
                "line_04_probe_type_needs_secondary_model",
                "line_04_probe_type_candidate_ints",
                "line_06_video_input",
                "line_06_video_input_code_0hdmi_1vga",
                "line_07_video_input_size_x",
                "line_08_video_input_size_y",
                "line_09_video_x_size",
                "line_10_video_y_size",
                "line_11_rect_echo",
                "line_11_top",
                "line_11_left",
                "line_11_bottom",
                "line_11_right",
                "line_11_source",
                "line_11_model_checkpoint",
                "line_11_method",
                "line_11_rect_red_margin_pct",
                "line_11_rect_red_winner_group",
                "line_12_group_orientation",
                "line_12_group_orientation_label",
                "line_12_source",
                "line_12_support",
                "su_giu_images_predicted",
                "su_giu_majority_label",
                "su_giu_majority_vote_ratio",
                "su_giu_mean_confidence",
                "su_giu_mean_prob_su",
                "su_giu_mean_prob_giu",
                "su_giu_source",
                "su_giu_checkpoint",
                "lr_marker_images_predicted",
                "lr_marker_majority_label",
                "lr_marker_majority_vote_ratio",
                "lr_marker_best_label",
                "lr_marker_best_score",
                "lr_marker_best_template_path",
                "lr_marker_best_search_strategy",
                "lr_marker_source",
                "lr_marker_method",
                "line_16_rect_orientation",
                "line_16_source",
                "line_16_marker_boxes_count",
                "line_16_groups_json",
                "rect_depth_status",
                "rect_depth_images_predicted",
                "rect_depth_accepted_count",
                "rect_depth_review_count",
                "rect_depth_reject_count",
                "rect_depth_missing_count",
                "rect_depth_acceptance_ratio",
                "rect_depth_majority_mode",
                "rect_depth_unique_depths_json",
                "rect_depth_source",
                "rect_depth_output_dir",
                "rect_depth_predictions_csv",
                "rect_depth_summary_json",
                "line_18_vect_depth",
                "line_19_pixel_ratio_x",
                "line_20_pixel_ratio_y",
                "line_21_scale_line",
                "scale_status",
                "scale_source",
                "scale_profile",
                "scale_frames_studied",
                "scale_depths_total",
                "scale_depths_accepted",
                "scale_depths_review",
                "scale_depths_reject",
                "scale_depths_interpolated",
                "scale_acceptance_ratio",
                "scale_ruler_x",
                "scale_output_dir",
                "scale_per_image_csv",
                "scale_per_depth_csv",
                "lt_images_predicted",
                "lt_majority_label",
                "lt_majority_vote_ratio",
                "lt_mean_confidence",
                "lt_mean_prob_l",
                "lt_mean_prob_t",
                "lt_source",
                "lt_checkpoint",
                "line_13_rect_name_echo",
                "line_13_top",
                "line_13_left",
                "line_13_bottom",
                "line_13_right",
                "line_13_source",
                "line_13_support",
                "line_13_model_checkpoint",
                "line_13_model_route",
                "line_13_pre_dark_trim_rect_name_echo",
                "line_14_rect_name_probe",
                "line_14_top",
                "line_14_left",
                "line_14_bottom",
                "line_14_right",
                "line_14_source",
                "line_14_support",
                "vendor_predicted",
                "vendor_confidence",
                "vendor_margin_top1_top2",
                "vendor_source",
                "vendor_decision_reason",
                "vendor_top3_json",
                "vendor_threshold",
                "vendor_ocr_used",
                "vendor_ocr_candidate",
                "vendor_ocr_candidate_score",
                "vendor_ocr_candidate_hits",
                "vendor_ocr_delta_score",
                "vendor_ocr_text",
                "vendor_ocr_top3_json",
                "vendor_ocr_samples_checked",
                "vendor_ocr_reason",
                "probe_source",
                "probe_decision_reason",
                "probe_top3_json",
                "probe_threshold",
                "capture_metadata_support_count",
                "status",
                "review_reasons",
            ],
        )
        writer.writeheader()
        for p in predictions:
            writer.writerow(
                {
                    "folder_path": p.folder_path.as_posix(),
                    "folder_name": p.folder_name,
                    "images_total": p.images_total,
                    "images_total_raw": p.images_total_raw,
                    "images_duplicates_removed": p.images_duplicates_removed,
                    "images_used_cls": p.images_used_cls,
                    "images_used_rect": p.images_used_rect,
                    "rotation_deg_clockwise": p.rotation_deg_clockwise,
                    "rotation_source": p.rotation_source,
                    "rotation_vote_ratio": f"{p.rotation_vote_ratio:.6f}",
                    "rotation_votes_total": p.rotation_votes_total,
                    "rotation_samples_checked": p.rotation_samples_checked,
                    "rotation_decision_reason": p.rotation_decision_reason,
                    "rotation_osd_debug_json": p.rotation_osd_debug_json,
                    "rotation_ocr_debug_json": p.rotation_ocr_debug_json,
                    "line_01_version": p.line_01_version,
                    "line_02_id_echo": p.line_02_id_echo,
                    "line_02_id_echo_source": p.line_02_source,
                    "line_02_id_echo_support": f"{p.line_02_support:.6f}",
                    "line_03_id_probe": p.line_03_id_probe,
                    "line_03_probe_confidence": f"{p.line_03_probe_conf:.6f}",
                    "line_04_probe_type": p.line_04_probe_type,
                    "line_04_probe_type_source": p.line_04_probe_type_source,
                    "line_04_probe_type_strategy": p.line_04_probe_type_strategy,
                    "line_04_probe_type_needs_secondary_model": p.line_04_probe_type_needs_secondary_model,
                    "line_04_probe_type_candidate_ints": p.line_04_probe_type_candidate_ints,
                    "line_06_video_input": p.line_06_video_input,
                    "line_06_video_input_code_0hdmi_1vga": p.line_06_video_input_code,
                    "line_07_video_input_size_x": p.line_07_video_input_size_x,
                    "line_08_video_input_size_y": p.line_08_video_input_size_y,
                    "line_09_video_x_size": p.line_09_video_x_size,
                    "line_10_video_y_size": p.line_10_video_y_size,
                    "line_11_rect_echo": p.line_11_rect_echo,
                    "line_11_top": p.line_11_top,
                    "line_11_left": p.line_11_left,
                    "line_11_bottom": p.line_11_bottom,
                    "line_11_right": p.line_11_right,
                    "line_11_source": p.line_11_source,
                    "line_11_model_checkpoint": p.line_11_model_checkpoint,
                    "line_11_method": p.line_11_method,
                    "line_11_rect_red_margin_pct": f"{p.line_11_rect_red_margin_pct:.4f}",
                    "line_11_rect_red_winner_group": p.line_11_rect_red_winner_group,
                    "line_12_group_orientation": p.line_12_group_orientation,
                    "line_12_group_orientation_label": p.line_12_group_orientation_label,
                    "line_12_source": p.line_12_source,
                    "line_12_support": f"{p.line_12_support:.6f}",
                    "su_giu_images_predicted": p.su_giu_images_predicted,
                    "su_giu_majority_label": p.su_giu_majority_label,
                    "su_giu_majority_vote_ratio": f"{p.su_giu_majority_vote_ratio:.6f}",
                    "su_giu_mean_confidence": f"{p.su_giu_mean_confidence:.6f}",
                    "su_giu_mean_prob_su": f"{p.su_giu_mean_prob_su:.6f}",
                    "su_giu_mean_prob_giu": f"{p.su_giu_mean_prob_giu:.6f}",
                    "su_giu_source": p.su_giu_source,
                    "su_giu_checkpoint": p.su_giu_checkpoint,
                    "lr_marker_images_predicted": p.lr_marker_images_predicted,
                    "lr_marker_majority_label": p.lr_marker_majority_label,
                    "lr_marker_majority_vote_ratio": f"{p.lr_marker_majority_vote_ratio:.6f}",
                    "lr_marker_best_label": p.lr_marker_best_label,
                    "lr_marker_best_score": f"{p.lr_marker_best_score:.6f}",
                    "lr_marker_best_template_path": p.lr_marker_best_template_path,
                    "lr_marker_best_search_strategy": p.lr_marker_best_search_strategy,
                    "lr_marker_source": p.lr_marker_source,
                    "lr_marker_method": p.lr_marker_method,
                    "line_16_rect_orientation": p.line_16_rect_orientation,
                    "line_16_source": p.line_16_source,
                    "line_16_marker_boxes_count": p.line_16_marker_boxes_count,
                    "line_16_groups_json": p.line_16_groups_json,
                    "rect_depth_status": p.rect_depth_status,
                    "rect_depth_images_predicted": p.rect_depth_images_predicted,
                    "rect_depth_accepted_count": p.rect_depth_accepted_count,
                    "rect_depth_review_count": p.rect_depth_review_count,
                    "rect_depth_reject_count": p.rect_depth_reject_count,
                    "rect_depth_missing_count": p.rect_depth_missing_count,
                    "rect_depth_acceptance_ratio": f"{p.rect_depth_acceptance_ratio:.6f}",
                    "rect_depth_majority_mode": p.rect_depth_majority_mode,
                    "rect_depth_unique_depths_json": p.rect_depth_unique_depths_json,
                    "rect_depth_source": p.rect_depth_source,
                    "rect_depth_output_dir": p.rect_depth_output_dir,
                    "rect_depth_predictions_csv": p.rect_depth_predictions_csv,
                    "rect_depth_summary_json": p.rect_depth_summary_json,
                    "line_18_vect_depth": p.line_18_vect_depth,
                    "line_19_pixel_ratio_x": p.line_19_pixel_ratio_x,
                    "line_20_pixel_ratio_y": p.line_20_pixel_ratio_y,
                    "line_21_scale_line": p.line_21_scale_line,
                    "scale_status": p.scale_status,
                    "scale_source": p.scale_source,
                    "scale_profile": p.scale_profile,
                    "scale_frames_studied": p.scale_frames_studied,
                    "scale_depths_total": p.scale_depths_total,
                    "scale_depths_accepted": p.scale_depths_accepted,
                    "scale_depths_review": p.scale_depths_review,
                    "scale_depths_reject": p.scale_depths_reject,
                    "scale_depths_interpolated": p.scale_depths_interpolated,
                    "scale_acceptance_ratio": f"{p.scale_acceptance_ratio:.6f}",
                    "scale_ruler_x": p.scale_ruler_x,
                    "scale_output_dir": p.scale_output_dir,
                    "scale_per_image_csv": p.scale_per_image_csv,
                    "scale_per_depth_csv": p.scale_per_depth_csv,
                    "lt_images_predicted": p.lt_images_predicted,
                    "lt_majority_label": p.lt_majority_label,
                    "lt_majority_vote_ratio": f"{p.lt_majority_vote_ratio:.6f}",
                    "lt_mean_confidence": f"{p.lt_mean_confidence:.6f}",
                    "lt_mean_prob_l": f"{p.lt_mean_prob_l:.6f}",
                    "lt_mean_prob_t": f"{p.lt_mean_prob_t:.6f}",
                    "lt_source": p.lt_source,
                    "lt_checkpoint": p.lt_checkpoint,
                    "line_13_rect_name_echo": p.line_13_rect_name_echo,
                    "line_13_top": p.line_13_top,
                    "line_13_left": p.line_13_left,
                    "line_13_bottom": p.line_13_bottom,
                    "line_13_right": p.line_13_right,
                    "line_13_source": p.line_13_source,
                    "line_13_support": f"{p.line_13_support:.6f}",
                    "line_13_model_checkpoint": p.line_13_model_checkpoint,
                    "line_13_model_route": p.line_13_model_route,
                    "line_13_pre_dark_trim_rect_name_echo": p.line_13_pre_dark_trim_rect_name_echo,
                    "line_14_rect_name_probe": p.line_14_rect_name_probe,
                    "line_14_top": p.line_14_top,
                    "line_14_left": p.line_14_left,
                    "line_14_bottom": p.line_14_bottom,
                    "line_14_right": p.line_14_right,
                    "line_14_source": p.line_14_source,
                    "line_14_support": f"{p.line_14_support:.6f}",
                    "vendor_predicted": p.vendor_pred,
                    "vendor_confidence": f"{p.vendor_conf:.6f}",
                    "vendor_margin_top1_top2": f"{p.vendor_margin:.6f}",
                    "vendor_source": p.vendor_source,
                    "vendor_decision_reason": p.vendor_decision_reason,
                    "vendor_top3_json": p.vendor_top3_json,
                    "vendor_threshold": f"{p.vendor_threshold:.6f}",
                    "vendor_ocr_used": p.vendor_ocr_used,
                    "vendor_ocr_candidate": p.vendor_ocr_candidate,
                    "vendor_ocr_candidate_score": p.vendor_ocr_candidate_score,
                    "vendor_ocr_candidate_hits": p.vendor_ocr_candidate_hits,
                    "vendor_ocr_delta_score": p.vendor_ocr_delta_score,
                    "vendor_ocr_text": p.vendor_ocr_text,
                    "vendor_ocr_top3_json": p.vendor_ocr_top3_json,
                    "vendor_ocr_samples_checked": p.vendor_ocr_samples_checked,
                    "vendor_ocr_reason": p.vendor_ocr_reason,
                    "probe_source": p.probe_source,
                    "probe_decision_reason": p.probe_decision_reason,
                    "probe_top3_json": p.probe_top3_json,
                    "probe_threshold": f"{p.probe_threshold:.6f}",
                    "capture_metadata_support_count": p.capture_metadata_support,
                    "status": p.status,
                    "review_reasons": p.review_reasons,
                }
            )

    su_giu_per_image_csv = output_dir / "su_giu_per_image_predictions.csv"
    with su_giu_per_image_csv.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(
            fh,
            fieldnames=[
                "folder_path",
                "folder_name",
                "image_index",
                "image_path",
                "pred_label",
                "pred_idx",
                "confidence",
                "prob_su",
                "prob_giu",
                "crop_top",
                "crop_left",
                "crop_bottom",
                "crop_right",
                "crop_source",
            ],
        )
        writer.writeheader()
        for row in su_giu_per_image_rows:
            writer.writerow(
                {
                    "folder_path": str(row.get("folder_path", "") or ""),
                    "folder_name": str(row.get("folder_name", "") or ""),
                    "image_index": int(row.get("image_index", 0) or 0),
                    "image_path": str(row.get("image_path", "") or ""),
                    "pred_label": str(row.get("pred_label", "") or ""),
                    "pred_idx": int(row.get("pred_idx", 0) or 0),
                    "confidence": f"{float(row.get('confidence', 0.0) or 0.0):.6f}",
                    "prob_su": f"{float(row.get('prob_su', 0.0) or 0.0):.6f}",
                    "prob_giu": f"{float(row.get('prob_giu', 0.0) or 0.0):.6f}",
                    "crop_top": int(row.get("crop_top", 0) or 0),
                    "crop_left": int(row.get("crop_left", 0) or 0),
                    "crop_bottom": int(row.get("crop_bottom", 0) or 0),
                    "crop_right": int(row.get("crop_right", 0) or 0),
                    "crop_source": str(row.get("crop_source", "") or ""),
                }
            )

    lr_marker_per_image_csv = output_dir / "lr_marker_per_image_predictions.csv"
    with lr_marker_per_image_csv.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(
            fh,
            fieldnames=[
                "folder_path",
                "folder_name",
                "image_index",
                "image_path",
                "image_width",
                "image_height",
                "canonical_image_size",
                "vendor",
                "lr_marker_method",
                "status",
                "review_reason",
                "orientation_group",
                "lr_label",
                "lr_label_it",
                "lr_binary",
                    "detected_marker_side",
                    "match_score",
                    "match_score_raw",
                    "match_score_adjusted",
                    "match_score_penalty_factor",
                    "match_score_penalty_reason",
                    "match_patch_is_blank",
                    "match_patch_dark_pct",
                    "match_patch_max_value",
                    "match_patch_std",
                "initial_match_score",
                "top_half_score",
                "bottom_half_score",
                "full_crop_fallback_score",
                "expanded_score",
                "search_strategy",
                "search_scope",
                "search_margin_px",
                "template_policy",
                "template_policy_requested",
                "template_policy_effective",
                "template_fallback_reason",
                "manual_seed_index",
                "folder_fixed_template_path",
                "folder_fixed_template_selection_score",
                    "folder_template_seed_image_path",
                    "folder_template_seed_darkness_pct",
                    "folder_template_seed_db_template_path",
                    "folder_template_rank_json",
                    "bundle_template_name",
                    "template_path",
                    "template_width",
                    "template_height",
                "su_giu_model_pred",
                "su_giu_pred",
                "su_giu_conf",
                "prob_su",
                "prob_giu",
                "dual_sugiu_search_used",
                "opposite_sugiu_match_score",
                "bundle_vertical_center",
                "bundle_vertical_source",
                "bundle_vertical_correction",
                "quadrant_valid",
                "quadrant_status",
                "quadrant_reason",
                "quadrant_expected",
                "quadrant_center",
                "quadrant_group",
                "quadrant_center_group",
                "echo_mid_x_abs",
                "echo_mid_y_abs",
                "echo_rect_top_abs",
                "echo_rect_left_abs",
                "echo_rect_bottom_abs",
                "echo_rect_right_abs",
                "marker_top_crop",
                "marker_left_crop",
                "marker_bottom_crop",
                "marker_right_crop",
                "marker_top_abs",
                "marker_left_abs",
                "marker_bottom_abs",
                    "marker_right_abs",
                    "marker_cx_crop_norm",
                    "marker_cy_crop_norm",
                    "marker_echo_overlap_ratio",
                    "marker_outside_echo_rect",
                    "marker_outside_echo_top_px",
                    "marker_outside_echo_left_px",
                "marker_outside_echo_bottom_px",
                "marker_outside_echo_right_px",
                "marker_outside_echo_max_px",
                "marker_outside_echo_gap_px",
                "marker_outside_echo_severity",
                "marker_top_left_outside_echo_rect",
                "marker_echo_near_margin_px",
                ],
            )
        writer.writeheader()
        for row in lr_marker_per_image_rows:
            writer.writerow(
                {
                    "folder_path": str(row.get("folder_path", "") or ""),
                    "folder_name": str(row.get("folder_name", "") or ""),
                    "image_index": int(row.get("image_index", 0) or 0),
                    "image_path": str(row.get("image_path", "") or ""),
                    "image_width": int(row.get("image_width", 0) or 0),
                    "image_height": int(row.get("image_height", 0) or 0),
                    "canonical_image_size": int(row.get("canonical_image_size", 0) or 0),
                    "vendor": str(row.get("vendor", "") or ""),
                    "lr_marker_method": str(row.get("lr_marker_method", "classical") or "classical"),
                    "status": str(row.get("status", "") or ""),
                    "review_reason": str(row.get("review_reason", "") or ""),
                    "orientation_group": str(row.get("orientation_group", "") or ""),
                    "lr_label": str(row.get("lr_label", "") or ""),
                    "lr_label_it": str(row.get("lr_label_it", "") or ""),
                    "lr_binary": int(row.get("lr_binary", -1) or -1),
                    "detected_marker_side": str(row.get("detected_marker_side", "") or ""),
                    "match_score": f"{float(row.get('match_score', 0.0) or 0.0):.6f}",
                    "match_score_raw": f"{float(row.get('match_score_raw', row.get('match_score', 0.0)) or 0.0):.6f}",
                    "match_score_adjusted": f"{float(row.get('match_score_adjusted', row.get('match_score', 0.0)) or 0.0):.6f}",
                    "match_score_penalty_factor": f"{float(row.get('match_score_penalty_factor', 1.0) or 1.0):.6f}",
                    "match_score_penalty_reason": str(row.get("match_score_penalty_reason", "") or ""),
                    "match_patch_is_blank": int(row.get("match_patch_is_blank", 0) or 0),
                    "match_patch_dark_pct": f"{float(row.get('match_patch_dark_pct', 0.0) or 0.0):.6f}",
                    "match_patch_max_value": int(row.get("match_patch_max_value", 0) or 0),
                    "match_patch_std": f"{float(row.get('match_patch_std', 0.0) or 0.0):.6f}",
                    "initial_match_score": row.get("initial_match_score", ""),
                    "top_half_score": row.get("top_half_score", ""),
                    "bottom_half_score": row.get("bottom_half_score", ""),
                    "full_crop_fallback_score": row.get("full_crop_fallback_score", ""),
                    "expanded_score": row.get("expanded_score", ""),
                    "search_strategy": str(row.get("search_strategy", "") or ""),
                    "search_scope": str(row.get("search_scope", "") or ""),
                    "search_margin_px": int(row.get("search_margin_px", 0) or 0),
                    "template_policy": str(row.get("template_policy", "") or ""),
                    "template_policy_requested": str(row.get("template_policy_requested", "") or ""),
                    "template_policy_effective": str(row.get("template_policy_effective", "") or ""),
                    "template_fallback_reason": str(row.get("template_fallback_reason", "") or ""),
                    "manual_seed_index": str(row.get("manual_seed_index", "") or ""),
                    "folder_fixed_template_path": str(row.get("folder_fixed_template_path", "") or ""),
                    "folder_fixed_template_selection_score": (
                        ""
                        if row.get("folder_fixed_template_selection_score", "") == ""
                        else f"{float(row.get('folder_fixed_template_selection_score', 0.0) or 0.0):.6f}"
                    ),
                    "folder_template_seed_image_path": str(row.get("folder_template_seed_image_path", "") or ""),
                    "folder_template_seed_darkness_pct": (
                        ""
                        if float(row.get("folder_template_seed_darkness_pct", -1.0) or -1.0) < 0.0
                        else f"{float(row.get('folder_template_seed_darkness_pct', 0.0) or 0.0):.6f}"
                    ),
                    "folder_template_seed_db_template_path": str(row.get("folder_template_seed_db_template_path", "") or ""),
                    "folder_template_rank_json": str(row.get("folder_template_rank_json", "") or ""),
                    "bundle_template_name": str(row.get("bundle_template_name", "") or ""),
                    "template_path": str(row.get("template_path", "") or ""),
                    "template_width": int(row.get("template_width", 0) or 0),
                    "template_height": int(row.get("template_height", 0) or 0),
                    "su_giu_model_pred": str(row.get("su_giu_model_pred", "") or ""),
                    "su_giu_pred": str(row.get("su_giu_pred", "") or ""),
                    "su_giu_conf": f"{float(row.get('su_giu_conf', 0.0) or 0.0):.6f}",
                    "prob_su": f"{float(row.get('prob_su', 0.0) or 0.0):.6f}",
                    "prob_giu": f"{float(row.get('prob_giu', 0.0) or 0.0):.6f}",
                    "dual_sugiu_search_used": int(row.get("dual_sugiu_search_used", 0) or 0),
                    "opposite_sugiu_match_score": (
                        ""
                        if row.get("opposite_sugiu_match_score", "") == ""
                        else f"{float(row.get('opposite_sugiu_match_score', 0.0) or 0.0):.6f}"
                    ),
                    "bundle_vertical_center": str(row.get("bundle_vertical_center", "") or ""),
                    "bundle_vertical_source": str(row.get("bundle_vertical_source", "") or ""),
                    "bundle_vertical_correction": str(row.get("bundle_vertical_correction", "") or ""),
                    "quadrant_valid": int(row.get("quadrant_valid", 0) or 0),
                    "quadrant_status": str(row.get("quadrant_status", "") or ""),
                    "quadrant_reason": str(row.get("quadrant_reason", "") or ""),
                    "quadrant_expected": str(row.get("quadrant_expected", "") or ""),
                    "quadrant_center": str(row.get("quadrant_center", "") or ""),
                    "quadrant_group": str(row.get("quadrant_group", "") or ""),
                    "quadrant_center_group": str(row.get("quadrant_center_group", "") or ""),
                    "echo_mid_x_abs": row.get("echo_mid_x_abs", ""),
                    "echo_mid_y_abs": row.get("echo_mid_y_abs", ""),
                    "echo_rect_top_abs": int(row.get("echo_rect_top_abs", 0) or 0),
                    "echo_rect_left_abs": int(row.get("echo_rect_left_abs", 0) or 0),
                    "echo_rect_bottom_abs": int(row.get("echo_rect_bottom_abs", 0) or 0),
                    "echo_rect_right_abs": int(row.get("echo_rect_right_abs", 0) or 0),
                    "marker_top_crop": int(row.get("marker_top_crop", 0) or 0),
                    "marker_left_crop": int(row.get("marker_left_crop", 0) or 0),
                    "marker_bottom_crop": int(row.get("marker_bottom_crop", 0) or 0),
                    "marker_right_crop": int(row.get("marker_right_crop", 0) or 0),
                    "marker_top_abs": int(row.get("marker_top_abs", 0) or 0),
                    "marker_left_abs": int(row.get("marker_left_abs", 0) or 0),
                    "marker_bottom_abs": int(row.get("marker_bottom_abs", 0) or 0),
                    "marker_right_abs": int(row.get("marker_right_abs", 0) or 0),
                    "marker_cx_crop_norm": f"{float(row.get('marker_cx_crop_norm', 0.0) or 0.0):.6f}",
                    "marker_cy_crop_norm": f"{float(row.get('marker_cy_crop_norm', 0.0) or 0.0):.6f}",
                    "marker_echo_overlap_ratio": f"{float(row.get('marker_echo_overlap_ratio', 0.0) or 0.0):.6f}",
                    "marker_outside_echo_rect": int(row.get("marker_outside_echo_rect", 0) or 0),
                    "marker_outside_echo_top_px": int(row.get("marker_outside_echo_top_px", 0) or 0),
                    "marker_outside_echo_left_px": int(row.get("marker_outside_echo_left_px", 0) or 0),
                    "marker_outside_echo_bottom_px": int(row.get("marker_outside_echo_bottom_px", 0) or 0),
                    "marker_outside_echo_right_px": int(row.get("marker_outside_echo_right_px", 0) or 0),
                    "marker_outside_echo_max_px": int(row.get("marker_outside_echo_max_px", 0) or 0),
                    "marker_outside_echo_gap_px": int(row.get("marker_outside_echo_gap_px", 0) or 0),
                    "marker_outside_echo_severity": f"{float(row.get('marker_outside_echo_severity', 0.0) or 0.0):.6f}",
                    "marker_top_left_outside_echo_rect": int(row.get("marker_top_left_outside_echo_rect", 0) or 0),
                    "marker_echo_near_margin_px": f"{float(row.get('marker_echo_near_margin_px', 0.0) or 0.0):.6f}",
                }
            )

    lt_per_image_csv = output_dir / "lt_per_image_predictions.csv"
    with lt_per_image_csv.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(
            fh,
            fieldnames=[
                "folder_path",
                "folder_name",
                "image_index",
                "image_path",
                "pred_label",
                "pred_idx",
                "confidence",
                "prob_l",
                "prob_t",
                "crop_top",
                "crop_left",
                "crop_bottom",
                "crop_right",
                "crop_source",
            ],
        )
        writer.writeheader()
        for row in lt_per_image_rows:
            writer.writerow(
                {
                    "folder_path": str(row.get("folder_path", "") or ""),
                    "folder_name": str(row.get("folder_name", "") or ""),
                    "image_index": int(row.get("image_index", 0) or 0),
                    "image_path": str(row.get("image_path", "") or ""),
                    "pred_label": str(row.get("pred_label", "") or ""),
                    "pred_idx": int(row.get("pred_idx", 0) or 0),
                    "confidence": f"{float(row.get('confidence', 0.0) or 0.0):.6f}",
                    "prob_l": f"{float(row.get('prob_l', 0.0) or 0.0):.6f}",
                    "prob_t": f"{float(row.get('prob_t', 0.0) or 0.0):.6f}",
                    "crop_top": int(row.get("crop_top", 0) or 0),
                    "crop_left": int(row.get("crop_left", 0) or 0),
                    "crop_bottom": int(row.get("crop_bottom", 0) or 0),
                    "crop_right": int(row.get("crop_right", 0) or 0),
                    "crop_source": str(row.get("crop_source", "") or ""),
                }
            )

    rect_depth_per_image_csv = output_dir / "rect_depth_autonomous_predictions.csv"
    with rect_depth_per_image_csv.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(
            fh,
            fieldnames=[
                "folder_path",
                "folder_name",
                "folder_index",
                "image_index",
                "image_path",
                "status",
                "score",
                "ranker_score",
                "mode",
                "depth_mm",
                "left",
                "top",
                "right",
                "bottom",
                "ocr_text",
                "reason",
                "best_direct_score",
                "best_direct_text",
                "best_scale_score",
                "best_scale_value_mm",
                "best_scale_text",
                "candidates",
                "rect_depth_run_dir",
            ],
            extrasaction="ignore",
        )
        writer.writeheader()
        for row in rect_depth_per_image_rows:
            writer.writerow(row)

    scale_per_image_csv_path = output_dir / "scale_per_image_predictions.csv"
    if scale_per_image_rows:
        # The stage's own columns plus the folder identity: the union of the keys, so a new
        # column added in the stage travels here without touching this list.
        scale_fields = ["folder_path", "folder_name", "folder_index", "scale_run_dir"]
        for row in scale_per_image_rows:
            for key in row:
                if key not in scale_fields:
                    scale_fields.append(str(key))
        with scale_per_image_csv_path.open("w", encoding="utf-8", newline="") as fh:
            writer = csv.DictWriter(fh, fieldnames=scale_fields, extrasaction="ignore")
            writer.writeheader()
            for row in scale_per_image_rows:
                writer.writerow(row)

    preview_txt = output_dir / "folder_fss_head_preview.txt"
    lines: List[str] = []
    for p in predictions:
        lines.append(f"[{p.folder_name}]")
        lines.append(f"#01 {p.line_01_version}")
        lines.append(f"#02 {p.line_02_id_echo}")
        lines.append(f"#03 {p.line_03_id_probe}")
        lines.append(f"#04 {p.line_04_probe_type}")
        lines.append(f"#06 {p.line_06_video_input_code or ''}")
        lines.append(f"#07 {p.line_07_video_input_size_x}")
        lines.append(f"#08 {p.line_08_video_input_size_y}")
        lines.append(f"#09 {p.line_09_video_x_size}")
        lines.append(f"#10 {p.line_10_video_y_size}")
        lines.append(f"#11 {p.line_11_rect_echo}")
        lines.append(
            "line11_method="
            f"{p.line_11_method} margin_pct={p.line_11_rect_red_margin_pct:.2f} "
            f"winner_group={p.line_11_rect_red_winner_group or '-'} "
            f"routing={p.line_11_source}"
        )
        lines.append(f"#12 {p.line_12_group_orientation} ({p.line_12_group_orientation_label})")
        lines.append(
            "su_giu_majority="
            f"{p.su_giu_majority_label or '-'} "
            f"vote_ratio={p.su_giu_majority_vote_ratio:.3f} "
            f"mean_conf={p.su_giu_mean_confidence:.3f} "
            f"mean_p_su={p.su_giu_mean_prob_su:.3f} "
            f"mean_p_giu={p.su_giu_mean_prob_giu:.3f} "
            f"n={p.su_giu_images_predicted} "
            f"source={p.su_giu_source}"
        )
        lines.append(
            "lr_marker="
            f"majority={p.lr_marker_majority_label or '-'} "
            f"vote_ratio={p.lr_marker_majority_vote_ratio:.3f} "
            f"best={p.lr_marker_best_label or '-'} "
            f"best_score={p.lr_marker_best_score:.3f} "
            f"best_strategy={p.lr_marker_best_search_strategy or '-'} "
            f"n={p.lr_marker_images_predicted} "
            f"source={p.lr_marker_source}"
        )
        lines.append(
            "lt_majority="
            f"{p.lt_majority_label or '-'} "
            f"vote_ratio={p.lt_majority_vote_ratio:.3f} "
            f"mean_conf={p.lt_mean_confidence:.3f} "
            f"mean_p_l={p.lt_mean_prob_l:.3f} "
            f"mean_p_t={p.lt_mean_prob_t:.3f} "
            f"n={p.lt_images_predicted} "
            f"source={p.lt_source}"
        )
        lines.append(f"#13 {p.line_13_rect_name_echo}")
        lines.append(
            "line13_model="
            f"{p.line_13_model_route or '-'} "
            f"checkpoint={p.line_13_model_checkpoint or '-'} "
            f"source={p.line_13_source or '-'}"
        )
        lines.append(f"#14 {p.line_14_rect_name_probe}")
        lines.append(f"#16 {p.line_16_rect_orientation}")
        lines.append(
            "line16_orientation="
            f"source={p.line_16_source or '-'} "
            f"marker_boxes={p.line_16_marker_boxes_count} "
            "thresholds=pending"
        )
        lines.append(
            "rect_depth="
            f"status={p.rect_depth_status or '-'} "
            f"accepted={p.rect_depth_accepted_count}/{p.rect_depth_images_predicted} "
            f"review={p.rect_depth_review_count} reject={p.rect_depth_reject_count} "
            f"ratio={p.rect_depth_acceptance_ratio:.3f} "
            f"mode={p.rect_depth_majority_mode or '-'} "
            f"depths={p.rect_depth_unique_depths_json or '[]'} "
            f"source={p.rect_depth_source or '-'}"
        )
        lines.append(f"#18 {p.line_18_vect_depth}")
        lines.append(f"#19 {p.line_19_pixel_ratio_x}")
        lines.append(f"#20 {p.line_20_pixel_ratio_y}")
        lines.append(f"#21 {p.line_21_scale_line}")
        lines.append(
            "scala="
            f"status={p.scale_status or '-'} "
            f"accepted={p.scale_depths_accepted}/{p.scale_depths_total} "
            f"review={p.scale_depths_review} reject={p.scale_depths_reject} "
            f"interpolate={p.scale_depths_interpolated} "
            f"ratio={p.scale_acceptance_ratio:.3f} "
            f"frame={p.scale_frames_studied} "
            f"profilo={p.scale_profile or '-'} "
            f"righello_x={p.scale_ruler_x or '-'} "
            f"source={p.scale_source or '-'}"
        )
        lines.append(
            "rotation_cw="
            f"{p.rotation_deg_clockwise} source={p.rotation_source} "
            f"votes={p.rotation_votes_total}/{p.rotation_samples_checked} ratio={p.rotation_vote_ratio:.3f}"
        )
        lines.append(f"rotation_reason={p.rotation_decision_reason}")
        lines.append(f"status={p.status} reasons={p.review_reasons}")
        lines.append("")
    preview_txt.write_text("\n".join(lines), encoding="utf-8")

    rect_red_json = output_dir / "rect_red_pipeline_by_folder.json"
    rect_red_json.write_text(
        json.dumps(rect_red_by_folder, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    summary = {
        "dataset_root": dataset_root.as_posix(),
        "vendor_checkpoint": vendor_ckpt_path.as_posix(),
        "probe_checkpoint": probe_ckpt_path.as_posix(),
        "rect_checkpoint": rect_ckpt_path.as_posix(),
        "rect_vendor_map_path": rect_vendor_map_path.as_posix() if rect_vendor_map_path else "",
        "rect_vendor_map_loaded": rect_vendor_map_loaded,
        "rect_vendor_min_confidence": args.rect_vendor_min_confidence,
        "line13_model_enabled": line13_model is not None,
        "line13_checkpoint_global": line13_ckpt_path.as_posix() if line13_model is not None else "",
        "line13_image_size": int(line13_image_size) if line13_model is not None else 0,
        "line13_vendor_map_path": line13_vendor_map_path.as_posix() if line13_vendor_map_path else "",
        "line13_vendor_map_loaded": line13_vendor_map_loaded,
        "line13_vendor_min_confidence": float(args.line13_vendor_min_confidence),
        "line13_vendor_models_required_when_routing_enabled": not args.disable_line13_vendor_routing,
        "line13_postprocess_enabled": not args.disable_line13_postprocess,
        "line13_postprocess_iou_threshold": float(args.line13_postprocess_iou_threshold),
        "line13_postprocess_min_keep": int(args.line13_postprocess_min_keep),
        "line13_template_postprocess_enabled": not args.disable_line13_template_postprocess,
        "line13_template_max_images": int(args.line13_template_max_images),
        "line13_template_search_margin_pct": float(args.line13_template_search_margin_pct),
        "line13_template_min_std": float(args.line13_template_min_std),
        "line13_template_min_nonblack_ratio": float(args.line13_template_min_nonblack_ratio),
        "line13_template_min_score": float(args.line13_template_min_score),
        "line13_template_min_consensus_iou": float(args.line13_template_min_consensus_iou),
        "line13_template_min_iou_with_model": float(args.line13_template_min_iou_with_model),
        "line13_dark_trim_postprocess_enabled": not args.disable_line13_dark_trim_postprocess,
        "line13_dark_trim_threshold": float(args.line13_dark_trim_threshold),
        "line13_dark_trim_max_std": float(args.line13_dark_trim_max_std),
        "line13_dark_trim_max_trim_pct": float(args.line13_dark_trim_max_trim_pct),
        "line13_dark_trim_min_keep_ratio": float(args.line13_dark_trim_min_keep_ratio),
        "line13_tail_policy": "disabled_after_coords",
        "rect_red_enabled": not args.disable_rect_red_line11,
        "rect_red_margin_pct": float(args.rect_red_margin_pct),
        "rect_red_bright_thr": float(args.rect_red_bright_thr),
        "reference_manifest": manifest_ref_path.as_posix(),
        "fss_version_line_01": args.fss_version,
        "device": str(device),
        "low_confidence_policy": args.low_confidence_policy,
        "vendor_min_confidence": args.vendor_min_confidence,
        "vendor_min_margin": float(args.vendor_min_margin),
        "vendor_ocr_fallback_enabled": not args.disable_vendor_ocr_fallback,
        "vendor_ocr_samples": int(args.vendor_ocr_samples),
        "vendor_ocr_timeout_sec": float(args.vendor_ocr_timeout_sec),
        "vendor_ocr_lang": str(args.vendor_ocr_lang),
        "vendor_ocr_psm": int(args.vendor_ocr_psm),
        "vendor_ocr_min_word_conf": float(args.vendor_ocr_min_word_conf),
        "vendor_ocr_min_score": float(args.vendor_ocr_min_score),
        "vendor_ocr_min_score_delta": float(args.vendor_ocr_min_score_delta),
        "vendor_ocr_min_hits": int(args.vendor_ocr_min_hits),
        "probe_min_confidence": args.probe_min_confidence,
        "probe_type_router_enabled": not args.disable_probe_type_router,
        "probe_type_router_loaded": probe_type_router is not None,
        "probe_type_router_entries": int(probe_type_router_entries),
        "probe_type_summary_csv": probe_type_summary_csv.as_posix(),
        "probe_type_evidence_csv": probe_type_evidence_csv.as_posix(),
        "su_giu_rect_classifier_enabled": su_giu_model is not None,
        "su_giu_rect_checkpoint": su_giu_ckpt_path.as_posix() if su_giu_model is not None else "",
        "su_giu_rect_image_size": int(su_giu_image_size) if su_giu_model is not None else 0,
        "su_giu_rect_batch_size": int(su_giu_batch_size),
        "su_giu_rect_class_names": list(su_giu_class_names) if su_giu_model is not None else [],
        "lr_marker_enabled": not bool(args.disable_lr_marker_classical),
        "lr_marker_method": str(args.lr_marker_method),
        "lr_marker_classical_enabled": (not bool(args.disable_lr_marker_classical)) and str(args.lr_marker_method) == "classical",
        "lr_marker_bundle_enabled": (not bool(args.disable_lr_marker_classical)) and str(args.lr_marker_method) == "bundle",
        "lr_marker_bundle_dir": args.lr_marker_bundle_dir.expanduser().resolve().as_posix(),
        "lr_marker_bundle_zip": args.lr_marker_bundle_zip.expanduser().resolve().as_posix(),
        "lr_marker_bundle_library_root": args.lr_marker_bundle_library_root.expanduser().resolve().as_posix(),
        "lr_marker_bundle_vertical_delta": float(args.lr_marker_bundle_vertical_delta),
        "lr_marker_bundle_match_max_side": int(args.lr_marker_bundle_match_max_side),
        "lr_marker_bundle_selection_images": int(args.lr_marker_bundle_selection_images),
        "lr_marker_template_roots": [p.as_posix() for p in lr_marker_template_roots],
        "lr_marker_exclude_vendors": [str(v) for v in args.lr_marker_exclude_vendors],
        "lr_marker_min_match_score": float(args.lr_marker_min_match_score),
        "lr_marker_full_crop_fallback_threshold": float(args.lr_marker_full_crop_fallback_threshold),
        "lr_marker_expanded_search_threshold": float(args.lr_marker_expanded_search_threshold),
        "lr_marker_expanded_search_steps": [float(x) for x in lr_marker_expanded_search_steps],
        "lr_marker_blank_template_max_value": int(args.lr_marker_blank_template_max_value),
        "lr_marker_min_sugiu_confidence": float(args.lr_marker_min_sugiu_confidence),
        "lr_marker_dual_sugiu_search_confidence": float(LR_MARKER_DUAL_SUGIU_SEARCH_CONFIDENCE),
        "lr_marker_template_policy": str(args.lr_marker_template_policy),
        "lr_marker_review_file": (
            args.lr_marker_review_file.expanduser().resolve().as_posix()
            if args.lr_marker_review_file is not None
            else ""
        ),
        "lr_marker_manual_seeds_file": (
            args.lr_marker_manual_seeds_file.expanduser().resolve().as_posix()
            if args.lr_marker_manual_seeds_file is not None
            else ""
        ),
        "lr_marker_manual_seed_rows_loaded": int(len(lr_marker_manual_seed_rows)),
        "lr_marker_manual_seed_templates": lr_marker_manual_seed_template_meta,
        "lr_marker_template_cache_meta": {
            key: meta for key, (_templates, meta) in sorted(lr_marker_templates_cache.items())
        },
        "lt_rect_classifier_enabled": lt_model is not None,
        "lt_rect_checkpoint": lt_ckpt_path.as_posix() if lt_model is not None else "",
        "lt_rect_image_size": int(lt_image_size) if lt_model is not None else 0,
        "lt_rect_batch_size": int(lt_batch_size),
        "lt_rect_class_names": list(lt_class_names) if lt_model is not None else [],
        "lt_min_confidence": float(args.lt_min_confidence),
        "rect_depth_autonomous_enabled": not bool(args.disable_rect_depth_autonomous),
        "rect_depth_autonomous_max_images": int(args.rect_depth_max_images),
        "rect_depth_autonomous_max_candidates_per_sample": int(args.rect_depth_max_candidates_per_sample),
        "rect_depth_autonomous_ocr_timeout": float(args.rect_depth_ocr_timeout),
        "rect_depth_autonomous_subprocess_timeout_sec": float(args.rect_depth_subprocess_timeout_sec),
        "rect_depth_autonomous_scale_side_preference": str(args.rect_depth_scale_side_preference),
        "rect_depth_autonomous_min_accepted_ratio": float(args.rect_depth_min_accepted_ratio),
        "line13_min_support": args.line13_min_support,
        "line14_min_support": args.line14_min_support,
        "interactive_topk": args.interactive_topk,
        "rotation_normalization_enabled": not args.disable_rotation_normalization,
        "rotation_max_samples": int(args.rotation_max_samples),
        "rotation_osd_timeout_sec": float(args.rotation_osd_timeout_sec),
        "rotation_osd_min_confidence": float(args.rotation_osd_min_confidence),
        "rotation_min_votes": int(args.rotation_min_votes),
        "rotation_min_ratio": float(args.rotation_min_ratio),
        "rotation_ocr_validation_enabled": not args.disable_rotation_ocr_validation,
        "rotation_ocr_samples": int(args.rotation_ocr_samples),
        "rotation_ocr_timeout_sec": float(args.rotation_ocr_timeout_sec),
        "rotation_ocr_lang": str(args.rotation_ocr_lang),
        "rotation_ocr_psm": int(args.rotation_ocr_psm),
        "rotation_ocr_min_word_conf": float(args.rotation_ocr_min_word_conf),
        "rotation_ocr_score_word_bonus": float(args.rotation_ocr_score_word_bonus),
        "rotation_ocr_min_words": int(args.rotation_ocr_min_words),
        "rotation_ocr_min_score_delta": float(args.rotation_ocr_min_score_delta),
        "encoding_struct_path": encoding_struct_path.as_posix(),
        "manual_entries_written": manual_entries_written,
        "pipeline_stage_line11_rect_echo": (
            "segment_top_red_rect_with_median_fallback"
            if not args.disable_rect_red_line11
            else "vendor_routing_with_global_fallback"
        ),
        "pipeline_stage_rotation": "ocr_osd_majority_plus_ocr_validation_post_dedup",
        "pipeline_stage_su_giu_rect": (
            "per_frame_rect_crop_classifier"
            if su_giu_model is not None
            else "disabled"
        ),
        "pipeline_stage_lr_marker": (
            "bundle_vendor_template_library_after_official_rect_su_giu"
            if not bool(args.disable_lr_marker_classical) and str(args.lr_marker_method) == "bundle"
            else "classical_vendor_template_library_after_su_giu"
            if not bool(args.disable_lr_marker_classical)
            else "disabled"
        ),
        "pipeline_stage_lt_rect": (
            "per_frame_rect_crop_classifier"
            if lt_model is not None
            else "disabled"
        ),
        "pipeline_stage_rect_depth": (
            "autonomous_ocr_classical_ranker_after_official_rect"
            if not bool(args.disable_rect_depth_autonomous)
            else "disabled"
        ),
        "rect_sampling_mode": (
            "all_images_per_folder"
            if int(args.rect_sample_per_folder) <= 0
            else "uniform_subset_per_folder"
        ),
        "rect_sample_per_folder": int(args.rect_sample_per_folder),
        "pipeline_stage_probe": "classifier_temp_with_manual_fallback",
        "pipeline_stage_line04_probe_type": (
            "probe_type_router_plus_lt_split_after_probe_classifier"
            if lt_model is not None
            else "probe_type_router_after_probe_classifier"
        ),
        "pipeline_stage_line13_rect_name_echo": (
            (
                "vendor_specialized_line13_required_plus_template_gate_plus_dark_trim_tail_disabled"
                if not args.disable_line13_vendor_routing
                else "global_line13_override_plus_template_gate_plus_dark_trim_tail_disabled"
            )
            if line13_model is not None
            else "historical_resolver_line13_tail_disabled"
        ),
        "pipeline_stage_line14_rect_name_probe": "historical_resolver_temp_with_manual_fallback",
        "folders_scanned": len(folders),
        "folders_predicted": len(predictions),
        "images_total_raw": total_images_raw,
        "images_total_unique": total_images_unique,
        "images_duplicates_removed": total_duplicates_removed,
        "status_counts": dict(status_counter),
        "rotation_source_counts": dict(rotation_source_counter),
        "rotation_degree_counts": {str(k): int(v) for k, v in sorted(rotation_degree_counter.items())},
        "rect_routing_counts": dict(rect_route_counter),
        "rect_vendor_specialized_usage_by_vendor": dict(sorted(rect_vendor_specialized_usage.items())),
        "line13_route_counts": dict(line13_route_counter),
        "line13_vendor_specialized_usage_by_vendor": dict(sorted(line13_vendor_specialized_usage.items())),
        "line13_postprocess_mode_counts": dict(line13_postprocess_mode_counter),
        "line13_postprocess_boxes_total": int(line13_postprocess_total_boxes),
        "line13_postprocess_boxes_kept_total": int(line13_postprocess_total_kept),
        "line13_postprocess_boxes_dropped_total": int(line13_postprocess_total_dropped),
        "line13_template_postprocess_mode_counts": dict(line13_template_postprocess_mode_counter),
        "line13_template_postprocess_attempted_total": int(line13_template_postprocess_attempted_total),
        "line13_template_postprocess_applied_total": int(line13_template_postprocess_applied_total),
        "line13_dark_trim_mode_counts": dict(line13_dark_trim_mode_counter),
        "line13_dark_trim_attempted_total": int(line13_dark_trim_attempted_total),
        "line13_dark_trim_applied_total": int(line13_dark_trim_applied_total),
        "su_giu_source_counts": dict(su_giu_source_counter),
        "su_giu_majority_counts": dict(sorted(su_giu_majority_counter.items())),
        "su_giu_per_image_predictions_csv": su_giu_per_image_csv.as_posix(),
        "su_giu_per_image_predictions_count": len(su_giu_per_image_rows),
        "lr_marker_source_counts": dict(lr_marker_source_counter),
        "lr_marker_majority_counts": dict(sorted(lr_marker_majority_counter.items())),
        "lr_marker_per_image_predictions_csv": lr_marker_per_image_csv.as_posix(),
        "lr_marker_per_image_predictions_count": len(lr_marker_per_image_rows),
        "lt_source_counts": dict(lt_source_counter),
        "lt_majority_counts": dict(sorted(lt_majority_counter.items())),
        "lt_per_image_predictions_csv": lt_per_image_csv.as_posix(),
        "lt_per_image_predictions_count": len(lt_per_image_rows),
        "rect_depth_status_counts": dict(rect_depth_status_counter),
        "rect_depth_mode_counts": dict(rect_depth_mode_counter),
        "rect_depth_autonomous_predictions_csv": rect_depth_per_image_csv.as_posix(),
        "rect_depth_autonomous_predictions_count": len(rect_depth_per_image_rows),
        "rect_red_pipeline_json": rect_red_json.as_posix(),
        "rect_red_pipeline_folders": len(rect_red_by_folder),
        "scale_stage_enabled": not bool(args.disable_scale_stage),
        "scale_status_counts": dict(scale_status_counter),
        "scale_source_counts": dict(scale_source_counter),
        "scale_max_frames": int(args.scale_max_frames),
        "scale_min_accepted_ratio": float(args.scale_min_accepted_ratio),
        "scale_per_image_predictions_csv": scale_per_image_csv_path.as_posix(),
        "scale_per_image_predictions_count": len(scale_per_image_rows),
        "scale_depths_total": sum(p.scale_depths_total for p in predictions),
        "scale_depths_accepted": sum(p.scale_depths_accepted for p in predictions),
        "scale_depths_interpolated": sum(p.scale_depths_interpolated for p in predictions),
        "scale_line21_complete_folders": sum(
            1 for p in predictions if str(p.line_21_scale_line or "").strip()
        ),
        "warnings_count": len(warnings),
        "warnings": warnings[:200],
        "output_csv": csv_path.as_posix(),
        "output_preview": preview_txt.as_posix(),
    }
    summary_json = output_dir / "summary.json"
    summary_json.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"Predictions CSV: {csv_path}", flush=True)
    print(f"SU/GIU per-image CSV: {su_giu_per_image_csv}", flush=True)
    print(f"LR marker per-image CSV: {lr_marker_per_image_csv}", flush=True)
    print(f"L/T per-image CSV: {lt_per_image_csv}", flush=True)
    print(f"RECT_DEPTH autonomous CSV: {rect_depth_per_image_csv}", flush=True)
    print(f"Preview TXT: {preview_txt}", flush=True)
    print(f"Rect Red JSON: {rect_red_json}", flush=True)
    print(f"Summary JSON: {summary_json}", flush=True)
    print(f"Encoding struct updates: {encoding_struct_path}", flush=True)
    print(
        f"Folders predicted: {len(predictions)} | ok={status_counter['ok']} "
        f"review={status_counter['review']}",
        flush=True,
    )
    if warnings:
        print(f"Warnings: {len(warnings)} (vedi summary.json)", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
